"""Integration tests for the ingest -> query -> delete pipeline."""

import json
import tempfile
import zlib
from collections.abc import Generator
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from queue import Queue
from threading import Barrier, Event, RLock, get_ident
from unittest.mock import MagicMock, PropertyMock

import numpy as np
import pytest
from docx import Document as DocxDocument

from scinexusrag.config import (
    EmbeddingSettings,
    IngestionSettings,
    LLMSettings,
    Settings,
    StorageSettings,
)
from scinexusrag.ingestion import Chunk
from scinexusrag.ingestion.embedder import EmbeddingConfig, EmbeddingConfigError
from scinexusrag.pipeline import IngestResult, NexusRAG, SystemStats, get_scinexusrag
from scinexusrag.storage import VectorStoreError


@pytest.fixture
def temp_data_dir() -> Generator[Path, None, None]:
    with tempfile.TemporaryDirectory() as tmpdir:
        yield Path(tmpdir)


@pytest.fixture
def test_settings(temp_data_dir: Path) -> Settings:
    return Settings(
        data_dir=temp_data_dir,
        storage=StorageSettings(
            lancedb_path=temp_data_dir / "lancedb",
            table_name="test_documents",
        ),
        embedding=EmbeddingSettings(
            model="test-model",
            revision="a" * 40,
            device="cpu",
            batch_size=8,
        ),
        llm=LLMSettings(
            model="test-llm",
            base_url="http://localhost:11434",
            timeout=30,
        ),
        ingestion=IngestionSettings(
            chunk_size=200,
            chunk_overlap=50,
            min_chunk_size=20,
        ),
    )


@pytest.fixture
def mock_embedder():
    embedder = MagicMock()
    embedder.dimension = 384
    embedder.embedding_config = EmbeddingConfig("test-model", "a" * 40, "", "", True, 384)

    def mock_embed(texts, batch_size=None, show_progress=False):
        if isinstance(texts, str):
            texts = [texts]
        np.random.seed(zlib.crc32("".join(texts).encode()))
        return np.random.rand(len(texts), 384).astype(np.float32)

    embedder.embed = mock_embed
    return embedder


@pytest.fixture
def mock_llm():
    llm = MagicMock()
    llm.model = "test-llm"

    def mock_generate(prompt, **kwargs):
        if "What is" in prompt or "?" in prompt:
            return "The fixture records temperature at thirty-second intervals."
        return "Sensor A is the reference."

    llm.generate = mock_generate
    llm.is_available = MagicMock(return_value=True)
    return llm


@pytest.fixture
def scinexusrag_instance(test_settings: Settings, mock_embedder, mock_llm) -> NexusRAG:
    rag = NexusRAG(settings=test_settings)

    rag._embedder = mock_embedder
    rag._llm = mock_llm

    return rag


@pytest.fixture
def sample_text_file(temp_data_dir: Path) -> Path:
    file_path = temp_data_dir / "test_document.txt"
    content = """INTRODUCTION

This research document is a software test fixture about temperature sensors.
The reference sensor is labeled A, and the comparison sensor is labeled B.

Each reading has a timestamp, a sensor label and a temperature in degrees Celsius.

METHODS

The methodology reads both sensors at thirty-second intervals.
Readings with the same timestamp belong to the same measurement pair.

Store each measurement pair on one row, with separate columns for sensors A and B.

RESULTS

The fixture contains paired readings with matching timestamps.
Its values are synthetic test data, not measured research results.

Missing readings are represented by an empty sensor column.

DISCUSSION

Compare sensor readings only when both values are present.
Keep missing values separate from numeric temperature differences.

The reference sensor provides the label used in retrieval queries.

CONCLUSION

Both sensors use the same units and sampling interval.

"""
    file_path.write_text(content, encoding="utf-8")
    return file_path


@pytest.fixture
def sample_markdown_file(temp_data_dir: Path) -> Path:
    file_path = temp_data_dir / "test_paper.md"
    content = """# Scientific Paper on Testing

## Abstract

This is a Markdown fixture for document ingestion and retrieval tests.

## Introduction

### Background

The fixture uses nested headings to test section extraction.
Each section contains text that can be indexed as a passage.

### Objectives

1. Parse headings and paragraphs.
2. Index the resulting passages.
3. Retrieve a passage by its text.

## Methodology

### Data Preparation

Save the fixture as UTF-8 Markdown.
Keep headings on separate lines from their paragraphs.

### Testing Framework

The tests use real parsing and storage with fixed model responses.

## Results

Test assertions check stored passages, query sources and deletion.

## Conclusion

This fixture contains no benchmark measurements.
"""
    file_path.write_text(content, encoding="utf-8")
    return file_path


@pytest.fixture
def docs_directory(temp_data_dir: Path) -> Path:
    docs_dir = temp_data_dir / "documents"
    docs_dir.mkdir()

    (docs_dir / "doc1.txt").write_text(
        "Document 1 content. This is the first test document.\n"
        "It contains multiple sentences for proper chunking.\n"
        "Testing the ingestion of text files.\n"
    )

    (docs_dir / "doc2.txt").write_text(
        "Document 2 content. This is the second test document.\n"
        "It also has enough content for chunking.\n"
        "Multiple ingestion operations are tested.\n"
    )

    (docs_dir / "paper.md").write_text(
        "# Test Paper\n\n"
        "## Introduction\n"
        "This is a markdown document for testing.\n\n"
        "## Methods\n"
        "Testing with multiple formats.\n\n"
        "## Results\n"
        "The fixture has one Methods section.\n"
    )

    return docs_dir


class TestBasicPipelineWorkflow:
    def test_uppercase_evidence_is_ingested_with_exact_source_offsets(
        self, scinexusrag_instance: NexusRAG
    ):
        rag = scinexusrag_instance
        content = "THE MEASURED CURRENT IS 42 MA.\n\nCALIBRATION PASSED."

        result = rag.ingest_bytes(content.encode(), "measurements.txt", ".txt")

        assert result.success, result.error
        assert result.chunk_count > 0
        chunks = rag.vector_store.get_chunks_by_document(result.document_id)
        assert all("[" not in chunk.content for chunk in chunks)
        for line in content.splitlines():
            if line:
                assert sum(line in chunk.content for chunk in chunks) == 1
        for chunk in chunks:
            assert chunk.start_offset is not None and chunk.end_offset is not None
            assert content[chunk.start_offset : chunk.end_offset] == chunk.content
        sources = rag.bm25.retrieve("42 MA", top_k=1)
        assert len(sources) == 1
        assert "THE MEASURED CURRENT IS 42 MA." in sources[0].chunk.content

    def test_ingest_single_text_file(self, scinexusrag_instance: NexusRAG, sample_text_file: Path):
        result = scinexusrag_instance.ingest(sample_text_file)

        assert result.success is True
        assert result.filename == "test_document.txt"
        assert result.chunk_count > 0
        assert result.word_count > 0
        assert result.document_id
        assert result.error is None

    def test_ingest_markdown_file(self, scinexusrag_instance: NexusRAG, sample_markdown_file: Path):
        result = scinexusrag_instance.ingest(sample_markdown_file)

        assert result.success is True
        assert result.filename == "test_paper.md"
        assert result.chunk_count > 0
        assert result.word_count > 0

    @pytest.mark.parametrize("interleaved", [False, True])
    def test_docx_table_rows_remain_retrievable_after_reload(
        self, scinexusrag_instance: NexusRAG, interleaved: bool
    ):
        document = DocxDocument()
        before = "A paragraph before the measurement table."
        after = "A paragraph after the measurement table."
        if interleaved:
            document.add_heading("Measurements", level=1)
            document.add_paragraph(before)
        rows = [
            ["Sensor", "Reading", "Run"],
            ["calibration", "42.50", "007"],
            ["comparison", "", "-3.25"],
        ]
        table = document.add_table(rows=3, cols=3)
        for row, values in zip(table.rows, rows, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
        if interleaved:
            document.add_paragraph(after)
        upload = BytesIO()
        document.save(upload)
        rag = scinexusrag_instance

        result = rag.ingest_bytes(upload.getvalue(), "measurements.docx", ".docx")

        assert result.success, result.error
        assert result.chunk_count > 0
        stored = json.loads(rag.document_store._doc_path(result.document_id).read_text())
        content = stored["content"]
        numeric_row = "calibration | 42.50 | 007"
        assert numeric_row in content
        empty_cell_row = next(
            line for line in content.splitlines() if line.startswith("comparison")
        )
        assert [cell.strip() for cell in empty_cell_row.split("|")] == ["comparison", "", "-3.25"]
        if interleaved:
            assert content.index(before) < content.index(numeric_row) < content.index(after)
        assert stored["sections"]
        for section in stored["sections"]:
            assert content[section["start_offset"] : section["end_offset"]] == section["content"]

        restarted = NexusRAG(settings=rag.settings)
        restarted._embedder = rag.embedder
        restarted._llm = rag.llm
        matches = restarted.bm25.retrieve("42.50", top_k=1)
        assert len(matches) == 1
        match = matches[0].chunk
        assert numeric_row in match.content
        persisted = {
            chunk.id: chunk
            for chunk in restarted.vector_store.get_chunks_by_document(result.document_id)
        }[match.id]
        assert persisted.full_context == match.full_context
        start, end = persisted.start_offset, persisted.end_offset
        assert start is not None and end is not None
        assert numeric_row in content[start:end]
        assert persisted.context_before == content[max(0, start - 150) : start].strip()
        assert persisted.context_after == content[end : end + 150].strip()

        prepared = restarted.orchestrator.synthesizer.prepare(
            "What reading was recorded?", matches, {result.document_id: "measurements.docx"}
        )
        assert prepared.error is None
        assert prepared.sources[0].chunk_id == match.id
        assert (
            prepared.sources[0].full_content or prepared.sources[0].content
        ) == match.full_context
        assert numeric_row in prepared.prompt

    def test_full_workflow_ingest_and_query(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path
    ):

        ingest_result = scinexusrag_instance.ingest(sample_text_file)
        assert ingest_result.success is True
        document_id = ingest_result.document_id

        docs = scinexusrag_instance.list_documents()
        assert len(docs) == 1
        assert docs[0]["id"] == document_id

        stats = scinexusrag_instance.get_stats()
        assert stats.total_documents == 1
        assert stats.total_chunks == ingest_result.chunk_count

    def test_delete_document_workflow(self, scinexusrag_instance: NexusRAG, sample_text_file: Path):

        ingest_result = scinexusrag_instance.ingest(sample_text_file)
        document_id = ingest_result.document_id

        assert len(scinexusrag_instance.list_documents()) == 1

        deleted = scinexusrag_instance.delete_document(document_id)
        assert deleted is True

        assert len(scinexusrag_instance.list_documents()) == 0

        stats = scinexusrag_instance.get_stats()
        assert stats.total_documents == 0
        assert stats.total_chunks == 0

    def test_delete_nonexistent_document(self, scinexusrag_instance: NexusRAG):
        deleted = scinexusrag_instance.delete_document("nonexistent_id")
        assert deleted is False

    def test_empty_stream_ends_with_the_authoritative_answer(self, scinexusrag_instance: NexusRAG):
        output = list(scinexusrag_instance.query_streaming("What was measured?"))
        events = [json.loads(part.removeprefix("\ndata: ")) for part in output[1:]]

        assert len(output) == 3
        assert events[0] == {
            "event": "verification",
            "citations_valid": [],
            "faithfulness": None,
            "warnings": [output[0]],
        }
        assert events[1] == {"event": "verified_answer", "answer": output[0]}

    @pytest.mark.parametrize("streaming", [False, True])
    def test_query_propagates_failed_storage_reads_before_returning_an_answer(
        self, scinexusrag_instance: NexusRAG, monkeypatch, streaming
    ):
        rag = scinexusrag_instance
        failure = OSError("Index read failed")
        monkeypatch.setattr(rag.vector_store.table, "count_rows", MagicMock(side_effect=failure))
        rag._orchestrator = MagicMock()

        with pytest.raises(VectorStoreError, match="row count") as raised:
            if streaming:
                list(rag.query_streaming("What was measured?"))
            else:
                rag.query("What was measured?")

        assert raised.value.__cause__ is failure
        rag._orchestrator.query.assert_not_called()
        rag._orchestrator.query_streaming.assert_not_called()

    @pytest.mark.parametrize("lazy", [False, True])
    def test_failed_sparse_read_can_retry_without_publishing_an_empty_index(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path, monkeypatch, lazy
    ):
        original = scinexusrag_instance
        result = original.ingest(sample_text_file)
        assert result.success
        rag = NexusRAG(settings=original.settings)
        rag._embedder = original.embedder
        get_all_chunks = rag.vector_store.get_all_chunks
        monkeypatch.setattr(
            rag.vector_store,
            "get_all_chunks",
            MagicMock(side_effect=VectorStoreError("Cannot read passages from the vector index.")),
        )

        with pytest.raises(VectorStoreError, match="read passages"):
            if lazy:
                _ = rag.bm25
            else:
                rag._rebuild_bm25_index()

        assert rag._bm25 is None
        monkeypatch.setattr(rag.vector_store, "get_all_chunks", get_all_chunks)
        assert rag.bm25.count() == result.chunk_count

    @pytest.mark.parametrize("stage", ["enumeration", "dense_before", "dense_after", "sparse"])
    def test_partial_deletion_reconciles_indexes_and_can_be_retried(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path, monkeypatch, stage
    ):
        rag = scinexusrag_instance
        result = rag.ingest(sample_text_file)
        other = rag.ingest_bytes(
            b"A separate document records voltage readings for sensor C.", "voltage.txt", ".txt"
        )
        assert result.success and other.success
        other_ids = rag.vector_store.get_chunk_ids_by_document(other.document_id)
        failure = OSError(f"Failed during {stage}")
        if stage == "enumeration":
            owner, method = rag.vector_store, "get_chunk_ids_by_document"
        elif stage.startswith("dense"):
            owner, method = rag.vector_store, "delete_by_document"
        else:
            owner, method = rag.bm25, "remove"
        operation = getattr(owner, method)

        def fail(*args, **kwargs):
            if stage == "dense_after":
                operation(*args, **kwargs)
            raise failure

        monkeypatch.setattr(owner, method, fail)
        with pytest.raises(OSError, match=f"Failed during {stage}"):
            rag.delete_document(result.document_id)
        monkeypatch.setattr(owner, method, operation)

        assert rag.document_store.exists(result.document_id)
        assert rag.document_store.exists(other.document_id)
        assert rag.vector_store.get_chunk_ids_by_document(other.document_id) == other_ids
        assert rag.bm25.count() == rag.vector_store.count()
        if stage in {"dense_after", "sparse"}:
            assert rag.vector_store.get_chunk_ids_by_document(result.document_id) == set()
            assert rag.bm25.retrieve("thirty-second") == []
        else:
            assert rag.bm25.count() == result.chunk_count + other.chunk_count

        assert rag.delete_document(result.document_id)
        assert not rag.document_store.exists(result.document_id)
        assert rag.bm25.count() == rag.vector_store.count() == other.chunk_count
        assert {chunk.id for chunk in rag.vector_store.get_all_chunks()} == other_ids
        restarted = NexusRAG(settings=rag.settings)
        restarted._embedder = rag.embedder
        assert restarted.bm25.count() == other.chunk_count

    def test_failed_deletion_recovery_discards_stale_sparse_handles(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path, monkeypatch
    ):
        rag = scinexusrag_instance
        result = rag.ingest(sample_text_file)
        assert result.success
        sparse = rag.bm25
        rag._orchestrator = MagicMock()
        get_all_chunks = rag.vector_store.get_all_chunks
        monkeypatch.setattr(
            sparse, "remove", MagicMock(side_effect=OSError("Sparse deletion failed"))
        )
        monkeypatch.setattr(
            rag.vector_store,
            "get_all_chunks",
            MagicMock(side_effect=VectorStoreError("Storage remains unavailable")),
        )

        with pytest.raises(OSError, match="Sparse deletion failed"):
            rag.delete_document(result.document_id)

        assert rag._bm25 is None
        assert rag._orchestrator is None
        assert rag.document_store.exists(result.document_id)
        assert rag.vector_store.count() == 0
        monkeypatch.setattr(rag.vector_store, "get_all_chunks", get_all_chunks)
        assert rag.bm25 is not sparse
        assert rag.bm25.count() == 0
        assert rag.delete_document(result.document_id)
        assert rag.document_store.count() == 0

    def test_delete_document_with_more_than_100000_chunks_removes_all_sparse_entries(
        self, scinexusrag_instance: NexusRAG
    ):
        from scinexusrag.ingestion import Chunk

        rag = scinexusrag_instance
        rag._embedder.embedding_config = replace(rag.embedder.embedding_config, dimension=4)
        document = rag.parser.parse_bytes(b"Current is 42 mA.", "large.txt", ".txt")
        count = 100_001
        chunks = [
            Chunk(id=f"large-{i}", content="Current is 42 mA.", document_id=document.id)
            for i in range(count)
        ]
        rag.document_store.add(document)
        rag.vector_store.add(chunks, np.ones((count, 4), dtype=np.float32))
        assert rag.bm25.count() == count

        assert rag.delete_document(document.id)

        assert rag.vector_store.count() == rag.bm25.count() == 0
        assert rag.bm25.retrieve("42 mA") == []
        assert rag.document_store.count() == 0

    @pytest.mark.parametrize("failure", ["unlink", "index_save"])
    def test_delete_metadata_failure_keeps_retrieval_indexes_synced(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path, monkeypatch, failure: str
    ):
        rag = scinexusrag_instance
        result = rag.ingest(sample_text_file)
        assert result.success
        store = rag.document_store
        error = OSError("document metadata unavailable")
        if failure == "unlink":
            monkeypatch.setattr(Path, "unlink", MagicMock(side_effect=error))
        else:
            monkeypatch.setattr(store, "_save_index", MagicMock(side_effect=error))

        with pytest.raises(OSError, match="document metadata unavailable"):
            rag.delete_document(result.document_id)

        assert rag.vector_store.count() == rag.bm25.count() == 0
        assert store.exists(result.document_id) is (failure == "unlink")
        assert rag.query("What methods were used?").sources == []


class TestMultiFileIngestion:
    def test_ingest_directory(self, scinexusrag_instance: NexusRAG, docs_directory: Path):
        results = scinexusrag_instance.ingest_directory(docs_directory)

        assert len(results) == 3
        assert all(r.success for r in results)

        docs = scinexusrag_instance.list_documents()
        assert len(docs) == 3

        for result in results:
            assert result.chunk_count > 0

    def test_ingest_directory_stats(self, scinexusrag_instance: NexusRAG, docs_directory: Path):
        results = scinexusrag_instance.ingest_directory(docs_directory)
        total_chunks = sum(r.chunk_count for r in results)
        total_words = sum(r.word_count for r in results)

        stats = scinexusrag_instance.get_stats()
        assert stats.total_documents == 3
        assert stats.total_chunks == total_chunks
        assert stats.total_words == total_words

    def test_ingest_directory_with_unsupported_format(
        self, scinexusrag_instance: NexusRAG, temp_data_dir: Path
    ):
        test_dir = temp_data_dir / "mixed_docs"
        test_dir.mkdir()

        (test_dir / "valid.txt").write_text("Valid content here.\n" * 10)
        (test_dir / "unsupported.xyz").write_text("This should be skipped.\n")

        results = scinexusrag_instance.ingest_directory(test_dir)

        assert len(results) == 2
        by_name = {r.filename: r for r in results}
        assert by_name["valid.txt"].success is True
        assert by_name["unsupported.xyz"].success is False
        assert "Unsupported file type" in (by_name["unsupported.xyz"].error or "")


class TestConcurrentWrites:
    @pytest.mark.parametrize("blocked_stage", ["vector_count", "llm_check"])
    def test_stats_snapshot_serializes_stores_without_locking_network_checks(
        self, scinexusrag_instance: NexusRAG, monkeypatch, blocked_stage
    ):
        rag = scinexusrag_instance
        store = rag.vector_store
        count_rows = store.count
        entered, release = Event(), Event()
        lock = _ObservedWriteLock()
        rag._write_lock = lock

        def hold(operation):
            entered.set()
            assert release.wait(timeout=5)
            return operation()

        if blocked_stage == "vector_count":
            monkeypatch.setattr(store, "count", lambda: hold(count_rows))
        else:
            monkeypatch.setattr(rag.llm, "is_available", lambda: hold(lambda: True))

        with ThreadPoolExecutor(max_workers=2) as workers:
            stats_result = workers.submit(rag.get_stats)
            try:
                assert entered.wait(timeout=5)
                ingest_result = workers.submit(
                    rag.ingest_bytes,
                    b"A new document contains current readings from sensor A.",
                    "current.txt",
                    ".txt",
                )
                if blocked_stage == "vector_count":
                    lock.wait_for_threads(2)
                    assert not ingest_result.done()
                else:
                    assert ingest_result.result(timeout=5).success
            finally:
                release.set()
            snapshot = stats_result.result(timeout=5)
            assert ingest_result.result(timeout=5).success

        assert snapshot.total_documents == snapshot.total_chunks == snapshot.total_words == 0
        assert rag.document_store.count() == 1
        assert rag.vector_store.count() > 0

    @pytest.mark.parametrize("first_operation", ["ingest", "list_documents"])
    def test_ingest_during_cold_index_load(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path, monkeypatch, first_operation
    ):
        import threading
        from concurrent.futures import ThreadPoolExecutor

        rag = scinexusrag_instance
        rag.vector_store.count()
        store = rag.document_store
        load_index = store._load_index
        index_loaded = threading.Event()
        release_index = threading.Event()
        second_started = threading.Event()

        def delayed_load():
            index = load_index()
            if not index_loaded.is_set():
                index_loaded.set()
                assert release_index.wait(timeout=10)
            return index

        def second_ingest():
            second_started.set()
            return rag.ingest(sample_text_file)

        monkeypatch.setattr(store, "_load_index", delayed_load)
        with ThreadPoolExecutor(max_workers=2) as pool:
            first = (
                pool.submit(rag.ingest, sample_text_file)
                if first_operation == "ingest"
                else pool.submit(rag.list_documents)
            )
            try:
                assert index_loaded.wait(timeout=5)
                second = pool.submit(second_ingest)
                assert second_started.wait(timeout=5)
                with pytest.raises(TimeoutError):
                    second.result(timeout=1)
            finally:
                release_index.set()
            first_result = first.result(timeout=10)
            results = [second.result(timeout=10)]
            if first_operation == "ingest":
                results.append(first_result)

        successful = [result for result in results if result.success]
        rejected = [result for result in results if not result.success]
        assert len(successful) == 1
        if first_operation == "ingest":
            assert len(rejected) == 1
            assert rejected[0].error == "Document already exists"
        else:
            assert not rejected
        assert store.count() == 1
        assert rag.vector_store.count() == successful[0].chunk_count
        assert rag.bm25.count() == successful[0].chunk_count

    def test_parallel_duplicate_ingest_preserves_document(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path, monkeypatch
    ):
        import threading
        from concurrent.futures import ThreadPoolExecutor

        rag = scinexusrag_instance
        rag.vector_store.count()
        rag.document_store.count()
        embed = rag.embedder.embed
        barrier = threading.Barrier(2)

        def simultaneous_embed(*args, **kwargs):
            result = embed(*args, **kwargs)
            barrier.wait(timeout=10)
            return result

        monkeypatch.setattr(rag.embedder, "embed", simultaneous_embed)
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(rag.ingest, sample_text_file) for _ in range(2)]
            results = [future.result(timeout=20) for future in futures]

        successful = [result for result in results if result.success]
        rejected = [result for result in results if not result.success]
        assert len(successful) == len(rejected) == 1
        assert rejected[0].error == "Document already exists"
        assert rag.document_store.count() == 1
        assert rag.vector_store.count() == successful[0].chunk_count
        assert rag.bm25.count() == successful[0].chunk_count

    def test_parallel_ingest_keeps_indexes_consistent(
        self, scinexusrag_instance: NexusRAG, temp_data_dir: Path
    ):

        import threading

        files = []
        for i in range(4):
            p = temp_data_dir / f"doc_{i}.txt"
            p.write_text(f"Document number {i}.\n" + f"Unique sentence {i} repeated often.\n" * 30)
            files.append(p)

        barrier = threading.Barrier(len(files))
        results: list[IngestResult] = []

        def worker(fp: Path) -> None:
            barrier.wait()
            results.append(scinexusrag_instance.ingest(fp))

        threads = [threading.Thread(target=worker, args=(f,)) for f in files]
        for t in threads:
            t.start()
        for t in threads:
            t.join()

        assert all(r.success for r in results)
        total_chunks = sum(r.chunk_count for r in results)
        assert scinexusrag_instance.vector_store.count() == total_chunks
        assert scinexusrag_instance.bm25.count() == total_chunks


class TestDuplicateDetection:
    def test_duplicate_file_detection(self, scinexusrag_instance: NexusRAG, sample_text_file: Path):

        result1 = scinexusrag_instance.ingest(sample_text_file)
        assert result1.success is True

        result2 = scinexusrag_instance.ingest(sample_text_file)
        assert result2.success is False
        assert result2.error == "Document already exists"
        assert result2.chunk_count == 0

        docs = scinexusrag_instance.list_documents()
        assert len(docs) == 1

    def test_duplicate_in_batch_ingest(self, scinexusrag_instance: NexusRAG, temp_data_dir: Path):
        docs_dir = temp_data_dir / "dup_test"
        docs_dir.mkdir()

        content = "Test content for duplicate detection.\n" * 5
        (docs_dir / "doc_v1.txt").write_text(content)
        (docs_dir / "doc_v2.txt").write_text(content)

        results = scinexusrag_instance.ingest_directory(docs_dir)

        assert len(results) == 2

        successful = [r for r in results if r.success]
        assert len(successful) >= 1


class TestClearOperations:
    def test_clear_metadata_failure_keeps_retrieval_indexes_synced(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path, monkeypatch
    ):
        rag = scinexusrag_instance
        assert rag.ingest(sample_text_file).success
        monkeypatch.setattr(
            rag.document_store,
            "clear",
            MagicMock(side_effect=OSError("document metadata unavailable")),
        )

        with pytest.raises(OSError, match="document metadata unavailable"):
            rag.clear_all()

        assert rag.vector_store.count() == rag.bm25.count() == 0
        assert rag.query("What methods were used?").sources == []

    def test_clear_then_reingest_updates_existing_retriever(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path
    ):
        rag = scinexusrag_instance
        assert rag.ingest(sample_text_file).success
        sparse = rag.orchestrator.retriever.base.sparse
        assert sparse.retrieve("research")

        rag.clear_all()
        assert not sparse.retrieve("research")
        result = rag.ingest(sample_text_file)

        assert result.success
        assert sparse.count() == result.chunk_count
        assert sparse.retrieve("research")

    def test_clear_all_documents(self, scinexusrag_instance: NexusRAG, docs_directory: Path):

        scinexusrag_instance.ingest_directory(docs_directory)
        assert scinexusrag_instance.get_stats().total_documents == 3

        scinexusrag_instance.clear_all()

        assert scinexusrag_instance.get_stats().total_documents == 0
        assert scinexusrag_instance.get_stats().total_chunks == 0
        assert len(scinexusrag_instance.list_documents()) == 0

    def test_clear_all_then_reingest(self, scinexusrag_instance: NexusRAG, sample_text_file: Path):

        result1 = scinexusrag_instance.ingest(sample_text_file)
        assert result1.success is True

        scinexusrag_instance.clear_all()
        assert scinexusrag_instance.get_stats().total_documents == 0

        result2 = scinexusrag_instance.ingest(sample_text_file)
        assert result2.success is True


class TestStatisticsAndMetadata:
    def test_get_stats_empty_system(self, scinexusrag_instance: NexusRAG):
        stats = scinexusrag_instance.get_stats()

        assert stats.total_documents == 0
        assert stats.total_chunks == 0
        assert stats.total_words == 0
        assert stats.llm_available is True
        assert stats.embedding_model == "test-model"
        assert stats.llm_model == "test-llm"

    def test_get_stats_after_ingestion(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path
    ):
        result = scinexusrag_instance.ingest(sample_text_file)

        stats = scinexusrag_instance.get_stats()
        assert stats.total_documents == 1
        assert stats.total_chunks == result.chunk_count
        assert stats.total_words > 0
        assert "tmp" in str(stats.storage_path).lower()

    def test_list_documents_empty(self, scinexusrag_instance: NexusRAG):
        docs = scinexusrag_instance.list_documents()
        assert docs == []

    def test_list_documents_with_metadata(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path
    ):
        ingest_result = scinexusrag_instance.ingest(sample_text_file)

        docs = scinexusrag_instance.list_documents()
        assert len(docs) == 1

        doc = docs[0]
        assert doc["id"] == ingest_result.document_id
        assert "filename" in doc or "original_filename" in doc

    def test_stats_with_multiple_documents(
        self, scinexusrag_instance: NexusRAG, docs_directory: Path
    ):
        results = scinexusrag_instance.ingest_directory(docs_directory)

        stats = scinexusrag_instance.get_stats()
        assert stats.total_documents == len(results)

        total_chunks_expected = sum(r.chunk_count for r in results)
        assert stats.total_chunks == total_chunks_expected


class TestErrorHandling:
    def test_ingest_empty_file(self, scinexusrag_instance: NexusRAG, temp_data_dir: Path):
        empty_file = temp_data_dir / "empty.txt"
        empty_file.write_text("")

        result = scinexusrag_instance.ingest(empty_file)

        assert result.success is False
        assert result.error == "No content extracted"
        assert result.chunk_count == 0

    def test_ingest_nonexistent_file(self, scinexusrag_instance: NexusRAG):
        result = scinexusrag_instance.ingest("/nonexistent/path/file.txt")

        assert result.success is False
        assert "Ingestion failed" in result.error or "FileNotFoundError" in result.error

    def test_ingest_unsupported_format(self, scinexusrag_instance: NexusRAG, temp_data_dir: Path):
        unsupported = temp_data_dir / "file.xyz"
        unsupported.write_text("Some content")

        result = scinexusrag_instance.ingest(unsupported)

        assert result.success is False
        assert "Ingestion failed" in result.error

    def test_ingest_directory_nonexistent(self, scinexusrag_instance: NexusRAG):
        with pytest.raises(ValueError, match="Not a directory"):
            scinexusrag_instance.ingest_directory("/nonexistent/directory")

    def test_ingest_bytes_empty(self, scinexusrag_instance: NexusRAG):
        result = scinexusrag_instance.ingest_bytes(b"", "empty.txt", ".txt")

        assert result.success is False
        assert result.chunk_count == 0

    def test_ingest_bytes_valid(self, scinexusrag_instance: NexusRAG):
        content = b"This is test content for bytes ingestion.\n" * 5
        result = scinexusrag_instance.ingest_bytes(content, "test.txt", ".txt")

        assert result.success is True
        assert result.filename == "test.txt"
        assert result.chunk_count > 0

    def test_ingest_bytes_without_extension(self, scinexusrag_instance: NexusRAG):
        content = b"Test content.\n" * 5

        result = scinexusrag_instance.ingest_bytes(content, "test.txt", "txt")

        assert result.success is True


class TestIngestResult:
    def test_ingest_result_success(self, sample_text_file: Path):
        result = IngestResult(
            document_id="doc_123",
            filename="test.txt",
            chunk_count=5,
            word_count=150,
            success=True,
        )

        assert result.document_id == "doc_123"
        assert result.filename == "test.txt"
        assert result.chunk_count == 5
        assert result.word_count == 150
        assert result.success is True
        assert result.error is None

    def test_ingest_result_failure(self):
        result = IngestResult(
            document_id="",
            filename="bad.xyz",
            chunk_count=0,
            word_count=0,
            success=False,
            error="Unsupported format",
        )

        assert result.success is False
        assert result.error == "Unsupported format"
        assert result.chunk_count == 0


class TestSystemStats:
    def test_system_stats_creation(self, test_settings: Settings):
        stats = SystemStats(
            total_documents=5,
            total_chunks=42,
            total_words=15000,
            storage_path="/data/storage",
            llm_model="llama3",
            embedding_model="all-MiniLM-L6-v2",
            llm_available=True,
        )

        assert stats.total_documents == 5
        assert stats.total_chunks == 42
        assert stats.total_words == 15000
        assert stats.llm_available is True

    def test_system_stats_no_llm(self, test_settings: Settings):
        stats = SystemStats(
            total_documents=2,
            total_chunks=10,
            total_words=5000,
            storage_path="/data/storage",
            llm_model="offline-model",
            embedding_model="offline-embeddings",
            llm_available=False,
        )

        assert stats.llm_available is False


class TestLazyLoading:
    def test_components_not_loaded_initially(self, scinexusrag_instance: NexusRAG):

        assert scinexusrag_instance._parser is None
        assert scinexusrag_instance._chunker is None
        assert scinexusrag_instance._vector_store is None
        assert scinexusrag_instance._document_store is None
        assert scinexusrag_instance._bm25 is None

    def test_parser_lazy_loads_on_access(self, scinexusrag_instance: NexusRAG):
        assert scinexusrag_instance._parser is None
        parser = scinexusrag_instance.parser
        assert parser is not None
        assert scinexusrag_instance._parser is parser

    def test_chunker_lazy_loads_on_access(self, scinexusrag_instance: NexusRAG):
        assert scinexusrag_instance._chunker is None
        chunker = scinexusrag_instance.chunker
        assert chunker is not None
        assert scinexusrag_instance._chunker is chunker

    def test_document_store_lazy_loads_on_access(self, scinexusrag_instance: NexusRAG):
        assert scinexusrag_instance._document_store is None
        store = scinexusrag_instance.document_store
        assert store is not None
        assert scinexusrag_instance._document_store is store


class TestSingletonBehavior:
    def test_get_scinexusrag_returns_same_instance(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            settings = Settings(data_dir=Path(tmpdir))
            instance1 = get_scinexusrag(settings)
            instance2 = get_scinexusrag()

            assert instance1 is instance2

    def test_get_scinexusrag_with_different_settings(self):
        with tempfile.TemporaryDirectory() as tmpdir1, tempfile.TemporaryDirectory() as tmpdir2:
            settings1 = Settings(data_dir=Path(tmpdir1))
            settings2 = Settings(data_dir=Path(tmpdir2))

            instance1 = get_scinexusrag(settings1)
            instance2 = get_scinexusrag(settings2)

            assert instance1 is instance2


class TestIngestBytes:
    def test_ingest_bytes_text(self, scinexusrag_instance: NexusRAG):
        content = b"""
        # Test Document

        Sensor A records temperature in degrees Celsius.
        Each reading includes a timestamp.

        Sensor B uses the same sampling interval.
        Readings with matching timestamps form a measurement pair.
        """

        result = scinexusrag_instance.ingest_bytes(content, "bytes_doc.txt", ".txt")

        assert result.success is True
        assert result.filename == "bytes_doc.txt"
        assert result.chunk_count > 0

    def test_ingest_bytes_md_format(self, scinexusrag_instance: NexusRAG):
        content = b"""# Markdown Test

## Section 1
Content here for testing.

## Section 2
More content for proper chunking.
"""

        result = scinexusrag_instance.ingest_bytes(content, "test.md", ".md")

        assert result.success is True
        assert result.chunk_count > 0

    def test_ingest_bytes_duplicate_detection(self, scinexusrag_instance: NexusRAG):
        content = b"Unique test content for duplicate detection.\n" * 5

        result1 = scinexusrag_instance.ingest_bytes(content, "dup.txt", ".txt")
        assert result1.success is True

        result2 = scinexusrag_instance.ingest_bytes(content, "dup.txt", ".txt")
        assert result2.success is False
        assert result2.document_id == result1.document_id
        assert "already exists" in (result2.error or "")


class TestEdgeCases:
    def test_ingest_very_short_document(self, scinexusrag_instance: NexusRAG, temp_data_dir: Path):
        short_file = temp_data_dir / "short.txt"
        short_file.write_text("Short content.")

        result = scinexusrag_instance.ingest(short_file)

        assert isinstance(result.success, bool)

    def test_ingest_large_document(self, scinexusrag_instance: NexusRAG, temp_data_dir: Path):
        large_file = temp_data_dir / "large.txt"

        large_content = "This is a test sentence that will be repeated many times. " * 200
        large_file.write_text(large_content)

        result = scinexusrag_instance.ingest(large_file)

        assert result.success is True
        assert result.chunk_count > 1

    def test_concurrent_operations_sequence(
        self, scinexusrag_instance: NexusRAG, temp_data_dir: Path
    ):
        file1 = temp_data_dir / "doc1.txt"
        file2 = temp_data_dir / "doc2.txt"

        content1 = "Document 1 content. " * 20
        content2 = "Document 2 content. " * 20

        file1.write_text(content1)
        file2.write_text(content2)

        result1 = scinexusrag_instance.ingest(file1)
        result2 = scinexusrag_instance.ingest(file2)

        assert result1.success is True
        assert result2.success is True

        docs = scinexusrag_instance.list_documents()
        assert len(docs) == 2

        deleted = scinexusrag_instance.delete_document(result1.document_id)
        assert deleted is True

        docs = scinexusrag_instance.list_documents()
        assert len(docs) == 1


class TestIntegrationWithRealComponents:
    def test_restart_rebuilds_sparse_index_from_persisted_chunks(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path
    ):
        rag = scinexusrag_instance
        result = rag.ingest(sample_text_file)
        assert result.success

        restarted = NexusRAG(settings=rag.settings)
        restarted._embedder = rag.embedder
        restarted._llm = rag.llm

        assert restarted.bm25.count() == result.chunk_count
        assert restarted.bm25.retrieve("research")
        assert restarted.vector_store.list_documents() == [result.document_id]

    def test_with_real_parser_and_chunker(
        self, test_settings: Settings, sample_text_file: Path, mock_embedder, mock_llm
    ):
        rag = NexusRAG(settings=test_settings)
        rag._embedder = mock_embedder
        rag._llm = mock_llm

        result = rag.ingest(sample_text_file)

        assert result.success is True
        assert result.chunk_count > 0

        stats = rag.get_stats()
        assert stats.total_chunks == result.chunk_count

    def test_document_store_persistence(
        self, test_settings: Settings, sample_text_file: Path, mock_embedder, mock_llm
    ):
        rag = NexusRAG(settings=test_settings)
        rag._embedder = mock_embedder
        rag._llm = mock_llm

        ingest_result = rag.ingest(sample_text_file)
        assert ingest_result.success is True

        docs = rag.list_documents()
        assert len(docs) == 1

        doc = docs[0]
        assert doc["id"] == ingest_result.document_id


class TestParametrized:
    @pytest.mark.parametrize(
        "filename,content",
        [
            ("test1.txt", "Test content one.\n" * 10),
            ("test2.txt", "Test content two.\n" * 10),
            ("test3.txt", "Different content.\n" * 10),
        ],
    )
    def test_ingest_multiple_variations(
        self, scinexusrag_instance: NexusRAG, temp_data_dir: Path, filename: str, content: str
    ):
        file_path = temp_data_dir / filename
        file_path.write_text(content)

        result = scinexusrag_instance.ingest(file_path)

        assert result.success is True
        assert result.filename == filename

    @pytest.mark.parametrize(
        "format_ext,format_name",
        [
            (".txt", "text"),
            (".md", "markdown"),
        ],
    )
    def test_various_formats(
        self, scinexusrag_instance: NexusRAG, temp_data_dir: Path, format_ext: str, format_name: str
    ):
        file_path = temp_data_dir / f"doc{format_ext}"

        if format_ext == ".txt":
            content = "Test content for text format.\n" * 10
        else:
            content = "# Test\n\nContent for markdown format.\n" * 5

        file_path.write_text(content)

        result = scinexusrag_instance.ingest(file_path)

        assert result.success is True


class TestIngestRollback:
    def test_failed_count_during_ingest_keeps_dense_rows_attached_to_a_document(
        self, scinexusrag_instance: NexusRAG, monkeypatch
    ):
        rag = scinexusrag_instance
        table = rag.vector_store.table
        count_rows = table.count_rows
        failure = OSError("Row count temporarily unavailable")
        monkeypatch.setattr(table, "count_rows", MagicMock(side_effect=failure))

        with pytest.raises(VectorStoreError, match="row count") as raised:
            rag.ingest_bytes(b"The measured current is forty-two milliamps.", "current.txt", ".txt")

        assert raised.value.__cause__ is failure
        monkeypatch.setattr(table, "count_rows", count_rows)
        chunks = rag.vector_store.get_all_chunks()
        documents = rag.document_store.list_with_metadata()
        assert chunks and len(documents) == 1
        assert {chunk.document_id for chunk in chunks} == set(documents)
        assert rag.bm25.count() == rag.vector_store.count() == len(chunks)
        document_id = chunks[0].document_id
        restarted = NexusRAG(settings=rag.settings)
        restarted._embedder = rag.embedder
        assert restarted.document_store.exists(document_id)
        assert restarted.bm25.count() == len(chunks)
        assert restarted.delete_document(document_id)
        assert restarted.document_store.count() == 0
        assert restarted.bm25.count() == restarted.vector_store.count() == 0

    @pytest.mark.parametrize(
        "stage", ["dense_before", "dense_partial", "dense_after", "sparse_read", "document_file"]
    )
    def test_failed_ingest_rollback_preserves_the_original_error_and_can_be_recovered(
        self, scinexusrag_instance: NexusRAG, monkeypatch, stage
    ):
        rag = scinexusrag_instance
        other = rag.ingest_bytes(
            b"A separate document records voltage for sensor B.", "voltage.txt", ".txt"
        )
        assert other.success
        document = rag.parser.parse_bytes(
            b"The measured current is forty-two milliamps.", "current.txt", ".txt"
        )
        chunks = [
            Chunk(
                id=f"current-{index}",
                content=f"Current measurement {index}",
                document_id=document.id,
            )
            for index in range(3)
        ]
        original_error = OSError("Document metadata update failed")
        monkeypatch.setattr(
            rag.document_store, "update_metadata", MagicMock(side_effect=original_error)
        )
        if stage.startswith("dense"):
            owner, method = rag.vector_store, "delete_by_document"
        elif stage == "sparse_read":
            owner, method = rag.vector_store, "get_all_chunks"
        else:
            owner, method = rag.document_store, "delete"
        cleanup = getattr(owner, method)

        def fail_cleanup(*args, **kwargs):
            if stage == "dense_partial":
                rag.vector_store.delete([chunks[0].id])
            elif stage == "dense_after":
                cleanup(*args, **kwargs)
            raise OSError(f"Rollback failed during {stage}")

        monkeypatch.setattr(owner, method, fail_cleanup)
        with pytest.raises(OSError, match="Document metadata update failed") as raised:
            rag._persist(document, chunks, np.ones((3, 384), dtype=np.float32))

        assert raised.value is original_error
        assert rag.document_store.exists(document.id)
        assert rag.document_store.exists(other.document_id)
        if stage == "sparse_read":
            assert rag._bm25 is None
            assert rag._orchestrator is None
        monkeypatch.setattr(owner, method, cleanup)
        expected_remaining = 3 if stage == "dense_before" else 2 if stage == "dense_partial" else 0
        assert len(rag.vector_store.get_chunks_by_document(document.id)) == expected_remaining
        assert (
            rag.bm25.count() == rag.vector_store.count() == other.chunk_count + expected_remaining
        )
        assert all(
            rag.document_store.exists(chunk.document_id)
            for chunk in rag.vector_store.get_all_chunks()
        )
        restarted = NexusRAG(settings=rag.settings)
        restarted._embedder = rag.embedder
        assert restarted.bm25.count() == other.chunk_count + expected_remaining
        assert restarted.delete_document(document.id)
        assert restarted.document_store.count() == 1
        assert restarted.document_store.exists(other.document_id)
        assert restarted.bm25.count() == restarted.vector_store.count() == other.chunk_count

    def test_document_index_failure_rolls_back_document_file(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path, monkeypatch
    ):
        rag = scinexusrag_instance
        store = rag.document_store
        store.count()
        save_index = store._save_index
        failed = False

        def fail_first_save():
            nonlocal failed
            if not failed:
                failed = True
                raise OSError("simulated index write failure")
            save_index()

        monkeypatch.setattr(store, "_save_index", fail_first_save)
        result = rag.ingest(sample_text_file)

        assert not result.success
        assert store.count() == 0
        assert rag.vector_store.count() == 0
        assert rag.bm25.count() == 0
        assert list(store.path.glob("*.json")) == [store._index_path]
        assert rag.ingest(sample_text_file).success

    def test_vector_failure_rolls_back_document_store(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path, monkeypatch
    ):
        def boom(chunks, embeddings):
            raise RuntimeError("simulated storage failure")

        monkeypatch.setattr(scinexusrag_instance.vector_store, "add", boom)

        result = scinexusrag_instance.ingest(sample_text_file)

        assert result.success is False
        assert scinexusrag_instance.document_store.count() == 0
        assert scinexusrag_instance.bm25.count() == 0

    def test_bm25_failure_rolls_back_all_stores(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path, monkeypatch
    ):
        def boom(chunks):
            raise RuntimeError("simulated index failure")

        monkeypatch.setattr(scinexusrag_instance.bm25, "add_incremental", boom)

        result = scinexusrag_instance.ingest(sample_text_file)

        assert result.success is False
        assert scinexusrag_instance.document_store.count() == 0
        assert scinexusrag_instance.vector_store.count() == 0


class TestUnloadModels:
    def test_unload_resets_lazy_handles(self, scinexusrag_instance: NexusRAG):

        assert scinexusrag_instance.embedder is not None
        _ = scinexusrag_instance.llm
        _ = scinexusrag_instance.orchestrator

        scinexusrag_instance.unload_models()

        assert scinexusrag_instance._embedder is None
        assert scinexusrag_instance._llm is None
        assert scinexusrag_instance._orchestrator is None
        assert scinexusrag_instance._vector_store is None
        assert scinexusrag_instance._bm25 is None


class _ObservedWriteLock:
    def __init__(self):
        self.lock = RLock()
        self.attempts = Queue()

    def __enter__(self):
        self.attempts.put(get_ident())
        if not self.lock.acquire(timeout=5):
            raise TimeoutError("Pipeline lock was not released")
        return self

    def __exit__(self, *_args):
        self.lock.release()

    def wait_for_attempts(self, count):
        return [self.attempts.get(timeout=5) for _ in range(count)]

    def wait_for_threads(self, count):
        threads = set()
        while len(threads) < count:
            threads.add(self.attempts.get(timeout=5))
        return threads


class TestLazyResourceLifecycle:
    @pytest.mark.parametrize("resource", ["bm25", "orchestrator"])
    @pytest.mark.parametrize("fail_first", [False, True])
    def test_lazy_retrieval_handles_wait_for_a_complete_sparse_index(
        self, scinexusrag_instance: NexusRAG, monkeypatch, resource, fail_first
    ):
        rag = scinexusrag_instance
        store = rag.vector_store
        get_all_chunks = store.get_all_chunks
        lock = _ObservedWriteLock()
        rag._write_lock = lock
        entered, release = Event(), Event()
        reads = 0

        def hold_first_read():
            nonlocal reads
            reads += 1
            if reads == 1:
                entered.set()
                assert release.wait(timeout=5)
                if fail_first:
                    raise VectorStoreError("Sparse initialization failed")
            return get_all_chunks()

        monkeypatch.setattr(store, "get_all_chunks", hold_first_read)
        with ThreadPoolExecutor(max_workers=2) as workers:
            first = workers.submit(lambda: getattr(rag, resource))
            try:
                assert entered.wait(timeout=5)
                assert rag._bm25 is None
                assert rag._orchestrator is None
                second = workers.submit(lambda: getattr(rag, resource))
                lock.wait_for_threads(2)
                assert not second.done()
            finally:
                release.set()
            if fail_first:
                with pytest.raises(VectorStoreError, match="Sparse initialization failed"):
                    first.result(timeout=5)
                assert second.result(timeout=5) is getattr(rag, resource)
                assert reads == 2
            else:
                assert first.result(timeout=5) is second.result(timeout=5)
                assert reads == 1
        assert rag.bm25.count() == 0

    @pytest.mark.parametrize(
        "resource,factory_name",
        [("embedder", "Embedder"), ("vector_store", "VectorStore"), ("llm", "LLMClient")],
    )
    def test_concurrent_first_access_constructs_one_handle(
        self, test_settings, mock_embedder, monkeypatch, resource, factory_name
    ):
        rag = NexusRAG(test_settings)
        if resource == "vector_store":
            rag._embedder = mock_embedder
        lock = _ObservedWriteLock()
        rag._write_lock = lock
        start = Barrier(4)
        constructing = Event()
        release = Event()
        handle = MagicMock()

        def construct(**_kwargs):
            constructing.set()
            assert release.wait(timeout=5)
            return handle

        def load():
            start.wait(timeout=5)
            return getattr(rag, resource)

        factory = MagicMock(side_effect=construct)
        monkeypatch.setattr("scinexusrag.pipeline." + factory_name, factory)
        with ThreadPoolExecutor(max_workers=4) as workers:
            futures = [workers.submit(load) for _ in range(4)]
            try:
                assert constructing.wait(timeout=5)
                assert len(lock.wait_for_threads(4)) == 4
            finally:
                release.set()
            assert all(future.result(timeout=5) is handle for future in futures)

        if resource == "embedder":
            factory.assert_called_once_with(
                model_name="test-model", device="cpu", revision="a" * 40
            )
        elif resource == "vector_store":
            factory.assert_called_once_with(
                path=test_settings.storage.lancedb_path,
                table_name=test_settings.storage.table_name,
                embedding_config=mock_embedder.embedding_config,
            )
        else:
            factory.assert_called_once_with(
                model=test_settings.llm.model,
                base_url=test_settings.llm.base_url,
                timeout=test_settings.llm.timeout,
            )

    @pytest.mark.parametrize(
        "resource,factory_name",
        [("embedder", "Embedder"), ("vector_store", "VectorStore"), ("llm", "LLMClient")],
    )
    def test_failed_constructor_can_retry(
        self, test_settings, mock_embedder, monkeypatch, resource, factory_name
    ):
        rag = NexusRAG(test_settings)
        if resource == "vector_store":
            rag._embedder = mock_embedder
        rag._write_lock = _ObservedWriteLock()
        handle = MagicMock()
        factory = MagicMock(side_effect=[RuntimeError("Constructor failed"), handle])
        monkeypatch.setattr("scinexusrag.pipeline." + factory_name, factory)

        with pytest.raises(RuntimeError, match="Constructor failed"):
            getattr(rag, resource)
        assert getattr(rag, "_" + resource) is None
        with ThreadPoolExecutor(max_workers=1) as workers:
            assert workers.submit(lambda: getattr(rag, resource)).result(timeout=5) is handle
        assert getattr(rag, resource) is handle
        assert factory.call_count == 2

    @pytest.mark.parametrize(
        "resource,factory_name", [("embedder", "Embedder"), ("llm", "LLMClient")]
    )
    def test_unload_waits_for_lazy_handle_construction(
        self, test_settings, monkeypatch, resource, factory_name
    ):
        rag = NexusRAG(test_settings)
        lock = _ObservedWriteLock()
        rag._write_lock = lock
        constructing = Event()
        release = Event()
        original_handle, fresh_handle = MagicMock(), MagicMock()

        def construct(**_kwargs):
            constructing.set()
            assert release.wait(timeout=5)
            return original_handle

        factory = MagicMock(side_effect=construct)
        monkeypatch.setattr("scinexusrag.pipeline." + factory_name, factory)
        with ThreadPoolExecutor(max_workers=2) as workers:
            loading = workers.submit(lambda: getattr(rag, resource))
            try:
                assert constructing.wait(timeout=5)
                lock.wait_for_attempts(1)
                unloading = workers.submit(rag.unload_models)
                lock.wait_for_attempts(1)
                assert not unloading.done()
            finally:
                release.set()
            assert loading.result(timeout=5) is original_handle
            unloading.result(timeout=5)

        if resource == "embedder":
            original_handle.unload.assert_called_once()
        else:
            original_handle.close.assert_called_once()
        assert rag._embedder is None
        assert rag._llm is None
        assert rag._orchestrator is None
        assert rag._vector_store is None
        assert rag._bm25 is None
        factory.side_effect = None
        factory.return_value = fresh_handle
        assert getattr(rag, resource) is fresh_handle
        assert getattr(rag, resource) is fresh_handle
        assert factory.call_count == 2


class TestStoredEmbeddingConfiguration:
    def test_explicit_reset_does_not_require_a_valid_embedder(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path
    ):
        original = scinexusrag_instance
        assert original.ingest(sample_text_file).success
        restarted = NexusRAG(settings=original.settings)
        invalid_embedder = MagicMock()
        config = PropertyMock(side_effect=EmbeddingConfigError("Unpinned model revision"))
        type(invalid_embedder).embedding_config = config
        restarted._embedder = invalid_embedder

        restarted.clear_all()

        config.assert_not_called()
        invalid_embedder.embed.assert_not_called()
        invalid_embedder.embed_query.assert_not_called()
        assert restarted.document_store.count() == 0
        assert (
            restarted.settings.storage.table_name
            not in original.vector_store.db.list_tables().tables
        )
        assert restarted._vector_store is None
        assert restarted._bm25 is None

    def test_changed_model_rejects_query_before_embedding_or_generation(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path
    ):
        original = scinexusrag_instance
        ingested = original.ingest(sample_text_file)
        assert ingested.success
        changed_embedder = MagicMock()
        changed_embedder.dimension = 384
        changed_embedder.embedding_config = replace(
            original.embedder.embedding_config, model_name="another-model"
        )
        restarted = NexusRAG(settings=original.settings)
        restarted._embedder = changed_embedder
        restarted._llm = MagicMock()

        with pytest.raises(EmbeddingConfigError):
            restarted.query("What measurements were recorded?")
        with pytest.raises(EmbeddingConfigError):
            list(restarted.query_streaming("What measurements were recorded?"))

        changed_embedder.embed_query.assert_not_called()
        restarted._llm.generate.assert_not_called()
        assert original.vector_store.count() == ingested.chunk_count
        assert original.document_store.count() == 1

    def test_changed_revision_rejects_ingest_before_embedding_writes_or_rollback(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path, monkeypatch
    ):
        original = scinexusrag_instance
        ingested = original.ingest(sample_text_file)
        assert ingested.success
        changed_embedder = MagicMock()
        changed_embedder.dimension = 384
        changed_embedder.embedding_config = replace(
            original.embedder.embedding_config, revision="b" * 40
        )
        restarted = NexusRAG(settings=original.settings)
        restarted._embedder = changed_embedder
        document = restarted.parser.parse_bytes(
            b"A distinct upload records pressure readings from sensor C.", "new.txt", ".txt"
        )
        writes = MagicMock()
        deletes = MagicMock()
        monkeypatch.setattr(restarted.document_store, "add", writes)
        monkeypatch.setattr(restarted.document_store, "delete", deletes)

        with pytest.raises(EmbeddingConfigError):
            restarted._ingest_document(document, "new.txt", show_progress=False)
        with pytest.raises(EmbeddingConfigError):
            restarted._persist(document, [], np.empty((0, 384), dtype=np.float32))
        new_path = sample_text_file.with_name("new.txt")
        new_path.write_text(document.content)
        with pytest.raises(EmbeddingConfigError):
            restarted.ingest(new_path)
        with pytest.raises(EmbeddingConfigError):
            restarted.ingest_bytes(document.content.encode(), "new.txt", ".txt")

        changed_embedder.embed.assert_not_called()
        writes.assert_not_called()
        deletes.assert_not_called()
        assert original.vector_store.count() == ingested.chunk_count
        assert original.document_store.count() == 1

    def test_unload_reloads_configuration_instead_of_reusing_old_store(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path
    ):
        rag = scinexusrag_instance
        ingested = rag.ingest(sample_text_file)
        assert ingested.success
        assert rag.bm25.count() == ingested.chunk_count
        changed_config = replace(rag.embedder.embedding_config, revision="b" * 40)

        rag.unload_models()
        changed_embedder = MagicMock()
        changed_embedder.dimension = 384
        changed_embedder.embedding_config = changed_config
        rag._embedder = changed_embedder

        assert rag._vector_store is None
        assert rag._bm25 is None
        with pytest.raises(EmbeddingConfigError):
            rag.query("What measurements were recorded?")
        changed_embedder.embed_query.assert_not_called()

    def test_explicit_reset_clears_documents_and_allows_reingestion(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path
    ):
        original = scinexusrag_instance
        assert original.ingest(sample_text_file).success
        changed_config = replace(original.embedder.embedding_config, model_name="another-model")
        original.unload_models()
        original._embedder = MagicMock()
        original._embedder.dimension = 384
        original._embedder.embedding_config = changed_config
        original._embedder.embed.return_value = np.ones((1, 384), dtype=np.float32)

        with pytest.raises(EmbeddingConfigError):
            original.query("What measurements were recorded?")

        original.clear_all()
        assert original.document_store.count() == 0
        assert original.vector_store.count() == 0
        assert original.bm25.count() == 0
        original._embedder.embed.side_effect = lambda texts, **kwargs: np.ones(
            (len(texts), 384), dtype=np.float32
        )
        result = original.ingest(sample_text_file)

        assert result.success
        assert original.document_store.count() == 1
        assert original.vector_store.count() == original.bm25.count() == result.chunk_count


class _ImmediateVisibilityStore:
    """Expose newly added chunks before lazy BM25 initialization."""

    def __init__(self, real):
        self._real = real
        self._seen: list = []

    def add(self, chunks, embeddings):
        n = self._real.add(chunks, embeddings)
        self._seen.extend(chunks)
        return n

    def get_all_chunks(self):
        return list(self._seen)

    def __getattr__(self, name):
        return getattr(self._real, name)


class TestBM25NoDoubleCount:
    def test_ingest_does_not_double_count_bm25_on_lazy_init(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path
    ):
        rag = scinexusrag_instance

        rag._vector_store = _ImmediateVisibilityStore(rag.vector_store)

        result = rag.ingest(sample_text_file)

        assert result.success
        assert rag.bm25.count() == result.chunk_count

    def test_first_ingest_keeps_bm25_and_vector_in_sync(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path
    ):

        r = scinexusrag_instance.ingest(sample_text_file)
        assert r.success
        assert scinexusrag_instance.bm25.count() == r.chunk_count
        assert scinexusrag_instance.vector_store.count() == r.chunk_count


class TestEndToEnd:
    def test_ingest_query_cites_a_real_source(
        self, scinexusrag_instance: NexusRAG, sample_text_file: Path
    ):

        rag = scinexusrag_instance
        rag._llm.generate = lambda prompt, **kw: "The sensors are read every thirty seconds [1]."
        rag._embedder.embed_query = lambda q: np.random.rand(384).astype(np.float32)

        assert rag.ingest(sample_text_file).success
        resp = rag.query("What methodology was used?")

        assert resp.sources, "query returned no sources"
        assert "[1]" in resp.answer, "answer did not cite a source"
        assert resp.confidence > 0

        assert resp.sources[0].chunk_id
