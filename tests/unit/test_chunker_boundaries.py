"""Keep all input text while respecting split and overlap boundaries."""

from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument

from scinexusrag.ingestion import (
    DocumentParser,
    FixedSizeChunker,
    ParsedDocument,
    Section,
    SemanticChunker,
)


@pytest.mark.parametrize("extension", [".txt", ".md", ".docx", ".pdf"])
def test_heading_only_documents_produce_chunks_with_exact_source_spans(extension, monkeypatch):
    headings = ["CALIBRATION PASSED.", "THE CURRENT IS 42 MA.", "CALIBRATION PASSED."]
    if extension == ".docx":
        source = DocxDocument()
        for heading in headings:
            source.add_heading(heading, level=1)
        buffer = BytesIO()
        source.save(buffer)
        data = buffer.getvalue()
        expected = headings
    elif extension == ".pdf":
        pages = [SimpleNamespace(extract_text=lambda text=text: text) for text in headings]
        reader = SimpleNamespace(is_encrypted=False, pages=pages)
        monkeypatch.setattr("scinexusrag.ingestion.parser.PdfReader", lambda path: reader)
        data = b"%PDF-1.4 fixture"
        expected = headings
    else:
        expected = [f"# {heading}" for heading in headings] if extension == ".md" else headings
        data = "\n\n".join(expected).encode()
    document = DocumentParser().parse_bytes(data, "readings" + extension, extension)

    chunks = SemanticChunker(context_chars=12).chunk(document)

    assert [chunk.content for chunk in chunks] == expected
    assert [chunk.section_title for chunk in chunks] == headings
    for chunk, section in zip(chunks, document.sections, strict=True):
        assert section.heading_only
        assert document.content[chunk.start_offset : chunk.end_offset] == chunk.content
        _assert_context_span(chunk, document.content, section.start_offset, section.end_offset, 12)
    assert chunks[0].start_offset < chunks[2].start_offset
    if extension == ".pdf":
        assert [chunk.page_number for chunk in chunks] == [1, 2, 3]


def test_body_and_empty_final_heading_are_chunked_once_with_their_own_metadata():
    source = "# Methods\nMeasured current: 42 mA.\n\n## CALIBRATION PASSED."
    document = DocumentParser().parse_bytes(source.encode(), "readings.md", ".md")

    chunks = SemanticChunker(context_chars=16).chunk(document)

    assert [chunk.content for chunk in chunks] == [
        "[Methods]\n\nMeasured current: 42 mA.",
        "## CALIBRATION PASSED.",
    ]
    assert [chunk.section_title for chunk in chunks] == ["Methods", "CALIBRATION PASSED."]
    for chunk, section in zip(chunks, document.sections, strict=True):
        _assert_context_span(chunk, document.content, section.start_offset, section.end_offset, 16)


@pytest.mark.parametrize("section_title", ["", "Legacy heading"])
def test_empty_sections_fall_back_to_original_text_with_exact_source_offsets(section_title):
    content = "  THE CURRENT IS 42 MA.\n\nCALIBRATION PASSED.  "
    document = ParsedDocument(
        id="empty-sections",
        content=content,
        sections=[Section(title=section_title, content=""), Section(title="", content="  ")],
    )

    chunks = SemanticChunker(context_chars=10).chunk(document)

    assert len(chunks) == 1
    chunk = chunks[0]
    assert chunk.content == content.strip()
    assert chunk.section_title == ""
    _assert_context_span(chunk, content, 2, len(content) - 2, 10)


@pytest.mark.parametrize("sectioned", [False, True])
@pytest.mark.parametrize("content", ["word " * 330, "long " * 299 + "\n\nTAILWORD"])
def test_short_tail_does_not_overflow_full_chunk(content, sectioned):
    prefix = "[Methods]\n\n" if sectioned else ""
    document = ParsedDocument(id="tail", content=content)
    if sectioned:
        document.sections = [Section(title="Methods", content=content, level=1, page_number=4)]
    chunker = SemanticChunker(include_context=False, overlap_size=0)

    chunks = chunker.chunk(document)

    assert len(chunks) >= 2
    assert all(len(chunk.content) <= chunker.max_chunk_size for chunk in chunks)
    recovered = " ".join(chunk.content.removeprefix(prefix) for chunk in chunks)
    assert recovered.split() == content.split()
    if sectioned:
        assert all(chunk.section_title == "Methods" and chunk.page_number == 4 for chunk in chunks)


@pytest.mark.parametrize("sectioned", [False, True])
def test_short_tail_merge_does_not_duplicate_overlap(sectioned):
    content = "\n\n".join(
        [("alpha " * 185).strip(), ("beta " * 10).strip(), ("gamma " * 10).strip()]
    )
    document = ParsedDocument(id="overlap-tail", content=content)
    if sectioned:
        document.sections = [Section(title="Methods", content=content, page_number=4)]
    chunker = SemanticChunker(min_chunk_size=200, include_context=False)

    chunks = chunker.chunk(document)

    assert len(chunks) == 1
    assert chunks[0].content.removeprefix("[Methods]\n\n") == content
    assert len(chunks[0].content) <= chunker.max_chunk_size
    if sectioned:
        assert chunks[0].section_title == "Methods"
        assert chunks[0].page_number == 4


def test_context_matches_full_passage_when_prefixes_repeat():
    prefix = "The shared introduction explains the study design and measurement procedure. "
    prefix += "Both groups followed this procedure. "
    first = prefix + "The first group used the original instrument."
    second = prefix + "The second group used the replacement instrument."
    content = f"First\n\n{first}\n\nSecond\n\n{second}\n\nEnd of the document."
    document = ParsedDocument(
        id="shared-prefix",
        content=content,
        sections=[Section(title="First", content=first), Section(title="Second", content=second)],
    )
    chunker = SemanticChunker(context_chars=20)

    chunks = chunker.chunk(document)

    second_start = content.index(second)
    second_end = second_start + len(second)
    assert chunks[1].context_before == content[second_start - 20 : second_start].strip()
    assert chunks[1].context_after == content[second_end : second_end + 20].strip()


def test_identical_passages_get_context_from_their_document_order():
    passage = "The instrument was calibrated before collecting the measurements."
    content = f"First marker.\n\n{passage}\n\nMiddle marker.\n\n{passage}\n\nLast marker."
    first_start = content.index(passage)
    second_start = content.rindex(passage)
    document = ParsedDocument(
        id="repeated-passage",
        content=content,
        sections=[
            Section(
                title="First",
                content=passage,
                start_offset=first_start,
                end_offset=first_start + len(passage),
            ),
            Section(
                title="Second",
                content=passage,
                start_offset=second_start,
                end_offset=second_start + len(passage),
            ),
        ],
    )

    chunks = SemanticChunker(context_chars=16).chunk(document)

    assert chunks[0].context_before == "First marker."
    assert chunks[1].context_before == "Middle marker."
    assert chunks[1].context_after == "Last marker."


def test_full_overlap_chunks_keep_context_from_the_same_source_position():
    first = ("alpha " * 40).strip()
    second = ("beta " * 210).strip()
    methods = first + "\n\n" + second
    results = "Following measurements were performed after calibration."
    content = f"Introduction marker.\n\n{methods}\n\nResults marker.\n\n{results}"
    document = ParsedDocument(
        id="full-overlap-context",
        content=content,
        sections=[
            Section(title="Methods", content=methods),
            Section(title="Results", content=results),
        ],
    )
    chunker = SemanticChunker(min_chunk_size=200, context_chars=40)

    chunks = chunker.chunk(document)

    assert len(chunks) == 3
    assert chunks[1].content == "[Methods]\n\n" + methods
    assert chunks[1].context_before == "Introduction marker."
    methods_end = content.index(methods) + len(methods)
    assert chunks[1].context_after == content[methods_end : methods_end + 40].strip()


def test_shorter_repeated_prefix_gets_context_from_its_later_occurrence():
    passage = "The instrument was calibrated before collecting the measurements."
    first = passage + " The first group used the original instrument."
    content = f"First marker.\n\n{first}\n\nMiddle marker.\n\n{passage}\n\nLast marker."
    second_start = content.rindex(passage)
    document = ParsedDocument(
        id="shorter-prefix",
        content=content,
        sections=[
            Section(title="First", content=first),
            Section(
                title="Second",
                content=passage,
                start_offset=second_start,
                end_offset=second_start + len(passage),
            ),
        ],
    )

    chunks = SemanticChunker(context_chars=16).chunk(document)

    assert chunks[1].context_before == "Middle marker."
    assert chunks[1].context_after == "Last marker."


def test_unmatched_passage_does_not_borrow_context_from_a_shared_prefix():
    prefix = "This shared passage prefix describes the methods used to collect the measurements. "
    document = ParsedDocument(
        id="unmatched-passage",
        content="Beginning.\n\n" + prefix + "The original result.\n\nEnding.",
        sections=[Section(title="Methods", content=prefix + "A different result.")],
    )

    chunk = SemanticChunker().chunk(document)[0]

    assert chunk.context_before == ""
    assert chunk.context_after == ""


def test_literal_bracketed_text_is_not_treated_as_an_added_section_header():
    content = "[Literal source label]\n\nA short document with a label in its source text."

    chunk = SemanticChunker().chunk(ParsedDocument(id="literal-label", content=content))[0]

    assert chunk.full_context == content
    assert chunk.context_before == ""
    assert chunk.context_after == ""


def _assert_context_span(chunk, content, start, end, context_chars):
    assert chunk.start_offset == start
    assert chunk.end_offset == end
    assert chunk.metadata["start_offset"] == start
    assert chunk.metadata["end_offset"] == end
    assert chunk.context_before == content[max(0, start - context_chars) : start].strip()
    assert chunk.context_after == content[end : end + context_chars].strip()


def test_repeated_wrapped_words_keep_their_original_whitespace_positions():
    body = "alpha\tbeta   alpha\nbeta alpha\tbeta"
    content = "Before marker.\n\n" + body + "\n\nAfter marker."
    start = len("Before marker.\n\n")
    document = ParsedDocument(
        id="wrapped-positions",
        content=content,
        sections=[
            Section(
                title="",
                content=body,
                start_offset=start,
                end_offset=start + len(body),
            )
        ],
    )
    chunker = SemanticChunker(
        min_chunk_size=1,
        target_chunk_size=12,
        max_chunk_size=12,
        overlap_size=0,
        context_chars=16,
    )

    chunks = chunker.chunk(document)

    assert [chunk.content for chunk in chunks] == ["alpha beta"] * 3
    for chunk, relative_start in zip(chunks, [0, 13, 24], strict=True):
        _assert_context_span(
            chunk, content, start + relative_start, start + relative_start + 10, 16
        )


def test_sentence_packing_keeps_source_offsets_when_separator_whitespace_changes():
    first = "First sentence."
    second = "Second sentence."
    third = "Third sentence."
    content = first + "\t  " + second + "\n " + third
    chunker = SemanticChunker(
        min_chunk_size=1,
        target_chunk_size=35,
        max_chunk_size=35,
        overlap_size=0,
        context_chars=18,
    )

    chunks = chunker.chunk(ParsedDocument(id="sentence-positions", content=content))

    assert [chunk.content for chunk in chunks] == [first + " " + second, third]
    second_end = len(first) + 3 + len(second)
    _assert_context_span(chunks[0], content, 0, second_end, 18)
    _assert_context_span(chunks[1], content, second_end + 2, len(content), 18)


def test_paragraph_overlap_keeps_source_offsets_with_headers_and_irregular_blank_lines():
    first = "alpha beta gamma."
    second = "delta epsilon zeta."
    third = "eta theta iota."
    body = first + "\n \n" + second + "\n\t\n" + third
    content = "Introduction.\n\n" + body + "\n\nConclusion."
    start = len("Introduction.\n\n")
    document = ParsedDocument(
        id="overlap-positions",
        content=content,
        sections=[
            Section(
                title="Methods",
                content=body,
                start_offset=start,
                end_offset=start + len(body),
            )
        ],
    )
    chunker = SemanticChunker(
        min_chunk_size=1,
        target_chunk_size=35,
        max_chunk_size=70,
        overlap_size=20,
        context_chars=16,
    )

    chunks = chunker.chunk(document)

    prefix = "[Methods]\n\n"
    assert [chunk.content for chunk in chunks] == [
        prefix + first,
        prefix + first + "\n\n" + second,
        prefix + second + "\n\n" + third,
    ]
    second_start = start + len(first) + 3
    _assert_context_span(chunks[0], content, start, start + len(first), 16)
    _assert_context_span(chunks[1], content, start, second_start + len(second), 16)
    _assert_context_span(chunks[2], content, second_start, start + len(body), 16)


def test_explicit_section_offsets_distinguish_identical_passages():
    passage = "The instrument was calibrated."
    content = passage + "\n\nMiddle marker.\n\n" + passage + "\n\nEnd marker."
    start = len(passage + "\n\nMiddle marker.\n\n")
    document = ParsedDocument(
        id="explicit-section-position",
        content=content,
        sections=[
            Section(
                title="Methods",
                content=passage,
                start_offset=start,
                end_offset=start + len(passage),
            )
        ],
    )

    chunk = SemanticChunker(context_chars=16).chunk(document)[0]

    assert chunk.content == "[Methods]\n\n" + passage
    _assert_context_span(chunk, content, start, start + len(passage), 16)


@pytest.mark.parametrize(
    "extension,headers", [(".txt", ["FIRST", "SECOND"]), (".md", ["# First", "# Second"])]
)
def test_parsed_repeated_sections_keep_context_after_whitespace_cleaning(extension, headers):
    body = "The instrument\twas calibrated.\n\n  Measurements were recorded."
    source = headers[0] + "\n" + body + "\n\n" + headers[1] + "\n" + body
    document = DocumentParser().parse_bytes(source.encode(), "study" + extension, extension)
    chunker = SemanticChunker(context_chars=20)

    chunks = chunker.chunk(document)

    assert len(chunks) == 2
    for chunk, section in zip(chunks, document.sections, strict=True):
        _assert_context_span(chunk, document.content, section.start_offset, section.end_offset, 20)
    assert chunks[0].start_offset < chunks[1].start_offset


def test_invalid_section_offsets_do_not_borrow_context_from_matching_text_elsewhere():
    passage = "The instrument was calibrated."
    content = "Before marker.\n\n" + passage + "\n\nAfter marker."
    document = ParsedDocument(
        id="invalid-section-position",
        content=content,
        sections=[
            Section(title="Methods", content=passage, start_offset=0, end_offset=len(passage))
        ],
    )

    chunk = SemanticChunker().chunk(document)[0]

    assert chunk.start_offset is None
    assert chunk.end_offset is None
    assert chunk.context_before == ""
    assert chunk.context_after == ""


def test_legacy_section_with_repeated_body_has_no_guessed_context():
    passage = "The instrument was calibrated."
    content = "First marker.\n\n" + passage + "\n\nSecond marker.\n\n" + passage
    document = ParsedDocument(
        id="ambiguous-legacy-section",
        content=content,
        sections=[Section(title="Methods", content=passage)],
    )

    chunk = SemanticChunker().chunk(document)[0]

    assert chunk.start_offset is None
    assert chunk.end_offset is None
    assert chunk.context_before == ""
    assert chunk.context_after == ""


def test_short_tail_merge_extends_source_end_without_repeating_overlap():
    first = ("alpha " * 185).strip()
    second = ("beta " * 10).strip()
    third = ("gamma " * 10).strip()
    content = " \t" + first + "\n \n" + second + "\n\t\n" + third + " \n"
    chunker = SemanticChunker(min_chunk_size=200, context_chars=20)

    chunks = chunker.chunk(ParsedDocument(id="merged-tail-position", content=content))

    assert len(chunks) == 1
    assert chunks[0].content == first + "\n\n" + second + "\n\n" + third
    _assert_context_span(chunks[0], content, 2, len(content) - 2, 20)


@pytest.mark.parametrize("length_function", ["chars", "words"])
def test_negative_overlap_fails_before_it_can_skip_content(length_function):
    with pytest.raises(ValueError, match="nonnegative"):
        FixedSizeChunker(chunk_size=2, chunk_overlap=-1, length_function=length_function)


@pytest.mark.parametrize("chunk_size", [0, -1])
def test_nonpositive_chunk_size_rejected(chunk_size):
    with pytest.raises(ValueError, match="positive"):
        FixedSizeChunker(chunk_size=chunk_size, chunk_overlap=-2)


@pytest.mark.parametrize(
    "content",
    [
        " ".join(f"word{i}" for i in range(60)),
        "A sentence fills this chunk. Another sentence also has enough words to overflow.",
    ],
)
def test_fixed_character_budget_preserves_every_word_without_overlap(content):
    chunker = FixedSizeChunker(chunk_size=30, chunk_overlap=0)
    chunks = chunker.chunk(ParsedDocument(id="fixed", content=content))

    assert len(chunks) > 1
    assert all(len(chunk.content) <= 30 for chunk in chunks)
    assert " ".join(chunk.content for chunk in chunks).split() == content.split()


def test_fixed_character_overlap_keeps_whole_words_and_respects_the_budget():
    content = "alpha beta gamma delta epsilon zeta"
    chunks = FixedSizeChunker(chunk_size=22, chunk_overlap=10).chunk(
        ParsedDocument(id="overlap", content=content)
    )

    assert [chunk.content for chunk in chunks] == [
        "alpha beta gamma delta",
        "delta epsilon zeta",
    ]


def test_fixed_character_overlap_shrinks_to_fit_the_next_word():
    chunks = FixedSizeChunker(chunk_size=10, chunk_overlap=8).chunk(
        ParsedDocument(id="budget", content="one two abcdefghi end")
    )

    assert [chunk.content for chunk in chunks] == ["one two", "abcdefghi", "end"]


def test_fixed_character_chunker_rejects_an_indivisible_oversized_word():
    with pytest.raises(ValueError, match="word length exceeds available chunk size"):
        FixedSizeChunker(chunk_size=10, chunk_overlap=0).chunk(
            ParsedDocument(id="word", content="valid " + "x" * 11)
        )


def test_fixed_chunker_rejects_an_unknown_length_unit():
    with pytest.raises(ValueError, match="length_function"):
        FixedSizeChunker(length_function="tokens")


@pytest.mark.parametrize("word_count,expected_chunks", [(20, 1), (35, 2), (36, 3)])
def test_fixed_word_chunks_do_not_emit_a_tail_made_only_of_overlap(word_count, expected_chunks):
    content = " ".join(f"word{i}" for i in range(word_count))
    chunks = FixedSizeChunker(chunk_size=20, chunk_overlap=5, length_function="words").chunk(
        ParsedDocument(id="word-tail", content=content)
    )

    assert len(chunks) == expected_chunks
    assert chunks[-1].content.split()[-1] == f"word{word_count - 1}"
