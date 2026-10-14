"""Tests for DocumentParser."""

import pytest

from scinexusrag.ingestion import DocumentParseError, DocumentParser, ParsedDocument


def _assert_section_spans(document):
    previous_end = 0
    for section in document.sections:
        assert section.start_offset is not None
        assert section.end_offset is not None
        assert previous_end <= section.start_offset <= section.end_offset <= len(document.content)
        assert document.content[section.start_offset : section.end_offset] == section.content
        previous_end = section.end_offset


class TestDocumentParser:
    @pytest.fixture
    def parser(self):
        return DocumentParser()

    def test_parse_txt_file(self, parser, sample_text_file):
        doc = parser.parse(sample_text_file)

        assert isinstance(doc, ParsedDocument)
        assert doc.id is not None
        assert len(doc.id) == 16
        assert "introduction" in doc.content.lower()
        assert "methodology" in doc.content.lower()
        assert "results" in doc.content.lower()
        assert doc.metadata["filename"] == "sample.txt"
        assert doc.metadata["extension"] == ".txt"

    def test_parse_md_file(self, parser, sample_markdown_file):
        doc = parser.parse(sample_markdown_file)

        assert isinstance(doc, ParsedDocument)
        assert "research paper title" in doc.content.lower()
        assert doc.metadata["extension"] == ".md"

        assert len(doc.sections) > 0
        section_titles = [s.title.lower() for s in doc.sections]
        assert any("abstract" in t for t in section_titles)
        assert any("introduction" in t for t in section_titles)

    def test_parse_nonexistent_file(self, parser, temp_dir):
        with pytest.raises(FileNotFoundError):
            parser.parse(temp_dir / "nonexistent.txt")

    def test_parse_unsupported_format(self, parser, temp_dir):
        unsupported = temp_dir / "file.xyz"
        unsupported.write_text("content")

        with pytest.raises(ValueError, match="Unsupported format"):
            parser.parse(unsupported)

    def test_non_utf8_text_raises_parse_error(self, parser, temp_dir):

        bad = temp_dir / "latin1.txt"
        bad.write_bytes("café résumé".encode("latin-1"))
        with pytest.raises(DocumentParseError, match="UTF-8"):
            parser.parse(bad)

    def test_non_utf8_markdown_raises_parse_error(self, parser, temp_dir):
        bad = temp_dir / "latin1.md"
        bad.write_bytes("# café".encode("latin-1"))
        with pytest.raises(DocumentParseError, match="UTF-8"):
            parser.parse(bad)

    def test_documents_sharing_long_prefix_are_not_duplicates(self, parser):

        prefix = "alpha " * 200
        a = parser.parse_bytes((prefix + "ending one").encode(), "paper.txt", ".txt")
        b = parser.parse_bytes((prefix + "ending two").encode(), "paper.txt", ".txt")
        assert a.id != b.id

    def test_generate_document_id(self, parser, sample_text_file):
        doc1 = parser.parse(sample_text_file)
        doc2 = parser.parse(sample_text_file)

        assert doc1.id == doc2.id
        assert len(doc1.id) == 16
        assert doc1.id.isalnum()

    def test_different_files_different_ids(self, parser, sample_text_file, sample_markdown_file):
        doc1 = parser.parse(sample_text_file)
        doc2 = parser.parse(sample_markdown_file)

        assert doc1.id != doc2.id

    def test_section_extraction_txt(self, parser, sample_text_file):
        doc = parser.parse(sample_text_file)

        assert len(doc.sections) >= 1

    def test_section_extraction_md(self, parser, sample_markdown_file):
        doc = parser.parse(sample_markdown_file)

        assert len(doc.sections) >= 4

        levels = [s.level for s in doc.sections]
        assert 1 in levels
        assert 2 in levels

    def test_section_content_preserved(self, parser, sample_markdown_file):
        doc = parser.parse(sample_markdown_file)

        abstract_sections = [s for s in doc.sections if "abstract" in s.title.lower()]
        assert len(abstract_sections) == 1
        assert "two sensor readings per timestamp" in abstract_sections[0].content.lower()

    def test_metadata_extraction(self, parser, sample_text_file):
        doc = parser.parse(sample_text_file)

        assert "filename" in doc.metadata
        assert "extension" in doc.metadata
        assert "size_bytes" in doc.metadata
        assert "created_at" in doc.metadata
        assert "modified_at" in doc.metadata
        assert doc.metadata["size_bytes"] > 0

    def test_word_count_property(self, parser, sample_text_file):
        doc = parser.parse(sample_text_file)

        assert doc.word_count > 0
        assert doc.word_count == len(doc.content.split())

    def test_empty_file(self, parser, empty_file):
        doc = parser.parse(empty_file)

        assert doc.content == ""
        assert doc.word_count == 0

    def test_parse_bytes(self, parser, sample_text_file):
        content = sample_text_file.read_bytes()
        doc = parser.parse_bytes(content, "uploaded.txt", ".txt")

        assert doc.metadata["original_filename"] == "uploaded.txt"
        assert "introduction" in doc.content.lower()

    def test_text_cleaning(self, parser, temp_dir):
        messy_file = temp_dir / "messy.txt"
        messy_file.write_text("Multiple   spaces\n\n\n\nToo many newlines")

        doc = parser.parse(messy_file)

        assert "   " not in doc.content
        assert "\n\n\n" not in doc.content

    @pytest.mark.parametrize("extension", [".txt", ".md"])
    def test_repeated_section_bodies_keep_their_own_normalized_source_spans(
        self, parser, temp_dir, extension
    ):
        first_heading = "METHODS" if extension == ".txt" else "## Methods"
        second_heading = "RESULTS" if extension == ".txt" else "## Results"
        passage = "Shared reading.\nSecond line.\n\nAnother paragraph."
        source = (
            "  Preamble\twith   spacing.\n\n\n"
            f"{first_heading}\n\n{passage}\n\n\n"
            f"{second_heading}\n\n{passage}\n\n"
        )
        path = temp_dir / f"repeated{extension}"
        path.write_text(source)

        document = parser.parse(path)

        _assert_section_spans(document)
        assert [section.content for section in document.sections] == [
            "Preamble with spacing.",
            passage,
            passage,
        ]
        assert document.sections[1].start_offset == document.content.index(first_heading) + len(
            first_heading + "\n\n"
        )
        assert document.sections[2].start_offset == document.content.index(second_heading) + len(
            second_heading + "\n\n"
        )
        assert document.sections[1].start_offset != document.sections[2].start_offset

    @pytest.mark.parametrize("extension", [".txt", ".md"])
    def test_headerless_text_has_one_exact_span(self, parser, temp_dir, extension):
        path = temp_dir / f"plain{extension}"
        path.write_text("\n\n First\t line.\n\n\n Second   paragraph. \n\n")

        document = parser.parse(path)

        _assert_section_spans(document)
        assert len(document.sections) == 1
        assert document.sections[0].title == ""
        assert document.sections[0].content == "First line.\n\n Second paragraph."
        assert document.sections[0].start_offset == 0
        assert document.sections[0].end_offset == len(document.content)

    def test_markdown_empty_headings_have_valid_empty_spans(self, parser, temp_dir):
        path = temp_dir / "headings.md"
        path.write_text("# Empty\n\n## Methods\n\nA measurement.\n\n### Empty tail\n")

        document = parser.parse(path)

        _assert_section_spans(document)
        assert [section.level for section in document.sections] == [1, 2, 3]
        assert [section.content for section in document.sections] == ["", "A measurement.", ""]
        for section in (document.sections[0], document.sections[2]):
            assert section.start_offset == section.end_offset

    def test_docx_sections_preserve_paragraph_joins_and_repeated_source_positions(
        self, parser, temp_dir
    ):
        from docx import Document as DocxDocument

        source = DocxDocument()
        source.add_paragraph("  First\t preamble.  ")
        source.add_paragraph("Second preamble.")
        source.add_heading("Methods", level=1)
        source.add_paragraph("  Shared   passage.  ")
        source.add_paragraph("Line one.\nLine two.")
        source.add_heading("Results", level=2)
        source.add_paragraph("Shared\tpassage.")
        source.add_paragraph("Line one.\nLine two.")
        path = temp_dir / "repeated.docx"
        source.save(path)

        document = parser.parse(path)

        _assert_section_spans(document)
        assert [section.content for section in document.sections] == [
            "First preamble.\n\nSecond preamble.",
            "Shared passage.\n\nLine one.\nLine two.",
            "Shared passage.\n\nLine one.\nLine two.",
        ]
        assert [section.level for section in document.sections] == [0, 1, 2]
        assert document.sections[1].start_offset == document.content.index("Methods\n\n") + len(
            "Methods\n\n"
        )
        assert document.sections[2].start_offset == document.content.index("Results\n\n") + len(
            "Results\n\n"
        )


class TestPdfEdgeCases:
    def test_page_offsets_follow_cleaned_joins_and_retain_page_numbers(self, temp_dir, monkeypatch):
        from types import SimpleNamespace

        page_texts = [
            "\n  METHODS\n\nShared   passage.\n\n\nFirst page tail.  \n",
            " \n\t ",
            "\nRESULTS\n\nShared\tpassage.\n\nLast page tail.\n",
            "  Shared passage.\n\nA fourth-page note. \n",
        ]
        reader = SimpleNamespace(
            is_encrypted=False,
            pages=[SimpleNamespace(extract_text=lambda text=text: text) for text in page_texts],
        )
        monkeypatch.setattr("scinexusrag.ingestion.parser.PdfReader", lambda path: reader)
        path = temp_dir / "repeated.pdf"
        path.write_bytes(b"%PDF-1.4 fake")

        document = DocumentParser().parse(path)

        _assert_section_spans(document)
        assert [section.page_number for section in document.sections] == [1, 3, 4]
        assert [section.title for section in document.sections] == ["METHODS", "RESULTS", ""]
        assert [section.content for section in document.sections] == [
            "Shared passage.\n\nFirst page tail.",
            "Shared passage.\n\nLast page tail.",
            "Shared passage.\n\nA fourth-page note.",
        ]
        first_page = "METHODS\n\nShared passage.\n\nFirst page tail."
        third_page = "RESULTS\n\nShared passage.\n\nLast page tail."
        assert document.sections[1].start_offset == len(first_page + "\n\nRESULTS\n\n")
        assert document.sections[2].start_offset == len(first_page + "\n\n" + third_page + "\n\n")

    def test_headerless_pages_reach_chunks(self, temp_dir, monkeypatch):
        from scinexusrag.ingestion import SemanticChunker

        class FakePage:
            def __init__(self, text):
                self.text = text

            def extract_text(self):
                return self.text

        class FakeReader:
            def __init__(self, path):
                self.pages = [
                    FakePage("METHODS\nThe sample was measured."),
                    FakePage("Critical measurements on a page without a heading."),
                ]
                self.is_encrypted = False

        monkeypatch.setattr("scinexusrag.ingestion.parser.PdfReader", FakeReader)
        pdf_path = temp_dir / "mixed.pdf"
        pdf_path.write_bytes(b"%PDF-1.4 fake")
        chunks = SemanticChunker().chunk(DocumentParser().parse(pdf_path))

        assert any("Critical measurements" in chunk.content for chunk in chunks)
        assert (
            next(chunk for chunk in chunks if "Critical measurements" in chunk.content).page_number
            == 2
        )

    def test_encrypted_pdf_actionable_error(self, temp_dir):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.encrypt("secret")
        pdf_path = temp_dir / "locked.pdf"
        with open(pdf_path, "wb") as f:
            writer.write(f)

        with pytest.raises(DocumentParseError, match="password-protected"):
            DocumentParser().parse(pdf_path)

    def test_scanned_pdf_actionable_error(self, temp_dir, monkeypatch):
        class FakePage:
            images = [object()]

            def extract_text(self):
                return ""

        class FakeReader:
            def __init__(self, path):
                self.pages = [FakePage()]
                self.is_encrypted = False

        monkeypatch.setattr("scinexusrag.ingestion.parser.PdfReader", FakeReader)
        pdf_path = temp_dir / "scan.pdf"
        pdf_path.write_bytes(b"%PDF-1.4 fake")

        with pytest.raises(DocumentParseError, match="scanned"):
            DocumentParser().parse(pdf_path)


class TestDuplicateNormalization:
    def test_whitespace_only_edits_keep_same_id(self, temp_dir):
        parser = DocumentParser()
        doc = temp_dir / "doc.txt"
        doc.write_text("Alpha beta. Gamma delta.")
        original_id = parser.parse(doc).id

        doc.write_text("Alpha beta.\n\nGamma  delta. ")

        assert parser.parse(doc).id == original_id

    def test_content_edits_change_id(self, temp_dir):
        parser = DocumentParser()
        doc = temp_dir / "doc.txt"
        doc.write_text("Alpha beta. Gamma delta.")
        original_id = parser.parse(doc).id

        doc.write_text("Alpha beta. Gamma delta. Epsilon.")

        assert parser.parse(doc).id != original_id


class TestPreamblePreserved:
    def test_text_preamble_before_first_heading_reaches_chunks(self, temp_dir):
        from scinexusrag.ingestion import SemanticChunker

        txt = temp_dir / "doc.txt"
        txt.write_text("Critical preamble sentence.\n\nMETHODS\nBody under heading.\n")
        chunks = SemanticChunker().chunk(DocumentParser().parse(txt))

        assert any("Critical preamble sentence." in chunk.content for chunk in chunks)
        assert any("Body under heading." in chunk.content for chunk in chunks)

    def test_markdown_preamble_before_first_heading_kept(self, temp_dir):
        md = temp_dir / "doc.md"
        md.write_text("Critical preamble sentence.\n\n# Heading One\n\nBody under heading.\n")
        doc = DocumentParser().parse(md)
        assert "critical preamble" in doc.content.lower()
        assert any("critical preamble" in s.content.lower() for s in doc.sections)

    def test_docx_preamble_before_first_heading_kept(self, temp_dir):
        from docx import Document as DocxDocument

        d = DocxDocument()
        d.add_paragraph("Critical preamble sentence.")
        d.add_heading("Heading One", level=1)
        d.add_paragraph("Body under heading.")
        path = temp_dir / "doc.docx"
        d.save(path)

        doc = DocumentParser().parse(path)
        assert any("critical preamble" in s.content.lower() for s in doc.sections)
