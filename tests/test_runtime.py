"""Runtime checks."""

from __future__ import annotations

import asyncio
import gc
import io
import json
import logging
import math
import os
import re
import subprocess
import sys
import threading
import traceback
import xml.etree.ElementTree as ET
import zipfile
import zlib
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from dataclasses import FrozenInstanceError, asdict, replace
from io import BytesIO
from pathlib import Path
from queue import Queue
from threading import Barrier, Event, RLock, get_ident
from types import MappingProxyType, ModuleType, SimpleNamespace
from typing import Any
from unittest.mock import MagicMock, Mock, PropertyMock, call, patch
from urllib.parse import quote
from zipfile import ZIP_DEFLATED, ZipFile

import httpx
import lancedb
import numpy as np
import pyarrow as pa
import pytest
from docx import Document as DocxDocument
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from pydantic import ValidationError
from scipy import sparse
from starlette.datastructures import Headers
from starlette.requests import Request

import scinexusrag.api.app as api
from scinexusrag.api import metrics, routes, security
from scinexusrag.api.app import create_app
from scinexusrag.api.dependencies import get_scinexusrag
from scinexusrag.api.metrics import MetricsCollector
from scinexusrag.api.routes import (
    MAX_FILENAME_LENGTH,
    QueryRequest,
    router,
)
from scinexusrag.application import (
    DocumentInventory,
    IngestResult,
    NexusRAG,
    SystemStats,
)
from scinexusrag.eval import datasets, run, systems
from scinexusrag.eval.metrics import holm_correction, pr_auc, risk_coverage_auc, roc_auc
from scinexusrag.generation.citations import extract_citations, strip_citations
from scinexusrag.generation.grounding import GroundingReport, GroundingVerifier
from scinexusrag.generation.llm import LLMClient, LLMError
from scinexusrag.generation.orchestrator import Orchestrator, RAGResponse
from scinexusrag.generation.query_analyzer import QueryAnalyzer, QueryType
from scinexusrag.generation.synthesizer import Source, Synthesizer
from scinexusrag.generation.verifier import AnswerVerifier
from scinexusrag.ingestion.chunker import (
    Chunk,
    FixedSizeChunker,
    HierarchicalChunker,
    SemanticChunker,
    get_chunker,
)
from scinexusrag.ingestion.embedder import (
    Embedder,
    EmbeddingConfig,
    EmbeddingConfigError,
)
from scinexusrag.ingestion.parser import (
    DocumentParseError,
    DocumentParser,
    ParsedDocument,
    Section,
)
from scinexusrag.retrieval.corrective import CorrectiveRetriever
from scinexusrag.retrieval.dense import DenseRetriever, RetrievalResult
from scinexusrag.retrieval.hybrid import (
    AdaptiveHybridRetriever,
    HybridRetriever,
    rrf_fuse,
)
from scinexusrag.retrieval.reranker import Reranker
from scinexusrag.retrieval.sparse import BM25Retriever
from scinexusrag.retrieval.splade import DEFAULT_MODEL, SpladeRetriever
from scinexusrag.settings import (
    HF_REVISIONS,
    APISettings,
    EmbeddingSettings,
    IngestionSettings,
    LLMSettings,
    RetrievalSettings,
    SelfCorrectionSettings,
    Settings,
    StorageSettings,
    get_settings,
)
from scinexusrag.storage.document_store import (
    DocumentExistsError,
    DocumentStore,
    DocumentStoreError,
)
from scinexusrag.storage.vector_store import SearchResult, VectorStore, VectorStoreError
from scinexusrag.utils.filenames import resolve_display_name
from scinexusrag.utils.text import SENTENCE_BOUNDARY, split_sentences
from scinexusrag.version import __version__


@pytest.fixture
def test_settings(tmp_path: Path) -> Settings:
    return Settings(
        data_dir=tmp_path,
        storage=StorageSettings(
            lancedb_path=tmp_path / "lancedb",
            table_name="test_documents",
        ),
        embedding=EmbeddingSettings(
            model="test-model",
            revision="a" * 40,
            device="cpu",
            batch_size=8,
        ),
        llm=LLMSettings(
            model="test-llm",
            base_url="http://localhost:11434",
            timeout=30,
        ),
        ingestion=IngestionSettings(
            chunk_size=200,
            chunk_overlap=50,
            min_chunk_size=20,
        ),
    )


@pytest.fixture
def mock_embedder():
    embedder = MagicMock()
    embedder.dimension = 384
    embedder.embedding_config = EmbeddingConfig("test-model", "a" * 40, "", "", True, 384)

    def mock_embed(texts, batch_size=None, show_progress=False):
        if isinstance(texts, str):
            texts = [texts]
        rng = np.random.default_rng(zlib.crc32("".join(texts).encode()))
        return rng.random((len(texts), 384), dtype=np.float32)

    embedder.embed = mock_embed
    return embedder


@pytest.fixture
def mock_llm():
    llm = MagicMock()
    llm.model = "test-llm"

    def mock_generate(prompt, **kwargs):
        if "What is" in prompt or "?" in prompt:
            return "The fixture records temperature at thirty-second intervals."
        return "Sensor A is the reference."

    llm.generate = mock_generate
    llm.is_available = MagicMock(return_value=True)
    return llm


@pytest.fixture
def pipeline_rag(test_settings: Settings, mock_embedder, mock_llm) -> NexusRAG:
    pipeline_rag = NexusRAG(settings=test_settings)

    pipeline_rag._embedder = mock_embedder
    pipeline_rag._llm = mock_llm

    return pipeline_rag


@pytest.fixture
def pipeline_text_file(tmp_path: Path) -> Path:
    file_path = tmp_path / "test_document.txt"
    content = """INTRODUCTION

This research document is a software test fixture about temperature sensors.
The reference sensor is labeled A, and the comparison sensor is labeled B.

Each reading has a timestamp, a sensor label and a temperature in degrees Celsius.

METHODS

The methodology reads both sensors at thirty-second intervals.
Readings with the same timestamp belong to the same measurement pair.

Store each measurement pair on one row, with separate columns for sensors A and B.

RESULTS

The fixture contains paired readings with matching timestamps.
Its values are synthetic test data, not measured research results.

Missing readings are represented by an empty sensor column.

DISCUSSION

Compare sensor readings only when both values are present.
Keep missing values separate from numeric temperature differences.

The reference sensor provides the label used in retrieval queries.

CONCLUSION

Both sensors use the same units and sampling interval.

"""
    file_path.write_text(content, encoding="utf-8")
    return file_path


@pytest.fixture
def pipeline_markdown_file(tmp_path: Path) -> Path:
    file_path = tmp_path / "test_paper.md"
    content = """# Scientific Paper on Testing

## Abstract

This is a Markdown fixture for document ingestion and retrieval tests.

## Introduction

### Background

The fixture uses nested headings to test section extraction.
Each section contains text that can be indexed as a passage.

### Objectives

1. Parse headings and paragraphs.
2. Index the resulting passages.
3. Retrieve a passage by its text.

## Methodology

### Data Preparation

Save the fixture as UTF-8 Markdown.
Keep headings on separate lines from their paragraphs.

### Testing Framework

The tests use real parsing and storage with fixed model responses.

## Results

Test assertions check stored passages, query sources and deletion.

## Conclusion

This fixture contains no benchmark measurements.
"""
    file_path.write_text(content, encoding="utf-8")
    return file_path


@pytest.fixture
def docs_directory(tmp_path: Path) -> Path:
    docs_dir = tmp_path / "documents"
    docs_dir.mkdir()

    (docs_dir / "doc1.txt").write_text(
        "Document 1 content. This is the first test document.\n"
        "It contains multiple sentences for proper chunking.\n"
        "Testing the ingestion of text files.\n"
    )

    (docs_dir / "doc2.txt").write_text(
        "Document 2 content. This is the second test document.\n"
        "It also has enough content for chunking.\n"
        "Multiple ingestion operations are tested.\n"
    )

    (docs_dir / "paper.md").write_text(
        "# Test Paper\n\n"
        "## Introduction\n"
        "This is a markdown document for testing.\n\n"
        "## Methods\n"
        "Testing with multiple formats.\n\n"
        "## Results\n"
        "The fixture has one Methods section.\n"
    )

    return docs_dir


class TestBasicPipelineWorkflow:
    def test_uppercase_evidence_is_ingested_with_exact_source_offsets(self, pipeline_rag: NexusRAG):
        content = "THE MEASURED CURRENT IS 42 MA.\n\nCALIBRATION PASSED."

        result = pipeline_rag.ingest_bytes(content.encode(), "measurements.txt", ".txt")

        assert result.success, result.error
        assert result.chunk_count > 0
        chunks = pipeline_rag.vector_store.get_chunks_by_document(result.document_id)
        assert all("[" not in chunk.content for chunk in chunks)
        for line in content.splitlines():
            if line:
                assert sum(line in chunk.content for chunk in chunks) == 1
        for chunk in chunks:
            assert chunk.start_offset is not None and chunk.end_offset is not None
            assert content[chunk.start_offset : chunk.end_offset] == chunk.content
        sources = pipeline_rag.bm25.retrieve("42 MA", top_k=1)
        assert len(sources) == 1
        assert "THE MEASURED CURRENT IS 42 MA." in sources[0].chunk.content

    def test_ingest_single_text_file(self, pipeline_rag: NexusRAG, pipeline_text_file: Path):
        result = pipeline_rag.ingest(pipeline_text_file)

        assert result.success is True
        assert result.filename == "test_document.txt"
        assert result.chunk_count > 0
        assert result.word_count > 0
        assert result.document_id
        assert result.error is None

    def test_ingest_markdown_file(self, pipeline_rag: NexusRAG, pipeline_markdown_file: Path):
        result = pipeline_rag.ingest(pipeline_markdown_file)

        assert result.success is True
        assert result.filename == "test_paper.md"
        assert result.chunk_count > 0
        assert result.word_count > 0

    @pytest.mark.parametrize("interleaved", [False, True])
    def test_docx_table_rows_remain_retrievable_after_reload(
        self, pipeline_rag: NexusRAG, interleaved: bool
    ):
        document = DocxDocument()
        before = "A paragraph before the measurement table."
        after = "A paragraph after the measurement table."
        if interleaved:
            document.add_heading("Measurements", level=1)
            document.add_paragraph(before)
        rows = [
            ["Sensor", "Reading", "Run"],
            ["calibration", "42.50", "007"],
            ["comparison", "", "-3.25"],
        ]
        table = document.add_table(rows=3, cols=3)
        for row, values in zip(table.rows, rows, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
        if interleaved:
            document.add_paragraph(after)
        upload = BytesIO()
        document.save(upload)

        result = pipeline_rag.ingest_bytes(upload.getvalue(), "measurements.docx", ".docx")

        assert result.success, result.error
        assert result.chunk_count > 0
        stored = json.loads(pipeline_rag.document_store._doc_path(result.document_id).read_text())
        content = stored["content"]
        numeric_row = "calibration | 42.50 | 007"
        assert numeric_row in content
        empty_cell_row = next(
            line for line in content.splitlines() if line.startswith("comparison")
        )
        assert [cell.strip() for cell in empty_cell_row.split("|")] == ["comparison", "", "-3.25"]
        if interleaved:
            assert content.index(before) < content.index(numeric_row) < content.index(after)
        assert stored["sections"]
        for section in stored["sections"]:
            assert content[section["start_offset"] : section["end_offset"]] == section["content"]

        restarted = NexusRAG(settings=pipeline_rag.settings)
        restarted._embedder = pipeline_rag.embedder
        restarted._llm = pipeline_rag.llm
        matches = restarted.bm25.retrieve("42.50", top_k=1)
        assert len(matches) == 1
        match = matches[0].chunk
        assert numeric_row in match.content
        persisted = {
            chunk.id: chunk
            for chunk in restarted.vector_store.get_chunks_by_document(result.document_id)
        }[match.id]
        assert persisted.full_context == match.full_context
        start, end = persisted.start_offset, persisted.end_offset
        assert start is not None and end is not None
        assert numeric_row in content[start:end]
        assert persisted.context_before == content[max(0, start - 150) : start].strip()
        assert persisted.context_after == content[end : end + 150].strip()

        prepared = restarted.orchestrator.synthesizer.prepare(
            "What reading was recorded?", matches, {result.document_id: "measurements.docx"}
        )
        assert prepared.error is None
        assert prepared.sources[0].chunk_id == match.id
        assert (
            prepared.sources[0].full_content or prepared.sources[0].content
        ) == match.full_context
        assert numeric_row in prepared.prompt

    def test_ingest_updates_inventory_and_statistics(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path
    ):

        ingest_result = pipeline_rag.ingest(pipeline_text_file)
        assert ingest_result.success is True
        document_id = ingest_result.document_id

        docs = pipeline_rag.list_documents()
        assert len(docs) == 1
        assert docs[0]["id"] == document_id

        stats = pipeline_rag.get_stats()
        assert stats.total_documents == 1
        assert stats.total_chunks == ingest_result.chunk_count

    def test_delete_document_workflow(self, pipeline_rag: NexusRAG, pipeline_text_file: Path):

        ingest_result = pipeline_rag.ingest(pipeline_text_file)
        document_id = ingest_result.document_id

        assert len(pipeline_rag.list_documents()) == 1

        deleted = pipeline_rag.delete_document(document_id)
        assert deleted is True

        assert len(pipeline_rag.list_documents()) == 0

        stats = pipeline_rag.get_stats()
        assert stats.total_documents == 0
        assert stats.total_chunks == 0

    def test_delete_nonexistent_document(self, pipeline_rag: NexusRAG):
        deleted = pipeline_rag.delete_document("nonexistent_id")
        assert deleted is False

    def test_empty_stream_ends_with_the_authoritative_answer(self, pipeline_rag: NexusRAG):
        output = list(pipeline_rag.query_streaming("What was measured?"))
        events = [json.loads(part.removeprefix("\ndata: ")) for part in output[1:]]

        assert len(output) == 3
        assert events[0] == {
            "event": "verification",
            "citations_valid": [],
            "faithfulness": None,
            "warnings": [output[0]],
        }
        assert events[1] == {"event": "verified_answer", "answer": output[0]}

    @pytest.mark.parametrize("streaming", [False, True])
    def test_query_propagates_failed_storage_reads_before_returning_an_answer(
        self, pipeline_rag: NexusRAG, monkeypatch, streaming
    ):
        failure = OSError("Index read failed")
        monkeypatch.setattr(
            pipeline_rag.vector_store.table, "count_rows", MagicMock(side_effect=failure)
        )
        pipeline_rag._orchestrator = MagicMock()

        with pytest.raises(VectorStoreError, match="row count") as raised:
            if streaming:
                list(pipeline_rag.query_streaming("What was measured?"))
            else:
                pipeline_rag.query("What was measured?")

        assert raised.value.__cause__ is failure
        pipeline_rag._orchestrator.query.assert_not_called()
        pipeline_rag._orchestrator.query_streaming.assert_not_called()

    @pytest.mark.parametrize("lazy", [False, True])
    def test_failed_sparse_read_can_retry_without_publishing_an_empty_index(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path, monkeypatch, lazy
    ):
        original = pipeline_rag
        result = original.ingest(pipeline_text_file)
        assert result.success
        pipeline_rag = NexusRAG(settings=original.settings)
        pipeline_rag._embedder = original.embedder
        get_all_chunks = pipeline_rag.vector_store.get_all_chunks
        monkeypatch.setattr(
            pipeline_rag.vector_store,
            "get_all_chunks",
            MagicMock(side_effect=VectorStoreError("Cannot read passages from the vector index.")),
        )

        with pytest.raises(VectorStoreError, match="read passages"):
            if lazy:
                _ = pipeline_rag.bm25
            else:
                pipeline_rag._rebuild_bm25_index()

        assert pipeline_rag._bm25 is None
        monkeypatch.setattr(pipeline_rag.vector_store, "get_all_chunks", get_all_chunks)
        assert pipeline_rag.bm25.count() == result.chunk_count

    @pytest.mark.parametrize("stage", ["enumeration", "dense_before", "dense_after", "sparse"])
    def test_partial_deletion_reconciles_indexes_and_can_be_retried(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path, monkeypatch, stage
    ):
        result = pipeline_rag.ingest(pipeline_text_file)
        other = pipeline_rag.ingest_bytes(
            b"A separate document records voltage readings for sensor C.", "voltage.txt", ".txt"
        )
        assert result.success and other.success
        other_ids = pipeline_rag.vector_store.get_chunk_ids_by_document(other.document_id)
        failure = OSError(f"Failed during {stage}")
        if stage == "enumeration":
            owner, method = pipeline_rag.vector_store, "get_chunk_ids_by_document"
        elif stage.startswith("dense"):
            owner, method = pipeline_rag.vector_store, "delete_by_document"
        else:
            owner, method = pipeline_rag.bm25, "remove"
        operation = getattr(owner, method)

        def fail(*args, **kwargs):
            if stage == "dense_after":
                operation(*args, **kwargs)
            raise failure

        monkeypatch.setattr(owner, method, fail)
        with pytest.raises(OSError, match=f"Failed during {stage}"):
            pipeline_rag.delete_document(result.document_id)
        monkeypatch.setattr(owner, method, operation)

        assert pipeline_rag.document_store.exists(result.document_id)
        assert pipeline_rag.document_store.exists(other.document_id)
        assert pipeline_rag.vector_store.get_chunk_ids_by_document(other.document_id) == other_ids
        assert pipeline_rag.bm25.count() == pipeline_rag.vector_store.count()
        if stage in {"dense_after", "sparse"}:
            assert pipeline_rag.vector_store.get_chunk_ids_by_document(result.document_id) == set()
            assert pipeline_rag.bm25.retrieve("thirty-second") == []
        else:
            assert pipeline_rag.bm25.count() == result.chunk_count + other.chunk_count

        assert pipeline_rag.delete_document(result.document_id)
        assert not pipeline_rag.document_store.exists(result.document_id)
        assert pipeline_rag.bm25.count() == pipeline_rag.vector_store.count() == other.chunk_count
        assert {chunk.id for chunk in pipeline_rag.vector_store.get_all_chunks()} == other_ids
        restarted = NexusRAG(settings=pipeline_rag.settings)
        restarted._embedder = pipeline_rag.embedder
        assert restarted.bm25.count() == other.chunk_count

    def test_failed_deletion_recovery_discards_stale_sparse_handles(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path, monkeypatch
    ):
        result = pipeline_rag.ingest(pipeline_text_file)
        assert result.success
        sparse = pipeline_rag.bm25
        pipeline_rag._orchestrator = MagicMock()
        get_all_chunks = pipeline_rag.vector_store.get_all_chunks
        monkeypatch.setattr(
            sparse, "remove", MagicMock(side_effect=OSError("Sparse deletion failed"))
        )
        monkeypatch.setattr(
            pipeline_rag.vector_store,
            "get_all_chunks",
            MagicMock(side_effect=VectorStoreError("Storage remains unavailable")),
        )

        with pytest.raises(OSError, match="Sparse deletion failed"):
            pipeline_rag.delete_document(result.document_id)

        assert pipeline_rag._bm25 is None
        assert pipeline_rag._orchestrator is None
        assert pipeline_rag.document_store.exists(result.document_id)
        assert pipeline_rag.vector_store.count() == 0
        monkeypatch.setattr(pipeline_rag.vector_store, "get_all_chunks", get_all_chunks)
        assert pipeline_rag.bm25 is not sparse
        assert pipeline_rag.bm25.count() == 0
        assert pipeline_rag.delete_document(result.document_id)
        assert pipeline_rag.document_store.count() == 0

    def test_delete_document_with_more_than_100000_chunks_removes_all_sparse_entries(
        self, pipeline_rag: NexusRAG
    ):
        pipeline_rag._embedder.embedding_config = replace(
            pipeline_rag.embedder.embedding_config, dimension=4
        )
        document = pipeline_rag.parser.parse_bytes(b"Current is 42 mA.", "large.txt", ".txt")
        count = 100_001
        chunks = [
            Chunk(id=f"large-{i}", content="Current is 42 mA.", document_id=document.id)
            for i in range(count)
        ]
        pipeline_rag.document_store.add(document)
        pipeline_rag.vector_store.add(chunks, np.ones((count, 4), dtype=np.float32))
        assert pipeline_rag.bm25.count() == count

        assert pipeline_rag.delete_document(document.id)

        assert pipeline_rag.vector_store.count() == pipeline_rag.bm25.count() == 0
        assert pipeline_rag.bm25.retrieve("42 mA") == []
        assert pipeline_rag.document_store.count() == 0

    @pytest.mark.parametrize("failure", ["unlink", "index_save"])
    def test_delete_metadata_failure_keeps_retrieval_indexes_synced(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path, monkeypatch, failure: str
    ):
        result = pipeline_rag.ingest(pipeline_text_file)
        assert result.success
        store = pipeline_rag.document_store
        error = OSError("document metadata unavailable")
        if failure == "unlink":
            monkeypatch.setattr(Path, "unlink", MagicMock(side_effect=error))
        else:
            monkeypatch.setattr(store, "_save_index", MagicMock(side_effect=error))

        with pytest.raises(OSError, match="document metadata unavailable"):
            pipeline_rag.delete_document(result.document_id)

        assert pipeline_rag.vector_store.count() == pipeline_rag.bm25.count() == 0
        assert store.exists(result.document_id) is (failure == "unlink")
        assert pipeline_rag.query("What methods were used?").sources == []


class TestMultiFileIngestion:
    def test_ingest_directory(self, pipeline_rag: NexusRAG, docs_directory: Path):
        results = pipeline_rag.ingest_directory(docs_directory)

        assert len(results) == 3
        assert all(r.success for r in results)

        docs = pipeline_rag.list_documents()
        assert len(docs) == 3

        for result in results:
            assert result.chunk_count > 0

    def test_ingest_directory_stats(self, pipeline_rag: NexusRAG, docs_directory: Path):
        results = pipeline_rag.ingest_directory(docs_directory)
        total_chunks = sum(r.chunk_count for r in results)
        total_words = sum(r.word_count for r in results)

        stats = pipeline_rag.get_stats()
        assert stats.total_documents == 3
        assert stats.total_chunks == total_chunks
        assert stats.total_words == total_words

    def test_ingest_directory_with_unsupported_format(self, pipeline_rag: NexusRAG, tmp_path: Path):
        test_dir = tmp_path / "mixed_docs"
        test_dir.mkdir()

        (test_dir / "valid.txt").write_text("Valid content here.\n" * 10)
        (test_dir / "unsupported.xyz").write_text("This should be skipped.\n")

        results = pipeline_rag.ingest_directory(test_dir)

        assert len(results) == 2
        by_name = {r.filename: r for r in results}
        assert by_name["valid.txt"].success is True
        assert by_name["unsupported.xyz"].success is False
        assert "Unsupported file type" in (by_name["unsupported.xyz"].error or "")


class TestConcurrentWrites:
    @pytest.mark.parametrize("blocked_stage", ["vector_count", "llm_check"])
    def test_stats_snapshot_serializes_stores_without_locking_network_checks(
        self, pipeline_rag: NexusRAG, monkeypatch, blocked_stage
    ):
        store = pipeline_rag.vector_store
        count_rows = store.count
        entered, release = Event(), Event()
        lock = _ObservedWriteLock()
        pipeline_rag._write_lock = lock

        def hold(operation):
            entered.set()
            assert release.wait(timeout=5)
            return operation()

        if blocked_stage == "vector_count":
            monkeypatch.setattr(store, "count", lambda: hold(count_rows))
        else:
            monkeypatch.setattr(pipeline_rag.llm, "is_available", lambda: hold(lambda: True))

        with ThreadPoolExecutor(max_workers=2) as workers:
            stats_result = workers.submit(pipeline_rag.get_stats)
            try:
                assert entered.wait(timeout=5)
                ingest_result = workers.submit(
                    pipeline_rag.ingest_bytes,
                    b"A new document contains current readings from sensor A.",
                    "current.txt",
                    ".txt",
                )
                if blocked_stage == "vector_count":
                    lock.wait_for_threads(2)
                    assert not ingest_result.done()
                else:
                    assert ingest_result.result(timeout=5).success
            finally:
                release.set()
            snapshot = stats_result.result(timeout=5)
            assert ingest_result.result(timeout=5).success

        assert snapshot.total_documents == snapshot.total_chunks == snapshot.total_words == 0
        assert pipeline_rag.document_store.count() == 1
        assert pipeline_rag.vector_store.count() > 0

    @pytest.mark.parametrize("first_operation", ["ingest", "list_documents"])
    def test_ingest_during_cold_index_load(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path, monkeypatch, first_operation
    ):
        pipeline_rag.vector_store.count()
        store = pipeline_rag.document_store
        load_index = store._load_index
        index_loaded = Event()
        release_index = Event()
        second_started = Event()

        def delayed_load():
            index = load_index()
            if not index_loaded.is_set():
                index_loaded.set()
                assert release_index.wait(timeout=10)
            return index

        def second_ingest():
            second_started.set()
            return pipeline_rag.ingest(pipeline_text_file)

        monkeypatch.setattr(store, "_load_index", delayed_load)
        with ThreadPoolExecutor(max_workers=2) as pool:
            first = (
                pool.submit(pipeline_rag.ingest, pipeline_text_file)
                if first_operation == "ingest"
                else pool.submit(pipeline_rag.list_documents)
            )
            try:
                assert index_loaded.wait(timeout=5)
                second = pool.submit(second_ingest)
                assert second_started.wait(timeout=5)
                with pytest.raises(TimeoutError):
                    second.result(timeout=1)
            finally:
                release_index.set()
            first_result = first.result(timeout=10)
            results = [second.result(timeout=10)]
            if first_operation == "ingest":
                results.append(first_result)

        successful = [result for result in results if result.success]
        rejected = [result for result in results if not result.success]
        assert len(successful) == 1
        if first_operation == "ingest":
            assert len(rejected) == 1
            assert rejected[0].error == "Document already exists"
        else:
            assert not rejected
        assert store.count() == 1
        assert pipeline_rag.vector_store.count() == successful[0].chunk_count
        assert pipeline_rag.bm25.count() == successful[0].chunk_count

    def test_parallel_duplicate_ingest_preserves_document(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path, monkeypatch
    ):
        pipeline_rag.vector_store.count()
        pipeline_rag.document_store.count()
        embed = pipeline_rag.embedder.embed
        barrier = Barrier(2)

        def simultaneous_embed(*args, **kwargs):
            result = embed(*args, **kwargs)
            barrier.wait(timeout=10)
            return result

        monkeypatch.setattr(pipeline_rag.embedder, "embed", simultaneous_embed)
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(pipeline_rag.ingest, pipeline_text_file) for _ in range(2)]
            results = [future.result(timeout=20) for future in futures]

        successful = [result for result in results if result.success]
        rejected = [result for result in results if not result.success]
        assert len(successful) == len(rejected) == 1
        assert rejected[0].error == "Document already exists"
        assert pipeline_rag.document_store.count() == 1
        assert pipeline_rag.vector_store.count() == successful[0].chunk_count
        assert pipeline_rag.bm25.count() == successful[0].chunk_count

    def test_parallel_ingest_keeps_indexes_consistent(self, pipeline_rag: NexusRAG, tmp_path: Path):
        files = []
        for i in range(4):
            p = tmp_path / f"doc_{i}.txt"
            p.write_text(f"Document number {i}.\n" + f"Unique sentence {i} repeated often.\n" * 30)
            files.append(p)

        barrier = Barrier(len(files))

        def worker(fp: Path) -> IngestResult:
            barrier.wait(timeout=5)
            return pipeline_rag.ingest(fp)

        with ThreadPoolExecutor(max_workers=len(files)) as workers:
            futures = [workers.submit(worker, path) for path in files]
            results = [future.result(timeout=20) for future in futures]

        assert len(results) == len(files)
        assert all(r.success for r in results)
        total_chunks = sum(r.chunk_count for r in results)
        assert pipeline_rag.vector_store.count() == total_chunks
        assert pipeline_rag.bm25.count() == total_chunks


class TestDuplicateDetection:
    def test_duplicate_file_detection(self, pipeline_rag: NexusRAG, pipeline_text_file: Path):

        result1 = pipeline_rag.ingest(pipeline_text_file)
        assert result1.success is True

        result2 = pipeline_rag.ingest(pipeline_text_file)
        assert result2.success is False
        assert result2.error == "Document already exists"
        assert result2.chunk_count == 0

        docs = pipeline_rag.list_documents()
        assert len(docs) == 1

    def test_duplicate_in_batch_ingest(self, pipeline_rag: NexusRAG, tmp_path: Path):
        docs_dir = tmp_path / "dup_test"
        docs_dir.mkdir()

        content = "Test content for duplicate detection.\n" * 5
        for name in ("first", "second"):
            directory = docs_dir / name
            directory.mkdir()
            (directory / "paper.txt").write_text(content)

        results = pipeline_rag.ingest_directory(docs_dir, recursive=True)

        assert len(results) == 2
        successful = [r for r in results if r.success]
        rejected = [r for r in results if not r.success]
        assert len(successful) == len(rejected) == 1
        assert rejected[0].error == "Document already exists"
        assert rejected[0].document_id == successful[0].document_id
        assert pipeline_rag.document_store.count() == 1
        assert pipeline_rag.vector_store.count() == successful[0].chunk_count
        assert pipeline_rag.bm25.count() == successful[0].chunk_count


class TestClearOperations:
    def test_clear_metadata_failure_keeps_retrieval_indexes_synced(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path, monkeypatch
    ):
        assert pipeline_rag.ingest(pipeline_text_file).success
        monkeypatch.setattr(
            pipeline_rag.document_store,
            "clear",
            MagicMock(side_effect=OSError("document metadata unavailable")),
        )

        with pytest.raises(OSError, match="document metadata unavailable"):
            pipeline_rag.clear_all()

        assert pipeline_rag.vector_store.count() == pipeline_rag.bm25.count() == 0
        assert pipeline_rag.query("What methods were used?").sources == []

    def test_clear_then_reingest_updates_existing_retriever(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path
    ):
        assert pipeline_rag.ingest(pipeline_text_file).success
        sparse = pipeline_rag.orchestrator.retriever.base.sparse
        assert sparse.retrieve("research")

        pipeline_rag.clear_all()
        assert not sparse.retrieve("research")
        result = pipeline_rag.ingest(pipeline_text_file)

        assert result.success
        assert sparse.count() == result.chunk_count
        assert sparse.retrieve("research")

    def test_clear_all_documents(self, pipeline_rag: NexusRAG, docs_directory: Path):

        pipeline_rag.ingest_directory(docs_directory)
        assert pipeline_rag.get_stats().total_documents == 3

        pipeline_rag.clear_all()

        assert pipeline_rag.get_stats().total_documents == 0
        assert pipeline_rag.get_stats().total_chunks == 0
        assert len(pipeline_rag.list_documents()) == 0

    def test_clear_all_then_reingest(self, pipeline_rag: NexusRAG, pipeline_text_file: Path):

        result1 = pipeline_rag.ingest(pipeline_text_file)
        assert result1.success is True

        pipeline_rag.clear_all()
        assert pipeline_rag.get_stats().total_documents == 0

        result2 = pipeline_rag.ingest(pipeline_text_file)
        assert result2.success is True


class TestStatisticsAndMetadata:
    def test_get_stats_reports_an_unavailable_llm(self, pipeline_rag: NexusRAG):
        pipeline_rag.llm.is_available.return_value = False

        stats = pipeline_rag.get_stats()

        assert stats.llm_available is False

    def test_get_stats_empty_system(self, pipeline_rag: NexusRAG):
        stats = pipeline_rag.get_stats()

        assert stats.total_documents == 0
        assert stats.total_chunks == 0
        assert stats.total_words == 0
        assert stats.llm_available is True
        assert stats.embedding_model == "test-model"
        assert stats.llm_model == "test-llm"

    def test_get_stats_after_ingestion(self, pipeline_rag: NexusRAG, pipeline_text_file: Path):
        result = pipeline_rag.ingest(pipeline_text_file)

        stats = pipeline_rag.get_stats()
        assert stats.total_documents == 1
        assert stats.total_chunks == result.chunk_count
        assert stats.total_words > 0
        assert stats.storage_path == str(pipeline_rag.settings.data_dir)

    def test_list_documents_empty(self, pipeline_rag: NexusRAG):
        docs = pipeline_rag.list_documents()
        assert docs == []

    def test_list_documents_with_metadata(self, pipeline_rag: NexusRAG, pipeline_text_file: Path):
        ingest_result = pipeline_rag.ingest(pipeline_text_file)

        docs = pipeline_rag.list_documents()
        assert len(docs) == 1

        doc = docs[0]
        assert doc["id"] == ingest_result.document_id
        assert "filename" in doc or "original_filename" in doc

    def test_stats_with_multiple_documents(self, pipeline_rag: NexusRAG, docs_directory: Path):
        results = pipeline_rag.ingest_directory(docs_directory)

        stats = pipeline_rag.get_stats()
        assert stats.total_documents == len(results)

        total_chunks_expected = sum(r.chunk_count for r in results)
        assert stats.total_chunks == total_chunks_expected


class TestPipelineErrors:
    def test_ingest_empty_file(self, pipeline_rag: NexusRAG, tmp_path: Path):
        empty_file = tmp_path / "empty.txt"
        empty_file.write_text("")

        result = pipeline_rag.ingest(empty_file)

        assert result.success is False
        assert result.error == "No content extracted"
        assert result.chunk_count == 0

    def test_ingest_nonexistent_file(self, pipeline_rag: NexusRAG):
        result = pipeline_rag.ingest("/nonexistent/path/file.txt")

        assert result.success is False
        assert "Ingestion failed" in result.error or "FileNotFoundError" in result.error

    def test_ingest_unsupported_format(self, pipeline_rag: NexusRAG, tmp_path: Path):
        unsupported = tmp_path / "file.xyz"
        unsupported.write_text("Some content")

        result = pipeline_rag.ingest(unsupported)

        assert result.success is False
        assert "Ingestion failed" in result.error

    def test_ingest_directory_nonexistent(self, pipeline_rag: NexusRAG):
        with pytest.raises(ValueError, match="Not a directory"):
            pipeline_rag.ingest_directory("/nonexistent/directory")

    def test_ingest_bytes_empty(self, pipeline_rag: NexusRAG):
        result = pipeline_rag.ingest_bytes(b"", "empty.txt", ".txt")

        assert result.success is False
        assert result.chunk_count == 0

    def test_ingest_bytes_valid(self, pipeline_rag: NexusRAG):
        content = b"This is test content for bytes ingestion.\n" * 5
        result = pipeline_rag.ingest_bytes(content, "test.txt", ".txt")

        assert result.success is True
        assert result.filename == "test.txt"
        assert result.chunk_count > 0

    def test_ingest_bytes_without_extension(self, pipeline_rag: NexusRAG):
        content = b"Test content.\n" * 5

        result = pipeline_rag.ingest_bytes(content, "test.txt", "txt")

        assert result.success is True


class TestLazyLoading:
    def test_components_not_loaded_initially(self, pipeline_rag: NexusRAG):

        assert pipeline_rag._parser is None
        assert pipeline_rag._chunker is None
        assert pipeline_rag._vector_store is None
        assert pipeline_rag._document_store is None
        assert pipeline_rag._bm25 is None

    @pytest.mark.parametrize("component", ["parser", "chunker", "document_store"])
    def test_component_lazy_loads_on_access(self, pipeline_rag, component):
        assert getattr(pipeline_rag, "_" + component) is None
        loaded = getattr(pipeline_rag, component)
        assert loaded is not None
        assert getattr(pipeline_rag, "_" + component) is loaded


class TestSingletonBehavior:
    def test_get_scinexusrag_returns_same_instance(self, tmp_path):
        first = get_scinexusrag(Settings(data_dir=tmp_path))
        assert get_scinexusrag() is first

    def test_get_scinexusrag_with_different_settings(self, tmp_path):
        first = get_scinexusrag(Settings(data_dir=tmp_path / "first"))
        assert get_scinexusrag(Settings(data_dir=tmp_path / "second")) is first


class TestIngestBytes:
    def test_ingest_bytes_text(self, pipeline_rag: NexusRAG):
        content = b"""
        # Test Document

        Sensor A records temperature in degrees Celsius.
        Each reading includes a timestamp.

        Sensor B uses the same sampling interval.
        Readings with matching timestamps form a measurement pair.
        """

        result = pipeline_rag.ingest_bytes(content, "bytes_doc.txt", ".txt")

        assert result.success is True
        assert result.filename == "bytes_doc.txt"
        assert result.chunk_count > 0

    def test_ingest_bytes_md_format(self, pipeline_rag: NexusRAG):
        content = b"""# Markdown Test

## Section 1
Content here for testing.

## Section 2
More content for proper chunking.
"""

        result = pipeline_rag.ingest_bytes(content, "test.md", ".md")

        assert result.success is True
        assert result.chunk_count > 0

    def test_ingest_bytes_duplicate_detection(self, pipeline_rag: NexusRAG):
        content = b"Unique test content for duplicate detection.\n" * 5

        result1 = pipeline_rag.ingest_bytes(content, "dup.txt", ".txt")
        assert result1.success is True

        result2 = pipeline_rag.ingest_bytes(content, "dup.txt", ".txt")
        assert result2.success is False
        assert result2.document_id == result1.document_id
        assert "already exists" in (result2.error or "")


class TestEdgeCases:
    def test_ingest_very_short_document(self, pipeline_rag: NexusRAG, tmp_path: Path):
        short_file = tmp_path / "short.txt"
        short_file.write_text("Short content.")

        result = pipeline_rag.ingest(short_file)

        assert result.success, result.error
        assert result.chunk_count == 1
        assert result.word_count == 2

    def test_ingest_large_document(self, pipeline_rag: NexusRAG, tmp_path: Path):
        large_file = tmp_path / "large.txt"

        large_content = "This is a test sentence that will be repeated many times. " * 200
        large_file.write_text(large_content)

        result = pipeline_rag.ingest(large_file)

        assert result.success is True
        assert result.chunk_count > 1

    def test_concurrent_operations_sequence(self, pipeline_rag: NexusRAG, tmp_path: Path):
        file1 = tmp_path / "doc1.txt"
        file2 = tmp_path / "doc2.txt"

        content1 = "Document 1 content. " * 20
        content2 = "Document 2 content. " * 20

        file1.write_text(content1)
        file2.write_text(content2)

        result1 = pipeline_rag.ingest(file1)
        result2 = pipeline_rag.ingest(file2)

        assert result1.success is True
        assert result2.success is True

        docs = pipeline_rag.list_documents()
        assert len(docs) == 2

        deleted = pipeline_rag.delete_document(result1.document_id)
        assert deleted is True

        docs = pipeline_rag.list_documents()
        assert len(docs) == 1


class TestIntegrationWithRealComponents:
    def test_restart_rebuilds_sparse_index_from_persisted_chunks(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path
    ):
        result = pipeline_rag.ingest(pipeline_text_file)
        assert result.success

        restarted = NexusRAG(settings=pipeline_rag.settings)
        restarted._embedder = pipeline_rag.embedder
        restarted._llm = pipeline_rag.llm

        assert restarted.bm25.count() == result.chunk_count
        assert restarted.bm25.retrieve("research")
        assert restarted.vector_store.list_documents() == [result.document_id]

    def test_with_real_parser_and_chunker(self, pipeline_rag: NexusRAG, pipeline_text_file: Path):
        result = pipeline_rag.ingest(pipeline_text_file)

        assert result.success is True
        assert result.chunk_count > 0

        stats = pipeline_rag.get_stats()
        assert stats.total_chunks == result.chunk_count

    def test_document_store_persistence(self, pipeline_rag: NexusRAG, pipeline_text_file: Path):
        ingest_result = pipeline_rag.ingest(pipeline_text_file)
        assert ingest_result.success is True

        docs = pipeline_rag.list_documents()
        assert len(docs) == 1

        doc = docs[0]
        assert doc["id"] == ingest_result.document_id


class TestParametrized:
    @pytest.mark.parametrize(
        "filename,content",
        [
            ("test1.txt", "Test content one.\n" * 10),
            ("test2.txt", "Test content two.\n" * 10),
            ("test3.txt", "Different content.\n" * 10),
        ],
    )
    def test_ingest_multiple_variations(
        self, pipeline_rag: NexusRAG, tmp_path: Path, filename: str, content: str
    ):
        file_path = tmp_path / filename
        file_path.write_text(content)

        result = pipeline_rag.ingest(file_path)

        assert result.success is True
        assert result.filename == filename

    @pytest.mark.parametrize(
        "format_ext,format_name",
        [
            (".txt", "text"),
            (".md", "markdown"),
        ],
    )
    def test_various_formats(
        self, pipeline_rag: NexusRAG, tmp_path: Path, format_ext: str, format_name: str
    ):
        file_path = tmp_path / f"doc{format_ext}"

        if format_ext == ".txt":
            content = "Test content for text format.\n" * 10
        else:
            content = "# Test\n\nContent for markdown format.\n" * 5

        file_path.write_text(content)

        result = pipeline_rag.ingest(file_path)

        assert result.success is True


class TestIngestRollback:
    @pytest.mark.parametrize("invalid_field", ["content", "metadata"])
    def test_prewrite_validation_preserves_preexisting_document_and_vectors(
        self, pipeline_rag: NexusRAG, invalid_field
    ):
        document = ParsedDocument("existing", "Original evidence.", {"filename": "current.txt"})
        chunks = pipeline_rag.chunker.chunk(document)
        embeddings = pipeline_rag.embedder.embed([chunk.content for chunk in chunks])
        assert pipeline_rag.document_store.count() == 0
        writer = DocumentStore(pipeline_rag.document_store.path)
        writer.add(document)
        writer.update_metadata(document.id, "notes", {"tags": ["original"]})
        pipeline_rag.vector_store.add(chunks, embeddings)
        files_before = {path.name: path.read_bytes() for path in writer.path.iterdir()}
        vectors_before = pipeline_rag.vector_store.table.to_arrow()
        invalid = replace(document, **{invalid_field: 42})

        with pytest.raises(ValueError):
            pipeline_rag._persist(invalid, chunks, embeddings)

        assert {path.name: path.read_bytes() for path in writer.path.iterdir()} == files_before
        assert pipeline_rag.vector_store.table.to_arrow().equals(vectors_before)
        assert DocumentStore(writer.path).list_with_metadata() == writer.list_with_metadata()
        assert pipeline_rag.bm25.count() == pipeline_rag.vector_store.count() == len(chunks)

    @pytest.mark.parametrize("entry_point", ["ingest", "ingest_bytes"])
    def test_duplicate_rejected_after_cached_inventory_preserves_existing_stores(
        self, pipeline_rag: NexusRAG, tmp_path: Path, entry_point
    ):
        source = tmp_path / "current.txt"
        source.write_text("The measured current is forty-two milliamps.", encoding="utf-8")
        document = pipeline_rag.parser.parse(source)
        chunks = pipeline_rag.chunker.chunk(document)
        assert pipeline_rag.document_store.count() == 0

        writer = DocumentStore(pipeline_rag.document_store.path)
        writer.add(document)
        writer.update_metadata(document.id, "notes", {"tags": ["original"]})
        writer.update_metadata(document.id, "chunk_count", len(chunks))
        pipeline_rag.vector_store.add(
            chunks, pipeline_rag.embedder.embed([chunk.content for chunk in chunks])
        )
        existing_metadata = writer.list_with_metadata()
        existing_files = {path.name: path.read_bytes() for path in writer.path.iterdir()}
        existing_vectors = pipeline_rag.vector_store.table.to_arrow()
        assert pipeline_rag.document_store.count() == 0

        if entry_point == "ingest":
            result = pipeline_rag.ingest(source)
        else:
            result = pipeline_rag.ingest_bytes(source.read_bytes(), source.name, ".txt")

        assert not result.success
        assert result.error == "Document already exists"
        assert result.document_id == document.id
        assert {path.name: path.read_bytes() for path in writer.path.iterdir()} == existing_files
        assert DocumentStore(writer.path).list_with_metadata() == existing_metadata
        assert pipeline_rag.document_store.list_with_metadata() == existing_metadata
        assert pipeline_rag.vector_store.table.to_arrow().equals(existing_vectors)
        assert pipeline_rag.bm25.count() == pipeline_rag.vector_store.count() == len(chunks)

    def test_failed_count_during_ingest_keeps_dense_rows_attached_to_a_document(
        self, pipeline_rag: NexusRAG, monkeypatch
    ):
        table = pipeline_rag.vector_store.table
        count_rows = table.count_rows
        failure = OSError("Row count temporarily unavailable")
        monkeypatch.setattr(table, "count_rows", MagicMock(side_effect=failure))

        with pytest.raises(VectorStoreError, match="row count") as raised:
            pipeline_rag.ingest_bytes(
                b"The measured current is forty-two milliamps.", "current.txt", ".txt"
            )

        assert raised.value.__cause__ is failure
        monkeypatch.setattr(table, "count_rows", count_rows)
        chunks = pipeline_rag.vector_store.get_all_chunks()
        documents = pipeline_rag.document_store.list_with_metadata()
        assert chunks and len(documents) == 1
        assert {chunk.document_id for chunk in chunks} == set(documents)
        assert pipeline_rag.bm25.count() == pipeline_rag.vector_store.count() == len(chunks)
        document_id = chunks[0].document_id
        restarted = NexusRAG(settings=pipeline_rag.settings)
        restarted._embedder = pipeline_rag.embedder
        assert restarted.document_store.exists(document_id)
        assert restarted.bm25.count() == len(chunks)
        assert restarted.delete_document(document_id)
        assert restarted.document_store.count() == 0
        assert restarted.bm25.count() == restarted.vector_store.count() == 0

    @pytest.mark.parametrize(
        "stage", ["dense_before", "dense_partial", "dense_after", "sparse_read", "document_file"]
    )
    def test_failed_ingest_rollback_preserves_the_original_error_and_can_be_recovered(
        self, pipeline_rag: NexusRAG, monkeypatch, stage
    ):
        other = pipeline_rag.ingest_bytes(
            b"A separate document records voltage for sensor B.", "voltage.txt", ".txt"
        )
        assert other.success
        document = pipeline_rag.parser.parse_bytes(
            b"The measured current is forty-two milliamps.", "current.txt", ".txt"
        )
        chunks = [
            Chunk(
                id=f"current-{index}",
                content=f"Current measurement {index}",
                document_id=document.id,
            )
            for index in range(3)
        ]
        original_error = OSError("Document metadata update failed")
        monkeypatch.setattr(
            pipeline_rag.document_store, "update_metadata", MagicMock(side_effect=original_error)
        )
        if stage.startswith("dense"):
            owner, method = pipeline_rag.vector_store, "delete_by_document"
        elif stage == "sparse_read":
            owner, method = pipeline_rag.vector_store, "get_all_chunks"
        else:
            owner, method = pipeline_rag.document_store, "delete"
        cleanup = getattr(owner, method)

        def fail_cleanup(*args, **kwargs):
            if stage == "dense_partial":
                pipeline_rag.vector_store.delete([chunks[0].id])
            elif stage == "dense_after":
                cleanup(*args, **kwargs)
            raise OSError(f"Rollback failed during {stage}")

        monkeypatch.setattr(owner, method, fail_cleanup)
        with pytest.raises(OSError, match="Document metadata update failed") as raised:
            pipeline_rag._persist(document, chunks, np.ones((3, 384), dtype=np.float32))

        assert raised.value is original_error
        assert pipeline_rag.document_store.exists(document.id)
        assert pipeline_rag.document_store.exists(other.document_id)
        if stage == "sparse_read":
            assert pipeline_rag._bm25 is None
            assert pipeline_rag._orchestrator is None
        monkeypatch.setattr(owner, method, cleanup)
        expected_remaining = 3 if stage == "dense_before" else 2 if stage == "dense_partial" else 0
        assert (
            len(pipeline_rag.vector_store.get_chunks_by_document(document.id)) == expected_remaining
        )
        assert (
            pipeline_rag.bm25.count()
            == pipeline_rag.vector_store.count()
            == other.chunk_count + expected_remaining
        )
        assert all(
            pipeline_rag.document_store.exists(chunk.document_id)
            for chunk in pipeline_rag.vector_store.get_all_chunks()
        )
        restarted = NexusRAG(settings=pipeline_rag.settings)
        restarted._embedder = pipeline_rag.embedder
        assert restarted.bm25.count() == other.chunk_count + expected_remaining
        assert restarted.delete_document(document.id)
        assert restarted.document_store.count() == 1
        assert restarted.document_store.exists(other.document_id)
        assert restarted.bm25.count() == restarted.vector_store.count() == other.chunk_count

    def test_document_index_failure_rolls_back_document_file(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path, monkeypatch
    ):
        store = pipeline_rag.document_store
        store.count()
        save_index = store._save_index
        failed = False

        def fail_first_save():
            nonlocal failed
            if not failed:
                failed = True
                raise OSError("simulated index write failure")
            save_index()

        monkeypatch.setattr(store, "_save_index", fail_first_save)
        result = pipeline_rag.ingest(pipeline_text_file)

        assert not result.success
        assert store.count() == 0
        assert pipeline_rag.vector_store.count() == 0
        assert pipeline_rag.bm25.count() == 0
        assert all(path == store._index_path for path in store.path.glob("*.json"))
        assert pipeline_rag.ingest(pipeline_text_file).success

    def test_vector_failure_rolls_back_document_store(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path, monkeypatch
    ):
        def boom(chunks, embeddings):
            raise RuntimeError("simulated storage failure")

        monkeypatch.setattr(pipeline_rag.vector_store, "add", boom)

        result = pipeline_rag.ingest(pipeline_text_file)

        assert result.success is False
        assert pipeline_rag.document_store.count() == 0
        assert pipeline_rag.bm25.count() == 0

    def test_bm25_failure_rolls_back_all_stores(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path, monkeypatch
    ):
        def boom(chunks):
            raise RuntimeError("simulated index failure")

        monkeypatch.setattr(pipeline_rag.bm25, "add_incremental", boom)

        result = pipeline_rag.ingest(pipeline_text_file)

        assert result.success is False
        assert pipeline_rag.document_store.count() == 0
        assert pipeline_rag.vector_store.count() == 0


class TestUnloadModels:
    def test_unload_resets_lazy_handles(self, pipeline_rag: NexusRAG):

        assert pipeline_rag.embedder is not None
        _ = pipeline_rag.llm
        _ = pipeline_rag.orchestrator

        pipeline_rag.unload_models()

        assert all(
            getattr(pipeline_rag, "_" + name) is None
            for name in ("embedder", "llm", "orchestrator", "vector_store", "bm25")
        )


class _ObservedWriteLock:
    def __init__(self):
        self.lock = RLock()
        self.attempts = Queue()

    def __enter__(self):
        self.attempts.put(get_ident())
        if not self.lock.acquire(timeout=5):
            raise TimeoutError("Pipeline lock was not released")
        return self

    def __exit__(self, *_args):
        self.lock.release()

    def wait_for_attempts(self, count):
        return [self.attempts.get(timeout=5) for _ in range(count)]

    def wait_for_threads(self, count):
        threads = set()
        while len(threads) < count:
            threads.add(self.attempts.get(timeout=5))
        return threads


class TestLazyResourceLifecycle:
    @pytest.mark.parametrize("resource", ["bm25", "orchestrator"])
    @pytest.mark.parametrize("fail_first", [False, True])
    def test_lazy_retrieval_handles_wait_for_a_complete_sparse_index(
        self, pipeline_rag: NexusRAG, monkeypatch, resource, fail_first
    ):
        store = pipeline_rag.vector_store
        get_all_chunks = store.get_all_chunks
        lock = _ObservedWriteLock()
        pipeline_rag._write_lock = lock
        entered, release = Event(), Event()
        reads = 0

        def hold_first_read():
            nonlocal reads
            reads += 1
            if reads == 1:
                entered.set()
                assert release.wait(timeout=5)
                if fail_first:
                    raise VectorStoreError("Sparse initialization failed")
            return get_all_chunks()

        monkeypatch.setattr(store, "get_all_chunks", hold_first_read)
        with ThreadPoolExecutor(max_workers=2) as workers:
            first = workers.submit(lambda: getattr(pipeline_rag, resource))
            try:
                assert entered.wait(timeout=5)
                assert pipeline_rag._bm25 is None
                assert pipeline_rag._orchestrator is None
                second = workers.submit(lambda: getattr(pipeline_rag, resource))
                lock.wait_for_threads(2)
                assert not second.done()
            finally:
                release.set()
            if fail_first:
                with pytest.raises(VectorStoreError, match="Sparse initialization failed"):
                    first.result(timeout=5)
                assert second.result(timeout=5) is getattr(pipeline_rag, resource)
                assert reads == 2
            else:
                assert first.result(timeout=5) is second.result(timeout=5)
                assert reads == 1
        assert pipeline_rag.bm25.count() == 0

    @pytest.mark.parametrize(
        "resource,factory_name",
        [("embedder", "Embedder"), ("vector_store", "VectorStore"), ("llm", "LLMClient")],
    )
    def test_concurrent_first_access_constructs_one_handle(
        self, test_settings, mock_embedder, monkeypatch, resource, factory_name
    ):
        pipeline_rag = NexusRAG(test_settings)
        if resource == "vector_store":
            pipeline_rag._embedder = mock_embedder
        lock = _ObservedWriteLock()
        pipeline_rag._write_lock = lock
        start = Barrier(4)
        constructing = Event()
        release = Event()
        handle = MagicMock()

        def construct(**_kwargs):
            constructing.set()
            assert release.wait(timeout=5)
            return handle

        def load():
            start.wait(timeout=5)
            return getattr(pipeline_rag, resource)

        factory = MagicMock(side_effect=construct)
        monkeypatch.setattr("scinexusrag.application." + factory_name, factory)
        with ThreadPoolExecutor(max_workers=4) as workers:
            futures = [workers.submit(load) for _ in range(4)]
            try:
                assert constructing.wait(timeout=5)
                assert len(lock.wait_for_threads(4)) == 4
            finally:
                release.set()
            assert all(future.result(timeout=5) is handle for future in futures)

        if resource == "embedder":
            factory.assert_called_once_with(
                model_name="test-model", device="cpu", revision="a" * 40
            )
        elif resource == "vector_store":
            factory.assert_called_once_with(
                path=test_settings.storage.lancedb_path,
                table_name=test_settings.storage.table_name,
                embedding_config=mock_embedder.embedding_config,
            )
        else:
            factory.assert_called_once_with(
                model=test_settings.llm.model,
                base_url=test_settings.llm.base_url,
                timeout=test_settings.llm.timeout,
            )

    @pytest.mark.parametrize(
        "resource,factory_name",
        [("embedder", "Embedder"), ("vector_store", "VectorStore"), ("llm", "LLMClient")],
    )
    def test_failed_constructor_can_retry(
        self, test_settings, mock_embedder, monkeypatch, resource, factory_name
    ):
        pipeline_rag = NexusRAG(test_settings)
        if resource == "vector_store":
            pipeline_rag._embedder = mock_embedder
        pipeline_rag._write_lock = _ObservedWriteLock()
        handle = MagicMock()
        factory = MagicMock(side_effect=[RuntimeError("Constructor failed"), handle])
        monkeypatch.setattr("scinexusrag.application." + factory_name, factory)

        with pytest.raises(RuntimeError, match="Constructor failed"):
            getattr(pipeline_rag, resource)
        assert getattr(pipeline_rag, "_" + resource) is None
        with ThreadPoolExecutor(max_workers=1) as workers:
            assert (
                workers.submit(lambda: getattr(pipeline_rag, resource)).result(timeout=5) is handle
            )
        assert getattr(pipeline_rag, resource) is handle
        assert factory.call_count == 2

    @pytest.mark.parametrize(
        "resource,factory_name", [("embedder", "Embedder"), ("llm", "LLMClient")]
    )
    def test_unload_waits_for_lazy_handle_construction(
        self, test_settings, monkeypatch, resource, factory_name
    ):
        pipeline_rag = NexusRAG(test_settings)
        lock = _ObservedWriteLock()
        pipeline_rag._write_lock = lock
        constructing = Event()
        release = Event()
        original_handle, fresh_handle = MagicMock(), MagicMock()

        def construct(**_kwargs):
            constructing.set()
            assert release.wait(timeout=5)
            return original_handle

        factory = MagicMock(side_effect=construct)
        monkeypatch.setattr("scinexusrag.application." + factory_name, factory)
        with ThreadPoolExecutor(max_workers=2) as workers:
            loading = workers.submit(lambda: getattr(pipeline_rag, resource))
            try:
                assert constructing.wait(timeout=5)
                lock.wait_for_attempts(1)
                unloading = workers.submit(pipeline_rag.unload_models)
                lock.wait_for_attempts(1)
                assert not unloading.done()
            finally:
                release.set()
            assert loading.result(timeout=5) is original_handle
            unloading.result(timeout=5)

        if resource == "embedder":
            original_handle.unload.assert_called_once()
        else:
            original_handle.close.assert_called_once()
        assert all(
            getattr(pipeline_rag, "_" + name) is None
            for name in ("embedder", "llm", "orchestrator", "vector_store", "bm25")
        )
        factory.side_effect = None
        factory.return_value = fresh_handle
        assert getattr(pipeline_rag, resource) is fresh_handle
        assert getattr(pipeline_rag, resource) is fresh_handle
        assert factory.call_count == 2


class TestStoredEmbeddingConfiguration:
    def test_explicit_reset_does_not_require_a_valid_embedder(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path
    ):
        original = pipeline_rag
        assert original.ingest(pipeline_text_file).success
        restarted = NexusRAG(settings=original.settings)
        invalid_embedder = MagicMock()
        config = PropertyMock(side_effect=EmbeddingConfigError("Unpinned model revision"))
        type(invalid_embedder).embedding_config = config
        restarted._embedder = invalid_embedder

        restarted.clear_all()

        config.assert_not_called()
        invalid_embedder.embed.assert_not_called()
        invalid_embedder.embed_query.assert_not_called()
        assert restarted.document_store.count() == 0
        assert (
            restarted.settings.storage.table_name
            not in original.vector_store.db.list_tables().tables
        )
        assert restarted._vector_store is None
        assert restarted._bm25 is None

    def test_changed_model_rejects_query_before_embedding_or_generation(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path
    ):
        original = pipeline_rag
        ingested = original.ingest(pipeline_text_file)
        assert ingested.success
        changed_embedder = MagicMock()
        changed_embedder.dimension = 384
        changed_embedder.embedding_config = replace(
            original.embedder.embedding_config, model_name="another-model"
        )
        restarted = NexusRAG(settings=original.settings)
        restarted._embedder = changed_embedder
        restarted._llm = MagicMock()

        with pytest.raises(EmbeddingConfigError):
            restarted.query("What measurements were recorded?")
        with pytest.raises(EmbeddingConfigError):
            list(restarted.query_streaming("What measurements were recorded?"))

        changed_embedder.embed_query.assert_not_called()
        restarted._llm.generate.assert_not_called()
        assert original.vector_store.count() == ingested.chunk_count
        assert original.document_store.count() == 1

    def test_changed_revision_rejects_ingest_before_embedding_writes_or_rollback(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path, monkeypatch
    ):
        original = pipeline_rag
        ingested = original.ingest(pipeline_text_file)
        assert ingested.success
        changed_embedder = MagicMock()
        changed_embedder.dimension = 384
        changed_embedder.embedding_config = replace(
            original.embedder.embedding_config, revision="b" * 40
        )
        restarted = NexusRAG(settings=original.settings)
        restarted._embedder = changed_embedder
        document = restarted.parser.parse_bytes(
            b"A distinct upload records pressure readings from sensor C.", "new.txt", ".txt"
        )
        writes = MagicMock()
        deletes = MagicMock()
        monkeypatch.setattr(restarted.document_store, "add", writes)
        monkeypatch.setattr(restarted.document_store, "delete", deletes)

        with pytest.raises(EmbeddingConfigError):
            restarted._ingest_document(document, "new.txt", show_progress=False)
        with pytest.raises(EmbeddingConfigError):
            restarted._persist(document, [], np.empty((0, 384), dtype=np.float32))
        new_path = pipeline_text_file.with_name("new.txt")
        new_path.write_text(document.content)
        with pytest.raises(EmbeddingConfigError):
            restarted.ingest(new_path)
        with pytest.raises(EmbeddingConfigError):
            restarted.ingest_bytes(document.content.encode(), "new.txt", ".txt")

        changed_embedder.embed.assert_not_called()
        writes.assert_not_called()
        deletes.assert_not_called()
        assert original.vector_store.count() == ingested.chunk_count
        assert original.document_store.count() == 1

    def test_unload_reloads_configuration_instead_of_reusing_old_store(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path
    ):
        ingested = pipeline_rag.ingest(pipeline_text_file)
        assert ingested.success
        assert pipeline_rag.bm25.count() == ingested.chunk_count
        changed_config = replace(pipeline_rag.embedder.embedding_config, revision="b" * 40)

        pipeline_rag.unload_models()
        changed_embedder = MagicMock()
        changed_embedder.dimension = 384
        changed_embedder.embedding_config = changed_config
        pipeline_rag._embedder = changed_embedder

        assert pipeline_rag._vector_store is None
        assert pipeline_rag._bm25 is None
        with pytest.raises(EmbeddingConfigError):
            pipeline_rag.query("What measurements were recorded?")
        changed_embedder.embed_query.assert_not_called()

    def test_explicit_reset_clears_documents_and_allows_reingestion(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path
    ):
        original = pipeline_rag
        assert original.ingest(pipeline_text_file).success
        changed_config = replace(original.embedder.embedding_config, model_name="another-model")
        original.unload_models()
        original._embedder = MagicMock()
        original._embedder.dimension = 384
        original._embedder.embedding_config = changed_config
        original._embedder.embed.return_value = np.ones((1, 384), dtype=np.float32)

        with pytest.raises(EmbeddingConfigError):
            original.query("What measurements were recorded?")

        original.clear_all()
        assert original.document_store.count() == 0
        assert original.vector_store.count() == 0
        assert original.bm25.count() == 0
        original._embedder.embed.side_effect = lambda texts, **kwargs: np.ones(
            (len(texts), 384), dtype=np.float32
        )
        result = original.ingest(pipeline_text_file)

        assert result.success
        assert original.document_store.count() == 1
        assert original.vector_store.count() == original.bm25.count() == result.chunk_count


class _ImmediateVisibilityStore:
    def __init__(self, real):
        self._real = real
        self._seen: list = []

    def add(self, chunks, embeddings):
        n = self._real.add(chunks, embeddings)
        self._seen.extend(chunks)
        return n

    def get_all_chunks(self):
        return list(self._seen)

    def __getattr__(self, name):
        return getattr(self._real, name)


class TestBM25NoDoubleCount:
    def test_ingest_does_not_double_count_bm25_on_lazy_init(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path
    ):

        pipeline_rag._vector_store = _ImmediateVisibilityStore(pipeline_rag.vector_store)

        result = pipeline_rag.ingest(pipeline_text_file)

        assert result.success
        assert pipeline_rag.bm25.count() == result.chunk_count

    def test_first_ingest_keeps_bm25_and_vector_in_sync(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path
    ):

        r = pipeline_rag.ingest(pipeline_text_file)
        assert r.success
        assert pipeline_rag.bm25.count() == r.chunk_count
        assert pipeline_rag.vector_store.count() == r.chunk_count


class TestEndToEnd:
    def test_ingest_query_cites_a_real_source(
        self, pipeline_rag: NexusRAG, pipeline_text_file: Path
    ):

        pipeline_rag._llm.generate = lambda prompt, **kw: (
            "The sensors are read every thirty seconds [1]."
        )
        query_embedding = np.random.default_rng(0).random(384, dtype=np.float32)
        pipeline_rag._embedder.embed_query = lambda q: query_embedding.copy()

        assert pipeline_rag.ingest(pipeline_text_file).success
        resp = pipeline_rag.query("What methodology was used?")

        assert resp.sources, "query returned no sources"
        assert "[1]" in resp.answer, "answer did not cite a source"
        assert resp.confidence > 0

        assert resp.sources[0].chunk_id


@pytest.mark.parametrize("threshold", [-0.1, 1.1, float("nan"), float("inf"), float("-inf")])
def test_grounding_rejects_invalid_threshold_before_loading(threshold):
    with pytest.raises(ValueError, match="threshold must be between 0 and 1"):
        GroundingVerifier(threshold=threshold)


@pytest.mark.parametrize("threshold", [0.0, 1.0])
def test_grounding_accepts_threshold_boundaries(threshold):
    verifier = GroundingVerifier(threshold=threshold)
    assert verifier.threshold == threshold
    assert verifier._model is None


def test_failed_label_resolution_does_not_leak_into_retry(monkeypatch):
    bad = SimpleNamespace(
        model=SimpleNamespace(
            config=SimpleNamespace(id2label={2: "entailment", "invalid": "contradiction"})
        )
    )
    good = SimpleNamespace(model=SimpleNamespace(config=SimpleNamespace(id2label={})))
    factory = Mock(side_effect=[bad, good])
    module = ModuleType("sentence_transformers")
    module.CrossEncoder = factory
    monkeypatch.setitem(sys.modules, "sentence_transformers", module)
    verifier = GroundingVerifier(model_name="fixture")

    with pytest.raises(ValueError):
        _ = verifier.model

    assert verifier._model is None
    assert (verifier.entail_idx, verifier.contra_idx) == (1, 0)
    assert verifier.model is good
    assert (verifier.entail_idx, verifier.contra_idx) == (1, 0)


@pytest.mark.parametrize("index", [-1, 3])
def test_entailment_cannot_fall_back_to_another_class(index):
    verifier = GroundingVerifier()
    verifier._model = SimpleNamespace(predict=Mock(return_value=np.array([[0.0, 1.0, 2.0]])))
    verifier.entail_idx = index

    with pytest.raises(ValueError, match="entailment label is outside"):
        verifier.entailment_scores([("evidence", "claim")])


def test_unnamed_model_preserves_explicit_class_indices(monkeypatch):
    model = SimpleNamespace(model=SimpleNamespace(config=SimpleNamespace(id2label={})))
    module = ModuleType("sentence_transformers")
    module.CrossEncoder = Mock(return_value=model)
    monkeypatch.setitem(sys.modules, "sentence_transformers", module)
    verifier = GroundingVerifier(model_name="fixture")
    verifier.entail_idx = 2

    assert verifier.model is model
    assert verifier.entail_idx == 2


def test_binary_probabilities_apply_the_sigmoid_once():
    logits = np.array([-2.0, 2.0])

    def predict(pairs, *, activation_fn=None, **kwargs):
        return activation_fn(logits) if activation_fn else 1 / (1 + np.exp(-logits))

    model = SimpleNamespace(predict=Mock(side_effect=predict))
    verifier = GroundingVerifier()
    verifier._model = model

    probabilities = verifier.class_probs([("first", "claim"), ("second", "claim")])

    expected = 1 / (1 + np.exp(-logits))
    np.testing.assert_allclose(probabilities, np.column_stack([1 - expected, expected]))
    assert model.predict.call_args.kwargs["apply_softmax"] is False


def test_multiclass_probabilities_preserve_raw_logit_softmax():
    logits = np.array([[1.2, -0.3, 0.6], [-5.1, 2.0, 1.5]])
    verifier = GroundingVerifier()
    verifier._model = SimpleNamespace(predict=Mock(return_value=logits))
    exp = np.exp(logits - logits.max(axis=1, keepdims=True))

    actual = verifier.class_probs([("first", "claim"), ("second", "claim")])

    np.testing.assert_array_equal(actual, exp / exp.sum(axis=1, keepdims=True))


@pytest.mark.parametrize(
    "prediction,pair_count",
    [
        (np.array([0.1, 0.2, 0.3]), 1),
        (np.ones((1, 3)), 2),
        (np.ones((2, 3)), 1),
        (np.empty((1, 0)), 1),
        (np.array(0.5), 1),
        (np.ones((1, 3, 1)), 1),
    ],
)
def test_grounding_rejects_misaligned_predictions(prediction, pair_count):
    verifier = GroundingVerifier()
    verifier._model = SimpleNamespace(predict=Mock(return_value=prediction))

    with pytest.raises(ValueError, match="one row per pair"):
        verifier.class_probs([("evidence", "claim")] * pair_count)


@pytest.mark.parametrize("number", [float("nan"), float("inf"), float("-inf")])
def test_grounding_rejects_nonfinite_predictions(number):
    verifier = GroundingVerifier()
    verifier._model = SimpleNamespace(predict=Mock(return_value=np.array([[0.0, number, 1.0]])))

    with pytest.raises(ValueError, match="NLI predictions must be finite"):
        verifier.class_probs([("evidence", "claim")])


def test_extreme_binary_logits_remain_finite_without_overflow():
    verifier = GroundingVerifier()
    verifier._model = SimpleNamespace(predict=Mock(return_value=np.array([-10000.0, 0.0, 10000.0])))

    with np.errstate(over="raise"):
        probabilities = verifier.class_probs([("evidence", "claim")] * 3)

    np.testing.assert_array_equal(probabilities, [[1.0, 0.0], [0.5, 0.5], [0.0, 1.0]])


def test_empty_nli_batch_does_not_load_a_model():
    verifier = GroundingVerifier()

    assert verifier.class_probs([]).shape == (0, 3)
    assert verifier._model is None


def test_blank_sources_cannot_ground_an_answer_or_load_a_model():
    verifier = GroundingVerifier()

    report = verifier.verify("The measurement is valid.", ["", None, " \n\t"])

    assert report.faithfulness == 0.0
    assert report.unsupported == ["The measurement is valid."]
    assert verifier._model is None


def test_grounding_preserves_original_source_positions_after_skipping_blanks():
    predict = Mock(return_value=np.array([[0.0, 4.0, -1.0]]))
    verifier = GroundingVerifier()
    verifier._model = SimpleNamespace(predict=predict)

    report = verifier.verify(
        "The measurement is valid.",
        ["", None, " \n\t", SimpleNamespace(content="Measured evidence.")],
    )

    assert report.faithfulness == 1.0
    assert report.sentences[0].source_index == 3
    assert predict.call_args.args == ([("Measured evidence.", "The measurement is valid.")],)


def test_invalid_source_content_fails_before_model_loading():
    verifier = GroundingVerifier()

    with pytest.raises(ValueError, match="source text must be a string"):
        verifier.verify("The measurement is valid.", [SimpleNamespace(content=123)])

    assert verifier._model is None


@pytest.mark.parametrize(
    "answer",
    [
        "  Claim [1].  \n\n\n    indented  evidence\n",
        "Claim [1].\n\n```python\n    values =  [1,  2]\n    print(values)\n```\n",
        "Claim [1].\n\n| Measure  | Value |\n| -------- | ----- |\n| pH       | 7.4   |\n",
        "Main finding [1].\n\nDetails:\n1. First [2]\n2. Second [3]",
    ],
)
def test_valid_citations_preserve_answer_layout_exactly(answer):
    assert strip_citations(answer, {1, 2, 3}) == answer


def test_invalid_citation_removal_does_not_reformat_other_lines():
    answer = "Claim [99] and evidence [1].\n\n```python\n    value =  42\n```\n"

    cleaned = strip_citations(answer, {1})

    assert cleaned == "Claim and evidence [1].\n\n```python\n    value =  42\n```\n"
    assert extract_citations(cleaned) == [1]


@pytest.mark.parametrize(
    "answer,expected",
    [
        ("a [9] b", "a b"),
        ("a[9] b", "a b"),
        ("[9] Claim", "Claim"),
        ("    [9] item", "    item"),
        ("a[9]b", "ab"),
    ],
)
def test_citation_removal_cleans_only_its_local_padding(answer, expected):
    assert strip_citations(answer, set()) == expected


def test_citation_cleanup_retains_unicode_numbers_and_original_spelling():
    answer = "Evidence [١] [００２] and invalid [9]."

    cleaned = strip_citations(answer, {1, 2})

    assert cleaned == "Evidence [١] [００２] and invalid ."
    assert extract_citations(cleaned) == [1, 2]


@pytest.mark.parametrize("index", [0, -1, True, 1.0, "1"])
def test_citation_verification_rejects_invalid_source_indices(index):
    source = Source(index=index, chunk_id="chunk", document_id="document", content="evidence")

    with pytest.raises(ValueError, match="unique positive integers"):
        AnswerVerifier().verify("Claim [1].", [source])


def test_citation_verification_rejects_ambiguous_source_indices():
    sources = [
        Source(index=1, chunk_id="first", document_id="first", content="alpha"),
        Source(index=1, chunk_id="second", document_id="second", content="beta"),
    ]

    with pytest.raises(ValueError, match="unique positive integers"):
        AnswerVerifier().verify("Claim [1].", sources)


def test_answer_verification_preserves_valid_noncontiguous_sources_and_layout():
    sources = [
        Source(index=2, chunk_id="second", document_id="second", content="alpha"),
        Source(index=4, chunk_id="fourth", document_id="fourth", content="beta"),
    ]
    answer = "  First [2].\n\n\n    Second [4].\n"

    result = AnswerVerifier().verify(answer, sources)

    assert result.verified_answer == answer
    assert result.citations_valid == [2, 4]
    assert result.citations_invalid == []
    assert result.warnings == []


@pytest.mark.parametrize("newline", ["\n", "\r\n"])
def test_invalid_citation_removal_preserves_markdown_hard_breaks(newline):
    answer = f"Claim [99]  {newline}Next claim [1]."

    cleaned = strip_citations(answer, {1})

    assert cleaned == f"Claim   {newline}Next claim [1]."
    assert extract_citations(cleaned) == [1]


def test_grounding_supports_legacy_cross_encoder_activation_parameter():
    logits = np.array([-2.0, 2.0])

    def predict(pairs, *, activation_fct, apply_softmax, convert_to_numpy):
        assert pairs == [("first", "claim"), ("second", "claim")]
        assert apply_softmax is False
        assert convert_to_numpy is True
        return activation_fct(logits)

    verifier = GroundingVerifier()
    verifier._model = SimpleNamespace(predict=predict)

    probabilities = verifier.class_probs([("first", "claim"), ("second", "claim")])

    expected = 1 / (1 + np.exp(-logits))
    np.testing.assert_allclose(probabilities, np.column_stack([1 - expected, expected]))


class TestHealthProbe:
    def test_health_returns_ok(self) -> None:
        app = create_app()
        client = TestClient(app, base_url="http://localhost")

        response = client.get("/health")

        assert response.status_code == 200
        assert response.json() == {"status": "ok"}


class TestPipelinePresentation:
    def test_pipeline_has_results_table_and_design_source(self) -> None:
        response = TestClient(create_app(), base_url="http://localhost").get("/pipeline")

        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/html")
        table = re.search(r"<table\b.*?</table>", response.text, re.DOTALL)
        assert table is not None
        assert len(re.findall(r"<th\b", table.group())) >= 3
        for value in ("SciFact", "NFCorpus", "0.708", "0.341", "0.704", "0.352", "0.703", "0.346"):
            assert value in table.group()
        assert "https://www.figma.com/design/bRUVlmkQKApYB2bWdHiblv?node-id=29-2" in response.text

    def test_main_page_links_to_pipeline(self) -> None:
        response = TestClient(create_app(), base_url="http://localhost").get("/")

        assert response.status_code == 200
        assert re.search(r'<a\b[^>]*href="/pipeline"[^>]*>', response.text)

    @pytest.mark.parametrize(
        ("path", "content_type"),
        [
            ("/css/pipeline.css", "text/css"),
            ("/images/pipeline.svg", "image/svg+xml"),
        ],
    )
    def test_pipeline_assets_are_served(self, path: str, content_type: str) -> None:
        response = TestClient(create_app(), base_url="http://localhost").get(path)

        assert response.status_code == 200
        assert response.headers["content-type"].startswith(content_type)
        assert response.content
        if content_type == "image/svg+xml":
            assert ET.fromstring(response.content).tag == "{http://www.w3.org/2000/svg}svg"

    @pytest.mark.parametrize(
        "path", ["/images/../settings.py", "/images/%2e%2e/%2e%2e/settings.py"]
    )
    def test_images_cannot_serve_python_source(self, path: str) -> None:
        assert TestClient(create_app(), base_url="http://localhost").get(path).status_code == 404


class TestLogLevel:
    def test_create_app_applies_log_level(self, monkeypatch) -> None:
        monkeypatch.setenv("LOG_LEVEL", "WARNING")
        logger = logging.getLogger("scinexusrag")
        original_level = logger.level
        get_settings.cache_clear()
        try:
            create_app()
            assert logger.level == logging.WARNING
        finally:
            get_settings.cache_clear()
            logger.setLevel(original_level)


def _write_frontend(directory: Path, label: str) -> Path:
    directory.mkdir()
    (directory / "index.html").write_text(f"<html>{label} index</html>", encoding="utf-8")
    (directory / "pipeline.html").write_text(f"<html>{label} pipeline</html>", encoding="utf-8")
    (directory / "css").mkdir()
    (directory / "css" / "style.css").write_text(
        f":root {{ --interface: {label}; }}", encoding="utf-8"
    )
    return directory


class TestFrontendConfiguration:
    def test_factory_reads_current_environment_and_keeps_apps_isolated(
        self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        first = _write_frontend(tmp_path / "first", "first")
        second = _write_frontend(tmp_path / "second", "second")
        monkeypatch.setenv("NEXUSRAG_FRONTEND_DIR", str(first))
        first_app = create_app()
        monkeypatch.setenv("NEXUSRAG_FRONTEND_DIR", str(second))
        second_app = create_app()

        for app, label in ((first_app, "first"), (second_app, "second")):
            with TestClient(app, base_url="http://localhost") as client:
                assert client.get("/").text == f"<html>{label} index</html>"
                assert client.get("/pipeline").text == f"<html>{label} pipeline</html>"
                assert client.get("/css/style.css").text == f":root {{ --interface: {label}; }}"

    @pytest.mark.parametrize("page,filename", [("/", "index.html"), ("/pipeline", "pipeline.html")])
    @pytest.mark.parametrize("remove_after_creation", [False, True])
    def test_missing_entry_page_returns_404(
        self,
        monkeypatch: pytest.MonkeyPatch,
        tmp_path: Path,
        page: str,
        filename: str,
        remove_after_creation: bool,
    ) -> None:
        frontend = _write_frontend(tmp_path / "frontend", "custom")
        monkeypatch.setenv("NEXUSRAG_FRONTEND_DIR", str(frontend))
        if not remove_after_creation:
            (frontend / filename).unlink()
        app = create_app()
        if remove_after_creation:
            (frontend / filename).unlink()

        with TestClient(app, base_url="http://localhost") as client:
            response = client.get(page)
            assert response.status_code == 404
            assert client.get("/health").json() == {"status": "ok"}

    @pytest.mark.parametrize("exists_as_file", [False, True])
    def test_non_directory_frontend_keeps_api_available(
        self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path, exists_as_file: bool
    ) -> None:
        frontend = tmp_path / "frontend"
        if exists_as_file:
            frontend.write_text("Not a directory", encoding="utf-8")
        monkeypatch.setenv("NEXUSRAG_FRONTEND_DIR", str(frontend))

        with TestClient(create_app(), base_url="http://localhost") as client:
            response = client.get("/health")
            assert response.status_code == 200
            assert response.json() == {"status": "ok"}
            assert client.get("/").status_code == 404
            assert client.get("/pipeline").status_code == 404

    def test_static_mount_ignores_regular_file(
        self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        frontend = tmp_path / "frontend"
        frontend.mkdir()
        (frontend / "index.html").write_text("<html>Custom interface</html>", encoding="utf-8")
        (frontend / "css").write_text("Not a directory", encoding="utf-8")
        monkeypatch.setenv("NEXUSRAG_FRONTEND_DIR", str(frontend))

        with TestClient(create_app(), base_url="http://localhost") as client:
            assert client.get("/").text == "<html>Custom interface</html>"
            assert client.get("/css/style.css").status_code == 404
            assert client.get("/health").status_code == 200

    def test_relative_frontend_stays_bound_after_working_directory_changes(
        self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        initial, later = tmp_path / "initial", tmp_path / "later"
        initial.mkdir()
        later.mkdir()
        _write_frontend(initial / "frontend", "initial")
        _write_frontend(later / "frontend", "later")
        monkeypatch.chdir(initial)
        monkeypatch.setenv("NEXUSRAG_FRONTEND_DIR", "frontend")
        app = create_app()
        monkeypatch.chdir(later)

        with TestClient(app, base_url="http://localhost") as client:
            assert client.get("/").text == "<html>initial index</html>"
            assert client.get("/pipeline").text == "<html>initial pipeline</html>"
            assert client.get("/css/style.css").text == ":root { --interface: initial; }"

    @pytest.mark.parametrize("page,filename", [("/", "index.html"), ("/pipeline", "pipeline.html")])
    def test_html_symlink_cannot_escape_frontend(
        self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path, page: str, filename: str
    ) -> None:
        frontend = _write_frontend(tmp_path / "frontend", "custom")
        secret = tmp_path / "private.html"
        secret.write_text("Private document contents", encoding="utf-8")
        (frontend / filename).unlink()
        (frontend / filename).symlink_to(secret)
        monkeypatch.setenv("NEXUSRAG_FRONTEND_DIR", str(frontend))

        with TestClient(create_app(), base_url="http://localhost") as client:
            response = client.get(page)
            assert response.status_code == 404
            assert "Private document contents" not in response.text


class TestFixedSizeChunker:
    def test_fixed_size_chunker_basic(self, parsed_document):
        chunker = FixedSizeChunker(chunk_size=100, chunk_overlap=20)
        chunks = chunker.chunk(parsed_document)

        assert len(chunks) > 0
        assert all(isinstance(c, Chunk) for c in chunks)

    def test_chunk_size_respected(self, parsed_document):
        chunk_size = 200
        chunker = FixedSizeChunker(chunk_size=chunk_size, chunk_overlap=20)
        chunks = chunker.chunk(parsed_document)

        assert chunks
        assert all(len(chunk.content) <= chunk_size for chunk in chunks)

    def test_chunk_overlap(self):
        content = " ".join(f"w{i}" for i in range(60))
        doc = ParsedDocument(id="test", content=content, metadata={}, sections=[])

        chunker = FixedSizeChunker(chunk_size=20, chunk_overlap=5, length_function="words")
        chunks = chunker.chunk(doc)

        assert len(chunks) >= 2

        assert chunks[0].content.split()[-5:] == chunks[1].content.split()[:5]

    def test_chunk_ids_unique(self, parsed_document):
        chunker = FixedSizeChunker(chunk_size=100, chunk_overlap=20)
        chunks = chunker.chunk(parsed_document)

        ids = [c.id for c in chunks]
        assert ids
        assert len(ids) == len(set(ids))

    def test_chunk_metadata_preserved(self, parsed_document):
        chunker = FixedSizeChunker(chunk_size=100, chunk_overlap=20)
        chunks = chunker.chunk(parsed_document)

        for i, chunk in enumerate(chunks):
            assert chunk.document_id == parsed_document.id
            assert chunk.metadata["chunk_index"] == i
            assert "filename" in chunk.metadata

    def test_word_based_chunking(self, parsed_document):
        chunker = FixedSizeChunker(chunk_size=20, chunk_overlap=5, length_function="words")
        chunks = chunker.chunk(parsed_document)

        assert len(chunks) > 0
        for chunk in chunks:
            word_count = len(chunk.content.split())
            assert 0 < word_count <= 20

    def test_overlap_validation(self):
        with pytest.raises(ValueError):
            FixedSizeChunker(chunk_size=100, chunk_overlap=100)

        with pytest.raises(ValueError):
            FixedSizeChunker(chunk_size=100, chunk_overlap=150)

    def test_empty_document(self):
        doc = ParsedDocument(id="empty", content="", metadata={}, sections=[])
        chunker = FixedSizeChunker(chunk_size=100, chunk_overlap=20)
        chunks = chunker.chunk(doc)

        assert chunks == []

    def test_whitespace_only_document(self):
        doc = ParsedDocument(id="ws", content="   \n\n   ", metadata={}, sections=[])
        chunker = FixedSizeChunker(chunk_size=100, chunk_overlap=20)
        chunks = chunker.chunk(doc)

        assert chunks == []


class TestSemanticChunker:
    @pytest.mark.parametrize(
        "overrides",
        [
            {"min_chunk_size": 0},
            {"target_chunk_size": 0},
            {"max_chunk_size": 0},
            {"min_chunk_size": -1},
            {"target_chunk_size": -1},
            {"max_chunk_size": -1},
            {"min_chunk_size": 101},
            {"max_chunk_size": 19},
            {"overlap_size": -1},
            {"overlap_size": 100},
            {"context_chars": -1},
        ],
    )
    def test_invalid_sizes_fail_before_chunking(self, overrides):
        kwargs = {
            "min_chunk_size": 20,
            "target_chunk_size": 100,
            "max_chunk_size": 150,
            "overlap_size": 0,
            **overrides,
        }
        with pytest.raises(ValueError):
            HierarchicalChunker(**kwargs)

    def test_zero_overlap_and_context_are_valid(self):
        chunker = HierarchicalChunker(overlap_size=0, context_chars=0)
        chunks = chunker.chunk(ParsedDocument(id="short", content="Short content."))

        assert len(chunks) == 1
        assert chunks[0].context_before == ""
        assert chunks[0].context_after == ""

    def test_semantic_chunker_basic(self, parsed_document):
        chunker = SemanticChunker(min_chunk_size=50, max_chunk_size=500)
        chunks = chunker.chunk(parsed_document)

        assert len(chunks) > 0
        assert all(isinstance(c, Chunk) for c in chunks)

    def test_respects_sections(self):
        sections = [
            Section(title="Introduction", content="Intro content here.", level=1),
            Section(title="Methods", content="Methods description.", level=1),
            Section(title="Results", content="Results summary.", level=1),
        ]
        doc = ParsedDocument(
            id="sectioned",
            content="Full document content",
            metadata={"filename": "test.txt"},
            sections=sections,
        )

        chunker = SemanticChunker(min_chunk_size=10, max_chunk_size=500)
        chunks = chunker.chunk(doc)

        assert len(chunks) >= len(sections)

        titles_found = sum(1 for c in chunks if any(s.title in c.content for s in sections))
        assert titles_found >= 2

    def test_section_metadata_preserved(self):
        sections = [
            Section(title="Methods", content="Detailed methods.", level=2, page_number=5),
        ]
        doc = ParsedDocument(id="test", content="Content", metadata={}, sections=sections)

        chunker = SemanticChunker(min_chunk_size=10, max_chunk_size=500)
        chunks = chunker.chunk(doc)

        methods_chunks = [c for c in chunks if c.metadata.get("section_title") == "Methods"]
        assert len(methods_chunks) >= 1
        assert methods_chunks[0].metadata.get("section_title") == "Methods"

    def test_large_section_split(self):
        large_content = "This is a sentence. " * 100
        sections = [Section(title="Large", content=large_content, level=1)]
        doc = ParsedDocument(id="large", content=large_content, metadata={}, sections=sections)

        chunker = SemanticChunker(min_chunk_size=50, max_chunk_size=200)
        chunks = chunker.chunk(doc)

        assert len(chunks) > 1

    def test_paragraph_based_chunking(self):
        content = """First paragraph with some content.

Second paragraph here.

Third paragraph follows."""

        doc = ParsedDocument(id="paras", content=content, metadata={}, sections=[])
        chunker = SemanticChunker(min_chunk_size=10, max_chunk_size=500)
        chunks = chunker.chunk(doc)

        assert len(chunks) >= 1

    def test_small_chunks_merged(self):
        content = "A.\n\nB.\n\nC.\n\nD.\n\nE."
        doc = ParsedDocument(id="small", content=content, metadata={}, sections=[])

        chunker = SemanticChunker(min_chunk_size=20, max_chunk_size=500)
        chunks = chunker.chunk(doc)

        assert [chunk.content for chunk in chunks] == [content]

    def test_empty_document(self):
        doc = ParsedDocument(id="empty", content="", metadata={}, sections=[])
        chunker = SemanticChunker()
        chunks = chunker.chunk(doc)

        assert chunks == []


class TestGetChunker:
    def test_get_fixed_chunker(self):
        chunker = get_chunker("fixed", chunk_size=200, chunk_overlap=30)
        assert isinstance(chunker, FixedSizeChunker)

    def test_get_semantic_chunker(self):
        chunker = get_chunker("semantic", max_chunk_size=500)
        assert isinstance(chunker, SemanticChunker)

    def test_default_is_semantic(self):
        chunker = get_chunker()
        assert isinstance(chunker, SemanticChunker)


class TestParagraphChunking:
    def _texts(self, paragraphs, section_title="", **overrides):
        params = {
            "min_chunk_size": 20,
            "target_chunk_size": 80,
            "max_chunk_size": 120,
            "overlap_size": 60,
            "include_context": False,
        }
        params.update(overrides)
        content = "\n\n".join(paragraphs)
        document = ParsedDocument(id="paragraphs", content=content)
        if section_title:
            document.sections = [Section(title=section_title, content=content)]
        return [chunk.content for chunk in HierarchicalChunker(**params).chunk(document)]

    def test_small_paragraphs_pack_into_one_chunk(self):
        assert self._texts(["one two", "three four"], min_chunk_size=5) == ["one two\n\nthree four"]

    def test_empty_paragraphs_skipped(self):
        assert self._texts(["", "  ", "real content here"], min_chunk_size=5) == [
            "real content here"
        ]

    def test_header_prefix_on_every_emitted_chunk(self):
        paras = ["a" * 50, "b" * 50, "c" * 50]
        texts = self._texts(paras, section_title="Methods", overlap_size=10)
        assert len(texts) > 1
        assert all(t.startswith("[Methods]\n\n") for t in texts)

    def test_target_overflow_starts_new_chunk_with_overlap(self):
        paras = ["a" * 50, "b" * 50, "c" * 50]
        texts = self._texts(paras)

        assert texts[0] == "a" * 50
        assert texts[1].startswith("a" * 50)
        assert "b" * 50 in texts[1]

    def test_oversized_paragraph_split_by_sentences(self):
        big = " ".join(f"Sentence number {i} is here." for i in range(20))
        assert len(big) > 120
        texts = self._texts([big])
        assert len(texts) > 1
        assert all(len(t) <= 120 for t in texts)
        assert all(len(t) >= 20 for t in texts)

    def test_boundaryless_oversized_paragraph_bounded_by_max(self):

        giant = "word " * 200
        texts = self._texts([giant.strip()])
        assert texts and all(len(t) <= 120 for t in texts)

    def test_short_tail_merges_into_same_batch_chunk(self):

        texts = self._texts(["a" * 100, "zz"])
        assert len(texts) == 1
        assert texts[0].startswith("a" * 100)
        assert texts[0].endswith("zz")

    def test_sole_short_content_still_emitted(self):

        assert self._texts(["tiny note"]) == ["tiny note"]

    def test_oversized_paragraph_flushes_pending_content_first(self):
        big = " ".join(f"Sentence number {i} is here." for i in range(20))
        texts = self._texts(["preamble text first", big])
        assert texts[0] == "preamble text first"

    def test_oversized_paragraph_short_trailing_sentence_not_dropped(self):

        big = "This is a long lead sentence that fills the budget nicely here. " * 3
        texts = self._texts(
            [big + "Short unique CONTENTX."], min_chunk_size=60, target_chunk_size=90
        )
        assert any("CONTENTX" in t for t in texts)


class TestChunkerNoDataLoss:
    def test_short_document_yields_a_chunk(self):
        doc = ParsedDocument(id="d", content="A short valid note.", metadata={})
        doc.sections = [Section(title="", content="A short valid note.", level=0)]
        chunks = SemanticChunker(include_context=False).chunk(doc)
        assert len(chunks) == 1
        assert "short valid note" in chunks[0].content

    def test_under_min_section_keeps_its_own_metadata(self):
        doc = ParsedDocument(id="d", content="x", metadata={})
        doc.sections = [
            Section(title="Alpha", content="alpha " * 300, level=0, page_number=1),
            Section(title="Beta", content="short beta note.", level=0, page_number=2),
        ]
        chunks = SemanticChunker(include_context=False).chunk(doc)
        beta = [c for c in chunks if c.section_title == "Beta"]
        assert beta and beta[0].page_number == 2

        assert not any(c.section_title == "Alpha" and "beta" in c.content.lower() for c in chunks)


@pytest.mark.parametrize("extension", [".txt", ".md", ".docx", ".pdf"])
def test_heading_only_documents_produce_chunks_with_exact_source_spans(extension, monkeypatch):
    headings = ["CALIBRATION PASSED.", "THE CURRENT IS 42 MA.", "CALIBRATION PASSED."]
    if extension == ".docx":
        source = DocxDocument()
        for heading in headings:
            source.add_heading(heading, level=1)
        buffer = BytesIO()
        source.save(buffer)
        data = buffer.getvalue()
        expected = headings
    elif extension == ".pdf":
        pages = [SimpleNamespace(extract_text=lambda text=text: text) for text in headings]
        reader = SimpleNamespace(is_encrypted=False, pages=pages)
        monkeypatch.setattr("scinexusrag.ingestion.parser.PdfReader", lambda path: reader)
        data = b"%PDF-1.4 fixture"
        expected = headings
    else:
        expected = [f"# {heading}" for heading in headings] if extension == ".md" else headings
        data = "\n\n".join(expected).encode()
    document = DocumentParser().parse_bytes(data, "readings" + extension, extension)

    chunks = SemanticChunker(context_chars=12).chunk(document)

    assert [chunk.content for chunk in chunks] == expected
    assert [chunk.section_title for chunk in chunks] == headings
    for chunk, section in zip(chunks, document.sections, strict=True):
        assert section.heading_only
        assert document.content[chunk.start_offset : chunk.end_offset] == chunk.content
        _assert_context_span(chunk, document.content, section.start_offset, section.end_offset, 12)
    assert chunks[0].start_offset < chunks[2].start_offset
    if extension == ".pdf":
        assert [chunk.page_number for chunk in chunks] == [1, 2, 3]


def test_body_and_empty_final_heading_are_chunked_once_with_their_own_metadata():
    source = "# Methods\nMeasured current: 42 mA.\n\n## CALIBRATION PASSED."
    document = DocumentParser().parse_bytes(source.encode(), "readings.md", ".md")

    chunks = SemanticChunker(context_chars=16).chunk(document)

    assert [chunk.content for chunk in chunks] == [
        "[Methods]\n\nMeasured current: 42 mA.",
        "## CALIBRATION PASSED.",
    ]
    assert [chunk.section_title for chunk in chunks] == ["Methods", "CALIBRATION PASSED."]
    for chunk, section in zip(chunks, document.sections, strict=True):
        _assert_context_span(chunk, document.content, section.start_offset, section.end_offset, 16)


@pytest.mark.parametrize("section_title", ["", "Legacy heading"])
def test_empty_sections_fall_back_to_original_text_with_exact_source_offsets(section_title):
    content = "  THE CURRENT IS 42 MA.\n\nCALIBRATION PASSED.  "
    document = ParsedDocument(
        id="empty-sections",
        content=content,
        sections=[Section(title=section_title, content=""), Section(title="", content="  ")],
    )

    chunks = SemanticChunker(context_chars=10).chunk(document)

    assert len(chunks) == 1
    chunk = chunks[0]
    assert chunk.content == content.strip()
    assert chunk.section_title == ""
    _assert_context_span(chunk, content, 2, len(content) - 2, 10)


@pytest.mark.parametrize("sectioned", [False, True])
@pytest.mark.parametrize("content", ["word " * 330, "long " * 299 + "\n\nTAILWORD"])
def test_short_tail_does_not_overflow_full_chunk(content, sectioned):
    prefix = "[Methods]\n\n" if sectioned else ""
    document = ParsedDocument(id="tail", content=content)
    if sectioned:
        document.sections = [Section(title="Methods", content=content, level=1, page_number=4)]
    chunker = SemanticChunker(include_context=False, overlap_size=0)

    chunks = chunker.chunk(document)

    assert len(chunks) >= 2
    assert all(len(chunk.content) <= chunker.max_chunk_size for chunk in chunks)
    recovered = " ".join(chunk.content.removeprefix(prefix) for chunk in chunks)
    assert recovered.split() == content.split()
    if sectioned:
        assert all(chunk.section_title == "Methods" and chunk.page_number == 4 for chunk in chunks)


@pytest.mark.parametrize("sectioned", [False, True])
def test_short_tail_merge_does_not_duplicate_overlap(sectioned):
    content = "\n\n".join(
        [("alpha " * 185).strip(), ("beta " * 10).strip(), ("gamma " * 10).strip()]
    )
    document = ParsedDocument(id="overlap-tail", content=content)
    if sectioned:
        document.sections = [Section(title="Methods", content=content, page_number=4)]
    chunker = SemanticChunker(min_chunk_size=200, include_context=False)

    chunks = chunker.chunk(document)

    assert len(chunks) == 1
    assert chunks[0].content.removeprefix("[Methods]\n\n") == content
    assert len(chunks[0].content) <= chunker.max_chunk_size
    if sectioned:
        assert chunks[0].section_title == "Methods"
        assert chunks[0].page_number == 4


def test_context_matches_full_passage_when_prefixes_repeat():
    prefix = "The shared introduction explains the study design and measurement procedure. "
    prefix += "Both groups followed this procedure. "
    first = prefix + "The first group used the original instrument."
    second = prefix + "The second group used the replacement instrument."
    content = f"First\n\n{first}\n\nSecond\n\n{second}\n\nEnd of the document."
    document = ParsedDocument(
        id="shared-prefix",
        content=content,
        sections=[Section(title="First", content=first), Section(title="Second", content=second)],
    )
    chunker = SemanticChunker(context_chars=20)

    chunks = chunker.chunk(document)

    second_start = content.index(second)
    second_end = second_start + len(second)
    assert chunks[1].context_before == content[second_start - 20 : second_start].strip()
    assert chunks[1].context_after == content[second_end : second_end + 20].strip()


def test_identical_passages_get_context_from_their_document_order():
    passage = "The instrument was calibrated before collecting the measurements."
    content = f"First marker.\n\n{passage}\n\nMiddle marker.\n\n{passage}\n\nLast marker."
    first_start = content.index(passage)
    second_start = content.rindex(passage)
    document = ParsedDocument(
        id="repeated-passage",
        content=content,
        sections=[
            Section(
                title="First",
                content=passage,
                start_offset=first_start,
                end_offset=first_start + len(passage),
            ),
            Section(
                title="Second",
                content=passage,
                start_offset=second_start,
                end_offset=second_start + len(passage),
            ),
        ],
    )

    chunks = SemanticChunker(context_chars=16).chunk(document)

    assert chunks[0].context_before == "First marker."
    assert chunks[1].context_before == "Middle marker."
    assert chunks[1].context_after == "Last marker."


def test_full_overlap_chunks_keep_context_from_the_same_source_position():
    first = ("alpha " * 40).strip()
    second = ("beta " * 210).strip()
    methods = first + "\n\n" + second
    results = "Following measurements were performed after calibration."
    content = f"Introduction marker.\n\n{methods}\n\nResults marker.\n\n{results}"
    document = ParsedDocument(
        id="full-overlap-context",
        content=content,
        sections=[
            Section(title="Methods", content=methods),
            Section(title="Results", content=results),
        ],
    )
    chunker = SemanticChunker(min_chunk_size=200, context_chars=40)

    chunks = chunker.chunk(document)

    assert len(chunks) == 3
    assert chunks[1].content == "[Methods]\n\n" + methods
    assert chunks[1].context_before == "Introduction marker."
    methods_end = content.index(methods) + len(methods)
    assert chunks[1].context_after == content[methods_end : methods_end + 40].strip()


def test_shorter_repeated_prefix_gets_context_from_its_later_occurrence():
    passage = "The instrument was calibrated before collecting the measurements."
    first = passage + " The first group used the original instrument."
    content = f"First marker.\n\n{first}\n\nMiddle marker.\n\n{passage}\n\nLast marker."
    second_start = content.rindex(passage)
    document = ParsedDocument(
        id="shorter-prefix",
        content=content,
        sections=[
            Section(title="First", content=first),
            Section(
                title="Second",
                content=passage,
                start_offset=second_start,
                end_offset=second_start + len(passage),
            ),
        ],
    )

    chunks = SemanticChunker(context_chars=16).chunk(document)

    assert chunks[1].context_before == "Middle marker."
    assert chunks[1].context_after == "Last marker."


def test_unmatched_passage_does_not_borrow_context_from_a_shared_prefix():
    prefix = "This shared passage prefix describes the methods used to collect the measurements. "
    document = ParsedDocument(
        id="unmatched-passage",
        content="Beginning.\n\n" + prefix + "The original result.\n\nEnding.",
        sections=[Section(title="Methods", content=prefix + "A different result.")],
    )

    chunk = SemanticChunker().chunk(document)[0]

    assert chunk.context_before == ""
    assert chunk.context_after == ""


def test_literal_bracketed_text_is_not_treated_as_an_added_section_header():
    content = "[Literal source label]\n\nA short document with a label in its source text."

    chunk = SemanticChunker().chunk(ParsedDocument(id="literal-label", content=content))[0]

    assert chunk.full_context == content
    assert chunk.context_before == ""
    assert chunk.context_after == ""


def _assert_context_span(chunk, content, start, end, context_chars):
    assert chunk.start_offset == start
    assert chunk.end_offset == end
    assert chunk.metadata["start_offset"] == start
    assert chunk.metadata["end_offset"] == end
    assert chunk.context_before == content[max(0, start - context_chars) : start].strip()
    assert chunk.context_after == content[end : end + context_chars].strip()


def test_repeated_wrapped_words_keep_their_original_whitespace_positions():
    body = "alpha\tbeta   alpha\nbeta alpha\tbeta"
    content = "Before marker.\n\n" + body + "\n\nAfter marker."
    start = len("Before marker.\n\n")
    document = ParsedDocument(
        id="wrapped-positions",
        content=content,
        sections=[
            Section(
                title="",
                content=body,
                start_offset=start,
                end_offset=start + len(body),
            )
        ],
    )
    chunker = SemanticChunker(
        min_chunk_size=1,
        target_chunk_size=12,
        max_chunk_size=12,
        overlap_size=0,
        context_chars=16,
    )

    chunks = chunker.chunk(document)

    assert [chunk.content for chunk in chunks] == ["alpha beta"] * 3
    for chunk, relative_start in zip(chunks, [0, 13, 24], strict=True):
        _assert_context_span(
            chunk, content, start + relative_start, start + relative_start + 10, 16
        )


def test_sentence_packing_keeps_source_offsets_when_separator_whitespace_changes():
    first = "First sentence."
    second = "Second sentence."
    third = "Third sentence."
    content = first + "\t  " + second + "\n " + third
    chunker = SemanticChunker(
        min_chunk_size=1,
        target_chunk_size=35,
        max_chunk_size=35,
        overlap_size=0,
        context_chars=18,
    )

    chunks = chunker.chunk(ParsedDocument(id="sentence-positions", content=content))

    assert [chunk.content for chunk in chunks] == [first + " " + second, third]
    second_end = len(first) + 3 + len(second)
    _assert_context_span(chunks[0], content, 0, second_end, 18)
    _assert_context_span(chunks[1], content, second_end + 2, len(content), 18)


def test_paragraph_overlap_keeps_source_offsets_with_headers_and_irregular_blank_lines():
    first = "alpha beta gamma."
    second = "delta epsilon zeta."
    third = "eta theta iota."
    body = first + "\n \n" + second + "\n\t\n" + third
    content = "Introduction.\n\n" + body + "\n\nConclusion."
    start = len("Introduction.\n\n")
    document = ParsedDocument(
        id="overlap-positions",
        content=content,
        sections=[
            Section(
                title="Methods",
                content=body,
                start_offset=start,
                end_offset=start + len(body),
            )
        ],
    )
    chunker = SemanticChunker(
        min_chunk_size=1,
        target_chunk_size=35,
        max_chunk_size=70,
        overlap_size=20,
        context_chars=16,
    )

    chunks = chunker.chunk(document)

    prefix = "[Methods]\n\n"
    assert [chunk.content for chunk in chunks] == [
        prefix + first,
        prefix + first + "\n\n" + second,
        prefix + second + "\n\n" + third,
    ]
    second_start = start + len(first) + 3
    _assert_context_span(chunks[0], content, start, start + len(first), 16)
    _assert_context_span(chunks[1], content, start, second_start + len(second), 16)
    _assert_context_span(chunks[2], content, second_start, start + len(body), 16)


def test_explicit_section_offsets_distinguish_identical_passages():
    passage = "The instrument was calibrated."
    content = passage + "\n\nMiddle marker.\n\n" + passage + "\n\nEnd marker."
    start = len(passage + "\n\nMiddle marker.\n\n")
    document = ParsedDocument(
        id="explicit-section-position",
        content=content,
        sections=[
            Section(
                title="Methods",
                content=passage,
                start_offset=start,
                end_offset=start + len(passage),
            )
        ],
    )

    chunk = SemanticChunker(context_chars=16).chunk(document)[0]

    assert chunk.content == "[Methods]\n\n" + passage
    _assert_context_span(chunk, content, start, start + len(passage), 16)


@pytest.mark.parametrize(
    "extension,headers", [(".txt", ["FIRST", "SECOND"]), (".md", ["# First", "# Second"])]
)
def test_parsed_repeated_sections_keep_context_after_whitespace_cleaning(extension, headers):
    body = "The instrument\twas calibrated.\n\n  Measurements were recorded."
    source = headers[0] + "\n" + body + "\n\n" + headers[1] + "\n" + body
    document = DocumentParser().parse_bytes(source.encode(), "study" + extension, extension)
    chunker = SemanticChunker(context_chars=20)

    chunks = chunker.chunk(document)

    assert len(chunks) == 2
    for chunk, section in zip(chunks, document.sections, strict=True):
        _assert_context_span(chunk, document.content, section.start_offset, section.end_offset, 20)
    assert chunks[0].start_offset < chunks[1].start_offset


def test_invalid_section_offsets_do_not_borrow_context_from_matching_text_elsewhere():
    passage = "The instrument was calibrated."
    content = "Before marker.\n\n" + passage + "\n\nAfter marker."
    document = ParsedDocument(
        id="invalid-section-position",
        content=content,
        sections=[
            Section(title="Methods", content=passage, start_offset=0, end_offset=len(passage))
        ],
    )

    chunk = SemanticChunker().chunk(document)[0]

    assert chunk.start_offset is None
    assert chunk.end_offset is None
    assert chunk.context_before == ""
    assert chunk.context_after == ""


def test_legacy_section_with_repeated_body_has_no_guessed_context():
    passage = "The instrument was calibrated."
    content = "First marker.\n\n" + passage + "\n\nSecond marker.\n\n" + passage
    document = ParsedDocument(
        id="ambiguous-legacy-section",
        content=content,
        sections=[Section(title="Methods", content=passage)],
    )

    chunk = SemanticChunker().chunk(document)[0]

    assert chunk.start_offset is None
    assert chunk.end_offset is None
    assert chunk.context_before == ""
    assert chunk.context_after == ""


def test_short_tail_merge_extends_source_end_without_repeating_overlap():
    first = ("alpha " * 185).strip()
    second = ("beta " * 10).strip()
    third = ("gamma " * 10).strip()
    content = " \t" + first + "\n \n" + second + "\n\t\n" + third + " \n"
    chunker = SemanticChunker(min_chunk_size=200, context_chars=20)

    chunks = chunker.chunk(ParsedDocument(id="merged-tail-position", content=content))

    assert len(chunks) == 1
    assert chunks[0].content == first + "\n\n" + second + "\n\n" + third
    _assert_context_span(chunks[0], content, 2, len(content) - 2, 20)


@pytest.mark.parametrize("length_function", ["chars", "words"])
def test_negative_overlap_fails_before_it_can_skip_content(length_function):
    with pytest.raises(ValueError, match="nonnegative"):
        FixedSizeChunker(chunk_size=2, chunk_overlap=-1, length_function=length_function)


@pytest.mark.parametrize("chunk_size", [0, -1])
def test_nonpositive_chunk_size_rejected(chunk_size):
    with pytest.raises(ValueError, match="positive"):
        FixedSizeChunker(chunk_size=chunk_size, chunk_overlap=-2)


@pytest.mark.parametrize(
    "content",
    [
        " ".join(f"word{i}" for i in range(60)),
        "A sentence fills this chunk. Another sentence also has enough words to overflow.",
    ],
)
def test_fixed_character_budget_preserves_every_word_without_overlap(content):
    chunker = FixedSizeChunker(chunk_size=30, chunk_overlap=0)
    chunks = chunker.chunk(ParsedDocument(id="fixed", content=content))

    assert len(chunks) > 1
    assert all(len(chunk.content) <= 30 for chunk in chunks)
    assert " ".join(chunk.content for chunk in chunks).split() == content.split()


def test_fixed_character_overlap_keeps_whole_words_and_respects_the_budget():
    content = "alpha beta gamma delta epsilon zeta"
    chunks = FixedSizeChunker(chunk_size=22, chunk_overlap=10).chunk(
        ParsedDocument(id="overlap", content=content)
    )

    assert [chunk.content for chunk in chunks] == [
        "alpha beta gamma delta",
        "delta epsilon zeta",
    ]


def test_fixed_character_overlap_shrinks_to_fit_the_next_word():
    chunks = FixedSizeChunker(chunk_size=10, chunk_overlap=8).chunk(
        ParsedDocument(id="budget", content="one two abcdefghi end")
    )

    assert [chunk.content for chunk in chunks] == ["one two", "abcdefghi", "end"]


def test_fixed_character_chunker_rejects_an_indivisible_oversized_word():
    with pytest.raises(ValueError, match="word length exceeds available chunk size"):
        FixedSizeChunker(chunk_size=10, chunk_overlap=0).chunk(
            ParsedDocument(id="word", content="valid " + "x" * 11)
        )


def test_fixed_chunker_rejects_an_unknown_length_unit():
    with pytest.raises(ValueError, match="length_function"):
        FixedSizeChunker(length_function="tokens")


@pytest.mark.parametrize("word_count,expected_chunks", [(20, 1), (35, 2), (36, 3)])
def test_fixed_word_chunks_do_not_emit_a_tail_made_only_of_overlap(word_count, expected_chunks):
    content = " ".join(f"word{i}" for i in range(word_count))
    chunks = FixedSizeChunker(chunk_size=20, chunk_overlap=5, length_function="words").chunk(
        ParsedDocument(id="word-tail", content=content)
    )

    assert len(chunks) == expected_chunks
    assert chunks[-1].content.split()[-1] == f"word{word_count - 1}"


@pytest.mark.parametrize(
    "option",
    ["min_chunk_size", "target_chunk_size", "max_chunk_size", "overlap_size", "context_chars"],
)
@pytest.mark.parametrize("value", [True, 1.5, float("nan"), float("inf"), "100"])
def test_hierarchical_sizes_reject_nonintegers_at_construction(option, value):
    with pytest.raises(ValueError, match=rf"{option} must be an integer"):
        HierarchicalChunker(**{option: value})


@pytest.mark.parametrize("option", ["chunk_size", "chunk_overlap"])
@pytest.mark.parametrize("value", [False, 1.5, float("nan"), float("inf"), "100"])
def test_fixed_sizes_reject_nonintegers_before_splitting(option, value):
    with pytest.raises(ValueError, match=rf"{option} must be an integer"):
        FixedSizeChunker(**{option: value})


@pytest.mark.parametrize("value", [0, 1, "false", None])
def test_context_switch_requires_a_boolean(value):
    with pytest.raises(ValueError, match="include_context must be a boolean"):
        HierarchicalChunker(include_context=value)


@pytest.mark.parametrize("strategy", ["", "unknown", "Fixed", None, []])
def test_unknown_chunking_strategies_are_rejected(strategy):
    with pytest.raises(ValueError, match="strategy must be"):
        get_chunker(strategy)


def test_invalid_length_unit_reports_configuration_error():
    with pytest.raises(ValueError, match="length_function"):
        FixedSizeChunker(length_function=[])


@pytest.mark.parametrize("sectioned", [False, True])
@pytest.mark.parametrize("paragraphs", [["alpha", "beta"], ["alpha", "beta", "gamma"]])
def test_paragraphs_that_exactly_fill_the_target_stay_together(sectioned, paragraphs):
    body = "\n\n".join(paragraphs)
    prefix = "[Methods]\n\n" if sectioned else ""
    document = ParsedDocument(id="exact-target", content=body)
    if sectioned:
        document.sections = [Section(title="Methods", content=body)]
    size = len(prefix) + len(body)
    chunker = HierarchicalChunker(
        min_chunk_size=1,
        target_chunk_size=size,
        max_chunk_size=size + 10,
        overlap_size=0,
    )

    chunks = chunker.chunk(document)

    assert [chunk.content for chunk in chunks] == [prefix + body]
    assert chunks[0].start_offset == 0
    assert chunks[0].end_offset == len(body)


@pytest.mark.parametrize(
    "options", [{"query": ""}, {"query": " \t\n"}, {"top_k": 0}, {"top_k": -1}]
)
def test_empty_corrective_requests_do_not_call_base(options):
    base = Mock()
    retriever = CorrectiveRetriever(base)

    assert retriever.retrieve_traced(**{"query": "query", **options}) == ([], False)
    base.retrieve_with_dense_top.assert_not_called()
    base.retrieve.assert_not_called()


@pytest.mark.parametrize(
    "options", [{"query": None}, {"top_k": False}, {"top_k": 1.5}, {"depth": 0}, {"depth": True}]
)
def test_invalid_corrective_requests_do_not_call_base(options):
    base = Mock()
    retriever = CorrectiveRetriever(base)

    with pytest.raises(ValueError):
        retriever.retrieve_traced(**{"query": "query", **options})

    base.retrieve_with_dense_top.assert_not_called()
    base.retrieve.assert_not_called()


@pytest.mark.parametrize(
    "tau", [-0.1, 1.1, float("nan"), float("inf"), True, "0.5", 10**1000, -(10**1000)]
)
def test_invalid_corrective_thresholds_fail_during_construction(tau):
    with pytest.raises(ValueError, match="threshold"):
        CorrectiveRetriever(Mock(), tau=tau)


@pytest.mark.parametrize("option", ["feedback_docs", "feedback_terms"])
@pytest.mark.parametrize("value", [0, -1, True, 1.5, float("nan")])
def test_feedback_sizes_require_positive_integers(option, value):
    with pytest.raises(ValueError, match=option):
        CorrectiveRetriever(Mock(), **{option: value})


@pytest.mark.parametrize("enabled", [0, 1, "false", None])
def test_corrective_switch_requires_a_boolean(enabled):
    with pytest.raises(ValueError, match="enabled"):
        CorrectiveRetriever(Mock(), enabled=enabled)


@pytest.mark.parametrize(
    "query,content",
    [
        ("kinase inhibition", "kinase inhibition kinase"),
        (" kinase inhibition ", "kinase inhibition"),
        ("kinase inhibition", "the and with"),
    ],
)
def test_feedback_without_new_terms_does_not_repeat_the_query(query, content):
    base = Mock()
    first = [RetrievalResult(Chunk(id="first", content=content, document_id="paper"), 0.7)]
    base.retrieve_with_dense_top.return_value = first, 0.1
    base.sparse = BM25Retriever()

    results, triggered = CorrectiveRetriever(base).retrieve_traced(query)

    assert results == first
    assert results[0] is first[0]
    assert triggered is False
    base.retrieve.assert_not_called()


_ORIGIN = "http://localhost:3000"


@pytest.fixture
def cors_settings(monkeypatch, local_settings):
    cors_settings = local_settings
    cors_settings.api.cors_origins = [_ORIGIN]
    cors_settings.api.max_upload_mb = 1

    def unexpected_pipeline():
        pytest.fail("Rejected requests and CORS preflights must not load the pipeline.")

    monkeypatch.setattr(routes, "get_scinexusrag", unexpected_pipeline)
    return cors_settings


@pytest.mark.parametrize("content_length,status", [("+1", 400), ("65537", 413)])
def test_body_rejection_keeps_allowed_origin_headers(cors_settings, content_length, status):
    with TestClient(api.create_app(), base_url="http://localhost") as client:
        response = client.post(
            "/api/query",
            content=b"x" if status == 400 else b"x" * 65537,
            headers={"Origin": _ORIGIN, "Content-Length": content_length},
        )

    assert response.status_code == status
    assert response.json()["detail"] == (
        "Invalid Content-Length" if status == 400 else "Request body too large"
    )
    assert response.headers["access-control-allow-origin"] == _ORIGIN
    assert response.headers["access-control-allow-credentials"] == "true"
    assert "origin" in response.headers["vary"].lower()


async def test_streamed_body_rejection_keeps_cors_headers(cors_settings, asgi_request):
    chunks = [b'{"question":"', b"x" * 32768, b"x" * 32768, b'"}']
    messages, consumed = await asgi_request(
        api.create_app(),
        [
            (b"host", b"localhost"),
            (b"origin", _ORIGIN.encode()),
            (b"content-type", b"application/json"),
        ],
        chunks,
    )

    starts = [message for message in messages if message["type"] == "http.response.start"]
    assert len(starts) == 1
    assert starts[0]["status"] == 413
    assert Headers(raw=starts[0]["headers"])["access-control-allow-origin"] == _ORIGIN
    assert consumed < len(chunks)


def test_allowed_preflight_reaches_cors_without_pipeline_access(cors_settings):
    with TestClient(api.create_app(), base_url="http://localhost") as client:
        response = client.options(
            "/api/query",
            headers={
                "Origin": _ORIGIN,
                "Access-Control-Request-Method": "POST",
                "Access-Control-Request-Headers": "Content-Type, X-API-Key",
            },
        )

    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == _ORIGIN
    assert "POST" in response.headers["access-control-allow-methods"]
    allowed_headers = response.headers["access-control-allow-headers"].lower()
    assert "content-type" in allowed_headers
    assert "x-api-key" in allowed_headers


@pytest.mark.parametrize(
    "hosts",
    [
        [],
        [b"rebind.example"],
        [b"localhost:0"],
        [b"localhost", b"rebind.example"],
        [b"localhost"] * 2,
    ],
)
async def test_invalid_preflight_hosts_are_rejected_before_cors(cors_settings, hosts, asgi_request):
    headers = [(b"host", host) for host in hosts]
    headers.extend([(b"origin", _ORIGIN.encode()), (b"access-control-request-method", b"POST")])

    messages, consumed = await asgi_request(api.create_app(), headers, method="OPTIONS")

    starts = [message for message in messages if message["type"] == "http.response.start"]
    assert len(starts) == 1
    assert starts[0]["status"] == 400
    assert "access-control-allow-origin" not in Headers(raw=starts[0]["headers"])
    assert consumed == 0


@pytest.mark.parametrize("preflight", [False, True])
def test_wildcard_origins_do_not_enable_credentials(cors_settings, preflight):
    cors_settings.api.cors_origins = ["*"]
    with TestClient(api.create_app(), base_url="http://localhost") as client:
        if preflight:
            response = client.options(
                "/api/query",
                headers={
                    "Origin": "https://client.example",
                    "Access-Control-Request-Method": "POST",
                },
            )
        else:
            response = client.get("/health", headers={"Origin": "https://client.example"})

    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "*"
    assert "access-control-allow-credentials" not in response.headers


def test_body_errors_do_not_allow_unconfigured_origins(cors_settings):
    with TestClient(api.create_app(), base_url="http://localhost") as client:
        response = client.post(
            "/api/query", content=b"x" * 65537, headers={"Origin": "https://foreign.example"}
        )

    assert response.status_code == 413
    assert "access-control-allow-origin" not in response.headers


@pytest.mark.parametrize("host,key", [("0.0.0.0", ""), ("127.0.0.1", "test-only-key")])
def test_explicit_network_configuration_keeps_public_hosts(cors_settings, host, key):
    cors_settings.api.host = host
    cors_settings.api.api_key = key
    with TestClient(api.create_app(), base_url="https://research.example") as client:
        health = client.get("/health", headers={"Origin": _ORIGIN})
        preflight = client.options(
            "/api/query",
            headers={"Origin": _ORIGIN, "Access-Control-Request-Method": "POST"},
        )

    for response in (health, preflight):
        assert response.status_code == 200
        assert response.headers["access-control-allow-origin"] == _ORIGIN


@pytest.fixture
def retriever():
    embedder = MagicMock()
    embedder.embed_query.return_value = np.array([1, 0], dtype=np.float32)
    store = MagicMock()
    store.search.return_value = []
    store.search_by_document.return_value = []
    return DenseRetriever(embedder, store)


@pytest.mark.parametrize("document_id", ["", " \t", False, 0, []])
def test_invalid_document_filter_never_falls_back_to_unfiltered_search(retriever, document_id):
    with pytest.raises(ValueError, match="document_id"):
        retriever.retrieve("evidence", document_id=document_id)

    retriever.embedder.embed_query.assert_not_called()
    retriever.vector_store.search.assert_not_called()
    retriever.vector_store.search_by_document.assert_not_called()


def test_document_filter_uses_only_the_requested_document(retriever):
    chunk = Chunk("passage", "Voltage is 3.3 V.", "paper")
    retriever.vector_store.search_by_document.return_value = [SearchResult(chunk, 0.75)]

    results = retriever.retrieve("voltage", top_k=2, document_id="paper")

    assert [(result.chunk, result.score, result.source) for result in results] == [
        (chunk, 0.75, "dense")
    ]
    retriever.vector_store.search.assert_not_called()
    retriever.vector_store.search_by_document.assert_called_once_with(
        retriever.embedder.embed_query.return_value, "paper", 2
    )


@pytest.mark.parametrize("top_k", [True, 1.5, "2", None, float("nan")])
def test_invalid_result_limits_fail_before_model_or_storage_access(retriever, top_k):
    with pytest.raises(ValueError, match="top_k"):
        retriever.retrieve("evidence", top_k=top_k)

    retriever.embedder.embed_query.assert_not_called()
    retriever.vector_store.validate_embedding_config.assert_not_called()


@pytest.mark.parametrize("query", [None, 1, ["evidence"]])
def test_invalid_queries_fail_before_model_or_storage_access(retriever, query):
    with pytest.raises(ValueError, match="query"):
        retriever.retrieve(query)

    retriever.embedder.embed_query.assert_not_called()
    retriever.vector_store.validate_embedding_config.assert_not_called()


@pytest.mark.parametrize("query,top_k", [("evidence", 0), ("evidence", -1), ("", 5), (" \n", 5)])
def test_empty_requests_skip_model_and_storage_access(retriever, query, top_k):
    assert retriever.retrieve(query, top_k=top_k) == []

    retriever.embedder.embed_query.assert_not_called()
    retriever.vector_store.validate_embedding_config.assert_not_called()
    retriever.vector_store.search.assert_not_called()


def test_nonempty_query_and_result_order_are_preserved(retriever):
    chunks = [Chunk(str(i), f"Evidence {i}.", "paper") for i in range(2)]
    retriever.vector_store.search.return_value = [
        SearchResult(chunks[0], 0.8),
        SearchResult(chunks[1], 0.5),
    ]
    query = " evidence \n"

    results = retriever.retrieve(query, top_k=2)

    retriever.embedder.embed_query.assert_called_once_with(query)
    assert [(result.chunk, result.score, result.source) for result in results] == [
        (chunks[0], 0.8, "dense"),
        (chunks[1], 0.5, "dense"),
    ]
    retriever.vector_store.search_by_document.assert_not_called()


@pytest.mark.parametrize("error_type", [EmbeddingConfigError, VectorStoreError])
def test_index_validation_errors_propagate_without_encoding(retriever, error_type):
    failure = error_type("index unavailable")
    retriever.vector_store.validate_embedding_config.side_effect = failure

    with pytest.raises(error_type) as caught:
        retriever.retrieve("evidence")

    assert caught.value is failure
    retriever.embedder.embed_query.assert_not_called()
    retriever.vector_store.search.assert_not_called()


def test_search_failure_propagates(retriever):
    failure = VectorStoreError("storage unavailable")
    retriever.vector_store.search.side_effect = failure

    with pytest.raises(VectorStoreError) as caught:
        retriever.retrieve("evidence")

    assert caught.value is failure


@pytest.mark.parametrize(
    "min_score", [float("nan"), float("inf"), -float("inf"), True, "0.5", None]
)
def test_invalid_thresholds_fail_before_model_or_storage_access(retriever, min_score):
    with pytest.raises(ValueError, match="min_score"):
        retriever.retrieve_with_threshold("evidence", min_score=min_score)

    retriever.embedder.embed_query.assert_not_called()
    retriever.vector_store.validate_embedding_config.assert_not_called()


@pytest.mark.parametrize(
    "min_score,expected",
    [(-1, [0.8, 0.5]), (0.5, [0.8, 0.5]), (0.8, [0.8]), (2, []), (10**1000, [])],
)
def test_thresholds_filter_scores_without_changing_order(retriever, min_score, expected):
    chunk = Chunk("passage", "Evidence.", "paper")
    retriever.vector_store.search.return_value = [
        SearchResult(chunk, 0.8),
        SearchResult(chunk, 0.5),
    ]

    results = retriever.retrieve_with_threshold("evidence", min_score=min_score)

    assert [result.score for result in results] == expected


@pytest.fixture
def inventory_rag(tmp_path):
    instance = NexusRAG(
        Settings(
            data_dir=tmp_path,
            storage=StorageSettings(lancedb_path=tmp_path / "lancedb"),
        )
    )
    instance._vector_store = MagicMock(spec=VectorStore)
    instance._vector_store.count.return_value = 1
    instance.document_store.add(
        ParsedDocument("original", "Sensor A reads temperature.", {"filename": "original.txt"})
    )
    instance.document_store.update_metadata("original", "chunk_count", 1)
    return instance


def test_inventory_uses_vector_count_when_metadata_has_no_chunk_count(inventory_rag):
    inventory_rag.document_store.add(
        ParsedDocument("recovered", "Sensor B reads pressure.", {"filename": "recovered.txt"})
    )
    inventory_rag._vector_store.count.return_value = 7

    inventory = inventory_rag.get_document_inventory()

    assert inventory.total_documents == 2
    assert inventory.total_chunks == 7
    assert "chunk_count" not in inventory.documents[1]
    assert inventory_rag._llm is None
    assert inventory_rag._embedder is None


def test_inventory_metadata_is_detached_from_later_store_updates(inventory_rag):
    inventory = inventory_rag.get_document_inventory()
    inventory_rag.document_store.update_metadata("original", "chunk_count", 4)

    assert inventory.documents[0]["chunk_count"] == 1
    assert inventory_rag.get_document_inventory().documents[0]["chunk_count"] == 4
    with pytest.raises(FrozenInstanceError):
        inventory.total_chunks = 4


def test_inventory_uses_stored_identity_when_metadata_contains_an_id(inventory_rag):
    inventory_rag.document_store.update_metadata("original", "id", "different")

    inventory = inventory_rag.get_document_inventory()

    assert inventory.documents[0]["id"] == "original"
    assert inventory_rag.document_store.exists("original")
    assert not inventory_rag.document_store.exists("different")


@pytest.mark.parametrize("method", ["list_documents", "get_document_inventory"])
def test_document_reads_wait_for_complete_writes(inventory_rag, method, observed_write_lock):
    lock = observed_write_lock
    inventory_rag._write_lock = lock
    inserted, finish_write = Event(), Event()

    def write_document():
        with inventory_rag._write_lock:
            inventory_rag.document_store.add(
                ParsedDocument("pending", "Sensor C reads humidity.", {"filename": "pending.txt"})
            )
            inserted.set()
            assert finish_write.wait(timeout=5)
            inventory_rag.document_store.update_metadata("pending", "chunk_count", 4)
            inventory_rag._vector_store.count.return_value = 5

    with ThreadPoolExecutor(max_workers=2) as workers:
        write = workers.submit(write_document)
        try:
            assert inserted.wait(timeout=5)
            read = workers.submit(getattr(inventory_rag, method))
            lock.wait_for_threads(2)
            assert not read.done()
        finally:
            finish_write.set()
        write.result(timeout=5)
        snapshot = read.result(timeout=5)

    documents = snapshot.documents if isinstance(snapshot, DocumentInventory) else snapshot
    assert [document["id"] for document in documents] == ["original", "pending"]
    assert documents[1]["chunk_count"] == 4
    if isinstance(snapshot, DocumentInventory):
        assert snapshot.total_documents == 2
        assert snapshot.total_chunks == 5


def test_inventory_holds_write_lock_until_vector_count_finishes(inventory_rag, observed_write_lock):
    lock = observed_write_lock
    inventory_rag._write_lock = lock
    counting, finish_count = Event(), Event()

    def count_vectors():
        counting.set()
        assert finish_count.wait(timeout=5)
        return 1

    inventory_rag._vector_store.count.side_effect = count_vectors

    def delete_document():
        with inventory_rag._write_lock:
            inventory_rag.document_store.delete("original")

    with ThreadPoolExecutor(max_workers=2) as workers:
        read = workers.submit(inventory_rag.get_document_inventory)
        try:
            assert counting.wait(timeout=5)
            write = workers.submit(delete_document)
            lock.wait_for_threads(2)
            assert not write.done()
        finally:
            finish_count.set()
        inventory = read.result(timeout=5)
        write.result(timeout=5)

    assert inventory.total_documents == inventory.total_chunks == 1
    assert inventory_rag.list_documents() == []


@pytest.mark.parametrize(
    "error",
    [VectorStoreError("Storage unavailable"), EmbeddingConfigError("Incompatible embeddings")],
)
def test_inventory_preserves_index_errors(inventory_rag, error):
    inventory_rag._vector_store.count.side_effect = error

    with pytest.raises(type(error), match=str(error)):
        inventory_rag.get_document_inventory()

    assert inventory_rag._llm is None


def test_inventory_preserves_metadata_read_errors(inventory_rag, monkeypatch):
    failure = MagicMock(side_effect=OSError("Metadata unavailable"))
    monkeypatch.setattr(inventory_rag.document_store, "list_with_metadata", failure)

    with pytest.raises(OSError, match="Metadata unavailable"):
        inventory_rag.get_document_inventory()

    inventory_rag._vector_store.count.assert_not_called()


def test_document_page_uses_inventory_totals_without_separate_reads(monkeypatch):
    pipeline = MagicMock(spec=NexusRAG)
    pipeline.get_document_inventory.return_value = DocumentInventory(
        documents=[
            {"id": "first", "filename": "first.txt", "word_count": 4, "chunk_count": 2},
            {"id": "recovered", "filename": "recovered.txt", "word_count": 3},
            {"id": "last", "filename": "last.txt", "word_count": 6, "chunk_count": 3},
        ],
        total_chunks=7,
    )
    pipeline.get_stats.side_effect = AssertionError("Unexpected statistics read")
    pipeline.list_documents.side_effect = AssertionError("Unexpected document read")
    monkeypatch.setattr("scinexusrag.api.routes.get_scinexusrag", lambda: pipeline)
    app = FastAPI()
    app.include_router(router)

    response = TestClient(app, base_url="http://localhost").get("/api/documents?limit=1&offset=1")

    assert response.status_code == 200
    page = response.json()
    assert [document["id"] for document in page["documents"]] == ["recovered"]
    assert page["total_documents"] == 3
    assert page["total_chunks"] == 7
    pipeline.get_document_inventory.assert_called_once_with()
    pipeline.get_stats.assert_not_called()
    pipeline.list_documents.assert_not_called()


@pytest.fixture
def document_store(tmp_path):
    return DocumentStore(tmp_path)


def _doc(doc_id: str = "abc123def456", name: str = "paper.pdf") -> ParsedDocument:
    return ParsedDocument(
        id=doc_id,
        content="Some parsed content with several words.",
        metadata={"original_filename": name, "file_type": "pdf"},
    )


def test_orphaned_doc_file_recovered_on_load(tmp_path, document_store) -> None:

    document_store.add(_doc())
    (tmp_path / "_index.json").unlink()

    reloaded = DocumentStore(path=tmp_path)

    assert reloaded.exists("abc123def456")
    assert reloaded.index["abc123def456"]["filename"] == "paper.pdf"
    assert reloaded.index["abc123def456"]["word_count"] > 0
    assert (tmp_path / "_index.json").exists()


def test_dangling_index_entry_dropped_on_load(tmp_path, document_store) -> None:

    document_store.add(_doc())
    (tmp_path / "abc123def456.json").unlink()

    reloaded = DocumentStore(path=tmp_path)

    assert not reloaded.exists("abc123def456")
    assert reloaded.count() == 0


def test_delete_survives_reload(tmp_path, document_store) -> None:
    document_store.add(_doc())
    assert document_store.delete("abc123def456") is True

    reloaded = DocumentStore(path=tmp_path)

    assert not reloaded.exists("abc123def456")
    assert reloaded.count() == 0


def test_reserved_index_id_is_rejected(tmp_path, document_store) -> None:

    with pytest.raises(ValueError, match="reserved"):
        document_store.add(_doc(doc_id="_index"))


def test_trailing_newline_id_is_rejected_before_writing(tmp_path, document_store) -> None:

    with pytest.raises(ValueError, match="invalid characters"):
        document_store.add(_doc(doc_id="paper\n"))

    assert list(tmp_path.iterdir()) == []


def test_interrupted_index_update_cannot_restore_a_deleted_document(
    tmp_path, monkeypatch, document_store
) -> None:
    document_store.add(_doc())
    monkeypatch.setattr(
        document_store, "_save_index", Mock(side_effect=OSError("interrupted index write"))
    )

    with pytest.raises(OSError, match="interrupted index write"):
        document_store.delete("abc123def456")

    reloaded = DocumentStore(path=tmp_path)
    assert not reloaded.exists("abc123def456")
    assert reloaded.count() == 0


def test_failed_file_deletion_keeps_the_index_entry(tmp_path, monkeypatch, document_store) -> None:
    document_store.add(_doc())
    monkeypatch.setattr(Path, "unlink", Mock(side_effect=OSError("file is unavailable")))

    with pytest.raises(OSError, match="file is unavailable"):
        document_store.delete("abc123def456")

    assert document_store.exists("abc123def456")
    assert DocumentStore(path=tmp_path).exists("abc123def456")


@pytest.mark.parametrize("operation", ["add", "delete", "update"])
def test_stats_use_one_snapshot_during_concurrent_writes(
    tmp_path, operation, document_store
) -> None:
    reading = Event()
    written = Event()

    class PausingCount(int):
        def __radd__(self, other):
            reading.set()
            assert written.wait(5), "Writer could not finish while statistics were calculated"
            return other + int(self)

    document_store.add(_doc("first"))
    document_store.add(_doc("second"))
    document_store.update_metadata("first", "word_count", PausingCount(7))
    document_store.update_metadata("first", "section_count", 1)
    document_store.update_metadata("second", "word_count", 2)
    document_store.update_metadata("second", "section_count", 2)

    def write():
        try:
            assert reading.wait(5), "Statistics never reached aggregation"
            if operation == "add":
                document_store.add(_doc("third"))
            elif operation == "delete":
                document_store.delete("second")
            else:
                document_store.update_metadata("second", "word_count", 99)
                document_store.update_metadata("second", "section_count", 99)
        finally:
            written.set()

    with ThreadPoolExecutor(max_workers=1) as executor:
        writer = executor.submit(write)
        stats = document_store.get_stats()
        writer.result(timeout=5)

    assert stats == {
        "document_count": 2,
        "total_words": 9,
        "total_sections": 3,
        "storage_path": str(tmp_path.resolve()),
    }


def test_metadata_snapshot_does_not_share_mutable_entries(tmp_path, document_store) -> None:
    document_store.add(_doc("first"))
    snapshot = document_store.list_with_metadata()
    initial_count = snapshot["first"]["word_count"]

    document_store.update_metadata("first", "word_count", 99)
    snapshot["first"]["filename"] = "changed.txt"

    assert snapshot["first"]["word_count"] == initial_count
    assert document_store.list_with_metadata()["first"]["filename"] == "paper.pdf"
    assert document_store.get_stats()["total_words"] == 99


def document(doc_id="paper"):
    return ParsedDocument(doc_id, "Two words.", {"filename": "paper.txt", "file_type": "txt"})


@pytest.mark.parametrize("original_name", [" \t", {"invalid": "name"}])
@pytest.mark.parametrize("recover", [False, True])
def test_document_names_skip_invalid_candidates_when_added_or_recovered(
    tmp_path, original_name, recover, document_store
):
    parsed = document()
    parsed.metadata["original_filename"] = original_name
    if recover:
        (tmp_path / "paper.json").write_text(
            json.dumps({"id": parsed.id, "content": parsed.content, "metadata": parsed.metadata}),
            encoding="utf-8",
        )
    else:
        document_store.add(parsed)

    metadata = document_store.list_with_metadata()["paper"]

    assert metadata["filename"] == "paper.txt"
    assert metadata["original_filename"] == "paper.txt"
    assert metadata["display_name"] == "paper.txt"


@pytest.mark.parametrize("read", ["index", "list_with_metadata"])
def test_public_metadata_snapshots_do_not_expose_nested_store_state(tmp_path, read, document_store):
    document_store.add(document())
    document_store.update_metadata("paper", "notes", {"tags": ["original"]})

    snapshot = document_store.index if read == "index" else document_store.list_with_metadata()
    snapshot["paper"]["notes"]["tags"].append("changed")
    snapshot["paper"]["word_count"] = 99
    snapshot.pop("paper")

    assert document_store.count() == 1
    assert document_store.list_with_metadata()["paper"]["notes"] == {"tags": ["original"]}
    assert document_store.get_stats()["total_words"] == 2
    assert DocumentStore(tmp_path).list_with_metadata() == document_store.list_with_metadata()


def test_metadata_updates_copy_caller_owned_values(tmp_path, document_store):
    document_store.add(document())
    value = {"tags": ["original"]}
    document_store.update_metadata("paper", "notes", value)
    value["tags"].append("changed")

    assert document_store.list_with_metadata()["paper"]["notes"] == {"tags": ["original"]}
    assert DocumentStore(tmp_path).list_with_metadata() == document_store.list_with_metadata()


@pytest.mark.parametrize("doc_id", ["_index", "_INDEX", "_Index"])
def test_reserved_index_names_are_rejected_before_any_write(tmp_path, doc_id, document_store):

    with pytest.raises(ValueError, match="reserved"):
        document_store.add(document(doc_id))

    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("operation", ["add", "delete"])
def test_case_colliding_document_ids_cannot_overwrite_or_delete_a_document(
    tmp_path, operation, document_store
):
    document_store.add(document("paper"))
    before = {path.name: path.read_bytes() for path in tmp_path.iterdir()}

    with pytest.raises(ValueError, match="letter case"):
        if operation == "add":
            document_store.add(document("PAPER"))
        else:
            document_store.delete("PAPER")

    assert document_store.count() == 1
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == before


@pytest.mark.parametrize("indexed", [False, True])
def test_recovery_rejects_document_symlinks_without_reading_or_writing_targets(tmp_path, indexed):
    outside = tmp_path / "outside.json"
    outside.write_text(json.dumps({"id": "linked", "content": "Private body.", "metadata": {}}))
    directory = tmp_path / "documents"
    directory.mkdir()
    (directory / "linked.json").symlink_to(outside)
    if indexed:
        (directory / "_index.json").write_text(json.dumps({"linked": {"word_count": 2}}))
    before = outside.read_bytes()

    with pytest.raises(DocumentStoreError, match="Unsafe stored"):
        DocumentStore(directory).count()

    assert outside.read_bytes() == before
    assert (directory / "linked.json").is_symlink()
    assert (directory / "_index.json").exists() is indexed


def test_index_symlinks_are_rejected_without_touching_the_target(tmp_path):
    outside = tmp_path / "outside.json"
    outside.write_text("{}")
    directory = tmp_path / "documents"
    directory.mkdir()
    (directory / "_index.json").symlink_to(outside)

    with pytest.raises(DocumentStoreError, match="index cannot be a symbolic link"):
        DocumentStore(directory).count()

    assert outside.read_text() == "{}"
    assert (directory / "_index.json").is_symlink()


def test_document_symlinks_inside_the_store_cannot_be_overwritten(tmp_path, document_store):
    document_store.add(document("original"))
    (tmp_path / "alias.json").symlink_to(tmp_path / "original.json")
    before = (tmp_path / "original.json").read_bytes()

    with pytest.raises(ValueError, match="symbolic link"):
        document_store.add(document("alias"))

    assert (tmp_path / "original.json").read_bytes() == before


def test_recovery_rejects_case_colliding_index_entries_without_rewriting_files(tmp_path):
    index = {"paper": {}, "PAPER": {}}
    (tmp_path / "_index.json").write_text(json.dumps(index))
    before = (tmp_path / "_index.json").read_bytes()

    with pytest.raises(DocumentStoreError, match="letter case"):
        DocumentStore(tmp_path).count()

    assert (tmp_path / "_index.json").read_bytes() == before


@pytest.mark.parametrize("key,value", [("word_count", 99), ("notes", {"tags": ["new"]})])
def test_failed_metadata_updates_restore_the_previous_memory_and_disk_state(
    tmp_path, monkeypatch, key, value, document_store
):
    document_store.add(document())
    before = document_store.list_with_metadata()
    disk_before = (tmp_path / "_index.json").read_bytes()

    def fail():
        raise OSError("index unavailable")

    monkeypatch.setattr(document_store, "_save_index", fail)

    with pytest.raises(OSError, match="index unavailable"):
        document_store.update_metadata("paper", key, value)

    assert document_store.list_with_metadata() == before
    assert (tmp_path / "_index.json").read_bytes() == disk_before
    assert DocumentStore(tmp_path).list_with_metadata() == before


def test_failed_add_index_write_removes_only_the_new_document_file(
    tmp_path, monkeypatch, document_store
):
    document_store.add(document("original"))
    before = {path.name: path.read_bytes() for path in tmp_path.iterdir()}

    def fail():
        raise OSError("index unavailable")

    monkeypatch.setattr(document_store, "_save_index", fail)

    with pytest.raises(OSError, match="index unavailable"):
        document_store.add(document("new"))

    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == before
    assert set(document_store.list_with_metadata()) == {"original"}
    assert DocumentStore(tmp_path).list_with_metadata() == document_store.list_with_metadata()


def test_failed_add_cleanup_preserves_original_error_and_recovers_the_orphan(
    tmp_path, monkeypatch, document_store
):
    document_store.add(document("original"))
    failure = OSError("index unavailable")
    unlink = type(tmp_path).unlink

    def fail_index():
        raise failure

    def fail_new_unlink(path, *args, **kwargs):
        if path.name == "new.json":
            raise PermissionError("cleanup unavailable")
        return unlink(path, *args, **kwargs)

    monkeypatch.setattr(document_store, "_save_index", fail_index)
    monkeypatch.setattr(type(tmp_path), "unlink", fail_new_unlink)

    with pytest.raises(OSError, match="index unavailable") as raised:
        document_store.add(document("new"))

    assert raised.value is failure
    assert json.loads((tmp_path / "new.json").read_text())["id"] == "new"
    assert set(document_store.list_with_metadata()) == {"original", "new"}
    assert DocumentStore(tmp_path).list_with_metadata() == document_store.list_with_metadata()
    assert not list(tmp_path.glob("*.tmp"))


def test_concurrent_file_creation_cannot_be_overwritten_or_removed(
    tmp_path, monkeypatch, document_store
):
    assert document_store.count() == 0
    writer = DocumentStore(tmp_path)
    link = os.link
    existing_files = {}
    published = False

    def create_before_publication(source, target):
        nonlocal published
        if not published:
            published = True
            writer.add(ParsedDocument("paper", "Original evidence.", {"filename": "original.txt"}))
            writer.update_metadata("paper", "notes", "Preserved")
            existing_files.update(
                {path.name: path.read_bytes() for path in tmp_path.glob("*.json")}
            )
        return link(source, target)

    monkeypatch.setattr(os, "link", create_before_publication)

    with pytest.raises(DocumentExistsError, match="already exists"):
        document_store.add(document())

    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == existing_files
    assert DocumentStore(tmp_path).list_with_metadata() == writer.list_with_metadata()


def test_interrupted_publication_recovers_the_completed_document_file(
    tmp_path, monkeypatch, document_store
):
    assert document_store.count() == 0
    link = os.link

    def interrupt_after_publication(source, target):
        link(source, target)
        raise KeyboardInterrupt

    monkeypatch.setattr(os, "link", interrupt_after_publication)

    with pytest.raises(KeyboardInterrupt):
        document_store.add(document())

    assert document_store._index is None
    assert json.loads((tmp_path / "paper.json").read_text())["content"] == "Two words."
    assert set(document_store.list_with_metadata()) == {"paper"}
    assert DocumentStore(tmp_path).list_with_metadata() == document_store.list_with_metadata()
    assert not list(tmp_path.glob("*.tmp"))


def test_duplicate_add_preserves_the_original_document_and_metadata(tmp_path, document_store):
    document_store.add(document())
    before = {path.name: path.read_bytes() for path in tmp_path.iterdir()}
    replacement = ParsedDocument("paper", "Different document body.", {"filename": "other.txt"})

    with pytest.raises(ValueError, match="already exists"):
        document_store.add(replacement)

    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == before


def test_add_never_overwrites_a_file_created_after_the_index_was_cached(tmp_path, document_store):
    assert document_store.count() == 0
    payload = {"id": "paper", "content": "Existing text.", "metadata": {}, "sections": []}
    (tmp_path / "paper.json").write_text(json.dumps(payload))
    before = (tmp_path / "paper.json").read_bytes()

    with pytest.raises(ValueError, match="already exists"):
        document_store.add(document())

    assert (tmp_path / "paper.json").read_bytes() == before
    assert not (tmp_path / "_index.json").exists()


@pytest.mark.parametrize("operation", ["delete", "clear"])
def test_missing_document_files_are_removed_from_cached_metadata(
    tmp_path, operation, document_store
):
    document_store.add(document())
    (tmp_path / "paper.json").unlink()

    if operation == "delete":
        assert document_store.delete("paper") is True
    else:
        assert document_store.clear() == 1
    assert document_store.count() == 0
    assert not document_store.exists("paper")
    assert DocumentStore(tmp_path).count() == 0


def test_failed_missing_file_cleanup_reconciles_on_retry(tmp_path, monkeypatch, document_store):
    document_store.add(document())
    (tmp_path / "paper.json").unlink()

    def fail():
        raise OSError("index unavailable")

    monkeypatch.setattr(document_store, "_save_index", fail)

    with pytest.raises(OSError, match="index unavailable"):
        document_store.delete("paper")

    assert document_store.count() == 0
    assert DocumentStore(tmp_path).count() == 0


@pytest.mark.parametrize("data", [b"{broken", b"\xff", b"[]", b"null", b'{"paper": []}'])
def test_corrupt_index_is_reported_without_overwriting_it_or_adding_documents(
    tmp_path, data, document_store
):
    (tmp_path / "_index.json").write_bytes(data)

    with pytest.raises(DocumentStoreError):
        document_store.add(document())

    assert (tmp_path / "_index.json").read_bytes() == data
    assert not (tmp_path / "paper.json").exists()
    assert not list(tmp_path.glob("*.tmp"))


@pytest.mark.parametrize(
    "entry",
    [
        {"word_count": "2"},
        {"word_count": True},
        {"section_count": -1},
        {"chunk_count": None},
        {"filename": ["paper.txt"]},
        {"file_type": None},
    ],
)
def test_invalid_index_metadata_is_reported_before_reconciliation(tmp_path, entry):
    data = json.dumps({"paper": entry})
    (tmp_path / "_index.json").write_text(data)

    with pytest.raises(DocumentStoreError, match="Invalid document index entry"):
        DocumentStore(tmp_path).get_stats()

    assert (tmp_path / "_index.json").read_text() == data


@pytest.mark.parametrize(
    "data",
    [
        b"{broken",
        b"\xff",
        b"[]",
        b'{"id":"other","content":"Text."}',
        b'{"id":"paper","content":null}',
        b'{"id":"paper","content":"Text.","metadata":null}',
        b'{"id":"paper","content":"Text.","sections":{}}',
        b'{"id":"paper","content":"Text.","sections":["invalid"]}',
    ],
)
def test_invalid_orphan_documents_are_preserved_and_never_reported_as_empty(tmp_path, data):
    (tmp_path / "paper.json").write_bytes(data)

    with pytest.raises(DocumentStoreError):
        DocumentStore(tmp_path).count()

    assert (tmp_path / "paper.json").read_bytes() == data
    assert not (tmp_path / "_index.json").exists()


def test_recovery_read_errors_propagate_without_rewriting_files(tmp_path, monkeypatch):
    (tmp_path / "paper.json").write_text("{}")
    read_text = type(tmp_path).read_text

    def fail(path, *args, **kwargs):
        if path.name == "paper.json":
            raise PermissionError("document unavailable")
        return read_text(path, *args, **kwargs)

    monkeypatch.setattr(type(tmp_path), "read_text", fail)

    with pytest.raises(PermissionError, match="document unavailable"):
        DocumentStore(tmp_path).count()

    assert not (tmp_path / "_index.json").exists()


@pytest.mark.parametrize(
    "key,value", [(None, 2), (1, 2), ("", 2), ("word_count", -1), ("notes", float("nan"))]
)
def test_invalid_metadata_updates_leave_memory_and_disk_unchanged(
    tmp_path, key, value, document_store
):
    document_store.add(document())
    before = document_store.list_with_metadata()
    disk_before = (tmp_path / "_index.json").read_bytes()

    with pytest.raises(ValueError):
        document_store.update_metadata("paper", key, value)

    assert document_store.list_with_metadata() == before
    assert (tmp_path / "_index.json").read_bytes() == disk_before


@pytest.mark.parametrize("content", [None, 42, ["words"]])
def test_nontext_documents_are_rejected_before_creating_files(tmp_path, content, document_store):

    with pytest.raises(ValueError, match="content must be text"):
        document_store.add(ParsedDocument("paper", content, {}))

    assert list(tmp_path.iterdir()) == []


def test_nonfinite_document_metadata_is_rejected_before_writing_a_document(
    tmp_path, document_store
):

    with pytest.raises(ValueError):
        document_store.add(ParsedDocument("paper", "Some text.", {"score": float("inf")}))

    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("filename", ["_index.json", "paper.json"])
@pytest.mark.parametrize("value", [float("nan"), float("inf"), -float("inf")])
def test_nonfinite_json_values_are_rejected_when_reading_stored_data(tmp_path, filename, value):
    payload = (
        {"paper": {"notes": value}}
        if filename == "_index.json"
        else {"id": "paper", "content": "Words.", "metadata": {"score": value}}
    )
    data = json.dumps(payload)
    (tmp_path / filename).write_text(data)

    with pytest.raises(DocumentStoreError, match="Cannot read stored JSON"):
        DocumentStore(tmp_path).count()

    assert (tmp_path / filename).read_text() == data


def test_index_symlink_created_after_loading_is_never_replaced_by_a_write(tmp_path):
    directory = tmp_path / "documents"
    document_store = DocumentStore(directory)
    document_store.add(document())
    before = document_store.list_with_metadata()
    outside = tmp_path / "outside.json"
    outside.write_text("{}")
    (directory / "_index.json").unlink()
    (directory / "_index.json").symlink_to(outside)

    with pytest.raises(DocumentStoreError, match="symbolic link"):
        document_store.update_metadata("paper", "notes", "Changed")

    assert document_store.list_with_metadata() == before
    assert outside.read_text() == "{}"
    assert (directory / "_index.json").is_symlink()


@pytest.mark.parametrize("operation", ["exists", "update"])
@pytest.mark.parametrize("doc_id", [None, [], "", "../paper", "paper\n", "_INDEX"])
def test_metadata_access_rejects_invalid_document_ids_without_changing_state(
    tmp_path, operation, doc_id, document_store
):
    document_store.add(document())
    before = document_store.list_with_metadata()

    with pytest.raises(ValueError, match="Document ID"):
        if operation == "exists":
            document_store.exists(doc_id)
        else:
            document_store.update_metadata(doc_id, "notes", "changed")

    assert document_store.list_with_metadata() == before


@pytest.mark.parametrize("filename", ["_index.json", "paper.json"])
def test_overflowing_json_numbers_are_rejected_without_rewriting_files(tmp_path, filename):
    data = (
        '{"paper":{"notes":{"ratio":1e999}}}'
        if filename == "_index.json"
        else '{"id":"paper","content":"Words.","metadata":{"notes":{"ratio":1e999}}}'
    )
    (tmp_path / filename).write_text(data)

    with pytest.raises(DocumentStoreError, match="Cannot read stored JSON"):
        DocumentStore(tmp_path).count()

    assert (tmp_path / filename).read_text() == data


def test_finite_custom_metadata_numbers_survive_a_reload(tmp_path, document_store):
    document_store.add(document())
    document_store.update_metadata("paper", "notes", {"ratio": 0.125, "exponent": 1.25e25})

    assert DocumentStore(tmp_path).list_with_metadata() == document_store.list_with_metadata()


class TestEmbedder:
    @pytest.fixture(autouse=True)
    def sentence_transformers_module(self, monkeypatch):
        module = ModuleType("sentence_transformers")
        module.SentenceTransformer = MagicMock()
        monkeypatch.setitem(sys.modules, "sentence_transformers", module)

    @pytest.fixture
    def mock_sentence_transformer(self):
        with patch("sentence_transformers.SentenceTransformer") as mock_cls:
            mock_model = MagicMock()
            mock_model.get_embedding_dimension.return_value = 384
            mock_model.get_sentence_embedding_dimension.return_value = 384
            rng = np.random.default_rng(0)

            def mock_encode(texts, **kwargs):
                shape = (384,) if isinstance(texts, str) else (len(texts), 384)
                return rng.random(shape, dtype=np.float32)

            mock_model.encode.side_effect = mock_encode
            mock_cls.return_value = mock_model
            yield mock_model

    def test_embed_single_text(self, mock_sentence_transformer):
        embedder = Embedder(model_name="test-model")
        texts = ["This is a test sentence."]

        embeddings = embedder.embed(texts)

        assert embeddings.shape == (1, 384)
        assert embeddings.dtype == np.float32

    def test_embed_batch(self, mock_sentence_transformer):
        embedder = Embedder(model_name="test-model")
        texts = [
            "First sentence.",
            "Second sentence.",
            "Third sentence.",
        ]

        embeddings = embedder.embed(texts)

        assert embeddings.shape == (3, 384)
        assert embeddings.dtype == np.float32
        mock_sentence_transformer.encode.assert_called_once()

    def test_embed_query(self, mock_sentence_transformer):
        embedder = Embedder(model_name="test-model")
        query = "What is the meaning of life?"

        embedding = embedder.embed_query(query)

        assert embedding.shape == (384,)
        assert embedding.dtype == np.float32

    def test_embedding_dimension(self, mock_sentence_transformer):
        embedder = Embedder(model_name="test-model")

        assert embedder.dimension == 384
        mock_sentence_transformer.get_embedding_dimension.assert_called()

    def test_empty_list(self, mock_sentence_transformer):
        embedder = Embedder(model_name="test-model")

        embeddings = embedder.embed([])

        assert embeddings.shape == (0, 384)

    def test_batch_size_passed(self, mock_sentence_transformer):
        embedder = Embedder(model_name="test-model")
        texts = ["text"] * 10

        embedder.embed(texts, batch_size=16)

        call_kwargs = mock_sentence_transformer.encode.call_args[1]
        assert call_kwargs["batch_size"] == 16

    @pytest.mark.parametrize("normalize", [True, False])
    def test_normalization_option_reaches_encoder(self, mock_sentence_transformer, normalize):
        Embedder(model_name="test-model", normalize=normalize).embed(["Test text"])
        assert (
            mock_sentence_transformer.encode.call_args.kwargs["normalize_embeddings"] is normalize
        )

    def test_lazy_model_loading(self):
        with patch("sentence_transformers.SentenceTransformer") as mock_cls:
            embedder = Embedder(model_name="test-model")

            mock_cls.assert_not_called()

            _ = embedder.model

            mock_cls.assert_called_once_with("test-model", device=None, revision=None)

    def test_device_configuration(self):
        with patch("sentence_transformers.SentenceTransformer") as mock_cls:
            mock_cls.return_value = MagicMock()
            embedder = Embedder(model_name="test-model", device="cuda")

            _ = embedder.model

            mock_cls.assert_called_once_with("test-model", device="cuda", revision=None)

    def test_concurrent_model_first_access_constructs_once(self, observed_lock):
        embedder = Embedder(model_name="test-model", revision="a" * 40)
        lock = observed_lock
        embedder._model_lock = lock
        start = Barrier(4)
        constructing = Event()
        release = Event()
        model = object()

        def construct(*_args, **_kwargs):
            constructing.set()
            assert release.wait(timeout=5)
            return model

        def load():
            start.wait(timeout=5)
            return embedder.model

        with (
            patch("sentence_transformers.SentenceTransformer", side_effect=construct) as factory,
            ThreadPoolExecutor(max_workers=4) as workers,
        ):
            futures = [workers.submit(load) for _ in range(4)]
            try:
                assert constructing.wait(timeout=5)
                assert len(set(lock.wait_for_attempts(4))) == 4
            finally:
                release.set()
            assert all(future.result(timeout=5) is model for future in futures)

        factory.assert_called_once_with("test-model", device=None, revision="a" * 40)
        assert embedder._model is model
        assert embedder._loaded_model_source == ("test-model", "a" * 40)

    def test_failed_constructor_can_retry_without_deadlock(self, observed_lock):
        embedder = Embedder(model_name="test-model", revision="a" * 40)
        embedder._model_lock = observed_lock
        model = object()
        with patch(
            "sentence_transformers.SentenceTransformer",
            side_effect=[RuntimeError("Constructor failed"), model],
        ) as factory:
            with pytest.raises(RuntimeError, match="Constructor failed"):
                _ = embedder.model
            assert embedder._model is None
            assert embedder._loaded_model_source is None

            with ThreadPoolExecutor(max_workers=1) as workers:
                assert workers.submit(lambda: embedder.model).result(timeout=5) is model
            assert embedder.model is model

        assert factory.call_count == 2
        assert embedder._loaded_model_source == ("test-model", "a" * 40)

    def test_unload_waits_for_initial_load_and_next_access_loads_fresh_model(self, observed_lock):
        embedder = Embedder(model_name="test-model", revision="a" * 40)
        lock = observed_lock
        embedder._model_lock = lock
        constructing = Event()
        release = Event()
        original_model, fresh_model = object(), object()

        def construct(*_args, **_kwargs):
            constructing.set()
            assert release.wait(timeout=5)
            return original_model

        with (
            patch("sentence_transformers.SentenceTransformer", side_effect=construct) as factory,
            ThreadPoolExecutor(max_workers=2) as workers,
        ):
            loading = workers.submit(lambda: embedder.model)
            try:
                assert constructing.wait(timeout=5)
                lock.wait_for_attempts(1)
                unloading = workers.submit(embedder.unload)
                lock.wait_for_attempts(1)
                assert not unloading.done()
            finally:
                release.set()
            assert loading.result(timeout=5) is original_model
            unloading.result(timeout=5)
            assert embedder._model is None
            assert embedder._loaded_model_source is None

            factory.side_effect = None
            factory.return_value = fresh_model
            assert embedder.model is fresh_model
            assert embedder.model is fresh_model

        assert factory.call_count == 2
        assert embedder._loaded_model_source == ("test-model", "a" * 40)

    def test_similarity_normalized(self, mock_sentence_transformer):
        embedder = Embedder(model_name="test-model", normalize=True)

        query_emb = np.array([1.0, 0.0, 0.0], dtype=np.float32)
        doc_embs = np.array(
            [
                [1.0, 0.0, 0.0],
                [0.0, 1.0, 0.0],
                [0.5, 0.5, 0.0],
            ],
            dtype=np.float32,
        )

        doc_embs = doc_embs / np.linalg.norm(doc_embs, axis=1, keepdims=True)

        similarities = embedder.similarity(query_emb, doc_embs)

        assert similarities.shape == (3,)
        assert similarities[0] == pytest.approx(1.0, abs=0.01)
        assert similarities[1] == pytest.approx(0.0, abs=0.01)

    def test_similarity_unnormalized(self, mock_sentence_transformer):
        embedder = Embedder(model_name="test-model", normalize=False)

        query_emb = np.array([2.0, 0.0], dtype=np.float32)
        doc_embs = np.array(
            [
                [3.0, 0.0],
                [0.0, 1.0],
            ],
            dtype=np.float32,
        )

        similarities = embedder.similarity(query_emb, doc_embs)

        assert similarities.shape == (2,)
        assert similarities[0] == pytest.approx(1.0, abs=0.01)
        assert similarities[1] == pytest.approx(0.0, abs=0.01)

    def test_show_progress_flag(self, mock_sentence_transformer):
        embedder = Embedder(model_name="test-model")
        texts = ["text"] * 5

        embedder.embed(texts, show_progress=True)

        call_kwargs = mock_sentence_transformer.encode.call_args[1]
        assert call_kwargs["show_progress_bar"] is True

    def test_embedding_configuration_uses_effective_defaults(self, mock_sentence_transformer):
        embedder = Embedder()

        assert embedder.embedding_config == EmbeddingConfig(
            model_name="BAAI/bge-small-en-v1.5",
            revision=HF_REVISIONS["BAAI/bge-small-en-v1.5"],
            query_prefix="Represent this sentence for searching relevant passages: ",
            doc_prefix="",
            normalize=True,
            dimension=384,
        )
        mock_sentence_transformer.encode.assert_not_called()

    def test_embedding_configuration_preserves_exact_overrides(self, mock_sentence_transformer):
        embedder = Embedder(
            model_name="test-model",
            revision="a" * 40,
            query_prefix="  query: ",
            doc_prefix="passage:  ",
            normalize=False,
        )

        assert embedder.embedding_config == EmbeddingConfig(
            "test-model", "a" * 40, "  query: ", "passage:  ", False, 384
        )

    @pytest.mark.parametrize("revision", [None, "main", "v1.0", "a" * 39, "g" * 40])
    def test_unknown_revision_rejects_before_model_loading(self, revision):
        with patch("sentence_transformers.SentenceTransformer") as model:
            embedder = Embedder(model_name="test-model", revision=revision)

            with pytest.raises(EmbeddingConfigError):
                _ = embedder.embedding_config

            model.assert_not_called()

    def test_local_model_path_rejects_even_with_commit_shaped_revision(self, tmp_path):
        with patch("sentence_transformers.SentenceTransformer") as model:
            embedder = Embedder(model_name=str(tmp_path), revision="a" * 40)

            with pytest.raises(EmbeddingConfigError):
                _ = embedder.embedding_config

            model.assert_not_called()

    @pytest.mark.parametrize("changed_field", ["model_name", "revision"])
    def test_loaded_weights_cannot_be_relabelled_before_unload(
        self, mock_sentence_transformer, changed_field
    ):
        embedder = Embedder(model_name="test-model", revision="a" * 40)
        original = embedder.embedding_config
        changed = "another-model" if changed_field == "model_name" else "b" * 40
        setattr(embedder, changed_field, changed)

        with pytest.raises(EmbeddingConfigError):
            _ = embedder.embedding_config
        mock_sentence_transformer.encode.assert_not_called()

        embedder.unload()
        reloaded = embedder.embedding_config
        assert getattr(reloaded, changed_field) == changed
        assert reloaded != original


@pytest.fixture
def model_factory(monkeypatch):
    module = ModuleType("sentence_transformers")
    module.SentenceTransformer = MagicMock()
    monkeypatch.setitem(sys.modules, "sentence_transformers", module)
    return module.SentenceTransformer


@pytest.mark.parametrize(
    "options",
    [
        {"model_name": ""},
        {"model_name": "  "},
        {"model_name": None},
        {"normalize": 1},
        {"query_prefix": 1},
        {"doc_prefix": False},
        {"revision": ""},
        {"revision": "  "},
        {"revision": 123},
    ],
)
def test_invalid_settings_fail_without_loading(model_factory, options):
    with pytest.raises(EmbeddingConfigError):
        Embedder(**options)

    model_factory.assert_not_called()


def test_invalid_device_fails_without_loading(model_factory):
    with pytest.raises(ValueError, match="device"):
        Embedder(device="automatic")

    model_factory.assert_not_called()


@pytest.mark.parametrize(
    ("field", "value"),
    [("model_name", ""), ("query_prefix", 1), ("doc_prefix", False), ("normalize", 1)],
)
def test_changed_invalid_settings_fail_before_index_model_load(model_factory, field, value):
    embedder = Embedder()
    setattr(embedder, field, value)

    with pytest.raises(EmbeddingConfigError):
        _ = embedder.embedding_config

    model_factory.assert_not_called()


def test_explicit_empty_prefixes_override_model_defaults(model_factory):
    embedder = Embedder(query_prefix="", doc_prefix="")

    assert embedder.query_prefix == ""
    assert embedder.doc_prefix == ""
    model_factory.assert_not_called()


@pytest.mark.parametrize(
    "options",
    [
        {"texts": "sentence"},
        {"texts": ["sentence", None]},
        {"texts": None},
        {"texts": ["sentence"], "batch_size": 0},
        {"texts": [], "batch_size": -1},
        {"texts": ["sentence"], "batch_size": True},
        {"texts": ["sentence"], "batch_size": 1.5},
        {"texts": ["sentence"], "show_progress": 1},
    ],
)
def test_invalid_batch_inputs_fail_without_model_load(model_factory, options):
    with pytest.raises(ValueError):
        Embedder().embed(**options)

    model_factory.assert_not_called()


@pytest.mark.parametrize("query", [None, 3, ["sentence"]])
def test_invalid_query_fails_without_model_load(model_factory, query):
    with pytest.raises(ValueError, match="query"):
        Embedder().embed_query(query)

    model_factory.assert_not_called()


@pytest.mark.parametrize("dimension", [None, False, 0, -1, 3.5, "3", np.nan])
def test_invalid_model_dimension_never_reaches_encoder(model_factory, dimension):
    model = model_factory.return_value
    model.get_embedding_dimension.return_value = dimension

    with pytest.raises(EmbeddingConfigError, match="dimension"):
        Embedder().embed(["sentence"])

    model.encode.assert_not_called()


def test_sentence_dimension_api_is_supported(model_factory):
    model = model_factory.return_value
    model.get_embedding_dimension = None
    model.get_sentence_embedding_dimension.return_value = np.int64(3)

    assert Embedder().dimension == 3


def test_missing_dimension_api_has_clear_error(model_factory):
    model = model_factory.return_value
    model.get_embedding_dimension = None
    model.get_sentence_embedding_dimension = None

    with pytest.raises(EmbeddingConfigError, match="dimension"):
        _ = Embedder().dimension


@pytest.mark.parametrize(
    "values",
    [
        np.ones(3),
        np.ones((1, 2)),
        np.ones((2, 3)),
        np.array([[1, np.nan, 0]]),
        np.array([[1, np.inf, 0]]),
        np.array([[1, 1e100, 0]]),
        np.array([["1", "2", "3"]]),
        np.ones((1, 3), dtype=complex),
    ],
)
def test_invalid_batch_result_rejected(model_factory, values):
    model = model_factory.return_value
    model.get_embedding_dimension.return_value = 3
    model.encode.return_value = values

    with pytest.raises(ValueError, match="Embedding model"):
        Embedder().embed(["sentence"])


@pytest.mark.parametrize("values", [np.ones((1, 3)), np.ones(2), [1, np.nan, 0]])
def test_invalid_query_result_rejected(model_factory, values):
    model = model_factory.return_value
    model.get_embedding_dimension.return_value = 3
    model.encode.return_value = values

    with pytest.raises(ValueError, match="Embedding model"):
        Embedder().embed_query("question")


def test_valid_vectors_are_preserved_as_float32(model_factory):
    model = model_factory.return_value
    model.get_embedding_dimension.return_value = 3
    values = np.array([[0.1, -0.2, 0.3]], dtype=np.float64)
    model.encode.return_value = values
    embedder = Embedder(query_prefix="q: ", doc_prefix="d: ", normalize=False)

    result = embedder.embed(["sentence"], batch_size=2)

    assert result.dtype == np.float32
    np.testing.assert_array_equal(result, values.astype(np.float32))
    model.encode.assert_called_once_with(
        ["d: sentence"],
        batch_size=2,
        show_progress_bar=False,
        normalize_embeddings=False,
        convert_to_numpy=True,
    )
    model.encode.reset_mock()
    model.encode.return_value = values[0]

    query = embedder.embed_query("question")

    np.testing.assert_array_equal(query, values[0].astype(np.float32))
    model.encode.assert_called_once_with(
        "q: question",
        normalize_embeddings=False,
        convert_to_numpy=True,
    )


@pytest.mark.parametrize("scale", [1e-30, 1.0, 1e20, 1e38])
def test_cosine_similarity_handles_finite_float32_scales(scale):
    embedder = Embedder(normalize=False)
    query = np.array([scale, 0], dtype=np.float32)
    documents = np.array([[scale, 0], [0, scale], [-scale, 0], [0, 0]], dtype=np.float32)
    original_query = query.copy()
    original_documents = documents.copy()

    with np.errstate(all="raise"):
        scores = embedder.similarity(query, documents)

    np.testing.assert_array_equal(scores, [1, 0, -1, 0])
    assert scores.dtype == np.float32
    np.testing.assert_array_equal(query, original_query)
    np.testing.assert_array_equal(documents, original_documents)


def test_cosine_similarity_of_zero_query_is_zero():
    query = np.zeros(2, dtype=np.float32)
    documents = np.array([[1, 0], [0, 0]], dtype=np.float32)

    with np.errstate(all="raise"):
        scores = Embedder(normalize=False).similarity(query, documents)

    np.testing.assert_array_equal(scores, [0, 0])


@pytest.mark.parametrize(
    "metadata,expected",
    [
        (
            {
                "document_name": "document.pdf",
                "display_name": "display.pdf",
                "filename": "stored.pdf",
                "original_filename": "original.pdf",
            },
            "original.pdf",
        ),
        (
            {"filename": "stored.pdf", "display_name": "display.pdf", "document_name": "doc.pdf"},
            "stored.pdf",
        ),
        ({"display_name": "display.pdf", "document_name": "doc.pdf"}, "display.pdf"),
        ({"document_name": "doc.pdf"}, "doc.pdf"),
    ],
)
def test_available_names_keep_the_existing_field_priority(metadata, expected):
    assert resolve_display_name(metadata) == expected


@pytest.mark.parametrize(
    "invalid_name",
    [
        None,
        False,
        True,
        0,
        42,
        b"paper.pdf",
        [],
        ["paper.pdf"],
        {},
        {"name": "paper.pdf"},
        "",
        " \t\n\u2003",
    ],
)
def test_invalid_candidates_do_not_hide_a_valid_filename(invalid_name):
    metadata = {"original_filename": invalid_name, "filename": "paper.pdf"}

    assert resolve_display_name(metadata) == "paper.pdf"


@pytest.mark.parametrize("fallback", ["Unknown", "Document", ""])
def test_missing_valid_names_return_the_exact_requested_fallback(fallback):
    metadata = {"original_filename": " \t", "filename": None, "display_name": False}

    assert resolve_display_name(metadata, fallback=fallback) == fallback


def test_empty_metadata_uses_the_default_name():
    assert resolve_display_name({}) == "Unknown"


@pytest.mark.parametrize(
    "name",
    [
        " \treport.pdf ",
        "résumé 数据 α.md",
        "folder/report.pdf",
        r"C:\papers\report.pdf",
        "file  name.pdf",
    ],
)
def test_valid_names_keep_their_text_and_whitespace(name):
    assert resolve_display_name({"original_filename": name}) == name


def test_read_only_metadata_is_supported():
    metadata = MappingProxyType({"original_filename": " \t", "filename": "paper.pdf"})

    assert resolve_display_name(metadata) == "paper.pdf"


def test_non_string_candidates_do_not_invoke_conversion_or_truthiness():
    class InvalidName:
        def __bool__(self):
            raise AssertionError("metadata value must not be interpreted as a Boolean")

        def __str__(self):
            raise AssertionError("metadata value must not become a fabricated filename")

    assert (
        resolve_display_name({"original_filename": InvalidName(), "filename": "paper.pdf"})
        == "paper.pdf"
    )


@pytest.mark.parametrize("chunker_type", [HierarchicalChunker, FixedSizeChunker])
@pytest.mark.parametrize("invalid_name", [" \t", 42])
def test_chunks_use_a_valid_filename_when_the_original_name_is_invalid(chunker_type, invalid_name):
    document = ParsedDocument(
        "paper",
        "The measured voltage is 3.3 V.",
        {"original_filename": invalid_name, "filename": "paper.pdf"},
    )

    chunks = chunker_type().chunk(document)

    assert chunks
    assert all(chunk.document_name == "paper.pdf" for chunk in chunks)


@pytest.mark.parametrize("stage", ["construction", "labels"])
def test_concurrent_first_requests_share_model_and_resolved_labels(
    monkeypatch, stage, observed_lock
):
    verifier = GroundingVerifier(model_name="test-nli", revision="a" * 40)
    lock = observed_lock
    verifier._model_lock = lock
    start = Barrier(4)
    initializing = Event()
    release = Event()

    def pause():
        initializing.set()
        assert release.wait(timeout=5)

    class Config:
        @property
        def id2label(self):
            if stage == "labels":
                pause()
            return {0: "neutral", 1: "contradiction", 2: "entailment"}

    model = SimpleNamespace(model=SimpleNamespace(config=Config()))

    def construct(*_args, **_kwargs):
        if stage == "construction":
            pause()
        return model

    factory = Mock(side_effect=construct)
    module = ModuleType("sentence_transformers")
    module.CrossEncoder = factory
    monkeypatch.setitem(sys.modules, "sentence_transformers", module)

    def load():
        start.wait(timeout=5)
        loaded = verifier.model
        return loaded, verifier.entail_idx, verifier.contra_idx

    with ThreadPoolExecutor(max_workers=4) as workers:
        futures = [workers.submit(load) for _ in range(4)]
        try:
            assert initializing.wait(timeout=5)
            assert len({lock.attempts.get(timeout=5) for _ in range(4)}) == 4
            assert not any(future.done() for future in futures)
        finally:
            release.set()
        assert all(future.result(timeout=5) == (model, 2, 1) for future in futures)

    factory.assert_called_once_with("test-nli", device=None, revision="a" * 40)


def test_failed_grounding_model_load_can_retry(monkeypatch):
    verifier = GroundingVerifier(model_name="test-nli")
    model = SimpleNamespace(model=SimpleNamespace(config=SimpleNamespace(id2label={})))
    factory = Mock(side_effect=[RuntimeError("Load failed"), model])
    module = ModuleType("sentence_transformers")
    module.CrossEncoder = factory
    monkeypatch.setitem(sys.modules, "sentence_transformers", module)

    with pytest.raises(RuntimeError, match="Load failed"):
        _ = verifier.model
    assert verifier._model is None

    with ThreadPoolExecutor(max_workers=1) as workers:
        assert workers.submit(lambda: verifier.model).result(timeout=5) is model
    assert verifier.model is model
    assert factory.call_count == 2


def test_concurrent_first_requests_share_http_client(monkeypatch, observed_lock):
    llm = LLMClient()
    lock = observed_lock
    llm._client_lock = lock
    start = Barrier(4)
    constructing = Event()
    release = Event()
    client = Mock()

    def construct(**_kwargs):
        constructing.set()
        assert release.wait(timeout=5)
        return client

    factory = Mock(side_effect=construct)
    monkeypatch.setattr("scinexusrag.generation.llm.httpx.Client", factory)

    def load():
        start.wait(timeout=5)
        return llm.client

    with ThreadPoolExecutor(max_workers=4) as workers:
        futures = [workers.submit(load) for _ in range(4)]
        try:
            assert constructing.wait(timeout=5)
            assert len({lock.attempts.get(timeout=5) for _ in range(4)}) == 4
        finally:
            release.set()
        assert all(future.result(timeout=5) is client for future in futures)

    assert factory.call_count == 1
    assert factory.call_args.kwargs["base_url"] == llm.base_url
    assert factory.call_args.kwargs["timeout"].read == llm.timeout
    llm.close()
    client.close.assert_called_once()


def test_failed_http_client_construction_can_retry(monkeypatch):
    llm = LLMClient()
    client = Mock()
    factory = Mock(side_effect=[RuntimeError("Client construction failed"), client])
    monkeypatch.setattr("scinexusrag.generation.llm.httpx.Client", factory)

    with pytest.raises(RuntimeError, match="Client construction failed"):
        _ = llm.client
    assert llm._client is None

    with ThreadPoolExecutor(max_workers=1) as workers:
        assert workers.submit(lambda: llm.client).result(timeout=5) is client
    assert llm.client is client
    assert factory.call_count == 2
    llm.close()


def test_close_waits_for_http_client_construction(monkeypatch, observed_lock):
    llm = LLMClient()
    lock = observed_lock
    llm._client_lock = lock
    constructing = Event()
    release = Event()
    original, fresh = Mock(), Mock()

    def construct(**_kwargs):
        constructing.set()
        assert release.wait(timeout=5)
        return original

    factory = Mock(side_effect=construct)
    monkeypatch.setattr("scinexusrag.generation.llm.httpx.Client", factory)

    with ThreadPoolExecutor(max_workers=2) as workers:
        loading = workers.submit(lambda: llm.client)
        try:
            assert constructing.wait(timeout=5)
            lock.attempts.get(timeout=5)
            closing = workers.submit(llm.close)
            lock.attempts.get(timeout=5)
            assert not closing.done()
        finally:
            release.set()
        assert loading.result(timeout=5) is original
        closing.result(timeout=5)

    original.close.assert_called_once()
    assert llm._client is None
    factory.side_effect = None
    factory.return_value = fresh
    assert llm.client is fresh
    llm.close()
    fresh.close.assert_called_once()


@pytest.mark.parametrize(
    "query", ["What is CVS?", "What is VSV infection?", "What is reasonable evidence?"]
)
def test_query_words_do_not_trigger_substring_instructions(query):
    assert QueryAnalyzer().analyze(query).query_type is QueryType.FACTUAL


@pytest.mark.parametrize(
    "query,expected",
    [
        ("Compare the two methods.", QueryType.COMPARISON),
        ("A vs. B", QueryType.COMPARISON),
        ("A versus B", QueryType.COMPARISON),
        ("Why did the study stop?", QueryType.CAUSAL),
        ("Explain the reason for the change.", QueryType.CAUSAL),
        ("Explain the reasons for the change.", QueryType.CAUSAL),
    ],
)
def test_comparison_and_causal_words_keep_their_meaning(query, expected):
    assert QueryAnalyzer().analyze(query).query_type is expected


@pytest.mark.parametrize(
    "query", ["  TELL ME\nABOUT\tTHIS PAPER?  ", "tell me   about this paper."]
)
def test_vague_rewriting_normalizes_whitespace(query):
    analyzer = QueryAnalyzer()
    assert analyzer.rewrite_vague_query(query) == analyzer.rewrite_vague_query(
        "tell me about this paper"
    )
    assert analyzer.analyze(query).original == query


def test_specific_query_is_not_rewritten():
    query = "  What is the measured threshold?\n"
    assert QueryAnalyzer().rewrite_vague_query(query) == query


@pytest.fixture
def passage():
    return RetrievalResult(
        chunk=Chunk(id="chunk", document_id="doc", content="Measured evidence."), score=0.8
    )


@pytest.mark.parametrize(
    "names,chunk_name,metadata,expected",
    [
        ({"doc": "index.pdf"}, "chunk.pdf", {"filename": "stored.pdf"}, "index.pdf"),
        ({"doc": ""}, "chunk.pdf", {"filename": "stored.pdf"}, "chunk.pdf"),
        ({"doc": " \t"}, "chunk.pdf", {"filename": "stored.pdf"}, "chunk.pdf"),
        ({}, "", {"original_filename": "original.pdf", "filename": "stored.pdf"}, "original.pdf"),
        ({}, "", {}, "Document 1"),
    ],
)
def test_prompt_source_names_keep_available_metadata(
    passage, names, chunk_name, metadata, expected
):
    passage.chunk.document_name = chunk_name
    passage.chunk.metadata = metadata
    prepared = Synthesizer(Mock(spec=LLMClient)).prepare("What was measured?", [passage], names)
    assert prepared.error is None
    assert prepared.sources[0].document_name == expected
    assert f"[1] Source: {expected}" in prepared.prompt


@pytest.mark.parametrize(
    "page,section,expected_page,expected_section",
    [
        (7, "Methods", 7, "Methods"),
        (0, "", 0, ""),
        (True, 7, None, None),
        ("7", ["Methods"], None, None),
    ],
)
def test_prompt_metadata_fallbacks_have_source_types(
    passage, page, section, expected_page, expected_section
):
    passage.chunk.metadata = {"page_number": page, "section_title": section}
    prepared = Synthesizer(Mock(spec=LLMClient)).prepare("What was measured?", [passage])
    assert prepared.sources[0].page_number == expected_page
    assert prepared.sources[0].section_title == expected_section
    if expected_page is not None:
        assert f"Page {expected_page}" in prepared.prompt
    else:
        assert ", Page " not in prepared.prompt


def test_explicit_source_location_precedes_metadata(passage):
    passage.chunk.page_number = 0
    passage.chunk.section_title = "Results"
    passage.chunk.metadata = {"page_number": 9, "section_title": "Methods"}
    prepared = Synthesizer(Mock(spec=LLMClient)).prepare("What was measured?", [passage])
    assert prepared.sources[0].page_number == 0
    assert prepared.sources[0].section_title == "Results"
    assert "Section: Results, Page 0" in prepared.prompt


@pytest.mark.parametrize(
    "name,value,error",
    [
        ("max_sources", -1, ValueError),
        ("max_sources", True, TypeError),
        ("max_sources", 1.5, TypeError),
        ("max_context_length", -1, ValueError),
        ("max_context_length", False, TypeError),
        ("max_tokens", 0, ValueError),
        ("max_tokens", 1.5, TypeError),
        ("temperature", True, TypeError),
        ("temperature", "0.1", TypeError),
        ("temperature", -0.1, ValueError),
        ("temperature", float("nan"), ValueError),
        ("temperature", float("inf"), ValueError),
        ("temperature", 10**1000, ValueError),
    ],
)
def test_invalid_synthesis_limits_fail_before_generation(name, value, error):
    llm = Mock(spec=LLMClient)
    with pytest.raises(error, match=name):
        Synthesizer(llm, **{name: value})
    llm.generate.assert_not_called()
    llm.stream.assert_not_called()


@pytest.mark.parametrize("name", ["max_sources", "max_context_length"])
def test_zero_source_budget_keeps_no_fit_response(passage, name):
    llm = Mock(spec=LLMClient)
    result = Synthesizer(llm, **{name: 0}).synthesize("What was measured?", [passage])
    assert result.error == "No source passages fit the context limit."
    assert result.sources == []
    llm.generate.assert_not_called()


@pytest.mark.parametrize("score", [float("nan"), float("inf"), float("-inf")])
def test_nonfinite_source_scores_fail_before_generation(passage, score):
    llm = Mock(spec=LLMClient)
    passage.score = score
    with pytest.raises(ValueError, match="scores must be finite"):
        Synthesizer(llm).synthesize("What was measured?", [passage])
    llm.generate.assert_not_called()


@pytest.mark.parametrize("score", [sys.float_info.max, -sys.float_info.max, -0.25])
def test_finite_source_scores_are_preserved_without_overflow(passage, score):
    passage.score = score
    synth = Synthesizer(Mock(spec=LLMClient), max_context_length=10000)
    prepared = synth.prepare("What was measured?", [passage])
    assert prepared.error is None
    assert prepared.sources[0].score == score
    assert "relevance:" in prepared.prompt
    assert "inf%" not in prepared.prompt


def test_default_prompt_and_generation_options_are_preserved(passage):
    llm = Mock(spec=LLMClient)
    llm.generate.return_value = " Answer [1].\n"
    passage.chunk.document_name = "paper.pdf"
    passage.chunk.section_title = "Results"
    passage.chunk.page_number = 7
    response = Synthesizer(llm).synthesize("What was measured?", [passage])

    assert response.answer == " Answer [1].\n"
    assert response.raw_response == " Answer [1].\n"
    assert llm.generate.call_args.args == (
        "QUESTION: What was measured?\n\n"
        "SOURCES:\n"
        "====================\n"
        "[1] Source: paper.pdf, Section: Results, Page 7 (relevance: 80%)\n"
        '"Measured evidence."\n'
        "====================\n\n"
        "Answer using only the sources above. Cite with [1], [2], etc.",
    )
    options = llm.generate.call_args.kwargs
    assert options["temperature"] == 0.1
    assert options["max_tokens"] == 356
    assert 'say: "This information is not found in the provided documents."' in options["system"]
    assert "Use ONLY information from the sources" in options["system"]


@pytest.fixture
def integer_string_limit():
    previous = sys.get_int_max_str_digits()
    sys.set_int_max_str_digits(4300)
    try:
        yield
    finally:
        sys.set_int_max_str_digits(previous)


def _generation_result(idx: int, content: str, score: float = 0.9) -> RetrievalResult:
    chunk = Chunk(id=f"c{idx}", content=content, document_id=f"d{idx}")
    return RetrievalResult(chunk=chunk, score=score)


def _consume_stream(chunks: list[str]) -> tuple[str, list[dict[str, Any]]]:
    answer = ""
    events = []
    for chunk in chunks:
        if chunk.startswith("\ndata: "):
            event = json.loads(chunk.removeprefix("\ndata: "))
            events.append(event)
            if event["event"] == "verified_answer":
                answer = event["answer"]
        else:
            answer += chunk
    return answer, events


class StubLLM:
    """Fixed output."""

    def __init__(self, answer: str) -> None:
        self.answer = answer
        self.calls = 0

    def generate(self, prompt: str, **kwargs: Any) -> str:
        self.calls += 1
        return self.answer


class StubRetriever:
    """Fixed retrieval."""

    def __init__(self, results: list[RetrievalResult]) -> None:
        self.results = results

    def retrieve(self, query: str, top_k: int = 5) -> list[RetrievalResult]:
        return self.results

    def retrieve_traced(self, query: str, top_k: int = 5) -> tuple[list[RetrievalResult], bool]:
        return self.results, False


class StubGrounding:
    """Fixed grounding."""

    def __init__(self, faithfulness: float, unsupported: list[str]) -> None:
        self.report = GroundingReport(
            faithfulness=faithfulness, sentences=[], unsupported=unsupported
        )

    def verify(self, answer: str, sources: list[Any]) -> GroundingReport:
        return self.report


class TestCitationValidation:
    def test_out_of_range_stripped_and_flagged(self) -> None:
        sources = [Source(index=1, chunk_id="c1", document_id="d1", content="a")]
        result = AnswerVerifier().verify("Claim one [1] and bogus [5].", sources)

        assert result.citations_valid == [1]
        assert result.citations_invalid == [5]
        assert "[5]" not in result.verified_answer
        assert "[1]" in result.verified_answer
        assert any("out-of-range" in w for w in result.warnings)

    def test_no_citations_warns(self) -> None:
        result = AnswerVerifier().verify("Plain answer.", [])
        assert result.citations_found == []
        assert any("no sources" in w for w in result.warnings)

    def test_citations_match_actual_source_indices(self) -> None:
        sources = [
            Source(index=2, chunk_id="c2", document_id="d2", content="alpha"),
            Source(index=4, chunk_id="c4", document_id="d4", content="beta"),
        ]

        result = AnswerVerifier().verify("Claim [1] [2] [4] [5].", sources)

        assert result.citations_valid == [2, 4]
        assert result.citations_invalid == [1, 5]
        assert "[2]" in result.verified_answer and "[4]" in result.verified_answer
        assert "[1]" not in result.verified_answer and "[5]" not in result.verified_answer

    @pytest.mark.usefixtures("integer_string_limit")
    def test_unparseable_citations_are_removed_without_inventing_indices(self) -> None:
        citation = "[" + "9" * 5000 + "]"
        sources = [Source(index=1, chunk_id="c1", document_id="d1", content="evidence")]

        result = AnswerVerifier().verify(
            f"Claim {citation} and {citation}. Known [1], invalid [9].", sources
        )

        assert result.verified_answer == "Claim and . Known [1], invalid ."
        assert result.citations_found == [1, 9]
        assert result.citations_valid == [1]
        assert result.citations_invalid == [9]
        assert "removed 2 unparseable numeric citation(s)" in result.warnings
        assert "removed 1 out-of-range citation(s): [9]" in result.warnings


class TestSynthesizerFormatting:
    def test_format_sources_numbered(self) -> None:
        synth = Synthesizer(StubLLM("unused"))  # type: ignore[arg-type]
        sources = [
            Source(index=1, chunk_id="c1", document_id="d1", content="alpha", score=0.8),
            Source(index=2, chunk_id="c2", document_id="d2", content="beta", score=0.6),
        ]
        out = synth._format_sources_for_llm(sources)

        assert "[1]" in out and "[2]" in out
        assert "alpha" in out and "beta" in out
        assert "relevance: 80%" in out

    def test_synthesize_uses_stub_llm(self) -> None:
        stub = StubLLM("Answer with citation [1].")
        synth = Synthesizer(stub)  # type: ignore[arg-type]
        out = synth.synthesize("q?", [_generation_result(1, "evidence")])

        assert stub.calls == 1
        assert out.answer == "Answer with citation [1]."
        assert len(out.sources) == 1
        assert out.sources[0].index == 1

    def test_token_budget_capped_by_config(self) -> None:
        five = [_generation_result(i, f"s{i}") for i in range(1, 6)]
        sources = Synthesizer(StubLLM("x"))._build_sources(five, {})  # type: ignore[arg-type]

        assert Synthesizer(StubLLM("x"), max_tokens=2048)._token_budget(sources) == 756  # type: ignore[arg-type]

        assert Synthesizer(StubLLM("x"), max_tokens=300)._token_budget(sources) == 300  # type: ignore[arg-type]

    @pytest.mark.parametrize("shortfall, expected_count", [(0, 2), (1, 1)])
    def test_context_budget_keeps_only_complete_blocks(
        self, shortfall: int, expected_count: int
    ) -> None:
        synth = Synthesizer(Mock(spec=LLMClient))
        results = [_generation_result(i, f"Passage {i}.") for i in range(1, 4)]
        results[1].chunk.context_before = "Nearby evidence."
        sources = synth._build_sources(results, {})
        synth.max_context_length = len(synth._format_sources_for_llm(sources[:2])) - shortfall

        prepared = synth.prepare("What did the study find?", results)

        assert prepared.error is None
        assert prepared.sources == sources[:expected_count]
        for source in prepared.sources:
            assert synth._format_sources_for_llm([source]) in prepared.prompt
        assert results[expected_count].chunk.content not in prepared.prompt
        assert "[truncated]" not in prepared.prompt

    def test_source_limit_applies_after_deduplication(self) -> None:
        first = _generation_result(1, "First passage.")
        synth = Synthesizer(Mock(spec=LLMClient), max_sources=2)

        prepared = synth.prepare(
            "What did the study find?",
            [
                first,
                first,
                _generation_result(2, "Second passage."),
                _generation_result(3, "Omitted passage."),
            ],
        )

        assert [source.chunk_id for source in prepared.sources] == ["c1", "c2"]
        assert [source.index for source in prepared.sources] == [1, 2]
        assert prepared.prompt.count("First passage.") == 1
        assert "Omitted passage." not in prepared.prompt

    def test_regular_and_streaming_use_same_retained_sources_and_token_budget(self) -> None:
        results = [_generation_result(i, f"Passage {i}: " + "evidence " * 145) for i in range(1, 6)]
        llm = Mock(spec=LLMClient)
        llm.generate.return_value = "Answer [2]."
        llm.stream.return_value = iter(["Answer ", "[2]."])
        synth = Synthesizer(llm)

        response = synth.synthesize("What did the study find?", results)
        streamed = "".join(synth.synthesize_streaming("What did the study find?", results))

        assert streamed == response.answer
        assert [source.chunk_id for source in response.sources] == ["c1", "c2", "c3", "c4"]
        assert llm.generate.call_args == llm.stream.call_args
        assert llm.generate.call_args.kwargs["max_tokens"] == 656
        prompt = llm.generate.call_args.args[0]
        assert results[4].chunk.content not in prompt
        assert all(source.content in prompt for source in response.sources)


class TestOrchestratorWiring:
    @pytest.mark.usefixtures("integer_string_limit")
    @pytest.mark.parametrize("streaming", [False, True])
    @pytest.mark.parametrize(
        ("digits", "valid", "warning"),
        [
            ("9" * 5000, False, "removed 1 unparseable numeric citation(s)"),
            ("0" * 5000 + "1", True, None),
            ("٠" * 5000 + "١", True, None),
            ("0" * 5000, False, "removed 1 out-of-range citation(s): [0]"),
            ("0" * 5000 + "99", False, "removed 1 out-of-range citation(s): [99]"),
        ],
        ids=["unparseable", "leading-zeros", "unicode-leading-zeros", "zero", "invalid-index"],
    )
    def test_long_citations_do_not_interrupt_the_final_answer(
        self, streaming: bool, digits: str, valid: bool, warning: str | None
    ) -> None:
        citation = f"[{digits}]"
        raw_answer = f"Claim {citation}. Known [1]."
        expected = raw_answer if valid else "Claim . Known [1]."

        def reply(request: httpx.Request) -> httpx.Response:
            if json.loads(request.content)["stream"]:
                records = [{"response": raw_answer}, {"done": True}]
                return httpx.Response(200, text="\n".join(json.dumps(row) for row in records))
            return httpx.Response(200, json={"response": raw_answer})

        with LLMClient() as llm:
            llm._client = httpx.Client(base_url=llm.base_url, transport=httpx.MockTransport(reply))
            orch = Orchestrator(StubRetriever([_generation_result(1, "evidence")]), llm)  # type: ignore[arg-type]

            if streaming:
                chunks = list(orch.query_streaming("What was measured?"))
                answer, events = _consume_stream(chunks)
                assert chunks[:-2] == [raw_answer]
                assert events[-1] == {"event": "verified_answer", "answer": expected}
                assert events[0]["citations_valid"] == ([1, 1] if valid else [1])
                warnings = events[0]["warnings"]
            else:
                response = orch.query("What was measured?")
                answer, warnings = response.answer, response.warnings
                assert [source.index for source in response.sources] == [1]

        assert answer == expected
        assert warnings == ([] if warning is None else [warning])

    def test_returns_rag_response(self) -> None:
        retriever = StubRetriever([_generation_result(1, "the sky is blue")])
        llm = StubLLM("The sky is blue [1].")
        orch = Orchestrator(retriever, llm)  # type: ignore[arg-type]

        resp = orch.query("why is the sky blue?")

        assert isinstance(resp, RAGResponse)
        assert resp.sources
        assert 0.0 <= resp.confidence <= 1.0
        assert "[1]" in resp.answer

    def test_invalid_citations_handled(self) -> None:

        retriever = StubRetriever([_generation_result(1, "only source")])
        llm = StubLLM("Grounded [1] but invalid [9].")
        orch = Orchestrator(retriever, llm)  # type: ignore[arg-type]

        resp = orch.query("q?")

        assert "[9]" not in resp.answer
        assert "[1]" in resp.answer
        assert any("out-of-range" in w for w in resp.warnings)

    def test_grounding_verifier_sets_faithfulness(self) -> None:
        retriever = StubRetriever([_generation_result(1, "source text")])
        llm = StubLLM("Answer [1].")
        grounding = StubGrounding(faithfulness=0.5, unsupported=["Answer [1]."])
        orch = Orchestrator(retriever, llm, grounding_verifier=grounding)  # type: ignore[arg-type]

        resp = orch.query("q?")

        assert resp.faithfulness == 0.5
        assert resp.confidence == 0.5
        assert any("not grounded" in w for w in resp.warnings)

    @pytest.mark.parametrize("grounding_enabled", [False, True])
    def test_failed_generation_keeps_sources_without_scoring_error_text(
        self, grounding_enabled: bool
    ) -> None:
        retriever = StubRetriever([_generation_result(1, "Retrieved evidence", score=1.0)])
        llm = Mock(spec=StubLLM)
        llm.generate.side_effect = LLMError("model unavailable")
        grounding = Mock(spec=StubGrounding)
        grounding.verify.return_value = GroundingReport(faithfulness=0.8, sentences=[])
        orch = Orchestrator(  # type: ignore[arg-type]
            retriever, llm, grounding_verifier=grounding if grounding_enabled else None
        )

        response = orch.query("What did the study find?")

        assert response.answer == "Unable to generate response. Please try again."
        assert response.confidence == 0.0
        assert response.faithfulness is None
        assert response.sources[0].content == "Retrieved evidence"
        assert response.warnings == ["Answer generation failed."]
        grounding.verify.assert_not_called()

    @pytest.mark.parametrize(
        ("tokens", "expected_answer", "expected_citations"),
        [
            (["Fact ", "[1]."], "Fact [1].", [1]),
            (["Claim ", "[99]."], "Claim .", []),
            (["Fact ", "[1] and [9]."], "Fact [1] and .", [1]),
        ],
    )
    def test_streaming_replaces_provisional_text_with_verified_answer(
        self, tokens: list[str], expected_answer: str, expected_citations: list[int]
    ) -> None:
        retriever = StubRetriever([_generation_result(1, "fact")])
        llm = Mock(spec=LLMClient)
        llm.stream.return_value = iter(tokens)
        orch = Orchestrator(retriever, llm)  # type: ignore[arg-type]

        chunks = list(orch.query_streaming("q?"))
        answer, events = _consume_stream(chunks)

        assert chunks[:-2] == tokens
        assert [event["event"] for event in events] == ["verification", "verified_answer"]
        assert events[0]["citations_valid"] == expected_citations
        assert events[1]["answer"] == expected_answer
        assert answer == expected_answer
        assert "[99]" not in answer and "[9]" not in answer

    def test_empty_retrieval_stream_returns_final_error_without_answer_checks(self) -> None:
        llm = Mock(spec=LLMClient)
        grounding = Mock(spec=StubGrounding)
        orch = Orchestrator(StubRetriever([]), llm, grounding_verifier=grounding)  # type: ignore[arg-type]
        orch.verifier = Mock(spec=AnswerVerifier)
        error = "No relevant documents found. Please upload documents and try again."

        chunks = list(orch.query_streaming("What did the study find?"))
        answer, events = _consume_stream(chunks)

        assert chunks[:-2] == [error]
        assert answer == error
        assert events[-1] == {"event": "verified_answer", "answer": error}
        assert events[0]["citations_valid"] == []
        assert events[0]["faithfulness"] is None
        assert events[0]["warnings"] == [error]
        llm.stream.assert_not_called()
        orch.verifier.verify.assert_not_called()
        grounding.verify.assert_not_called()

    def test_omitted_sources_cannot_be_cited_or_used_for_grounding(self) -> None:
        results = [_generation_result(i, f"Passage {i}: " + "evidence " * 145) for i in range(1, 6)]
        llm = Mock(spec=LLMClient)
        llm.generate.return_value = "Visible claim [1]. Omitted claim [5]."
        grounding = Mock(spec=StubGrounding)
        grounding.verify.return_value = GroundingReport(faithfulness=1.0, sentences=[])
        orch = Orchestrator(StubRetriever(results), llm, grounding_verifier=grounding)  # type: ignore[arg-type]

        response = orch.query("What did the study find?")

        assert [source.index for source in response.sources] == [1, 2, 3, 4]
        assert "[1]" in response.answer and "[5]" not in response.answer
        assert any("out-of-range" in warning for warning in response.warnings)
        assert grounding.verify.call_args.args == (response.answer, response.sources)
        prompt = llm.generate.call_args.args[0]
        assert all((source.full_content or source.content) in prompt for source in response.sources)
        assert results[4].chunk.content not in prompt

    @pytest.mark.parametrize("streaming", [False, True])
    def test_oversized_context_stops_before_generation_and_verification(
        self, streaming: bool
    ) -> None:
        first = _generation_result(1, "Short main passage.")
        first.chunk.context_before = "Nearby context. " * 500
        llm = Mock(spec=LLMClient)
        grounding = Mock(spec=StubGrounding)
        orch = Orchestrator(  # type: ignore[arg-type]
            StubRetriever([first, _generation_result(2, "A later passage that would fit.")]),
            llm,
            grounding_verifier=grounding,
        )
        error = "No source passages fit the context limit."

        if streaming:
            chunks = list(orch.query_streaming("What did the study find?"))
            answer, events = _consume_stream(chunks)
            event = events[0]
            assert chunks[:-2] == [error]
            assert answer == error
            assert events[-1] == {"event": "verified_answer", "answer": error}
            assert event["citations_valid"] == []
            assert event["faithfulness"] is None
            assert event["warnings"] == [error]
        else:
            response = orch.query("What did the study find?")
            assert response.answer == error
            assert response.sources == []
            assert response.confidence == 0.0
            assert response.faithfulness is None
            assert response.warnings == [error]

        llm.generate.assert_not_called()
        llm.stream.assert_not_called()
        grounding.verify.assert_not_called()

    def test_stream_verifies_the_source_snapshot_sent_to_the_llm(self) -> None:
        results = [_generation_result(i, f"Passage {i}: " + "evidence " * 145) for i in range(1, 6)]
        original_contents = [result.chunk.content for result in results[:4]]
        llm = Mock(spec=LLMClient)
        grounding = Mock(spec=StubGrounding)
        grounding.verify.return_value = GroundingReport(faithfulness=1.0, sentences=[])
        orch = Orchestrator(StubRetriever(results), llm, grounding_verifier=grounding)  # type: ignore[arg-type]

        def mutate_after_preparation(*args: Any, **kwargs: Any) -> Any:
            results[0].chunk.content = "Changed after prompt construction."
            results.clear()
            orch.synthesizer.max_sources = 0
            orch.synthesizer.max_context_length = 0
            yield "Visible claim [1]. Omitted claim [5]."

        llm.stream.side_effect = mutate_after_preparation

        chunks = list(orch.query_streaming("What did the study find?"))
        answer, events = _consume_stream(chunks)
        event = events[0]

        assert event["citations_valid"] == [1]
        assert any("out-of-range" in warning for warning in event["warnings"])
        verified_answer, sources = grounding.verify.call_args.args
        assert "[5]" not in verified_answer
        assert answer == verified_answer
        assert [source.content for source in sources] == original_contents
        assert [source.index for source in sources] == [1, 2, 3, 4]
        prompt = llm.stream.call_args.args[0]
        assert all(source.content in prompt for source in sources)

    @pytest.mark.parametrize("partial", [False, True])
    def test_failed_stream_reports_failure_without_grounding_partial_text(
        self, partial: bool
    ) -> None:
        def failed_stream(*args: Any, **kwargs: Any) -> Any:
            if partial:
                yield "Partial answer [1]."
            raise LLMError("model unavailable")

        llm = Mock(spec=LLMClient)
        llm.stream.side_effect = failed_stream
        grounding = Mock(spec=StubGrounding)
        grounding.verify.return_value = GroundingReport(faithfulness=0.8, sentences=[])
        orch = Orchestrator(  # type: ignore[arg-type]
            StubRetriever([_generation_result(1, "evidence")]), llm, grounding_verifier=grounding
        )

        chunks = list(orch.query_streaming("What did the study find?"))
        answer, events = _consume_stream(chunks)
        event = events[0]

        assert chunks[-3] == "Unable to generate response."
        assert "Partial answer" not in answer
        assert answer == "Unable to generate response."
        assert events[-1] == {"event": "verified_answer", "answer": answer}
        assert event["warnings"] == ["Answer generation failed."]
        assert event["faithfulness"] is None
        assert event["citations_valid"] == []
        grounding.verify.assert_not_called()


def _client_with(behaviors: list[httpx.Response | Exception]) -> tuple[LLMClient, Mock]:
    client = LLMClient(max_retries=2, backoff=0.0)
    for behavior in behaviors:
        if isinstance(behavior, httpx.Response):
            behavior.request = httpx.Request("POST", f"{client.base_url}/api/generate")
    fake = Mock(spec=httpx.Client)
    fake.post.side_effect = behaviors
    client._client = fake
    return client, fake


class TestLLMRetry:
    def test_retries_then_succeeds(self) -> None:
        client, fake = _client_with(
            [httpx.ConnectError("down"), httpx.Response(200, json={"response": "recovered"})]
        )
        assert client.generate("hi") == "recovered"
        assert fake.post.call_count == 2

    def test_raises_llmerror_after_exhaustion(self) -> None:
        client, fake = _client_with([httpx.ConnectError("down") for _ in range(3)])
        with pytest.raises(LLMError) as exc:
            client.generate("hi")

        assert "base_url" in str(exc.value)
        assert fake.post.call_count == 3

    def test_does_not_retry_on_4xx(self) -> None:
        client, fake = _client_with([httpx.Response(404)])
        with pytest.raises(httpx.HTTPStatusError):
            client.generate("hi")

        assert fake.post.call_count == 1

    def test_retries_on_5xx(self) -> None:
        client, fake = _client_with(
            [httpx.Response(503), httpx.Response(200, json={"response": "ok"})]
        )
        assert client.generate("hi") == "ok"
        assert fake.post.call_count == 2

    def test_missing_response_field_raises_llmerror(self) -> None:
        client, fake = _client_with([httpx.Response(200, json={"done": True})])
        with pytest.raises(LLMError, match="response"):
            client.generate("hi")
        assert fake.post.call_count == 1


@contextmanager
def _llm_response(response: httpx.Response, **options):
    with LLMClient(**options) as client:
        client._client = httpx.Client(
            base_url=client.base_url, transport=httpx.MockTransport(lambda request: response)
        )
        yield client


class TestLLMResponseValidation:
    @pytest.mark.parametrize("answer", ["", " \n\t"])
    def test_generate_rejects_empty_answers(self, answer: str) -> None:
        with _llm_response(httpx.Response(200, json={"response": answer})) as client:
            with pytest.raises(LLMError, match="empty response"):
                client.generate("question")

    @pytest.mark.parametrize("payload", [[], None, {"response": None}, {"response": 42}])
    def test_generate_rejects_invalid_response_fields(self, payload: Any) -> None:
        with _llm_response(httpx.Response(200, text=json.dumps(payload))) as client:
            with pytest.raises(LLMError, match="string 'response' field"):
                client.generate("question")

    @pytest.mark.parametrize(
        ("body", "error"),
        [
            ("[]", "JSON object"),
            ('{"response": null}', "must be a string"),
            ('{"response": 42}', "must be a string"),
            ('{"error": "model was unloaded"}', "backend reported an error"),
            ('{"done": "false"}', "must be a boolean"),
            ("invalid JSON", "stream failed"),
        ],
    )
    def test_stream_rejects_invalid_records(self, body: str, error: str) -> None:
        with _llm_response(httpx.Response(200, text=body)) as client:
            with pytest.raises(LLMError, match=error):
                list(client.stream("question"))

    @pytest.mark.parametrize("body", ["", '{"response": "Partial answer"}\n'])
    def test_stream_requires_completion_record(self, body: str) -> None:
        with _llm_response(httpx.Response(200, text=body)) as client:
            with pytest.raises(LLMError, match="before completion"):
                list(client.stream("question"))

    def test_stream_accepts_completed_answer(self) -> None:
        body = '{"response": "Complete answer"}\n{"done": true}\n'
        with _llm_response(httpx.Response(200, text=body)) as client:
            assert list(client.stream("question")) == ["Complete answer"]

    @pytest.mark.parametrize("answer", ["", " \n\t"])
    def test_stream_rejects_empty_completed_answers(self, answer: str) -> None:
        with _llm_response(httpx.Response(200, json={"response": answer, "done": True})) as client:
            with pytest.raises(LLMError, match="empty response"):
                list(client.stream("question"))

    def test_stream_surfaces_an_error_after_partial_output(self) -> None:
        body = '{"response": "Partial answer"}\n{"error": "model was unloaded"}\n'
        with _llm_response(httpx.Response(200, text=body)) as client:
            tokens = client.stream("question")
            assert next(tokens) == "Partial answer"
            with pytest.raises(LLMError, match="backend reported an error"):
                next(tokens)

    @pytest.mark.parametrize("payload", [[], None])
    def test_readiness_rejects_non_object_json(self, payload: Any) -> None:
        with _llm_response(httpx.Response(200, text=json.dumps(payload))) as client:
            assert client.is_available() is False


@pytest.mark.parametrize(
    ("configured", "installed", "available"),
    [
        ("llama3.2:3b", "llama3.2:3b", True),
        ("llama3.2:3b", "llama3.2:3b-instruct", False),
        ("llama3.2", "llama3.2:3b", False),
        ("llama3.2", "llama3.2:latest", True),
        ("localhost:5000/model", "localhost:5000/model:latest", True),
    ],
)
def test_model_readiness_matches_exact_tag(configured, installed, available):
    with _llm_response(
        httpx.Response(200, json={"models": [{"name": installed}]}), model=configured
    ) as client:
        assert client.is_available() is available


def _fusion_result(chunk_id):
    return RetrievalResult(Chunk(id=chunk_id, content=chunk_id, document_id="paper"), 0.5)


@pytest.mark.parametrize(
    "weight", [-1, -0.1, float("nan"), float("inf"), True, "0.5", 10**1000, -(10**1000)]
)
def test_invalid_fusion_weights_are_rejected(weight):
    with pytest.raises(ValueError, match="weights"):
        rrf_fuse([[_fusion_result("first")]], [weight])
    with pytest.raises(ValueError, match="weights"):
        HybridRetriever(Mock(), Mock(), dense_weight=weight)


@pytest.mark.parametrize("k", [-1, True, 0.5, float("nan")])
def test_invalid_rrf_constants_are_rejected(k):
    with pytest.raises(ValueError, match="RRF k"):
        rrf_fuse([[_fusion_result("first")]], [1.0], k=k)
    with pytest.raises(ValueError, match="RRF k"):
        HybridRetriever(Mock(), Mock(), rrf_k=k)


def test_fusion_requires_one_weight_per_list():
    with pytest.raises(ValueError, match="one fusion weight"):
        rrf_fuse([[_fusion_result("first")]], [])


@pytest.mark.parametrize("top_k", [True, 1.5, "2"])
def test_fusion_rejects_noninteger_limits(top_k):
    with pytest.raises(ValueError, match="top_k"):
        rrf_fuse([[_fusion_result("first")]], [1.0], top_k=top_k)


@pytest.mark.parametrize("top_k", [0, -1])
def test_nonpositive_fusion_limits_return_no_results(top_k):
    assert rrf_fuse([[_fusion_result("first"), _fusion_result("second")]], [1.0], top_k=top_k) == []


def test_large_finite_weights_preserve_normalized_ranking():
    ranked = [[_fusion_result("first"), _fusion_result("second")]] * 2

    results = rrf_fuse(ranked, [1e308, 1e308], k=0)

    assert [result.chunk.id for result in results] == ["first", "second"]
    assert [result.score for result in results] == [1.0, 0.5]
    assert all(math.isfinite(result.score) for result in results)


def test_zero_rrf_constant_uses_reciprocal_rank():
    results = rrf_fuse([[_fusion_result("first"), _fusion_result("second")]], [1.0], k=0)

    assert [result.score for result in results] == [1.0, 0.5]


def test_duplicate_chunk_does_not_gain_extra_rank_votes():
    first = _fusion_result("first")
    second = _fusion_result("second")

    clean = rrf_fuse([[first, second], [second, first]], [0.5, 0.5])
    repeated = rrf_fuse([[first, first, second], [second, first]], [0.5, 0.5])

    assert repeated == clean


def test_first_occurrence_supplies_chunk_content_when_results_repeat():
    first = _fusion_result("first")
    conflicting = _fusion_result("first")
    conflicting.chunk.content = "Conflicting repeated content."

    result = rrf_fuse([[first, conflicting]], [1.0])[0]

    assert result.chunk is first.chunk


@pytest.mark.parametrize("query", ["", " \t\n"])
def test_empty_hybrid_queries_do_not_call_backends(query):
    dense, sparse = Mock(), Mock()
    retriever = HybridRetriever(dense, sparse)

    assert retriever.retrieve_with_dense_top(query) == ([], 0.0)
    assert retriever.retrieve(query) == []
    dense.retrieve.assert_not_called()
    sparse.retrieve.assert_not_called()


@pytest.mark.parametrize("top_k", [0, -1])
def test_nonpositive_hybrid_limit_does_not_call_backends(top_k):
    dense, sparse = Mock(), Mock()

    assert HybridRetriever(dense, sparse).retrieve_with_dense_top("query", top_k=top_k) == (
        [],
        0.0,
    )
    dense.retrieve.assert_not_called()
    sparse.retrieve.assert_not_called()


@pytest.mark.parametrize(
    "options", [{"query": None}, {"top_k": True}, {"top_k": 1.5}, {"depth": 0}, {"depth": 1.5}]
)
def test_invalid_hybrid_requests_do_not_call_backends(options):
    dense, sparse = Mock(), Mock()
    retriever = HybridRetriever(dense, sparse)

    with pytest.raises(ValueError):
        retriever.retrieve(**{"query": "query", **options})

    dense.retrieve.assert_not_called()
    sparse.retrieve.assert_not_called()


def test_explicit_candidate_depth_and_query_text_are_preserved():
    dense, sparse = Mock(), Mock()
    dense.retrieve.return_value = [_fusion_result("dense")]
    sparse.retrieve.return_value = [_fusion_result("sparse")]

    results = HybridRetriever(dense, sparse).retrieve(" query ", top_k=10, depth=1)

    assert len(results) == 2
    dense.retrieve.assert_called_once_with(" query ", 1)
    sparse.retrieve.assert_called_once_with(" query ", 1)


@pytest.mark.parametrize(
    "shift", [-0.1, float("nan"), float("inf"), True, "0.2", 10**1000, -(10**1000)]
)
def test_invalid_adaptive_shifts_fail_during_construction(shift):
    with pytest.raises(ValueError, match="shift"):
        AdaptiveHybridRetriever(Mock(), Mock(), shift=shift)


@pytest.mark.parametrize(
    "query,weights",
    [
        ("kinase inhibition", (0.3, 0.7)),
        ("how does the treatment change the observed rate of recovery over time", (0.7, 0.3)),
        ("kinase inhibition in human cells", (0.5, 0.5)),
    ],
)
def test_default_adaptive_weight_rules_remain_unchanged(query, weights):
    retriever = AdaptiveHybridRetriever(Mock(), Mock())

    assert retriever._weights(query) == weights


def test_subnormal_positive_weights_retain_their_relative_influence():
    ranked = [
        [_fusion_result("first"), _fusion_result("second")],
        [_fusion_result("second"), _fusion_result("first")],
    ]

    results = rrf_fuse(ranked, [5e-324, 1e-323])

    assert results == rrf_fuse(ranked, [0.5, 1.0])
    assert results[0].chunk.id == "second"


@pytest.mark.parametrize(
    "payload,message",
    [
        ({"response": "Partial answer", "done": False}, "before completion"),
        ({"response": "Partial answer", "done": "true"}, "must be a boolean"),
        ({"response": "Partial answer", "done": 1}, "must be a boolean"),
        (
            {"response": "Partial answer", "error": "model was unloaded"},
            "backend reported an error",
        ),
    ],
)
def test_generate_rejects_incomplete_or_failed_responses(payload, message):
    requests = []

    def respond(request):
        requests.append(request)
        return httpx.Response(200, json=payload)

    with LLMClient() as llm:
        llm._client = httpx.Client(base_url=llm.base_url, transport=httpx.MockTransport(respond))
        with pytest.raises(LLMError, match=message):
            llm.generate("Question")

    assert len(requests) == 1


@pytest.mark.parametrize("done", [True, None])
def test_generate_keeps_complete_response_and_legacy_metadata_contract(done):
    payload = {"response": "Answer.\n"}
    if done is not None:
        payload["done"] = done
    transport = httpx.MockTransport(lambda request: httpx.Response(200, json=payload))
    with LLMClient() as llm:
        llm._client = httpx.Client(base_url=llm.base_url, transport=transport)
        assert llm.generate("Question") == "Answer.\n"


def test_generate_wraps_invalid_unicode_body_as_backend_error():
    transport = httpx.MockTransport(lambda request: httpx.Response(200, content=b"\xff"))
    with LLMClient() as llm:
        llm._client = httpx.Client(base_url=llm.base_url, transport=transport)
        with pytest.raises(LLMError, match="non-JSON body"):
            llm.generate("Question")


def test_model_readiness_rejects_invalid_unicode_body():
    transport = httpx.MockTransport(lambda request: httpx.Response(200, content=b"\xff"))
    with LLMClient() as llm:
        llm._client = httpx.Client(base_url=llm.base_url, transport=transport)
        assert llm.is_available() is False


def test_generate_and_stream_share_generation_options():
    requests = []

    def respond(request):
        payload = json.loads(request.content)
        requests.append(payload)
        return httpx.Response(200, json={"response": "Answer.", "done": True})

    with LLMClient(model="custom-model") as llm:
        llm._client = httpx.Client(base_url=llm.base_url, transport=httpx.MockTransport(respond))
        assert (
            llm.generate("Question", system="Instructions", temperature=0.2, max_tokens=17)
            == "Answer."
        )
        assert list(
            llm.stream("Question", system="Instructions", temperature=0.2, max_tokens=17)
        ) == ["Answer."]

    assert [request.pop("stream") for request in requests] == [False, True]
    assert (
        requests[0]
        == requests[1]
        == {
            "model": "custom-model",
            "prompt": "Question",
            "system": "Instructions",
            "options": {"temperature": 0.2, "num_predict": 17, "top_p": 0.9},
        }
    )


def test_generate_preserves_stop_sequences():
    requests = []

    def respond(request):
        requests.append(json.loads(request.content))
        return httpx.Response(200, json={"response": "Answer.", "done": True})

    with LLMClient() as llm:
        llm._client = httpx.Client(base_url=llm.base_url, transport=httpx.MockTransport(respond))
        assert llm.generate("Question", stop=["END", "STOP"]) == "Answer."

    assert requests[0]["options"]["stop"] == ["END", "STOP"]


@pytest.mark.parametrize(
    "options",
    [
        {"max_retries": -1},
        {"max_retries": 1.5},
        {"max_retries": True},
        {"timeout": 0},
        {"timeout": float("nan")},
        {"timeout": float("inf")},
        {"timeout": True},
        {"timeout": False},
        {"timeout": 10**1000},
        {"backoff": -0.1},
        {"backoff": float("nan")},
        {"backoff": float("inf")},
        {"backoff": True},
        {"backoff": False},
        {"backoff": 10**1000},
    ],
)
def test_invalid_client_configuration_fails_before_http_creation(monkeypatch, options):
    factory = Mock()
    monkeypatch.setattr("scinexusrag.generation.llm.httpx.Client", factory)

    with pytest.raises(ValueError):
        LLMClient(**options)

    factory.assert_not_called()


class _FailingCloseTransport(httpx.MockTransport):
    def close(self):
        raise OSError("Transport cleanup failed")


def test_close_failure_does_not_reuse_closed_http_client(monkeypatch):
    with LLMClient() as llm:
        original = httpx.Client(
            base_url=llm.base_url,
            transport=_FailingCloseTransport(lambda request: httpx.Response(200)),
        )
        llm._client = original
        with pytest.raises(OSError, match="Transport cleanup failed"):
            llm.close()

        assert original.is_closed
        assert llm._client is None
        replacement = httpx.Client(
            base_url=llm.base_url,
            transport=httpx.MockTransport(
                lambda request: httpx.Response(200, json={"response": "Recovered.", "done": True})
            ),
        )
        factory = Mock(return_value=replacement)
        monkeypatch.setattr("scinexusrag.generation.llm.httpx.Client", factory)

        assert llm.generate("Question") == "Recovered."
        factory.assert_called_once()

    assert replacement.is_closed


def test_finalizer_does_not_report_transport_cleanup_errors(monkeypatch):
    errors = []
    monkeypatch.setattr(sys, "unraisablehook", errors.append)
    llm = LLMClient()
    llm._client = httpx.Client(
        base_url=llm.base_url,
        transport=_FailingCloseTransport(lambda request: httpx.Response(200)),
    )

    del llm
    gc.collect()

    assert errors == []


@pytest.fixture
def private_backend():
    secrets = (
        "fixture-password",
        "fixture-path-secret",
        "fixture-query-secret",
        "fixture-fragment",
    )
    password, path, query, fragment = secrets
    base_url = f"http://fixture-user:{password}@localhost:11434/{path}?token={query}#{fragment}"
    return base_url, (*secrets, "fixture-user")


def _failed_backend(base_url, failure):
    def respond(request):
        if failure == "timeout":
            raise httpx.ReadTimeout(base_url, request=request)
        if failure == "connect":
            raise httpx.ConnectError(base_url, request=request)
        if failure.startswith("http_"):
            return httpx.Response(int(failure[5:]), text=base_url)
        if failure == "unicode":
            return httpx.Response(200, content=base_url.encode() + b"\xff")
        if failure == "json":
            return httpx.Response(200, text=base_url)
        return httpx.Response(200, json={"error": base_url})

    client = LLMClient(base_url=base_url, max_retries=0)
    client._client = httpx.Client(base_url=base_url, transport=httpx.MockTransport(respond))
    return client


@pytest.mark.parametrize(
    "failure", ["timeout", "connect", "http_400", "http_503", "unicode", "json", "backend"]
)
@pytest.mark.parametrize("streaming", [False, True])
def test_backend_errors_and_tracebacks_omit_private_url_components(
    private_backend, failure, streaming
):
    base_url, secrets = private_backend
    with (
        _failed_backend(base_url, failure) as client,
        pytest.raises((LLMError, httpx.HTTPStatusError)) as raised,
    ):
        if streaming:
            list(client.stream("Question"))
        else:
            client.generate("Question")
    error = raised.value
    diagnostic = str(error) + repr(error) + "".join(traceback.format_exception(error))
    assert "base_url=http://localhost:11434" in diagnostic
    assert all(secret not in diagnostic for secret in secrets)
    assert error.__cause__ is None
    if failure.startswith("http_"):
        assert f"HTTP {failure[5:]}" in str(error)


@pytest.mark.parametrize(
    "failure", ["timeout", "connect", "http_400", "http_503", "unicode", "json", "backend"]
)
@pytest.mark.parametrize("streaming", [False, True])
def test_generation_logs_omit_private_backend_url(private_backend, failure, streaming, caplog):
    base_url, secrets = private_backend
    results = [RetrievalResult(Chunk(id="passage", document_id="paper", content="A finding."), 0.8)]
    with _failed_backend(base_url, failure) as client, caplog.at_level(logging.ERROR):
        if streaming:
            retriever = Mock()
            retriever.retrieve.return_value = results
            output = list(Orchestrator(retriever, client).query_streaming("Question"))
            assert "Unable to generate response." in "".join(output)
        else:
            result = Synthesizer(client).synthesize("Question", results)
            assert result.error == "Answer generation failed."
    assert caplog.records
    assert "base_url=http://localhost:11434" in caplog.text
    assert all(secret not in caplog.text for secret in secrets)


class TestMetricsCollector:
    def test_initial_state(self) -> None:
        collector = MetricsCollector()
        snap = collector.snapshot()

        assert snap["total_queries"] == 0
        assert snap["total_ingestions"] == 0
        assert snap["total_deletions"] == 0
        assert snap["avg_query_time_ms"] == 0.0
        assert snap["uptime_seconds"] >= 0
        assert set(snap) == {
            "total_queries",
            "total_ingestions",
            "total_deletions",
            "avg_query_time_ms",
            "uptime_seconds",
        }
        for key in ("total_queries", "total_ingestions", "total_deletions"):
            assert type(snap[key]) is int
        for key in ("avg_query_time_ms", "uptime_seconds"):
            assert type(snap[key]) is float

    def test_record_query(self) -> None:
        collector = MetricsCollector()
        collector.record_query(100.0)
        collector.record_query(200.0)

        snap = collector.snapshot()
        assert snap["total_queries"] == 2
        assert snap["avg_query_time_ms"] == 150.0

    def test_record_ingest(self) -> None:
        collector = MetricsCollector()
        collector.record_ingest()
        collector.record_ingest()
        collector.record_ingest()

        snap = collector.snapshot()
        assert snap["total_ingestions"] == 3

    def test_record_delete(self) -> None:
        collector = MetricsCollector()
        collector.record_delete()

        snap = collector.snapshot()
        assert snap["total_deletions"] == 1

    def test_zero_duration_is_recorded(self) -> None:
        collector = MetricsCollector()
        collector.record_query(0)
        snap = collector.snapshot()
        assert snap["total_queries"] == 1
        assert snap["avg_query_time_ms"] == 0.0

    @pytest.mark.parametrize(
        "duration,error",
        [
            (True, TypeError),
            (False, TypeError),
            ("12.5", TypeError),
            (None, TypeError),
            (complex(1, 1), TypeError),
            (float("nan"), ValueError),
            (float("inf"), ValueError),
            (float("-inf"), ValueError),
            (-1.0, ValueError),
            (10**1000, ValueError),
        ],
    )
    def test_invalid_duration_preserves_statistics(self, duration, error):
        collector = MetricsCollector()
        collector.record_query(100.0)

        with pytest.raises(error):
            collector.record_query(duration)

        snap = collector.snapshot()
        assert snap["total_queries"] == 1
        assert snap["avg_query_time_ms"] == 100.0
        collector.record_query(200)
        assert collector.snapshot()["avg_query_time_ms"] == 150.0

    @pytest.mark.parametrize(
        "durations,expected",
        [
            ([sys.float_info.max] * 4, sys.float_info.max),
            ([0.0, sys.float_info.max], sys.float_info.max / 2),
            ([sys.float_info.max, 0.0], sys.float_info.max / 2),
        ],
    )
    def test_finite_durations_do_not_overflow_average(self, durations, expected):
        collector = MetricsCollector()
        for duration in durations:
            collector.record_query(duration)
        snap = collector.snapshot()
        assert snap["total_queries"] == len(durations)
        assert math.isfinite(snap["avg_query_time_ms"])
        assert snap["avg_query_time_ms"] == pytest.approx(expected, rel=1e-15)

    def test_varied_durations_match_arithmetic_mean(self):
        collector = MetricsCollector()
        durations = [0.0, 1.25, 200, 5.5, 0.125, 750.0, 12.75] * 20
        for duration in durations:
            collector.record_query(duration)
        snap = collector.snapshot()
        assert snap["total_queries"] == len(durations)
        assert snap["avg_query_time_ms"] == pytest.approx(
            math.fsum(durations) / len(durations), rel=1e-14
        )

    def test_snapshots_are_detached(self):
        collector = MetricsCollector()
        before = collector.snapshot()
        before["total_queries"] = 999
        collector.record_query(25.0)
        after = collector.snapshot()

        assert before["avg_query_time_ms"] == 0.0
        assert after["total_queries"] == 1
        assert after["avg_query_time_ms"] == 25.0
        after["avg_query_time_ms"] = -1.0
        assert collector.snapshot()["avg_query_time_ms"] == 25.0

    def test_overlapping_writers_and_snapshot_readers(self):
        collector = MetricsCollector()
        writers, readers, iterations = 4, 2, 200
        start = threading.Barrier(writers + readers, timeout=5)

        def write():
            start.wait()
            for _ in range(iterations):
                collector.record_query(12.5)
                collector.record_ingest()
                collector.record_delete()

        def read():
            start.wait()
            previous = dict.fromkeys(("total_queries", "total_ingestions", "total_deletions"), 0)
            for _ in range(iterations):
                snap = collector.snapshot()
                for key in previous:
                    assert previous[key] <= snap[key] <= writers * iterations
                    previous[key] = snap[key]
                assert snap["avg_query_time_ms"] == (12.5 if snap["total_queries"] else 0.0)

        with ThreadPoolExecutor(max_workers=writers + readers) as pool:
            jobs = [pool.submit(write) for _ in range(writers)]
            jobs.extend(pool.submit(read) for _ in range(readers))
            for job in jobs:
                job.result(timeout=10)

        snap = collector.snapshot()
        assert snap["total_queries"] == writers * iterations
        assert snap["total_ingestions"] == writers * iterations
        assert snap["total_deletions"] == writers * iterations
        assert snap["avg_query_time_ms"] == 12.5


def test_concurrent_first_access_constructs_one_singleton(monkeypatch):
    workers = 4
    start = threading.Barrier(workers + 1, timeout=5)
    entered = threading.Event()
    release = threading.Event()
    construction = []

    def create_collector():
        construction.append(1)
        entered.set()
        assert release.wait(timeout=5)
        return MetricsCollector()

    monkeypatch.setattr(metrics, "_collector", None)
    monkeypatch.setattr(metrics, "MetricsCollector", create_collector)

    def get_collector():
        start.wait()
        return metrics.get_metrics_collector()

    with ThreadPoolExecutor(max_workers=workers) as pool:
        jobs = [pool.submit(get_collector) for _ in range(workers)]
        start.wait()
        try:
            assert entered.wait(timeout=5)
        finally:
            release.set()
        collectors = [job.result(timeout=10) for job in jobs]

    assert len(construction) == 1
    assert all(collector is collectors[0] for collector in collectors)
    collectors[0].record_query(4.0)
    assert metrics.get_metrics_collector().snapshot()["total_queries"] == 1


def test_metrics_http_response_preserves_counts_types_and_rounding(monkeypatch, local_settings):
    collector = MetricsCollector()
    collector.record_query(12.34)
    collector.record_query(56.78)
    collector.record_ingest()
    collector.record_delete()
    statistics = SystemStats(
        total_documents=3,
        total_chunks=12,
        total_words=1200,
        storage_path="unused",
        llm_model="fixture-generator",
        embedding_model="fixture-embedder",
        llm_available=False,
    )
    monkeypatch.setattr(routes, "get_metrics_collector", lambda: collector)
    monkeypatch.setattr(
        routes, "get_scinexusrag", lambda: SimpleNamespace(get_stats=lambda: statistics)
    )

    with TestClient(api.create_app(), base_url="http://localhost") as client:
        response = client.get("/api/metrics")

    assert response.status_code == 200
    data = response.json()
    expected = {
        "total_queries": 2,
        "total_ingestions": 1,
        "total_deletions": 1,
        "avg_query_time_ms": 34.6,
        "total_documents": 3,
        "total_chunks": 12,
        "version": __version__,
    }
    assert set(data) == {"uptime_seconds", *expected}
    for key, value in expected.items():
        assert data[key] == value
        assert type(data[key]) is type(value)
    assert type(data["uptime_seconds"]) is float
    assert data["uptime_seconds"] >= 0


@pytest.fixture
def reranker():
    instance = Reranker()
    instance._model = Mock()
    return instance


def test_concurrent_first_reranker_load_constructs_once(monkeypatch, observed_lock):
    reranker = Reranker()
    lock = observed_lock
    reranker._model_lock = lock
    start = Barrier(3)
    constructing, release = Event(), Event()
    model = object()

    def construct(*_args, **_kwargs):
        constructing.set()
        assert release.wait(timeout=5)
        return model

    factory = Mock(side_effect=construct)
    monkeypatch.setitem(sys.modules, "sentence_transformers", SimpleNamespace(CrossEncoder=factory))

    def load():
        start.wait(timeout=5)
        return reranker.model

    with ThreadPoolExecutor(max_workers=3) as workers:
        futures = [workers.submit(load) for _ in range(3)]
        try:
            assert constructing.wait(timeout=5)
            assert len({lock.attempts.get(timeout=5) for _ in range(3)}) == 3
        finally:
            release.set()
        assert all(future.result(timeout=5) is model for future in futures)

    factory.assert_called_once_with(reranker.model_name, device=None, revision=reranker.revision)


def test_failed_reranker_load_can_retry(monkeypatch):
    model = object()
    factory = Mock(side_effect=[RuntimeError("load failed"), model])
    monkeypatch.setitem(sys.modules, "sentence_transformers", SimpleNamespace(CrossEncoder=factory))
    reranker = Reranker()

    with pytest.raises(RuntimeError, match="load failed"):
        _ = reranker.model

    assert reranker._model is None
    assert reranker.model is model
    assert reranker.model is model
    assert factory.call_count == 2


@pytest.mark.parametrize("stage", ["to", "eval"])
def test_failed_splade_initialization_does_not_cache_partial_model(monkeypatch, stage):
    tokenizer = Mock()
    model = Mock()
    model.to.return_value = model
    getattr(model, stage).side_effect = [RuntimeError("model failed"), model]
    tokenizer_loader = Mock(return_value=tokenizer)
    model_loader = Mock(return_value=model)
    monkeypatch.setitem(
        sys.modules,
        "transformers",
        SimpleNamespace(
            AutoTokenizer=SimpleNamespace(from_pretrained=tokenizer_loader),
            AutoModelForMaskedLM=SimpleNamespace(from_pretrained=model_loader),
        ),
    )
    monkeypatch.setattr(SpladeRetriever, "_encode", lambda *_: sparse.csr_matrix((0, 0)))
    retriever = SpladeRetriever([])

    with pytest.raises(RuntimeError, match="model failed"):
        retriever._load()

    assert retriever._model is None
    assert retriever._tok is None
    retriever._load()
    retriever._load()
    assert retriever._model is model
    assert retriever._tok is tokenizer
    assert model_loader.call_count == 2


def candidates(count=3):
    return [
        RetrievalResult(Chunk(str(index), f"passage {index}", "doc"), 0.5, "hybrid")
        for index in range(count)
    ]


@pytest.mark.parametrize(
    "options",
    [
        {"model_name": ""},
        {"model_name": None},
        {"device": "automatic"},
        {"batch_size": 0},
        {"batch_size": False},
        {"batch_size": 1.5},
        {"revision": ""},
        {"revision": 1},
    ],
)
def test_invalid_reranker_settings_fail_without_load(monkeypatch, options):
    factory = Mock()
    monkeypatch.setitem(sys.modules, "sentence_transformers", SimpleNamespace(CrossEncoder=factory))

    with pytest.raises(ValueError):
        Reranker(**options)

    factory.assert_not_called()


@pytest.mark.parametrize("top_k", [0, -1, -99])
def test_nonpositive_reranking_depth_skips_model(reranker, top_k):
    assert reranker.rerank("query", candidates(), top_k=top_k) == []
    reranker._model.predict.assert_not_called()


@pytest.mark.parametrize("top_k", [True, False, 1.5, "1"])
def test_invalid_reranking_depth_fails_without_model(reranker, top_k):
    with pytest.raises(ValueError, match="top_k"):
        reranker.rerank("query", candidates(), top_k=top_k)

    reranker._model.predict.assert_not_called()


@pytest.mark.parametrize("query", ["", " \n\t"])
def test_blank_reranking_query_skips_model(reranker, query):
    assert reranker.rerank(query, candidates()) == []
    reranker._model.predict.assert_not_called()


@pytest.mark.parametrize("query", [None, 1, ["question"]])
def test_nontext_reranking_query_fails_without_model(reranker, query):
    with pytest.raises(ValueError, match="query"):
        reranker.rerank(query, candidates())

    reranker._model.predict.assert_not_called()


@pytest.mark.parametrize(
    "scores",
    [
        [0.1, 0.2],
        [0.1, np.nan, 0.3],
        [0.1, np.inf, 0.3],
        [[0.1, 0.2]] * 3,
        ["0.1", "0.2", "0.3"],
        [1j, 2j, 3j],
    ],
)
def test_malformed_reranker_scores_are_rejected(scores):
    reranker = Reranker()
    reranker._model = Mock(predict=Mock(return_value=scores))

    with pytest.raises(ValueError, match="reranker"):
        reranker.rerank("query", candidates())


def test_column_reranker_scores_have_scalar_outputs_and_stable_ties():
    reranker = Reranker(batch_size=2)
    reranker._model = Mock(predict=Mock(return_value=np.array([[0.2], [0.8], [0.8]])))
    results = candidates()

    reranked = reranker.rerank("query", results)

    assert [result.chunk.id for result in reranked] == ["1", "2", "0"]
    assert [result.score for result in reranked] == [1.0, 1.0, 0.0]
    assert all(type(result.score) is float for result in reranked)
    assert all(result.source == "hybrid+rerank" for result in reranked)
    assert [result.score for result in results] == [0.5] * 3
    reranker._model.predict.assert_called_once_with(
        [("query", f"passage {index}") for index in range(3)], batch_size=2
    )


def test_large_finite_reranker_scores_normalize_without_overflow():
    reranker = Reranker()
    reranker._model = Mock(predict=Mock(return_value=np.array([-1e308, 0, 1e308])))

    with np.errstate(all="raise"):
        reranked = reranker.rerank("query", candidates())

    assert [result.score for result in reranked] == [1.0, 0.5, 0.0]


def test_splade_index_keeps_a_snapshot_of_caller_chunks(monkeypatch):
    encode = Mock(
        side_effect=[
            sparse.csr_matrix([[1.0], [3.0]]),
            sparse.csr_matrix([[2.0]]),
            sparse.csr_matrix([[2.0]]),
        ]
    )
    monkeypatch.setattr(SpladeRetriever, "_encode", encode)
    chunks = [
        Chunk("first", "first passage", "first-doc"),
        Chunk("second", "second passage", "second-doc", metadata={"name": "original"}),
    ]
    retriever = SpladeRetriever(chunks)
    chunks[1].content = "replacement passage"
    chunks[1].metadata["name"] = "replacement"
    chunks.append(Chunk("third", "third passage", "third-doc"))

    results = retriever.retrieve("question")

    assert [result.chunk.id for result in results] == ["second", "first"]
    assert results[0].chunk.content == "second passage"
    assert results[0].chunk.metadata["name"] == "original"
    assert results[0].score == 6.0
    assert encode.call_args_list[0].args == (["first passage", "second passage"],)
    results[0].chunk.content = "changed result"
    results[0].chunk.metadata["name"] = "changed result"

    repeated = retriever.retrieve("question")

    assert repeated[0].chunk.content == "second passage"
    assert repeated[0].chunk.metadata["name"] == "original"


@pytest.mark.parametrize(
    "options",
    [
        {"model_name": ""},
        {"model_name": None},
        {"device": "automatic"},
        {"batch_size": False},
        {"batch_size": 1.5},
        {"max_length": True},
        {"max_length": 1.5},
        {"revision": ""},
        {"revision": 1},
        {"chunks": None},
        {"chunks": ["passage"]},
        {"chunks": [Chunk("doc", None, "doc")]},
    ],
)
def test_invalid_splade_settings_fail_before_encoding(monkeypatch, options):
    encode = Mock()
    monkeypatch.setattr(SpladeRetriever, "_encode", encode)
    settings = {"chunks": [], **options}

    with pytest.raises(ValueError):
        SpladeRetriever(**settings)

    encode.assert_not_called()


@pytest.mark.parametrize("top_k", [True, False, 1.5, "1", None])
def test_invalid_splade_depth_fails_without_encoding(monkeypatch, top_k):
    encode = Mock(return_value=sparse.csr_matrix([[1.0]]))
    monkeypatch.setattr(SpladeRetriever, "_encode", encode)
    retriever = SpladeRetriever([Chunk("doc", "passage", "doc")])
    encode.reset_mock()

    with pytest.raises(ValueError, match="top_k"):
        retriever.retrieve("question", top_k=top_k)

    encode.assert_not_called()


@pytest.mark.parametrize("query", ["", " \n\t"])
def test_blank_splade_query_skips_encoding(monkeypatch, query):
    encode = Mock(return_value=sparse.csr_matrix([[1.0]]))
    monkeypatch.setattr(SpladeRetriever, "_encode", encode)
    retriever = SpladeRetriever([Chunk("doc", "passage", "doc")])
    encode.reset_mock()

    assert retriever.retrieve(query) == []
    encode.assert_not_called()


@pytest.mark.parametrize("query", [None, 1, ["question"]])
def test_nontext_splade_query_fails_without_encoding(monkeypatch, query):
    encode = Mock(return_value=sparse.csr_matrix([[1.0]]))
    monkeypatch.setattr(SpladeRetriever, "_encode", encode)
    retriever = SpladeRetriever([Chunk("doc", "passage", "doc")])
    encode.reset_mock()

    with pytest.raises(ValueError, match="query"):
        retriever.retrieve(query)

    encode.assert_not_called()


def splade_with_logits(logits, mask):
    import torch

    class TensorBatch(dict):
        def to(self, device):
            return self

    batch = TensorBatch(attention_mask=torch.tensor(mask))
    retriever = SpladeRetriever([])
    retriever._tok = Mock(return_value=batch)
    retriever._model = Mock(return_value=SimpleNamespace(logits=torch.tensor(logits)))
    return retriever


def test_splade_pooling_excludes_padding_and_preserves_encoding_defaults():
    retriever = splade_with_logits([[[1.0, -1.0, 3.0], [100.0, 100.0, 100.0]]], [[1, 0]])

    vectors = retriever._encode(["passage"])

    np.testing.assert_allclose(vectors.toarray(), [[np.log(2), 0, np.log(4)]])
    assert sparse.isspmatrix_csr(vectors)
    retriever._tok.assert_called_once_with(
        ["passage"], padding=True, truncation=True, max_length=256, return_tensors="pt"
    )


@pytest.mark.parametrize("value", [np.nan, np.inf])
def test_nonfinite_splade_weights_are_rejected(value):
    retriever = splade_with_logits([[[1.0, value, 3.0]]], [[1]])

    with pytest.raises(ValueError, match="finite"):
        retriever._encode(["passage"])


@pytest.mark.parametrize(
    "logits,mask",
    [
        ([[1.0, 2.0, 3.0]], [[1]]),
        ([[[1.0], [2.0]]], [[1]]),
        ([[[1.0]], [[2.0]]], [[1], [1]]),
    ],
)
def test_splade_logits_must_match_tokenized_batch(logits, mask):
    retriever = splade_with_logits(logits, mask)

    with pytest.raises(ValueError, match="align"):
        retriever._encode(["passage"])


@pytest.mark.parametrize(
    "query_matrix", [sparse.csr_matrix([[1.0, 2.0]]), sparse.csr_matrix((0, 1))]
)
def test_splade_query_vectors_must_align_with_index(monkeypatch, query_matrix):
    encode = Mock(side_effect=[sparse.csr_matrix([[1.0]]), query_matrix])
    monkeypatch.setattr(SpladeRetriever, "_encode", encode)
    retriever = SpladeRetriever([Chunk("doc", "passage", "doc")])

    with pytest.raises(ValueError, match="align"):
        retriever.retrieve("question")


def test_nonfinite_splade_dot_products_are_rejected(monkeypatch):
    encode = Mock(side_effect=[sparse.csr_matrix([[1e30]], dtype=np.float32)] * 2)
    monkeypatch.setattr(SpladeRetriever, "_encode", encode)
    retriever = SpladeRetriever([Chunk("doc", "passage", "doc")])

    with pytest.raises(ValueError, match="finite"):
        retriever.retrieve("question")


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("model_name", "other/model"),
        ("revision", "different-revision"),
        ("device", "cuda"),
        ("max_length", 512),
    ],
)
def test_splade_queries_cannot_use_changed_encoding_settings(monkeypatch, field, value):
    encode = Mock(return_value=sparse.csr_matrix([[1.0]]))
    monkeypatch.setattr(SpladeRetriever, "_encode", encode)
    retriever = SpladeRetriever([Chunk("doc", "passage", "doc")])
    setattr(retriever, field, value)
    encode.reset_mock()

    with pytest.raises(ValueError, match="settings changed"):
        retriever.retrieve("question")

    encode.assert_not_called()


@pytest.fixture
def orchestrator():
    results = [
        RetrievalResult(
            Chunk(id="one", content="First evidence.", document_id="doc-one"), score=0.8
        ),
        RetrievalResult(
            Chunk(id="two", content="Second evidence.", document_id="doc-two"), score=0.8
        ),
    ]
    retriever = Mock(spec=CorrectiveRetriever)
    retriever.retrieve.return_value = results
    retriever.retrieve_traced.return_value = (results, False)
    llm = Mock(spec=LLMClient)
    llm.generate.return_value = "A claim [1]."
    llm.stream.return_value = iter(["A claim [1].", " Unsupported [99]."])
    return Orchestrator(retriever, llm)


@pytest.mark.parametrize("stage", ["citations", "grounding"])
def test_failed_stream_verification_replaces_provisional_answer(orchestrator, stage):
    if stage == "grounding":
        orchestrator.grounding_verifier = Mock(spec=GroundingVerifier)
        verifier = orchestrator.grounding_verifier
    else:
        orchestrator.verifier = Mock(spec=AnswerVerifier)
        verifier = orchestrator.verifier
    verifier.verify.side_effect = RuntimeError("Verification unavailable")

    chunks = list(orchestrator.query_streaming("What was measured?"))
    events = [json.loads(chunk.removeprefix("\ndata: ")) for chunk in chunks[-2:]]

    assert chunks[:-2] == ["A claim [1].", " Unsupported [99]."]
    assert events == [
        {
            "event": "verification",
            "citations_valid": [],
            "faithfulness": None,
            "warnings": ["Answer verification failed."],
        },
        {"event": "verified_answer", "answer": "Unable to verify response."},
    ]


@pytest.mark.parametrize("stage", ["citations", "grounding"])
def test_nonstreaming_verification_errors_still_propagate(orchestrator, stage):
    if stage == "grounding":
        orchestrator.grounding_verifier = Mock(spec=GroundingVerifier)
        verifier = orchestrator.grounding_verifier
    else:
        orchestrator.verifier = Mock(spec=AnswerVerifier)
        verifier = orchestrator.verifier
    verifier.verify.side_effect = RuntimeError("Verification unavailable")

    with pytest.raises(RuntimeError, match="Verification unavailable"):
        orchestrator.query("What was measured?")


@pytest.mark.parametrize(
    "scores,cited,expected",
    [
        ([-2.0, -1.0], False, 0.0),
        ([-2.0, -1.0], True, 0.25),
        ([0.4, 0.8], False, 0.3),
        ([0.4, 0.8], True, 0.55),
        ([2.0, 4.0], False, 0.5),
        ([2.0, 4.0], True, 0.75),
    ],
)
def test_answer_confidence_stays_bounded_for_finite_ranking_scores(
    orchestrator, scores, cited, expected
):
    results, _ = orchestrator.retriever.retrieve_traced.return_value
    for result, score in zip(results, scores, strict=True):
        result.score = score
    orchestrator.llm.generate.return_value = "A claim [1]." if cited else "A claim."

    response = orchestrator.query("What was measured?")

    assert response.confidence == expected
    assert 0 <= response.confidence <= 1


@pytest.mark.parametrize("top_k", [0, -1, True, 1.5, "5", None])
def test_invalid_retrieval_limit_fails_before_models(orchestrator, top_k):
    with pytest.raises(ValueError, match="top_k must be a positive integer"):
        Orchestrator(orchestrator.retriever, orchestrator.llm, top_k=top_k)

    orchestrator.retriever.retrieve.assert_not_called()
    orchestrator.retriever.retrieve_traced.assert_not_called()
    orchestrator.llm.generate.assert_not_called()
    orchestrator.llm.stream.assert_not_called()


@pytest.mark.parametrize("streaming", [False, True])
@pytest.mark.parametrize("question", ["", "  ", "\n\t", None, 1, b"Question"])
def test_invalid_question_fails_before_retrieval(orchestrator, question, streaming):
    error = ValueError if isinstance(question, str) else TypeError
    with pytest.raises(error, match="Question"):
        if streaming:
            list(orchestrator.query_streaming(question))
        else:
            orchestrator.query(question)

    orchestrator.retriever.retrieve.assert_not_called()
    orchestrator.retriever.retrieve_traced.assert_not_called()
    orchestrator.llm.generate.assert_not_called()
    orchestrator.llm.stream.assert_not_called()


@pytest.mark.parametrize("streaming", [False, True])
def test_question_whitespace_is_trimmed_before_retrieval(orchestrator, streaming):
    if streaming:
        list(orchestrator.query_streaming("  What was measured? \n"))
        retrieval = orchestrator.retriever.retrieve
    else:
        orchestrator.query("  What was measured? \n")
        retrieval = orchestrator.retriever.retrieve_traced

    retrieval.assert_called_once_with("What was measured?", top_k=5)


def test_falsey_metadata_store_still_supplies_document_names(orchestrator):
    class MetadataStore:
        def __bool__(self):
            return False

        def list_with_metadata(self):
            return {"doc-one": {"original_filename": "measurements.txt"}}

    orchestrator.document_store = MetadataStore()

    response = orchestrator.query("What was measured?")

    assert response.sources[0].document_name == "measurements.txt"


@pytest.mark.parametrize(
    "answer",
    [
        "    value =  42\n    print(value)\n\nClaim [1].\n",
        "\n\n```python\n    value =  42\n```\n\nClaim [1].\n",
        "  Claim [1].\n\n\n    Supporting detail.\n",
    ],
)
def test_complete_and_streamed_answers_preserve_the_same_layout(orchestrator, answer):
    orchestrator.llm.generate.return_value = answer
    orchestrator.llm.stream.return_value = iter([answer])

    response = orchestrator.query("What was measured?")
    chunks = list(orchestrator.query_streaming("What was measured?"))
    final = json.loads(chunks[-1].removeprefix("\ndata: "))

    assert response.answer == answer
    assert final == {"event": "verified_answer", "answer": answer}


def _assert_section_spans(document):
    previous_end = 0
    for section in document.sections:
        assert section.start_offset is not None
        assert section.end_offset is not None
        assert previous_end <= section.start_offset <= section.end_offset <= len(document.content)
        assert document.content[section.start_offset : section.end_offset] == section.content
        previous_end = section.end_offset


def _add_table(container, rows):
    table = container.add_table(rows=len(rows), cols=len(rows[0]))
    for row, values in zip(table.rows, rows, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    return table


class TestDocumentParser:
    @pytest.fixture
    def parser(self):
        return DocumentParser()

    def test_parse_txt_file(self, parser, sample_text_file):
        doc = parser.parse(sample_text_file)

        assert isinstance(doc, ParsedDocument)
        assert doc.id is not None
        assert len(doc.id) == 16
        assert "introduction" in doc.content.lower()
        assert "methodology" in doc.content.lower()
        assert "results" in doc.content.lower()
        assert doc.metadata["filename"] == "sample.txt"
        assert doc.metadata["extension"] == ".txt"

    def test_parse_md_file(self, parser, sample_markdown_file):
        doc = parser.parse(sample_markdown_file)

        assert isinstance(doc, ParsedDocument)
        assert "research paper title" in doc.content.lower()
        assert doc.metadata["extension"] == ".md"

        assert len(doc.sections) > 0
        section_titles = [s.title.lower() for s in doc.sections]
        assert any("abstract" in t for t in section_titles)
        assert any("introduction" in t for t in section_titles)

    def test_parse_nonexistent_file(self, parser, tmp_path):
        with pytest.raises(FileNotFoundError):
            parser.parse(tmp_path / "nonexistent.txt")

    def test_parse_unsupported_format(self, parser, tmp_path):
        unsupported = tmp_path / "file.xyz"
        unsupported.write_text("content")

        with pytest.raises(ValueError, match="Unsupported format"):
            parser.parse(unsupported)

    @pytest.mark.parametrize("extension,content", [(".txt", "café résumé"), (".md", "# café")])
    def test_non_utf8_files_raise_parse_error(self, parser, tmp_path, extension, content):
        path = tmp_path / ("latin1" + extension)
        path.write_bytes(content.encode("latin-1"))
        with pytest.raises(DocumentParseError, match="UTF-8"):
            parser.parse(path)

    def test_documents_sharing_long_prefix_are_not_duplicates(self, parser):

        prefix = "alpha " * 200
        a = parser.parse_bytes((prefix + "ending one").encode(), "paper.txt", ".txt")
        b = parser.parse_bytes((prefix + "ending two").encode(), "paper.txt", ".txt")
        assert a.id != b.id

    def test_generate_document_id(self, parser, sample_text_file):
        doc1 = parser.parse(sample_text_file)
        doc2 = parser.parse(sample_text_file)

        assert doc1.id == doc2.id
        assert len(doc1.id) == 16
        assert doc1.id.isalnum()

    def test_different_files_different_ids(self, parser, sample_text_file, sample_markdown_file):
        doc1 = parser.parse(sample_text_file)
        doc2 = parser.parse(sample_markdown_file)

        assert doc1.id != doc2.id

    def test_section_extraction_txt(self, parser, sample_text_file):
        doc = parser.parse(sample_text_file)

        assert len(doc.sections) >= 1

    def test_section_extraction_md(self, parser, sample_markdown_file):
        doc = parser.parse(sample_markdown_file)

        assert len(doc.sections) >= 4

        levels = [s.level for s in doc.sections]
        assert 1 in levels
        assert 2 in levels

    def test_section_content_preserved(self, parser, sample_markdown_file):
        doc = parser.parse(sample_markdown_file)

        abstract_sections = [s for s in doc.sections if "abstract" in s.title.lower()]
        assert len(abstract_sections) == 1
        assert "two sensor readings per timestamp" in abstract_sections[0].content.lower()

    def test_metadata_extraction(self, parser, sample_text_file):
        doc = parser.parse(sample_text_file)

        assert "filename" in doc.metadata
        assert "extension" in doc.metadata
        assert "size_bytes" in doc.metadata
        assert "created_at" in doc.metadata
        assert "modified_at" in doc.metadata
        assert doc.metadata["size_bytes"] > 0

    def test_word_count_property(self, parser, sample_text_file):
        doc = parser.parse(sample_text_file)

        assert doc.word_count > 0
        assert doc.word_count == len(doc.content.split())

    def test_empty_file(self, parser, empty_file):
        doc = parser.parse(empty_file)

        assert doc.content == ""
        assert doc.word_count == 0

    def test_parse_bytes(self, parser, sample_text_file):
        content = sample_text_file.read_bytes()
        doc = parser.parse_bytes(content, "uploaded.txt", ".txt")

        assert doc.metadata["original_filename"] == "uploaded.txt"
        assert "introduction" in doc.content.lower()

    def test_text_cleaning(self, parser, tmp_path):
        messy_file = tmp_path / "messy.txt"
        messy_file.write_text("Multiple   spaces\n\n\n\nToo many newlines")

        doc = parser.parse(messy_file)

        assert "   " not in doc.content
        assert "\n\n\n" not in doc.content

    @pytest.mark.parametrize("extension", [".txt", ".md"])
    def test_repeated_section_bodies_keep_their_own_normalized_source_spans(
        self, parser, tmp_path, extension
    ):
        first_heading = "METHODS" if extension == ".txt" else "## Methods"
        second_heading = "RESULTS" if extension == ".txt" else "## Results"
        passage = "Shared reading.\nSecond line.\n\nAnother paragraph."
        source = (
            "  Preamble\twith   spacing.\n\n\n"
            f"{first_heading}\n\n{passage}\n\n\n"
            f"{second_heading}\n\n{passage}\n\n"
        )
        path = tmp_path / f"repeated{extension}"
        path.write_text(source)

        document = parser.parse(path)

        _assert_section_spans(document)
        assert [section.content for section in document.sections] == [
            "Preamble with spacing.",
            passage,
            passage,
        ]
        assert document.sections[1].start_offset == document.content.index(first_heading) + len(
            first_heading + "\n\n"
        )
        assert document.sections[2].start_offset == document.content.index(second_heading) + len(
            second_heading + "\n\n"
        )
        assert document.sections[1].start_offset != document.sections[2].start_offset

    @pytest.mark.parametrize("extension", [".txt", ".md"])
    def test_headerless_text_has_one_exact_span(self, parser, tmp_path, extension):
        path = tmp_path / f"plain{extension}"
        path.write_text("\n\n First\t line.\n\n\n Second   paragraph. \n\n")

        document = parser.parse(path)

        _assert_section_spans(document)
        assert len(document.sections) == 1
        assert document.sections[0].title == ""
        assert document.sections[0].content == "First line.\n\n Second paragraph."
        assert document.sections[0].start_offset == 0
        assert document.sections[0].end_offset == len(document.content)

    def test_markdown_empty_heading_bodies_preserve_the_heading_spans(self, parser, tmp_path):
        path = tmp_path / "headings.md"
        path.write_text("# Empty\n\n## Methods\n\nA measurement.\n\n### Empty tail\n")

        document = parser.parse(path)

        _assert_section_spans(document)
        assert [section.level for section in document.sections] == [1, 2, 3]
        assert [section.content for section in document.sections] == [
            "# Empty",
            "A measurement.",
            "### Empty tail",
        ]
        assert [section.heading_only for section in document.sections] == [True, False, True]

    @pytest.mark.parametrize(
        "source,expected",
        [
            ("THE CURRENT IS 42 MA.", ["THE CURRENT IS 42 MA."]),
            (
                "  CALIBRATION PASSED.\n\nTHE CURRENT IS 42 MA.\n\nCALIBRATION PASSED.\n",
                ["CALIBRATION PASSED.", "THE CURRENT IS 42 MA.", "CALIBRATION PASSED."],
            ),
            (
                "RESULTS\nA measurement.\n\nCALIBRATION PASSED.",
                ["A measurement.", "CALIBRATION PASSED."],
            ),
        ],
    )
    def test_uppercase_sentences_keep_nonempty_source_spans(self, parser, source, expected):
        document = parser.parse_bytes(source.encode(), "readings.txt", ".txt")

        _assert_section_spans(document)
        assert [section.content for section in document.sections] == expected
        assert all(section.start_offset < section.end_offset for section in document.sections)
        assert document.sections[-1].heading_only

    def test_docx_sections_preserve_paragraph_joins_and_repeated_source_positions(
        self, parser, tmp_path
    ):
        source = DocxDocument()
        source.add_paragraph("  First\t preamble.  ")
        source.add_paragraph("Second preamble.")
        source.add_heading("Methods", level=1)
        source.add_paragraph("  Shared   passage.  ")
        source.add_paragraph("Line one.\nLine two.")
        source.add_heading("Results", level=2)
        source.add_paragraph("Shared\tpassage.")
        source.add_paragraph("Line one.\nLine two.")
        path = tmp_path / "repeated.docx"
        source.save(path)

        document = parser.parse(path)

        assert document.content == (
            "First preamble.\n\nSecond preamble.\n\nMethods\n\n"
            "Shared passage.\n\nLine one.\nLine two.\n\nResults\n\n"
            "Shared passage.\n\nLine one.\nLine two."
        )
        _assert_section_spans(document)
        assert [section.content for section in document.sections] == [
            "First preamble.\n\nSecond preamble.",
            "Shared passage.\n\nLine one.\nLine two.",
            "Shared passage.\n\nLine one.\nLine two.",
        ]
        assert [section.level for section in document.sections] == [0, 1, 2]
        assert document.sections[1].start_offset == document.content.index("Methods\n\n") + len(
            "Methods\n\n"
        )
        assert document.sections[2].start_offset == document.content.index("Results\n\n") + len(
            "Results\n\n"
        )

    def test_docx_tables_follow_paragraph_and_heading_order(self, parser, tmp_path):
        source = DocxDocument()
        source.add_paragraph("Preamble.")
        _add_table(source, [["Sensor", "Offset"], ["A", "-0.25 °C"]])
        source.add_heading("Results", level=1)
        source.add_paragraph("Recorded readings.")
        _add_table(source, [["Sensor", "Value"], ["B", "12.50%"]])
        source.add_paragraph("After table.")
        path = tmp_path / "ordered.docx"
        source.save(path)

        document = parser.parse(path)

        assert document.content == (
            "Preamble.\n\nSensor | Offset\nA | -0.25 °C\n\nResults\n\n"
            "Recorded readings.\n\nSensor | Value\nB | 12.50%\n\nAfter table."
        )
        assert [section.title for section in document.sections] == ["", "Results"]
        assert document.sections[0].content == "Preamble.\n\nSensor | Offset\nA | -0.25 °C"
        assert document.sections[1].content == (
            "Recorded readings.\n\nSensor | Value\nB | 12.50%\n\nAfter table."
        )
        _assert_section_spans(document)

    def test_table_only_docx_preserves_numeric_text_empty_columns_and_upload_metadata(
        self, parser, tmp_path
    ):
        rows = [
            ["-12.50 °C", "0.0001 mg/L", "98.6%", ""],
            ["", "+0.25 m/s", "−3.2 kPa", "1.0e-06 mol"],
        ]
        source = DocxDocument()
        source.add_table(rows=1, cols=4)
        _add_table(source, rows)
        source.add_table(rows=1, cols=4)
        path = tmp_path / "readings.docx"
        source.save(path)

        document = parser.parse_bytes(path.read_bytes(), "readings.docx", ".docx")

        extracted_rows = [
            [value.strip() for value in line.split("|")]
            for line in document.content.splitlines()
            if line.strip()
        ]
        assert extracted_rows == rows
        assert [section.title for section in document.sections] == [""]
        assert document.sections[0].content == document.content
        _assert_section_spans(document)
        assert document.metadata["filename"] == "readings.docx"
        assert document.metadata["original_filename"] == "readings.docx"
        assert document.metadata["extension"] == ".docx"
        assert document.metadata["file_type"] == "docx"
        assert document.metadata["display_name"] == "readings.docx"

    def test_repeated_docx_tables_have_distinct_section_offsets(self, parser, tmp_path):
        source = DocxDocument()
        for title in ("Methods", "Results"):
            source.add_heading(title, level=1)
            _add_table(source, [["Sensor A", "-0.25 °C"]])
        path = tmp_path / "repeated-tables.docx"
        source.save(path)

        document = parser.parse(path)

        table_text = "Sensor A | -0.25 °C"
        assert [section.content for section in document.sections] == [table_text, table_text]
        assert document.sections[0].start_offset == document.content.index(table_text)
        assert document.sections[1].start_offset == document.content.rindex(table_text)
        assert document.sections[0].start_offset != document.sections[1].start_offset
        _assert_section_spans(document)

    @pytest.mark.parametrize("direction", ["horizontal", "vertical"])
    def test_merged_docx_cells_do_not_duplicate_evidence(self, parser, tmp_path, direction):
        source = DocxDocument()
        table = source.add_table(rows=2, cols=3)
        table.cell(0, 0).text = "Shared result."
        table.cell(0, 0).merge(table.cell(0, 1) if direction == "horizontal" else table.cell(1, 0))
        table.cell(0, 2).text = "First row."
        table.cell(1, 2).text = "Second row."
        path = tmp_path / "merged.docx"
        source.save(path)

        document = parser.parse(path)

        assert document.content.count("Shared result.") == 1
        assert [
            [value.strip() for value in line.split("|")]
            for line in document.content.splitlines()
            if line.strip()
        ] == [["Shared result.", "", "First row."], ["", "", "Second row."]]
        _assert_section_spans(document)

    def test_docx_omitted_edge_cells_keep_their_column_positions(self, parser, tmp_path):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        source = DocxDocument()
        table = _add_table(source, [["", "-0.25 °C", ""], ["Left", "Center", "Right"]])
        row = table.rows[0]._tr
        first, _, last = row.tc_lst
        row.remove(first)
        row.remove(last)
        for name in ("gridBefore", "gridAfter"):
            omitted = OxmlElement("w:" + name)
            omitted.set(qn("w:val"), "1")
            row.get_or_add_trPr().append(omitted)
        path = tmp_path / "omitted.docx"
        source.save(path)

        document = parser.parse(path)

        assert [
            [value.strip() for value in line.split("|")]
            for line in document.content.splitlines()
            if line.strip()
        ] == [["", "-0.25 °C", ""], ["Left", "Center", "Right"]]
        _assert_section_spans(document)

    def test_nested_docx_tables_keep_cell_block_order(self, parser, tmp_path):
        source = DocxDocument()
        table = source.add_table(rows=1, cols=2)
        cell = table.cell(0, 0)
        cell.text = "Before nested."
        _add_table(cell, [["Nested A", "-0.5 °C"], ["Nested B", "2.0%"]])
        cell.add_paragraph("After nested.")
        table.cell(0, 1).text = "Adjacent cell."
        path = tmp_path / "nested.docx"
        source.save(path)

        document = parser.parse(path)

        assert document.content == (
            "Before nested.; [Nested A | -0.5 °C; Nested B | 2.0%]; After nested. | Adjacent cell."
        )
        _assert_section_spans(document)


class TestPdfEdgeCases:
    def test_page_offsets_follow_cleaned_joins_and_retain_page_numbers(self, tmp_path, monkeypatch):
        page_texts = [
            "\n  METHODS\n\nShared   passage.\n\n\nFirst page tail.  \n",
            " \n\t ",
            "\nRESULTS\n\nShared\tpassage.\n\nLast page tail.\n",
            "  Shared passage.\n\nA fourth-page note. \n",
        ]
        reader = SimpleNamespace(
            is_encrypted=False,
            pages=[SimpleNamespace(extract_text=lambda text=text: text) for text in page_texts],
        )
        monkeypatch.setattr("scinexusrag.ingestion.parser.PdfReader", lambda path: reader)
        path = tmp_path / "repeated.pdf"
        path.write_bytes(b"%PDF-1.4 fake")

        document = DocumentParser().parse(path)

        _assert_section_spans(document)
        assert [section.page_number for section in document.sections] == [1, 3, 4]
        assert [section.title for section in document.sections] == ["METHODS", "RESULTS", ""]
        assert [section.content for section in document.sections] == [
            "Shared passage.\n\nFirst page tail.",
            "Shared passage.\n\nLast page tail.",
            "Shared passage.\n\nA fourth-page note.",
        ]
        first_page = "METHODS\n\nShared passage.\n\nFirst page tail."
        third_page = "RESULTS\n\nShared passage.\n\nLast page tail."
        assert document.sections[1].start_offset == len(first_page + "\n\nRESULTS\n\n")
        assert document.sections[2].start_offset == len(first_page + "\n\n" + third_page + "\n\n")

    def test_headerless_pages_reach_chunks(self, tmp_path, monkeypatch):

        pages = [
            SimpleNamespace(extract_text=lambda: "METHODS\nThe sample was measured."),
            SimpleNamespace(
                extract_text=lambda: "Critical measurements on a page without a heading."
            ),
        ]
        reader = SimpleNamespace(pages=pages, is_encrypted=False)
        monkeypatch.setattr("scinexusrag.ingestion.parser.PdfReader", lambda path: reader)
        pdf_path = tmp_path / "mixed.pdf"
        pdf_path.write_bytes(b"%PDF-1.4 fake")
        chunks = SemanticChunker().chunk(DocumentParser().parse(pdf_path))

        assert any("Critical measurements" in chunk.content for chunk in chunks)
        assert (
            next(chunk for chunk in chunks if "Critical measurements" in chunk.content).page_number
            == 2
        )

    def test_encrypted_pdf_actionable_error(self, tmp_path):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.encrypt("secret")
        pdf_path = tmp_path / "locked.pdf"
        with open(pdf_path, "wb") as f:
            writer.write(f)

        with pytest.raises(DocumentParseError, match="password-protected"):
            DocumentParser().parse(pdf_path)

    def test_scanned_pdf_actionable_error(self, tmp_path, monkeypatch):
        page = SimpleNamespace(images=[object()], extract_text=lambda: "")
        reader = SimpleNamespace(pages=[page], is_encrypted=False)
        monkeypatch.setattr("scinexusrag.ingestion.parser.PdfReader", lambda path: reader)
        pdf_path = tmp_path / "scan.pdf"
        pdf_path.write_bytes(b"%PDF-1.4 fake")

        with pytest.raises(DocumentParseError, match="scanned"):
            DocumentParser().parse(pdf_path)


class TestDuplicateNormalization:
    def test_whitespace_only_edits_keep_same_id(self, tmp_path):
        parser = DocumentParser()
        doc = tmp_path / "doc.txt"
        doc.write_text("Alpha beta. Gamma delta.")
        original_id = parser.parse(doc).id

        doc.write_text("Alpha beta.\n\nGamma  delta. ")

        assert parser.parse(doc).id == original_id

    def test_content_edits_change_id(self, tmp_path):
        parser = DocumentParser()
        doc = tmp_path / "doc.txt"
        doc.write_text("Alpha beta. Gamma delta.")
        original_id = parser.parse(doc).id

        doc.write_text("Alpha beta. Gamma delta. Epsilon.")

        assert parser.parse(doc).id != original_id


class TestPreamblePreserved:
    def test_text_preamble_before_first_heading_reaches_chunks(self, tmp_path):

        txt = tmp_path / "doc.txt"
        txt.write_text("Critical preamble sentence.\n\nMETHODS\nBody under heading.\n")
        chunks = SemanticChunker().chunk(DocumentParser().parse(txt))

        assert any("Critical preamble sentence." in chunk.content for chunk in chunks)
        assert any("Body under heading." in chunk.content for chunk in chunks)

    def test_markdown_preamble_before_first_heading_kept(self, tmp_path):
        md = tmp_path / "doc.md"
        md.write_text("Critical preamble sentence.\n\n# Heading One\n\nBody under heading.\n")
        doc = DocumentParser().parse(md)
        assert "critical preamble" in doc.content.lower()
        assert any("critical preamble" in s.content.lower() for s in doc.sections)

    def test_docx_preamble_before_first_heading_kept(self, tmp_path):
        d = DocxDocument()
        d.add_paragraph("Critical preamble sentence.")
        d.add_heading("Heading One", level=1)
        d.add_paragraph("Body under heading.")
        path = tmp_path / "doc.docx"
        d.save(path)

        doc = DocumentParser().parse(path)
        assert any("critical preamble" in s.content.lower() for s in doc.sections)


@pytest.mark.parametrize(
    "extension,data",
    [(".pdf", b"%PDF-1.4\ntruncated"), (".docx", b"not a Word document")],
)
def test_damaged_documents_have_actionable_parse_errors(extension, data):
    with pytest.raises(DocumentParseError, match="Cannot read") as error:
        DocumentParser().parse_bytes(data, "damaged" + extension, extension)

    assert error.value.__cause__ is not None


@pytest.mark.parametrize("content_types", [None, b"<broken>"])
def test_docx_missing_or_malformed_package_entries_are_parse_errors(content_types):
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"<unused />")
        if content_types is not None:
            archive.writestr("[Content_Types].xml", content_types)

    with pytest.raises(DocumentParseError, match="Cannot read DOCX"):
        DocumentParser().parse_bytes(stream.getvalue(), "damaged.docx", ".docx")


@pytest.mark.parametrize("root", ["document", "invalid"])
def test_docx_requires_a_document_root_and_body(root):
    source = BytesIO()
    DocxDocument().save(source)
    source.seek(0)
    damaged = BytesIO()
    replacement = (
        f'<w:{root} xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" />'
    ).encode()
    with ZipFile(source) as original, ZipFile(damaged, "w") as rewritten:
        for entry in original.infolist():
            content = replacement if entry.filename == "word/document.xml" else original.read(entry)
            rewritten.writestr(entry, content)

    with pytest.raises(DocumentParseError, match="document body"):
        DocumentParser().parse_bytes(damaged.getvalue(), "damaged.docx", ".docx")


@pytest.mark.parametrize("error_type", [OSError, RuntimeError])
def test_parser_preserves_non_document_failures(tmp_path, monkeypatch, error_type):
    path = tmp_path / "paper.docx"
    path.write_bytes(b"placeholder")
    failure = error_type("unrelated failure")

    def fail(_path):
        raise failure

    monkeypatch.setattr("scinexusrag.ingestion.parser.DocxDocument", fail)

    with pytest.raises(error_type) as error:
        DocumentParser().parse(path)

    assert error.value is failure


@pytest.mark.parametrize("extension", [".unknown", ".bad/format", ""])
def test_upload_rejects_unsupported_extensions_before_creating_files(
    tmp_path, monkeypatch, extension
):
    monkeypatch.setattr("tempfile.tempdir", str(tmp_path))

    with pytest.raises(ValueError, match="Unsupported format"):
        DocumentParser().parse_bytes(b"contents", "paper", extension)

    assert not list(tmp_path.iterdir())


def test_upload_cleans_temporary_files_after_a_write_failure(tmp_path, monkeypatch):
    monkeypatch.setattr("tempfile.tempdir", str(tmp_path))

    def fail_write(_path, _data):
        raise OSError("disk full")

    monkeypatch.setattr(Path, "write_bytes", fail_write)

    with pytest.raises(OSError, match="disk full"):
        DocumentParser().parse_bytes(b"contents", "paper.txt", ".txt")

    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("data", [b"a reading", b"\xff"])
def test_upload_cleans_temporary_files_after_parsing(tmp_path, monkeypatch, data):
    monkeypatch.setattr("tempfile.tempdir", str(tmp_path))

    if data == b"\xff":
        with pytest.raises(DocumentParseError, match="UTF-8"):
            DocumentParser().parse_bytes(data, "paper.txt", ".txt")
    else:
        assert DocumentParser().parse_bytes(data, "paper.txt", ".txt").content == "a reading"

    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("marker", ["```", "~~~", " ````", "   ~~~~"])
def test_markdown_fenced_comments_remain_in_their_section(marker):
    source = (
        "# Methods\n\n"
        f"{marker}python\n# calibration comment\n\n\nprint(42)\n{marker}\n\n"
        "## Results\n\nValue   42."
    )

    document = DocumentParser().parse_bytes(source.encode(), "paper.md", ".md")

    assert [section.title for section in document.sections] == ["Methods", "Results"]
    assert "# calibration comment" in document.sections[0].content
    assert "print(42)" in document.sections[0].content
    assert document.sections[1].content == "Value 42."
    for section in document.sections:
        assert document.content[section.start_offset : section.end_offset] == section.content


@pytest.mark.parametrize("false_close", ["```", "~~~~", "```` trailing text"])
def test_markdown_fences_require_a_matching_close(false_close):
    source = f"# Methods\n\n````python\n{false_close}\n# inside code\n`````\n\n## Results\n\n42."

    document = DocumentParser().parse_bytes(source.encode(), "paper.md", ".md")

    assert [section.title for section in document.sections] == ["Methods", "Results"]
    assert "# inside code" in document.sections[0].content


def test_unclosed_markdown_fences_keep_later_hash_lines_as_code():
    source = "# Methods\n\n```python\n# calibration comment\n\n## still code\nprint(42)"

    document = DocumentParser().parse_bytes(source.encode(), "paper.md", ".md")

    assert [section.title for section in document.sections] == ["Methods"]
    assert "## still code" in document.sections[0].content


@pytest.mark.parametrize("literal_fence", ["    ```", "```not`a`fence", "\t~~~"])
def test_literal_fence_text_does_not_hide_markdown_headings(literal_fence):
    source = f"Preamble.\n\n{literal_fence}\n\n## Results\n\n42."

    document = DocumentParser().parse_bytes(source.encode(), "paper.md", ".md")

    assert [section.title for section in document.sections] == ["", "Results"]
    assert document.sections[1].content == "42."


def test_indented_markdown_heading_has_an_exact_normalized_span():
    source = "Preamble.\n \n\n   ## Methods\n\nMeasured   42.\n\n  ## Results\n\n42."

    document = DocumentParser().parse_bytes(source.encode(), "paper.md", ".md")

    assert [section.title for section in document.sections] == ["", "Methods", "Results"]
    assert document.sections[1].content == "Measured 42."
    for section in document.sections:
        assert document.content[section.start_offset : section.end_offset] == section.content


@pytest.mark.parametrize("extension,heading", [(".md", "# Methods"), (".txt", "METHODS")])
def test_utf8_bom_does_not_change_content_sections_or_document_identity(extension, heading):
    source = f"{heading}\n\nMeasured 42 mA."
    parser = DocumentParser()
    plain = parser.parse_bytes(source.encode(), "paper" + extension, extension)
    with_bom = parser.parse_bytes(source.encode("utf-8-sig"), "paper" + extension, extension)

    assert with_bom.content == plain.content
    assert with_bom.sections == plain.sections
    assert with_bom.id == plain.id
    assert with_bom.sections[0].title.lower() == "methods"
    assert with_bom.sections[0].content == "Measured 42 mA."


def test_utf8_bom_only_file_is_empty():
    document = DocumentParser().parse_bytes(b"\xef\xbb\xbf", "empty.txt", ".txt")

    assert document.content == ""
    assert document.sections == []
    assert document.word_count == 0


def test_interior_byte_order_mark_is_preserved():
    source = "Measured \ufeff42 mA."

    document = DocumentParser().parse_bytes(source.encode(), "paper.txt", ".txt")

    assert document.content == source


@pytest.fixture
def request_settings(local_settings):
    request_settings = local_settings
    request_settings.api.max_upload_mb = 1
    return request_settings


@pytest.mark.parametrize(
    "host",
    [
        "rebind.example",
        "localhost.evil.example",
        "evil@localhost",
        "localhost:not-a-port",
        "[::1",
        "[::1]evil.example",
        "[::1]foo:8000",
        "localhost:",
        "127.0.0.1:",
        "[::1]:",
        "localhost:0",
        "[::1]:65536",
    ],
)
def test_local_api_rejects_foreign_hosts(request_settings, monkeypatch, host):
    def unexpected_pipeline():
        pytest.fail("Rejected hosts must not read the document collection.")

    monkeypatch.setattr(routes, "get_scinexusrag", unexpected_pipeline)
    response = TestClient(api.create_app(), base_url="http://localhost").get(
        "/api/documents", headers={"Host": host, "Sec-Fetch-Site": "same-origin"}
    )
    assert response.status_code == 400


def test_loopback_host_rejects_unicode_case_folding():
    assert not security.is_loopback_host("localhoſt")


@pytest.mark.parametrize(
    "host",
    [
        "localhost",
        "LOCALHOST:8000",
        "127.0.0.1",
        "127.0.0.1:8000",
        "[::1]",
        "[::1]:65535",
    ],
)
def test_local_api_accepts_loopback_hosts(request_settings, host):
    response = TestClient(api.create_app(), base_url="http://localhost").get(
        "/health", headers={"Host": host}
    )
    assert response.status_code == 200


def test_configured_key_preserves_public_hosts(request_settings):
    request_settings.api.api_key = "test-only-key"
    response = TestClient(api.create_app(), base_url="https://research.example").get("/health")
    assert response.status_code == 200


def test_explicit_network_bind_preserves_public_hosts(request_settings):
    request_settings.api.host = "0.0.0.0"
    response = TestClient(api.create_app(), base_url="https://research.example").get("/health")
    assert response.status_code == 200


@pytest.mark.parametrize("content_length", [None, b"1", b"2097152"])
@pytest.mark.parametrize("mount_path", ["", "/research"])
def test_oversized_multipart_stops_before_reading_all_bytes(
    request_settings, monkeypatch, content_length, mount_path, asgi_request
):
    import starlette.formparsers

    temporary_files = []
    create_temporary_file = starlette.formparsers.SpooledTemporaryFile

    def tracked_temporary_file(*args, **kwargs):
        file = create_temporary_file(*args, **kwargs)
        temporary_files.append(file)
        return file

    monkeypatch.setattr(starlette.formparsers, "SpooledTemporaryFile", tracked_temporary_file)
    boundary = b"security-boundary"
    prefix = (
        b"--" + boundary + b'\r\nContent-Disposition: form-data; name="file"; '
        b'filename="large.txt"\r\nContent-Type: text/plain\r\n\r\n'
    )
    chunks = [prefix] + [b"x" * 65536] * 32 + [b"\r\n--" + boundary + b"--\r\n"]
    headers = [
        (b"host", b"localhost"),
        (b"content-type", b"multipart/form-data; boundary=" + boundary),
    ]
    if content_length is not None:
        headers.append((b"content-length", content_length))
    messages, consumed = asyncio.run(
        asgi_request(
            api.create_app(),
            headers,
            chunks,
            path=mount_path + "/api/ingest",
            root_path=mount_path,
        )
    )

    assert (
        next(message["status"] for message in messages if message["type"] == "http.response.start")
        == 413
    )
    assert consumed < len(chunks)
    assert all(file.closed for file in temporary_files)


def test_upload_trailing_slash_preserves_redirect(request_settings):
    response = TestClient(api.create_app(), base_url="http://localhost").post(
        "/api/ingest/",
        files={"file": ("normal.txt", b"x" * (128 * 1024), "text/plain")},
        follow_redirects=False,
    )
    assert response.status_code == 307
    assert response.headers["location"] == "http://localhost/api/ingest"


def test_normal_upload_reaches_ingestion(request_settings, monkeypatch):
    ingested = []

    def ingest(data, filename, extension):
        ingested.append((data, filename, extension))
        return SimpleNamespace(
            success=True,
            document_id="test-doc",
            filename=filename,
            chunk_count=1,
            word_count=2,
            error=None,
        )

    monkeypatch.setattr(routes, "get_scinexusrag", lambda: SimpleNamespace(ingest_bytes=ingest))
    response = TestClient(api.create_app(), base_url="http://localhost").post(
        "/api/ingest", files={"file": ("small.txt", b"safe text", "text/plain")}
    )
    assert response.status_code == 200
    assert ingested == [(b"safe text", "small.txt", ".txt")]


@pytest.mark.parametrize(
    "headers",
    [
        [],
        [(b"host", b"localhost"), (b"host", b"evil.example")],
        [(b"host", b"evil.example"), (b"host", b"localhost")],
        [(b"host", b"localhost"), (b"host", b"localhost")],
    ],
)
def test_local_api_rejects_missing_or_duplicate_hosts_before_reading_body(
    request_settings, monkeypatch, headers, asgi_request
):
    def unexpected_pipeline():
        pytest.fail("Invalid host headers must not reach the document pipeline.")

    monkeypatch.setattr(routes, "get_scinexusrag", unexpected_pipeline)
    messages, consumed = asyncio.run(asgi_request(api.create_app(), headers))
    assert [
        message["status"] for message in messages if message["type"] == "http.response.start"
    ] == [400]
    assert consumed == 0


@pytest.mark.parametrize(
    "content_lengths, expected_status",
    [
        ([b""], 400),
        ([b"-1"], 400),
        ([b"+1"], 400),
        ([b" 1"], 400),
        ([b"1 "], 400),
        ([b"1_000"], 400),
        ([b"1.0"], 400),
        ([b"\xb2"], 400),
        ([b"1, 1"], 400),
        ([b"1", b"1"], 400),
        ([b"1", b"65537"], 400),
        ([b"65537"], 413),
        ([b"9" * 5000], 413),
    ],
)
def test_api_rejects_invalid_or_oversized_declared_lengths_before_reading_body(
    request_settings, monkeypatch, content_lengths, expected_status, asgi_request
):
    def unexpected_pipeline():
        pytest.fail("Rejected request lengths must not reach the document pipeline.")

    monkeypatch.setattr(routes, "get_scinexusrag", unexpected_pipeline)
    headers = [(b"host", b"localhost")]
    headers.extend((b"content-length", value) for value in content_lengths)
    messages, consumed = asyncio.run(asgi_request(api.create_app(), headers))
    assert [
        message["status"] for message in messages if message["type"] == "http.response.start"
    ] == [expected_status]
    assert consumed == 0


async def _body_size_app(scope, receive, send):
    size = 0
    while True:
        message = await receive()
        size += len(message.get("body", b""))
        if not message.get("more_body", False):
            break
    await send({"type": "http.response.start", "status": 200, "headers": []})
    await send({"type": "http.response.body", "body": str(size).encode()})


@pytest.mark.parametrize("content_length", [None, b"1", b"00065536", b"0" * 5000])
@pytest.mark.parametrize("overflow", [False, True])
def test_api_stream_limit_counts_actual_bytes(content_length, overflow, asgi_request):
    headers = [(b"host", b"localhost")]
    if content_length is not None:
        headers.append((b"content-length", content_length))
    chunks = [b"x" * 32768, b"x" * 32768]
    if overflow:
        chunks.extend([b"x", b"unread"])
    app = security.LoopbackHostMiddleware(
        security.RequestBodyLimitMiddleware(_body_size_app, max_upload_bytes=1)
    )

    messages, consumed = asyncio.run(asgi_request(app, headers, chunks))

    assert [
        message["status"] for message in messages if message["type"] == "http.response.start"
    ] == [413 if overflow else 200]
    assert consumed == (3 if overflow else 2)
    if not overflow:
        assert messages[-1]["body"] == b"65536"


def test_concurrent_requests_have_independent_body_limits(asgi_request):
    app = security.LoopbackHostMiddleware(
        security.RequestBodyLimitMiddleware(_body_size_app, max_upload_bytes=1)
    )

    async def requests():
        return await asyncio.gather(
            *(
                asgi_request(app, [(b"host", b"localhost")], [b"x" * 32768, b"x" * 32768])
                for _ in range(2)
            )
        )

    for messages, consumed in asyncio.run(requests()):
        assert messages[0]["status"] == 200
        assert messages[-1]["body"] == b"65536"
        assert consumed == 2


@pytest.mark.parametrize("scope", [{"type": "lifespan"}, {"type": "websocket"}])
def test_non_http_scope_passes_through_unchanged(scope):
    calls = []

    async def receive():
        pytest.fail("Middleware must not consume non-HTTP messages.")

    async def send(message):
        pytest.fail("Middleware must not send non-HTTP messages.")

    async def app(received_scope, received_receive, received_send):
        calls.append((received_scope, received_receive, received_send))

    middleware = security.LoopbackHostMiddleware(
        security.RequestBodyLimitMiddleware(app, max_upload_bytes=1)
    )
    asyncio.run(middleware(scope, receive, send))
    assert len(calls) == 1
    assert calls[0][0] is scope
    assert calls[0][1] is receive
    assert calls[0][2] is send


@pytest.fixture
def sample_chunks():
    texts = [
        "Machine learning is a subset of artificial intelligence.",
        "Deep learning uses neural networks with many layers.",
        "Natural language processing enables computers to understand text.",
        "Computer vision allows machines to interpret images.",
        "Reinforcement learning trains agents through rewards.",
    ]
    docs = ["doc1", "doc1", "doc2", "doc2", "doc3"]
    return [
        Chunk(id=f"chunk{i + 1}", content=t, document_id=d)
        for i, (t, d) in enumerate(zip(texts, docs, strict=True))
    ]


class TestBM25Retriever:
    def test_add_indexes_all_chunks(self, sample_chunks):
        r = BM25Retriever()
        assert r.add(sample_chunks) == 5
        assert r.count() == 5

    def test_retrieve_ranks_the_matching_chunk_first(self, sample_chunks):
        r = BM25Retriever()
        r.add(sample_chunks)
        results = r.retrieve("neural networks deep learning", top_k=3)
        assert results[0].chunk.id == "chunk2"
        assert all(res.source == "sparse" for res in results)
        assert all(0 <= res.score <= 1 for res in results)

    def test_empty_index_and_empty_query(self, sample_chunks):
        assert BM25Retriever().retrieve("anything", top_k=5) == []
        r = BM25Retriever()
        r.add(sample_chunks)
        assert r.retrieve("", top_k=5) == []

    def test_single_matching_passage_with_negative_idf_is_returned(self):
        retriever = BM25Retriever()
        retriever.add([Chunk(id="match", document_id="paper", content="kinase tumor inhibition")])

        results = retriever.retrieve("kinase")

        assert [result.chunk.id for result in results] == ["match"]
        assert results[0].score == 0.0

    @pytest.mark.parametrize(
        "texts",
        [
            ["cardiac disease", "kinase inhibition"],
            ["shared pathway", "shared pathway kinase", "shared pathway kinase kinase"],
        ],
    )
    def test_nonpositive_match_is_ranked_before_unmatched_passage(self, texts):
        retriever = BM25Retriever()
        retriever.add(
            [
                Chunk(id=f"passage{i}", document_id="paper", content=text)
                for i, text in enumerate(texts)
            ]
        )

        results = retriever.retrieve("kinase", top_k=1)

        assert [result.chunk.id for result in results] == ["passage1"]
        assert results[0].score == 0.0

    def test_query_without_token_overlap_returns_no_passages(self):
        retriever = BM25Retriever()
        retriever.add([Chunk(id="one", document_id="paper", content="kinase tumor inhibition")])

        assert retriever.retrieve("cardiac disease") == []

    @pytest.mark.parametrize("top_k", [0, -1])
    def test_nonpositive_depth_returns_no_passages(self, sample_chunks, top_k):
        retriever = BM25Retriever()
        retriever.add(sample_chunks)

        assert retriever.retrieve("neural networks", top_k=top_k) == []

    def test_tokenize_drops_stopwords(self):
        tokens = BM25Retriever().tokenize("the quick brown fox is jumping")
        assert "the" not in tokens and "is" not in tokens
        assert "quick" in tokens and "brown" in tokens

    def test_explicit_empty_stopwords_keep_common_words(self):
        assert BM25Retriever(stopwords=set()).tokenize("the quick fox") == [
            "the",
            "quick",
            "fox",
        ]

    def test_incremental_add_and_remove(self, sample_chunks):
        r = BM25Retriever()
        r.add(sample_chunks[:3])
        r.add_incremental(sample_chunks[3:])
        assert r.count() == 5
        assert r.remove({"chunk1", "chunk2"}) == 2
        assert r.count() == 3

    def test_incremental_add_is_idempotent_by_id(self, sample_chunks):

        r = BM25Retriever()
        r.add(sample_chunks[:3])
        r.add_incremental(sample_chunks[:3])
        assert r.count() == 3
        r.add_incremental(sample_chunks[2:])
        assert r.count() == 5

    @pytest.mark.parametrize("preload", [False, True])
    def test_incremental_add_deduplicates_the_new_batch(self, sample_chunks, preload):
        retriever = BM25Retriever()
        if preload:
            retriever.add(sample_chunks[:1])

        assert retriever.add_incremental([*sample_chunks[:2], *sample_chunks[:2]]) == 2
        assert retriever.count() == 2
        assert retriever.remove({sample_chunks[0].id}) == 1
        assert retriever.count() == 1

    def test_add_survives_all_empty_tokenization(self):

        chunks = [
            Chunk(id=f"c{i}", content=c, document_id="d")
            for i, c in enumerate(["the a is", "of to in", "x y z"])
        ]
        r = BM25Retriever()
        assert r.add(chunks) == 3
        assert r.count() == 3
        assert r.retrieve("machine learning", top_k=5) == []


class TestReranker:
    def _results(self):
        return [
            RetrievalResult(
                chunk=Chunk(id=f"c{i}", content=t, document_id="d"), score=0.0, source="hybrid"
            )
            for i, t in enumerate(["alpha", "beta", "gamma"])
        ]

    def _reranker(self, scores):
        r = Reranker()
        model = MagicMock()
        model.predict.return_value = np.array(scores, dtype=np.float32)
        r._model = model
        return r

    def test_empty_results_returns_empty(self):
        assert self._reranker([]).rerank("q", []) == []

    def test_orders_by_cross_encoder_score_desc(self):
        out = self._reranker([0.1, 0.9, 0.5]).rerank("q", self._results())
        assert [x.chunk.id for x in out] == ["c1", "c2", "c0"]
        assert out[0].score == 1.0 and out[-1].score == 0.0
        assert all(x.source == "hybrid+rerank" for x in out)

    def test_all_equal_scores_map_to_neutral_one(self):
        out = self._reranker([0.4, 0.4, 0.4]).rerank("q", self._results())
        assert all(x.score == 1.0 for x in out)

    def test_top_k_truncates(self):
        out = self._reranker([0.1, 0.9, 0.5]).rerank("q", self._results(), top_k=2)
        assert [x.chunk.id for x in out] == ["c1", "c2"]


class TestHybridRetriever:
    @pytest.fixture
    def dense(self, sample_chunks):
        m = MagicMock()
        m.retrieve.return_value = [
            RetrievalResult(sample_chunks[0], 0.9, "dense"),
            RetrievalResult(sample_chunks[1], 0.8, "dense"),
            RetrievalResult(sample_chunks[2], 0.6, "dense"),
        ]
        return m

    @pytest.fixture
    def sparse(self, sample_chunks):
        m = MagicMock()
        m.retrieve.return_value = [
            RetrievalResult(sample_chunks[1], 0.85, "sparse"),
            RetrievalResult(sample_chunks[3], 0.7, "sparse"),
            RetrievalResult(sample_chunks[0], 0.5, "sparse"),
        ]
        return m

    def test_fusion_promotes_chunk_in_both_lists(self, dense, sparse):
        hybrid = HybridRetriever(dense, sparse, 0.5, 0.5)
        results = hybrid.retrieve("q", top_k=3)
        assert results[0].chunk.id == "chunk2"
        assert all(r.source == "hybrid" for r in results)

    def test_empty_results(self, dense, sparse):
        dense.retrieve.return_value = []
        sparse.retrieve.return_value = []
        assert HybridRetriever(dense, sparse).retrieve("q", top_k=5) == []

    def test_fusion_retains_lexical_match_with_zero_bm25_score(self):
        sparse = BM25Retriever()
        sparse.add([Chunk(id="match", document_id="paper", content="kinase tumor inhibition")])
        assert sparse.retrieve("kinase")[0].score == 0.0
        dense = MagicMock()
        dense.retrieve.return_value = []

        results = HybridRetriever(dense, sparse).retrieve("kinase", top_k=1)

        assert [result.chunk.id for result in results] == ["match"]
        assert results[0].source == "hybrid"
        assert results[0].score == 1.0


class TestSpladeEmptyCorpus:
    def test_empty_corpus_does_not_crash(self):

        r = SpladeRetriever([])
        assert r.retrieve("query") == []


def _ranked_result(cid: str, score: float = 0.0) -> RetrievalResult:
    return RetrievalResult(chunk=Chunk(id=cid, content=cid, document_id=cid), score=score)


def test_rrf_fuse_orders_by_summed_reciprocal_rank():
    a = [_ranked_result("c1"), _ranked_result("c2"), _ranked_result("c3")]
    b = [_ranked_result("c2"), _ranked_result("c3"), _ranked_result("c4")]
    fused = rrf_fuse([a, b], [1.0, 1.0], k=60)

    assert [r.chunk.id for r in fused] == ["c2", "c3", "c1", "c4"]
    assert fused[0].score == 1.0


def test_rrf_fuse_weights_shift_ranking():
    a = [_ranked_result("c1"), _ranked_result("c2")]
    b = [_ranked_result("c2"), _ranked_result("c1")]
    sparse_heavy = rrf_fuse([a, b], [0.1, 0.9], k=60)
    assert sparse_heavy[0].chunk.id == "c2"


def test_rrf_fuse_zero_weights_return_zero_scores_in_stable_order():
    first = [_ranked_result("c1"), _ranked_result("c2")]
    second = [_ranked_result("c2"), _ranked_result("c3")]

    fused = rrf_fuse([first, second], [0.0, 0.0], top_k=2)

    assert [result.chunk.id for result in fused] == ["c1", "c2"]
    assert all(result.score == 0.0 for result in fused)


def test_roc_auc_known_value():
    assert math.isclose(roc_auc([0.1, 0.4, 0.35, 0.8], [0, 0, 1, 1]), 0.75)
    assert roc_auc([0.1, 0.2, 0.3, 0.4], [0, 0, 1, 1]) == 1.0
    assert roc_auc([0.4, 0.3, 0.2, 0.1], [0, 0, 1, 1]) == 0.0


def test_roc_auc_handles_ties():
    assert roc_auc([0.5, 0.5, 0.5, 0.5], [0, 0, 1, 1]) == 0.5


def test_pr_auc_known_value():
    assert math.isclose(pr_auc([0.1, 0.4, 0.35, 0.8], [0, 0, 1, 1]), 0.8333333, rel_tol=1e-5)


def test_holm_is_monotone_and_scales_smallest():
    adj = holm_correction({"a": 0.01, "b": 0.02, "c": 0.5})
    assert math.isclose(adj["a"], 0.03)
    assert adj["a"] <= adj["b"] <= adj["c"]


def test_risk_coverage_auc_rewards_confident_correct():
    assert risk_coverage_auc([0.9, 0.8, 0.7], [1, 1, 1]) == 0.0
    assert risk_coverage_auc([0.9, 0.8, 0.7], [0, 0, 0]) == 1.0
    assert math.isclose(risk_coverage_auc([0.9, 0.8], [1, 0]), 0.25)


class _FakeDense:
    def __init__(self, top_score: float):
        self.top_score = top_score
        self.calls = 0

    def retrieve(self, query, top_k=5):
        self.calls += 1
        return [_ranked_result("d1", self.top_score)]


class _FakeBase:
    def __init__(self, top_score: float):
        self.dense = _FakeDense(top_score)
        self.rrf_k = 60
        self.calls: list[str] = []

    def tokenize_sparse(self, text):
        return [t for t in text.lower().split() if len(t) > 1]

    @property
    def sparse(self):
        return self

    def tokenize(self, text):
        return self.tokenize_sparse(text)

    def retrieve(self, query, top_k=10, depth=50):
        self.calls.append(query)
        return [_ranked_result("d1", 0.9), _ranked_result("d2", 0.8)]

    def retrieve_with_dense_top(self, query, top_k=10, depth=50):

        self.calls.append(query)
        return [_ranked_result("d1", 0.9), _ranked_result("d2", 0.8)], self.dense.retrieve(
            query, 1
        )[0].score


def test_corrective_skips_when_confident():
    base = _FakeBase(top_score=0.9)
    cr = CorrectiveRetriever(base, tau=0.55)
    _, triggered = cr.retrieve_traced("kinase inhibits tumor growth", top_k=2)
    assert triggered is False
    assert len(base.calls) == 1
    assert base.dense.calls == 1


def test_corrective_expands_when_weak():
    base = _FakeBase(top_score=0.2)
    cr = CorrectiveRetriever(base, tau=0.55, feedback_terms=3)
    _, triggered = cr.retrieve_traced("kinase inhibits tumor growth", top_k=2)
    assert triggered is True
    assert len(base.calls) == 2
    assert base.calls[1] != base.calls[0]


def test_corrective_disabled_skips_second_pass():

    base = _FakeBase(top_score=0.2)
    cr = CorrectiveRetriever(base, tau=0.55, enabled=False)
    _, triggered = cr.retrieve_traced("kinase inhibits tumor growth", top_k=2)
    assert triggered is False
    assert len(base.calls) == 1


def _content_result(text: str) -> RetrievalResult:
    return RetrievalResult(chunk=Chunk(id="x", content=text, document_id="d"), score=0.5)


def test_expand_appends_frequent_non_query_terms():

    base = _FakeBase(top_score=0.2)
    cr = CorrectiveRetriever(base, feedback_terms=2)
    results = [
        _content_result("apoptosis apoptosis apoptosis signaling pathway"),
        _content_result("signaling pathway pathway"),
    ]
    added = cr.expand("kinase inhibits", results).split()[2:]

    assert added == ["apoptosis", "pathway"]
    assert "kinase" not in added and "inhibits" not in added
    assert "signaling" not in added


def test_expand_respects_feedback_terms_cap():
    base = _FakeBase(top_score=0.2)
    cr = CorrectiveRetriever(base, feedback_terms=1)
    added = cr.expand("kinase", [_content_result("apoptosis apoptosis signaling")]).split()[1:]
    assert added == ["apoptosis"]


def test_hybrid_retrieve_with_dense_top_single_pass():

    class _Dense:
        def __init__(self):
            self.calls = 0

        def retrieve(self, q, k):
            self.calls += 1
            return [_ranked_result("d1", 0.72)]

    class _Sparse:
        def retrieve(self, q, k):
            return [_ranked_result("d1", 0.4)]

    dense = _Dense()
    fused, top = HybridRetriever(dense, _Sparse()).retrieve_with_dense_top("q", top_k=5, depth=10)
    assert dense.calls == 1
    assert top == 0.72
    assert fused


def test_looks_technical_ignores_trailing_punctuation():
    from scinexusrag.retrieval.hybrid import _looks_technical

    assert _looks_technical("classification.") is False
    assert _looks_technical("immunodeficiency,") is True
    assert _looks_technical("TF-IDF") is True


@pytest.fixture
def local_client(local_settings):
    with TestClient(api.create_app(), base_url="http://localhost") as local_client:
        yield local_client


@pytest.mark.parametrize(
    "document_id",
    [
        "bad id",
        "bad.id",
        "bad$id",
        "paper\\notes",
        "pápér",
        "bad\x00id",
        "a" * 65,
        "_index",
        "_INDEX",
        "_InDeX",
    ],
)
def test_invalid_document_ids_do_not_reach_pipeline(local_client, monkeypatch, document_id):
    def unexpected_pipeline():
        pytest.fail("Invalid document IDs must be rejected before loading the pipeline.")

    monkeypatch.setattr(routes, "get_scinexusrag", unexpected_pipeline)

    response = local_client.delete(f"/api/documents/{quote(document_id, safe='')}")

    assert response.status_code == 400
    assert "Invalid document ID" in response.json()["detail"]


async def test_empty_document_id_does_not_reach_pipeline(monkeypatch):
    def unexpected_pipeline():
        pytest.fail("An empty document ID must not load the pipeline.")

    monkeypatch.setattr(routes, "get_scinexusrag", unexpected_pipeline)

    with pytest.raises(HTTPException) as error:
        await routes.delete_document("")

    assert error.value.status_code == 400


def test_corrupt_metadata_is_reported_as_server_error(local_client, monkeypatch, tmp_path):
    settings = Settings(data_dir=tmp_path)
    settings.storage.lancedb_path = tmp_path / "lancedb"
    rag = NexusRAG(settings)
    index_path = rag.document_store.path / "_index.json"
    corrupt_index = '{"private_document_metadata":'
    index_path.write_text(corrupt_index, encoding="utf-8")
    monkeypatch.setattr(routes, "get_scinexusrag", lambda: rag)

    response = local_client.delete("/api/documents/valid_document")

    assert response.status_code == 500
    assert response.json() == {"detail": "Failed to delete document"}
    assert index_path.read_text(encoding="utf-8") == corrupt_index
    assert rag._embedder is None
    assert rag._vector_store is None
    assert rag._bm25 is None
    assert rag._llm is None
    assert not settings.storage.lancedb_path.exists()


@pytest.fixture
def failure_settings(tmp_path, local_settings):
    failure_settings = local_settings
    failure_settings.data_dir = tmp_path
    failure_settings.storage.lancedb_path = tmp_path / "lancedb"
    return failure_settings


@pytest.fixture
def failure_client(failure_settings):
    with TestClient(api.create_app(), base_url="http://localhost") as failure_client:
        yield failure_client


@pytest.mark.parametrize("path", ["/api/health", "/api/metrics"])
def test_corrupt_document_index_is_not_reported_as_available(
    failure_settings, failure_client, monkeypatch, path
):
    documents = failure_settings.data_dir / "documents"
    documents.mkdir()
    (documents / "_index.json").write_text('{"private_document":', encoding="utf-8")
    pipeline = NexusRAG(failure_settings)
    monkeypatch.setattr(routes, "get_scinexusrag", lambda: pipeline)

    response = failure_client.get(path)

    if path == "/api/health":
        assert response.status_code == 503
        assert response.json()["status"] == "error"
        assert response.json()["llm_available"] is False
    else:
        assert response.status_code == 500
        assert response.json() == {"detail": "Failed to retrieve metrics"}
    assert "private_document" not in response.text
    assert str(failure_settings.data_dir) not in response.text


@pytest.mark.parametrize("path", ["/api/health", "/api/metrics"])
@pytest.mark.parametrize("failure_stage", ["initialization", "statistics"])
def test_unreadable_pipeline_does_not_return_success(
    failure_client, monkeypatch, path, failure_stage
):
    def fail():
        raise OSError("private storage details")

    pipeline = SimpleNamespace(get_stats=fail)
    factory = fail if failure_stage == "initialization" else lambda: pipeline
    monkeypatch.setattr(routes, "get_scinexusrag", factory)

    response = failure_client.get(path)

    if path == "/api/health":
        assert response.status_code == 503
        assert response.json()["status"] == "error"
        assert response.json()["llm_available"] is False
    else:
        assert response.status_code == 500
        assert response.json() == {"detail": "Failed to retrieve metrics"}
    assert "private storage details" not in response.text


def test_metrics_collector_failure_returns_generic_error(failure_client, monkeypatch):
    def fail():
        raise RuntimeError("private collector details")

    monkeypatch.setattr(routes, "get_metrics_collector", lambda: SimpleNamespace(snapshot=fail))

    response = failure_client.get("/api/metrics")

    assert response.status_code == 500
    assert response.json() == {"detail": "Failed to retrieve metrics"}
    assert "private collector details" not in response.text


def test_unavailable_llm_keeps_valid_pipeline_health_visible(failure_client, monkeypatch):
    statistics = SystemStats(
        total_documents=3,
        total_chunks=12,
        total_words=1200,
        storage_path="unused",
        llm_model="fixture-generator",
        embedding_model="fixture-embedder",
        llm_available=False,
    )
    monkeypatch.setattr(
        routes, "get_scinexusrag", lambda: SimpleNamespace(get_stats=lambda: statistics)
    )

    response = failure_client.get("/api/health")

    assert response.status_code == 200
    assert response.json() == {
        "status": "ok",
        "llm_available": False,
        "llm_model": "fixture-generator",
        "embedding_model": "fixture-embedder",
        "total_documents": 3,
        "total_chunks": 12,
    }


@pytest.fixture
def pipeline(monkeypatch):
    pipeline = MagicMock(spec=NexusRAG)
    pipeline.query.return_value = RAGResponse(
        answer="Answer [2].",
        sources=[Source(index=2, chunk_id="chunk", document_id="doc", content="Evidence.")],
        confidence=0.625,
        processing_time_ms=12.36,
        warnings=["Check the cited passage."],
    )
    pipeline.list_documents.return_value = []
    monkeypatch.setattr(routes, "get_scinexusrag", lambda: pipeline)
    return pipeline


@pytest.fixture
def source_client(pipeline, local_settings):
    with TestClient(api.create_app(), base_url="http://localhost") as source_client:
        yield source_client


@pytest.mark.parametrize(
    "metadata,source_name,expected",
    [
        (
            [{"id": "doc", "original_filename": "papers/metadata.pdf", "filename": "stored.pdf"}],
            "embedded.pdf",
            "metadata.pdf",
        ),
        (
            [{"id": "doc", "filename": r"C:\papers\metadata.pdf"}],
            "embedded.pdf",
            "metadata.pdf",
        ),
        ([], "papers/embedded.pdf", "embedded.pdf"),
        ([], r"C:\papers/mixed\embedded.pdf", "embedded.pdf"),
        ([{"id": "doc"}], "embedded.pdf", "embedded.pdf"),
        ([{"id": "doc", "original_filename": ""}], "embedded.pdf", "embedded.pdf"),
        ([{"id": "doc", "original_filename": "/"}], "embedded.pdf", "embedded.pdf"),
        ([{"id": "doc", "filename": "\\"}], "/", "Unknown"),
        ([], "", "Unknown"),
    ],
)
def test_source_filename_uses_metadata_then_embedded_name(
    source_client, pipeline, metadata, source_name, expected
):
    pipeline.list_documents.return_value = metadata
    pipeline.query.return_value.sources[0].document_name = source_name

    response = source_client.post("/api/query", json={"question": "Evidence?"})

    assert response.status_code == 200
    assert response.json()["sources"][0]["filename"] == expected


@pytest.mark.parametrize("length", [500, 501])
def test_source_preview_preserves_wire_fields_and_truncation_boundary(
    source_client, pipeline, length
):
    source = pipeline.query.return_value.sources[0]
    source.content = "x" * length
    source.document_name = "paper.pdf"
    source.score = 0.87654

    response = source_client.post("/api/query", json={"question": "Evidence?"})

    assert response.status_code == 200
    preview = "x" * 500 + ("..." if length > 500 else "")
    assert response.json() == {
        "answer": "Answer [2].",
        "confidence": 0.625,
        "sources": [
            {
                "index": 2,
                "content": preview,
                "text": preview,
                "truncated": length > 500,
                "filename": "paper.pdf",
                "document_id": "doc",
                "section_title": "",
                "page": None,
                "page_number": None,
                "score": 0.877,
            }
        ],
        "processing_time_ms": 12.4,
        "warnings": ["Check the cited passage."],
    }


def test_sources_keep_sparse_citation_indices_and_page_aliases(source_client, pipeline):
    pipeline.query.return_value.answer = "Answer [2] [4]."
    pipeline.query.return_value.sources.append(
        Source(
            index=4,
            chunk_id="another-chunk",
            document_id="another-doc",
            content="Additional evidence.",
            document_name="another.pdf",
            section_title="Results",
            page_number=7,
        )
    )

    response = source_client.post("/api/query", json={"question": "Evidence?"})

    assert response.status_code == 200
    data = response.json()
    assert data["answer"] == "Answer [2] [4]."
    assert [source["index"] for source in data["sources"]] == [2, 4]
    assert data["sources"][1]["page"] == data["sources"][1]["page_number"] == 7
    assert data["sources"][1]["section_title"] == "Results"


def test_metadata_read_failure_keeps_answer_and_embedded_source_name(source_client, pipeline):
    pipeline.list_documents.side_effect = OSError("private metadata details")
    pipeline.query.return_value.sources[0].document_name = r"C:\papers\embedded.pdf"

    response = source_client.post("/api/query", json={"question": "Evidence?"})

    assert response.status_code == 200
    assert response.json()["answer"] == "Answer [2]."
    assert response.json()["sources"][0]["filename"] == "embedded.pdf"
    assert response.json()["warnings"] == ["Check the cited passage."]
    assert "private metadata details" not in response.text


@pytest.mark.parametrize(
    "error,status_code,detail",
    [
        (EmbeddingConfigError("private metadata details"), 409, "EMBEDDING_REVISION"),
        (VectorStoreError("private metadata details"), 503, "vector index is unavailable"),
    ],
)
def test_metadata_storage_failures_are_not_hidden(
    source_client, pipeline, error, status_code, detail
):
    pipeline.list_documents.side_effect = error

    response = source_client.post("/api/query", json={"question": "Evidence?"})

    assert response.status_code == status_code
    assert detail in response.json()["detail"]
    assert "answer" not in response.json()
    assert "private metadata details" not in response.text


@pytest.fixture
def upload_app(monkeypatch, local_settings):
    settings = local_settings
    settings.api.max_upload_mb = 1
    monkeypatch.setattr(security, "_MAGIC", None)

    def ingest(data, filename, extension):
        return IngestResult(
            document_id="uploaded_document",
            filename=filename,
            chunk_count=1,
            word_count=1,
            success=True,
        )

    pipeline = Mock(spec=NexusRAG)
    pipeline.ingest_bytes.side_effect = ingest
    factory = Mock(return_value=pipeline)
    monkeypatch.setattr(routes, "get_scinexusrag", factory)
    return SimpleNamespace(app=api.create_app(), pipeline=pipeline, factory=factory)


@pytest.mark.parametrize("control", [b"\x00", b"\t", b"\x1f", b"\x7f"])
def test_control_characters_in_upload_names_are_rejected(upload_app, control):
    body = (
        b'--boundary\r\nContent-Disposition: form-data; name="file"; filename="bad'
        + control
        + b'.txt"\r\nContent-Type: text/plain\r\n\r\nEvidence.\r\n--boundary--\r\n'
    )
    with TestClient(upload_app.app, base_url="http://localhost") as client:
        response = client.post(
            "/api/ingest",
            content=body,
            headers={"Content-Type": "multipart/form-data; boundary=boundary"},
        )

    assert response.status_code == 400
    upload_app.factory.assert_not_called()


@pytest.mark.parametrize(
    "filename",
    ["txt", "md", ".txt", "paper.exe", "folder/", "C:\\folder\\", "a" * 256 + ".txt"],
)
def test_invalid_upload_names_do_not_reach_pipeline(upload_app, filename):
    with TestClient(upload_app.app, base_url="http://localhost") as client:
        response = client.post(
            "/api/ingest", files={"file": (filename, b"Evidence.", "text/plain")}
        )

    assert response.status_code == 400
    upload_app.factory.assert_not_called()


@pytest.mark.parametrize(
    "filename, basename, extension",
    [
        ("../../papers/Study.TXT", "Study.TXT", ".txt"),
        (r"C:\papers\Study.MD", "Study.MD", ".md"),
    ],
)
def test_upload_paths_use_basename_and_case_insensitive_extension(
    upload_app, filename, basename, extension
):
    with TestClient(upload_app.app, base_url="http://localhost") as client:
        response = client.post(
            "/api/ingest", files={"file": (filename, b"Evidence.", "text/plain")}
        )

    assert response.status_code == 200
    assert response.json()["filename"] == basename
    upload_app.pipeline.ingest_bytes.assert_called_once_with(b"Evidence.", basename, extension)


@pytest.mark.parametrize("extra_bytes", [0, 1])
def test_upload_size_boundary(upload_app, extra_bytes):
    content = b"x" * (1024 * 1024 + extra_bytes)
    with TestClient(upload_app.app, base_url="http://localhost") as client:
        response = client.post("/api/ingest", files={"file": ("paper.txt", content, "text/plain")})

    if extra_bytes:
        assert response.status_code == 413
        upload_app.factory.assert_not_called()
    else:
        assert response.status_code == 200
        upload_app.pipeline.ingest_bytes.assert_called_once_with(content, "paper.txt", ".txt")


async def test_upload_validation_does_not_block_health_requests(upload_app, monkeypatch):
    entered = threading.Event()
    release = threading.Event()

    def validate(*args):
        entered.set()
        release.wait()

    monkeypatch.setattr(routes, "validate_upload", validate)
    fail_safe = threading.Timer(5, release.set)
    fail_safe.daemon = True
    fail_safe.start()
    transport = httpx.ASGITransport(app=upload_app.app)
    async with httpx.AsyncClient(transport=transport, base_url="http://localhost") as client:
        upload = asyncio.create_task(
            client.post("/api/ingest", files={"file": ("paper.txt", b"Evidence.", "text/plain")})
        )
        try:
            assert await asyncio.to_thread(entered.wait, 2)
            assert not release.is_set()

            health = await asyncio.wait_for(client.get("/health"), timeout=1)

            assert health.status_code == 200
            assert health.json() == {"status": "ok"}
            assert not upload.done()
            upload_app.factory.assert_not_called()
        finally:
            release.set()
            fail_safe.cancel()
            response = await asyncio.wait_for(upload, timeout=2)

    assert response.status_code == 200
    upload_app.pipeline.ingest_bytes.assert_called_once_with(b"Evidence.", "paper.txt", ".txt")


MAX_QUERY_LENGTH = get_settings().retrieval.max_query_length


def _docx_bytes() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("[Content_Types].xml", "<?xml version='1.0'?><Types/>")
    return buf.getvalue()


PDF_BYTES = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\ntest content"


DOCX_BYTES = _docx_bytes()


@pytest.fixture
def app():
    test_app = FastAPI()
    test_app.include_router(router)
    return test_app


@pytest.fixture
def route_client(app):
    return TestClient(app, base_url="http://localhost")


@pytest.fixture
def route_rag(monkeypatch):
    mock = MagicMock(spec=NexusRAG)

    mock.get_stats.return_value = SystemStats(
        total_documents=2,
        total_chunks=10,
        total_words=5000,
        storage_path="/tmp/scinexusrag",
        llm_model="llama3.2:3b",
        embedding_model="all-MiniLM-L6-v2",
        llm_available=True,
    )

    mock.list_documents.return_value = []

    def document_inventory():
        documents = mock.list_documents()
        return DocumentInventory(
            documents=documents,
            total_chunks=sum(doc.get("chunk_count", 0) for doc in documents),
        )

    mock.get_document_inventory.side_effect = document_inventory
    monkeypatch.setattr("scinexusrag.api.routes.get_scinexusrag", lambda: mock)

    return mock


@pytest.fixture
def upload_limit_bytes(monkeypatch):
    monkeypatch.setattr(get_settings().api, "max_upload_mb", 1)
    return 1024 * 1024


class TestHealthCheck:
    def test_health_check_success(self, route_client, route_rag):
        response = route_client.get("/api/health")

        assert response.status_code == 200
        data = response.json()

        assert data["status"] == "ok"
        assert data["llm_available"] is True
        assert data["llm_model"] == "llama3.2:3b"
        assert data["embedding_model"] == "all-MiniLM-L6-v2"
        assert data["total_documents"] == 2
        assert data["total_chunks"] == 10

    def test_health_check_with_error(self, route_client):
        with patch("scinexusrag.api.routes.get_scinexusrag", side_effect=Exception("Test error")):
            response = route_client.get("/api/health")

            assert response.status_code == 503
            data = response.json()

            assert data["status"] == "error"
            assert data["llm_available"] is False


class TestIngestDocument:
    @pytest.mark.parametrize(
        "extension,content_type,content",
        [
            (".pdf", "application/pdf", PDF_BYTES),
            (".txt", "text/plain", b"Document text."),
            (
                ".docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                DOCX_BYTES,
            ),
            (".md", "text/markdown", b"# Document\n\nDocument text."),
        ],
    )
    def test_supported_upload_reaches_ingestion(
        self, route_client, route_rag, extension, content_type, content
    ):
        filename = f"paper{extension}"
        route_rag.ingest_bytes.return_value = IngestResult(
            document_id="uploaded_document",
            filename=filename,
            chunk_count=5,
            word_count=1000,
            success=True,
        )

        response = route_client.post(
            "/api/ingest", files={"file": (filename, content, content_type)}
        )

        assert response.status_code == 200
        assert response.json() == {
            "document_id": "uploaded_document",
            "filename": filename,
            "chunk_count": 5,
            "word_count": 1000,
            "success": True,
            "error": None,
        }
        route_rag.ingest_bytes.assert_called_once_with(content, filename, extension)

    @pytest.mark.parametrize(
        "filename, content, content_type, status, message",
        [
            ("", b"content", "text/plain", 422, None),
            (
                "a" * (MAX_FILENAME_LENGTH + 1) + ".txt",
                b"content",
                "text/plain",
                400,
                "Filename too long",
            ),
            ("test.xyz", b"content", "application/octet-stream", 400, "Unsupported file type"),
            ("test.txt", b"", "text/plain", 400, "File is empty"),
        ],
        ids=["missing-name", "long-name", "unsupported-format", "empty-file"],
    )
    def test_invalid_upload_is_rejected(
        self, route_client, route_rag, filename, content, content_type, status, message
    ):
        response = route_client.post(
            "/api/ingest", files={"file": (filename, content, content_type)}
        )

        assert response.status_code == status
        if message is not None:
            assert message in response.json()["detail"]
        route_rag.ingest_bytes.assert_not_called()

    def test_ingest_file_too_large(self, route_client, route_rag, upload_limit_bytes):
        oversized_content = b"x" * (upload_limit_bytes + 1)

        response = route_client.post(
            "/api/ingest",
            files={"file": ("large.pdf", oversized_content, "application/pdf")},
        )

        assert response.status_code == 413
        assert "File too large" in response.json()["detail"]
        assert "1MB" in response.json()["detail"]

    def test_ingest_file_at_size_limit(self, route_client, route_rag, upload_limit_bytes):
        route_rag.ingest_bytes.return_value = IngestResult(
            document_id="doc_large",
            filename="large.pdf",
            chunk_count=100,
            word_count=50000,
            success=True,
        )

        content_at_limit = b"%PDF-1.4\n" + b"x" * (upload_limit_bytes - 9)

        response = route_client.post(
            "/api/ingest",
            files={"file": ("large.pdf", content_at_limit, "application/pdf")},
        )

        assert response.status_code == 200
        assert response.json()["success"] is True

    def test_ingest_path_traversal_in_filename(self, route_client, route_rag):
        route_rag.ingest_bytes.return_value = IngestResult(
            document_id="doc_traversal",
            filename="test.txt",
            chunk_count=1,
            word_count=100,
            success=True,
        )

        response = route_client.post(
            "/api/ingest",
            files={"file": ("../../etc/passwd.txt", b"content", "text/plain")},
        )

        assert response.status_code == 200

        route_rag.ingest_bytes.assert_called_once_with(b"content", "passwd.txt", ".txt")

    def test_ingest_ingestion_failure(self, route_client, route_rag):
        route_rag.ingest_bytes.return_value = IngestResult(
            document_id="",
            filename="test.txt",
            chunk_count=0,
            word_count=0,
            success=False,
            error="Failed to parse document",
        )

        response = route_client.post(
            "/api/ingest",
            files={"file": ("test.txt", b"invalid content", "text/plain")},
        )

        assert response.status_code == 200
        data = response.json()
        assert data["success"] is False
        assert data["error"] == "Failed to parse document"


class TestQueryDocuments:
    def test_query_success(self, route_client, route_rag):
        source = Source(
            index=1,
            chunk_id="chunk_1",
            document_id="doc_123",
            document_name="test.pdf",
            content="Relevant content about the topic",
            section_title="Introduction",
            page_number=1,
            score=0.95,
        )

        route_rag.query.return_value = RAGResponse(
            answer="Paris is the capital of France.",
            sources=[source],
            confidence=0.92,
            processing_time_ms=150.5,
        )
        route_rag.list_documents.return_value = [
            {"id": "doc_123", "filename": "test.pdf", "original_filename": "test.pdf"}
        ]

        response = route_client.post(
            "/api/query", json={"question": "What is the capital of France?"}
        )

        assert response.status_code == 200
        data = response.json()

        assert data["answer"] == "Paris is the capital of France."
        assert data["confidence"] == 0.92
        assert len(data["sources"]) == 1
        assert data["sources"][0]["filename"] == "test.pdf"
        assert data["processing_time_ms"] == 150.5
        route_rag.query.assert_called_once_with("What is the capital of France?")

    def test_query_surfaces_warnings(self, route_client, route_rag):
        route_rag.query.return_value = RAGResponse(
            answer="Ungrounded answer.",
            sources=[],
            confidence=0.3,
            processing_time_ms=12.0,
            warnings=["answer cites no sources"],
        )
        route_rag.list_documents.return_value = []

        data = route_client.post("/api/query", json={"question": "q?"}).json()

        assert data["warnings"] == ["answer cites no sources"]

    @pytest.mark.parametrize(
        "payload, message",
        [
            ({"question": ""}, "cannot be empty"),
            ({"question": "   \n\t   "}, None),
            ({"question": "a" * (MAX_QUERY_LENGTH + 1)}, "too long"),
            ({}, None),
        ],
        ids=["empty", "whitespace", "too-long", "missing"],
    )
    def test_invalid_query_is_rejected(self, route_client, route_rag, payload, message):
        response = route_client.post("/api/query", json=payload)

        assert response.status_code == 422
        if message is not None:
            assert message in response.json()["detail"][0]["msg"].lower()
        route_rag.query.assert_not_called()

    def test_query_at_max_length(self, route_client, route_rag):
        max_length_question = "a" * MAX_QUERY_LENGTH

        route_rag.query.return_value = RAGResponse(
            answer="Answer",
            sources=[],
            confidence=0.5,
            processing_time_ms=100.0,
        )

        response = route_client.post("/api/query", json={"question": max_length_question})

        assert response.status_code == 200

    def test_query_no_documents(self, route_client, route_rag):
        route_rag.query.return_value = RAGResponse(
            answer="No documents have been uploaded yet.",
            sources=[],
            confidence=0.0,
            processing_time_ms=10.0,
        )
        route_rag.list_documents.return_value = []

        response = route_client.post("/api/query", json={"question": "What is the answer?"})

        assert response.status_code == 200
        data = response.json()

        assert data["sources"] == []
        assert data["processing_time_ms"] == 10.0

    def test_query_multiple_sources(self, route_client, route_rag):
        sources = [
            Source(
                index=1,
                chunk_id="chunk_1",
                document_id="doc_1",
                document_name="paper1.pdf",
                content="Content from paper 1",
                section_title="Methods",
                page_number=5,
                score=0.95,
            ),
            Source(
                index=2,
                chunk_id="chunk_2",
                document_id="doc_2",
                document_name="paper2.pdf",
                content="Content from paper 2",
                section_title="Results",
                page_number=10,
                score=0.87,
            ),
        ]

        route_rag.query.return_value = RAGResponse(
            answer="Combined answer from both sources",
            sources=sources,
            confidence=0.91,
            processing_time_ms=200.0,
        )
        route_rag.list_documents.return_value = [
            {"id": "doc_1", "filename": "paper1.pdf", "original_filename": "paper1.pdf"},
            {"id": "doc_2", "filename": "paper2.pdf", "original_filename": "paper2.pdf"},
        ]

        response = route_client.post("/api/query", json={"question": "What did the research show?"})

        assert response.status_code == 200
        data = response.json()
        assert len(data["sources"]) == 2
        assert data["sources"][0]["score"] == 0.95
        assert data["sources"][1]["score"] == 0.87

    def test_query_preserves_source_indices(self, route_client, route_rag):
        sources = [
            Source(index=index, chunk_id=f"chunk_{index}", document_id="doc", content="Evidence.")
            for index in (2, 4)
        ]
        route_rag.query.return_value = RAGResponse("Answer [2] [4].", sources, 0.5)

        response = route_client.post("/api/query", json={"question": "What was measured?"})

        assert response.status_code == 200
        assert [source["index"] for source in response.json()["sources"]] == [2, 4]

    def test_query_long_source_content_truncated(self, route_client, route_rag):
        long_content = "x" * 1000
        source = Source(
            index=1,
            chunk_id="chunk_1",
            document_id="doc_123",
            document_name="test.pdf",
            content=long_content,
            section_title="Section",
            page_number=1,
            score=0.9,
        )

        route_rag.query.return_value = RAGResponse(
            answer="Answer",
            sources=[source],
            confidence=0.9,
            processing_time_ms=100.0,
        )
        route_rag.list_documents.return_value = []

        response = route_client.post("/api/query", json={"question": "What is this?"})

        assert response.status_code == 200
        data = response.json()
        assert len(data["sources"][0]["content"]) <= 503
        assert data["sources"][0]["content"].endswith("...")


class TestListDocuments:
    def test_list_documents_empty(self, route_client, route_rag):
        route_rag.list_documents.return_value = []

        response = route_client.get("/api/documents")

        assert response.status_code == 200
        data = response.json()

        assert data["documents"] == []
        assert data["total_documents"] == 0
        assert data["total_chunks"] == 0

    def test_list_documents_with_documents(self, route_client, route_rag):
        route_rag.list_documents.return_value = [
            {
                "id": "doc_1",
                "filename": "paper1.pdf",
                "original_filename": "paper1.pdf",
                "word_count": 1000,
                "chunk_count": 5,
                "file_type": "pdf",
                "uploaded_at": "2024-01-15T10:30:00",
            },
            {
                "id": "doc_2",
                "filename": "paper2.pdf",
                "original_filename": "paper2.pdf",
                "word_count": 800,
                "chunk_count": 4,
                "file_type": "pdf",
                "uploaded_at": "2024-01-15T11:00:00",
            },
        ]

        response = route_client.get("/api/documents")

        assert response.status_code == 200
        data = response.json()

        assert len(data["documents"]) == 2
        assert data["total_documents"] == 2
        assert data["total_chunks"] == 9
        assert data["documents"][0]["id"] == "doc_1"
        assert data["documents"][0]["word_count"] == 1000
        assert data["documents"][1]["id"] == "doc_2"

    def test_list_documents_pagination(self, route_client, route_rag):
        route_rag.list_documents.return_value = [
            {"id": f"doc_{i}", "filename": f"p{i}.pdf", "word_count": 10} for i in range(5)
        ]

        page = route_client.get("/api/documents?limit=2&offset=2").json()

        assert [d["id"] for d in page["documents"]] == ["doc_2", "doc_3"]
        assert page["total_documents"] == 5

        assert route_client.get("/api/documents?limit=0").status_code == 422
        assert route_client.get("/api/documents?offset=-1").status_code == 422


class TestGetStatus:
    def test_get_status_success(self, route_client, route_rag):
        response = route_client.get("/api/status")

        assert response.status_code == 200
        data = response.json()

        assert data["total_documents"] == 2
        assert data["total_chunks"] == 10
        assert data["total_words"] == 5000
        assert data["llm_available"] is True
        assert data["llm_model"] == "llama3.2:3b"
        assert data["embedding_model"] == "all-MiniLM-L6-v2"


class TestDeleteDocument:
    def test_delete_document_success(self, route_client, route_rag):
        route_rag.delete_document.return_value = True

        response = route_client.delete("/api/documents/doc_123")

        assert response.status_code == 200
        data = response.json()

        assert data["success"] is True
        assert data["message"] == "Document deleted"
        route_rag.delete_document.assert_called_once_with("doc_123")

    def test_delete_document_not_found(self, route_client, route_rag):
        route_rag.delete_document.return_value = False

        response = route_client.delete("/api/documents/nonexistent")

        assert response.status_code == 404
        assert "Document not found" in response.json()["detail"]

    def test_trailing_slash_uses_clear_all_route(self, route_client, route_rag):
        response = route_client.delete("/api/documents/")

        assert response.status_code == 200
        route_rag.clear_all.assert_called_once()

    def test_delete_document_id_too_long(self, route_client, route_rag):
        long_id = "a" * 65

        response = route_client.delete(f"/api/documents/{long_id}")

        assert response.status_code == 400
        assert "Invalid document ID" in response.json()["detail"]


class TestClearAllDocuments:
    def test_clear_all_documents_success(self, route_client, route_rag):
        response = route_client.delete("/api/documents")

        assert response.status_code == 200
        data = response.json()

        assert data["success"] is True
        assert data["message"] == "All documents cleared"
        route_rag.clear_all.assert_called_once()


class TestPydanticValidation:
    @pytest.mark.parametrize(
        "question, message",
        [
            ("", "Question cannot be empty"),
            ("   \n\t   ", "Question cannot be empty"),
            ("a" * (MAX_QUERY_LENGTH + 1), "too long"),
        ],
        ids=["empty", "whitespace", "too-long"],
    )
    def test_invalid_query_request(self, question, message):
        with pytest.raises(ValidationError, match=message):
            QueryRequest(question=question)

    @pytest.mark.parametrize(
        "question, expected",
        [
            ("What is the meaning of life?", "What is the meaning of life?"),
            ("  What is this?  \n", "What is this?"),
        ],
        ids=["unchanged", "trimmed"],
    )
    def test_valid_query_request(self, question, expected):
        assert QueryRequest(question=question).question == expected


class TestRouteErrors:
    @pytest.mark.parametrize(
        "method, path, operation, error, message",
        [
            (
                "POST",
                "/api/ingest",
                "ingest_bytes",
                Exception("Database error"),
                "Failed to process document",
            ),
            ("POST", "/api/query", "query", Exception("LLM error"), "Failed to process query"),
            (
                "GET",
                "/api/documents",
                "list_documents",
                Exception("Storage error"),
                "Failed to retrieve documents",
            ),
            (
                "GET",
                "/api/status",
                "get_stats",
                Exception("Database error"),
                "Failed to retrieve status",
            ),
            (
                "DELETE",
                "/api/documents/doc_123",
                "delete_document",
                Exception("Deletion error"),
                "Failed to delete document",
            ),
            (
                "DELETE",
                "/api/documents/invalid-id-format",
                "delete_document",
                ValueError("Invalid stored metadata"),
                "Failed to delete document",
            ),
            (
                "DELETE",
                "/api/documents",
                "clear_all",
                Exception("Clear error"),
                "Failed to clear documents",
            ),
        ],
        ids=["ingest", "query", "inventory", "status", "delete", "invalid-metadata", "clear"],
    )
    def test_backend_errors_are_reported(
        self, route_client, route_rag, method, path, operation, error, message
    ):
        getattr(route_rag, operation).side_effect = error
        kwargs = {}
        if operation == "ingest_bytes":
            kwargs["files"] = {"file": ("test.txt", b"content", "text/plain")}
        elif operation == "query":
            kwargs["json"] = {"question": "What is the answer?"}

        response = route_client.request(method, path, **kwargs)

        assert response.status_code == 500
        assert response.json()["detail"] == message

    def test_missing_content_type_header(self, route_client, route_rag):
        route_rag.ingest_bytes.return_value = IngestResult(
            document_id="doc_1",
            filename="test.txt",
            chunk_count=1,
            word_count=100,
            success=True,
        )

        response = route_client.post(
            "/api/ingest",
            content=(
                b'--boundary\r\nContent-Disposition: form-data; name="file"; '
                b'filename="test.txt"\r\n\r\ncontent\r\n--boundary--\r\n'
            ),
            headers={"Content-Type": "multipart/form-data; boundary=boundary"},
        )

        assert response.status_code == 200
        route_rag.ingest_bytes.assert_called_once_with(b"content", "test.txt", ".txt")

    def test_invalid_json_in_query(self, route_client, route_rag):
        response = route_client.post(
            "/api/query",
            content="{invalid json}",
            headers={"Content-Type": "application/json"},
        )

        assert response.status_code == 422


@pytest.fixture
def application_client(route_rag):
    return TestClient(create_app(), base_url="http://localhost")


class TestEmbeddingConfigErrors:
    @staticmethod
    def assert_conflict(response):
        assert response.status_code == 409
        detail = response.json()["detail"]
        assert "EMBEDDING_REVISION" in detail
        assert "DATA_DIR" in detail
        assert "LANCEDB_PATH" in detail
        assert "original documents" in detail
        assert "private embedding metadata" not in detail

    @pytest.mark.parametrize(
        "path", ["/api/health", "/api/status", "/api/documents", "/api/metrics"]
    )
    def test_count_conflict_is_not_reported_as_empty_or_healthy(
        self, application_client, route_rag, path
    ):
        method = "get_document_inventory" if path == "/api/documents" else "get_stats"
        getattr(route_rag, method).side_effect = EmbeddingConfigError("private embedding metadata")

        self.assert_conflict(application_client.get(path))

    @pytest.mark.parametrize("failed_method", ["query", "list_documents"])
    def test_query_conflict_is_actionable_and_does_not_expose_metadata(
        self, application_client, route_rag, failed_method
    ):
        route_rag.query.return_value = RAGResponse("Answer.", [], 0.0)
        getattr(route_rag, failed_method).side_effect = EmbeddingConfigError(
            "private embedding metadata"
        )

        response = application_client.post("/api/query", json={"question": "What was measured?"})

        self.assert_conflict(response)

    def test_ingest_conflict_returns_409(self, application_client, route_rag):
        route_rag.ingest_bytes.side_effect = EmbeddingConfigError("private embedding metadata")

        response = application_client.post(
            "/api/ingest", files={"file": ("paper.txt", b"Sensor readings.", "text/plain")}
        )

        self.assert_conflict(response)

    def test_real_ingest_rejects_same_dimension_model_change(self, tmp_path):
        settings = Settings(
            data_dir=tmp_path,
            storage=StorageSettings(lancedb_path=tmp_path / "lancedb"),
            embedding=EmbeddingSettings(model="test-model", revision="a" * 40),
        )
        config = EmbeddingConfig("test-model", "a" * 40, "", "", True, 4)
        original_embedder = MagicMock(spec=Embedder)
        original_embedder.embedding_config = config
        original_embedder.embed.return_value = np.array([[1, 0, 0, 0]], dtype=np.float32)
        original = NexusRAG(settings)
        original._embedder = original_embedder
        result = original.ingest_bytes(
            b"Sensor A records temperature every thirty seconds.", "original.txt", ".txt"
        )
        assert result.success
        assert original.vector_store.count() == 1

        changed = NexusRAG(settings)
        changed_embedder = MagicMock(spec=Embedder)
        changed_embedder.embedding_config = replace(config, model_name="different-model")
        changed._embedder = changed_embedder
        with (
            patch("scinexusrag.api.routes.get_scinexusrag", return_value=changed),
            TestClient(create_app(), base_url="http://localhost") as route_client,
        ):
            self.assert_conflict(
                route_client.post(
                    "/api/ingest",
                    files={"file": ("new.txt", b"Sensor B records pressure.", "text/plain")},
                )
            )
            self.assert_conflict(
                route_client.post("/api/query", json={"question": "What did the sensor record?"})
            )

        changed_embedder.embed.assert_not_called()
        changed_embedder.embed_query.assert_not_called()
        assert changed._llm is None
        reopened = NexusRAG(settings)
        reopened._embedder = original_embedder
        assert reopened.vector_store.count() == 1
        assert [document["id"] for document in reopened.list_documents()] == [result.document_id]

    def test_delete_conflict_returns_409(self, application_client, route_rag):
        route_rag.delete_document.side_effect = EmbeddingConfigError("private embedding metadata")

        self.assert_conflict(application_client.delete("/api/documents/paper"))

    def test_explicit_clear_does_not_require_compatible_stats(self, application_client, route_rag):
        route_rag.get_stats.side_effect = EmbeddingConfigError("private embedding metadata")

        response = application_client.delete("/api/documents")

        assert response.status_code == 200
        route_rag.clear_all.assert_called_once()
        route_rag.get_stats.assert_not_called()

    def test_auth_runs_before_index_checks(self, application_client, route_rag):
        route_rag.get_stats.side_effect = EmbeddingConfigError("private embedding metadata")
        with patch(
            "scinexusrag.api.security.get_settings", return_value=_route_settings(api_key="secret")
        ):
            response = application_client.get("/api/health")

        assert response.status_code == 401
        route_rag.get_stats.assert_not_called()


class TestStorageErrors:
    @staticmethod
    def assert_unavailable(response):
        assert response.status_code == 503
        detail = response.json()["detail"]
        assert "vector index is unavailable" in detail
        assert "server logs" in detail
        assert "private database details" not in detail
        assert "No documents" not in detail

    @pytest.mark.parametrize(
        ("method", "path", "failed_method"),
        [
            ("GET", "/api/health", "get_stats"),
            ("GET", "/api/status", "get_stats"),
            ("GET", "/api/documents", "get_document_inventory"),
            ("GET", "/api/metrics", "get_stats"),
            ("POST", "/api/query", "query"),
            ("DELETE", "/api/documents/paper", "delete_document"),
            ("DELETE", "/api/documents", "clear_all"),
        ],
    )
    def test_storage_failure_returns_503(
        self, application_client, route_rag, method, path, failed_method
    ):
        getattr(route_rag, failed_method).side_effect = VectorStoreError("private database details")

        response = application_client.request(
            method, path, json={"question": "What was measured?"} if method == "POST" else None
        )

        self.assert_unavailable(response)

    def test_ingest_storage_failure_returns_503(self, application_client, route_rag):
        route_rag.ingest_bytes.side_effect = VectorStoreError("private database details")

        response = application_client.post(
            "/api/ingest", files={"file": ("paper.txt", b"Sensor readings.", "text/plain")}
        )

        self.assert_unavailable(response)

    def test_failed_real_index_count_is_distinct_from_an_empty_corpus(self, tmp_path, monkeypatch):
        route_rag = NexusRAG(
            Settings(
                data_dir=tmp_path,
                storage=StorageSettings(lancedb_path=tmp_path / "lancedb"),
            )
        )
        embedder = MagicMock(spec=Embedder)
        embedder.embedding_config = EmbeddingConfig("fixture", "a" * 40, "", "", True, 4)
        route_rag._embedder = embedder
        with (
            patch("scinexusrag.api.routes.get_scinexusrag", return_value=route_rag),
            TestClient(create_app(), base_url="http://localhost") as route_client,
        ):
            empty = route_client.post("/api/query", json={"question": "What was measured?"})
            assert empty.status_code == 200
            assert "No documents" in empty.json()["answer"]

            def fail_count():
                raise OSError("private database details")

            monkeypatch.setattr(route_rag.vector_store.table, "count_rows", fail_count)
            self.assert_unavailable(
                route_client.post("/api/query", json={"question": "What was measured?"})
            )

        embedder.embed_query.assert_not_called()
        assert route_rag._orchestrator is None


def _route_settings(**api_overrides) -> Settings:
    settings = Settings()
    for key, value in api_overrides.items():
        setattr(settings.api, key, value)
    return settings


class TestApiKeyAuth:
    @pytest.mark.parametrize("key", [None, "wrong"], ids=["missing", "wrong"])
    def test_invalid_key_is_rejected(self, route_client, route_rag, key):
        with patch(
            "scinexusrag.api.security.get_settings", return_value=_route_settings(api_key="secret")
        ):
            response = route_client.get("/api/status", headers={"X-API-Key": key} if key else {})

        assert response.status_code == 401
        route_rag.get_stats.assert_not_called()

    def test_correct_key_allowed(self, route_client, route_rag):
        with patch(
            "scinexusrag.api.security.get_settings", return_value=_route_settings(api_key="secret")
        ):
            response = route_client.get("/api/status", headers={"X-API-Key": "secret"})
            assert response.status_code == 200


class TestUploadValidationThroughRoute:
    def test_wrong_magic_bytes_rejected(self, route_client, route_rag):
        files = {"file": ("bad.pdf", b"this is not a pdf at all", "application/pdf")}
        response = route_client.post("/api/ingest", files=files)
        assert response.status_code == 415

    def test_zip_bomb_rejected(self, route_client, route_rag):

        with patch(
            "scinexusrag.api.security.get_settings",
            return_value=_route_settings(max_uncompressed_mb=0),
        ):
            files = {
                "file": (
                    "doc.docx",
                    DOCX_BYTES,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            }
            response = route_client.post("/api/ingest", files=files)
            assert response.status_code == 413


def _security_settings(**api_overrides: object) -> Settings:
    settings = Settings()
    for key, value in api_overrides.items():
        setattr(settings.api, key, value)
    return settings


class _MagicDatabaseError(Exception):
    pass


class TestMimeDetectorLoading:
    def test_missing_optional_package_returns_none(self, monkeypatch):
        monkeypatch.setitem(sys.modules, "magic", None)
        assert security._load_magic() is None

    @pytest.mark.parametrize("error", [OSError, _MagicDatabaseError])
    def test_unavailable_native_detector_returns_none(self, monkeypatch, error):
        def create_detector(**kwargs):
            raise error("Native MIME detector unavailable")

        monkeypatch.setitem(
            sys.modules,
            "magic",
            SimpleNamespace(Magic=create_detector, MagicException=_MagicDatabaseError),
        )
        assert security._load_magic() is None

    def test_successful_detector_is_returned(self, monkeypatch):
        detector = SimpleNamespace(from_buffer=lambda data: "text/plain")
        monkeypatch.setitem(
            sys.modules,
            "magic",
            SimpleNamespace(Magic=lambda **kwargs: detector, MagicException=_MagicDatabaseError),
        )
        assert security._load_magic() is detector

    def test_unexpected_initialization_error_is_not_suppressed(self, monkeypatch):
        def create_detector(**kwargs):
            raise RuntimeError("Detector initialization failed")

        monkeypatch.setitem(
            sys.modules,
            "magic",
            SimpleNamespace(Magic=create_detector, MagicException=_MagicDatabaseError),
        )
        with pytest.raises(RuntimeError, match="Detector initialization failed"):
            security._load_magic()


class TestApiKey:
    async def test_no_key_configured_is_open(self, monkeypatch) -> None:
        monkeypatch.setattr(security, "get_settings", lambda: _security_settings(api_key=""))
        assert await security.require_api_key(None) is None

    @pytest.mark.parametrize("key", [None, "nope"], ids=["missing", "wrong"])
    async def test_invalid_key_is_rejected(self, monkeypatch, key):
        monkeypatch.setattr(
            security, "get_settings", lambda: _security_settings(api_key="test-only-key")
        )
        with pytest.raises(HTTPException) as exc:
            await security.require_api_key(key)
        assert exc.value.status_code == 401
        assert exc.value.headers == {"WWW-Authenticate": "ApiKey"}

    async def test_correct_key_accepted(self, monkeypatch) -> None:
        monkeypatch.setattr(
            security, "get_settings", lambda: _security_settings(api_key="test-only-key")
        )
        assert await security.require_api_key("test-only-key") is None

    async def test_non_ascii_key_rejected_not_crash(self, monkeypatch) -> None:
        monkeypatch.setattr(
            security, "get_settings", lambda: _security_settings(api_key="test-only-key")
        )
        with pytest.raises(HTTPException) as exc:
            await security.require_api_key("sécrét-ключ-🔑")
        assert exc.value.status_code == 401


class TestRateLimiter:
    def test_limit_strings_track_settings(self, monkeypatch) -> None:
        monkeypatch.setattr(
            security,
            "get_settings",
            lambda: _security_settings(query_rate_per_minute=42, upload_rate_per_minute=7),
        )
        assert security.query_limit() == "42/minute"
        assert security.upload_limit() == "7/minute"

    def test_slowapi_returns_429_over_limit(self) -> None:
        from slowapi import Limiter, _rate_limit_exceeded_handler
        from slowapi.errors import RateLimitExceeded
        from slowapi.util import get_remote_address

        limiter = Limiter(key_func=get_remote_address)
        app = FastAPI()
        app.state.limiter = limiter
        app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)

        @app.get("/x")
        @limiter.limit("2/minute")
        async def _x(request: Request) -> dict[str, bool]:
            return {"ok": True}

        client = TestClient(app, base_url="http://localhost")
        assert [client.get("/x").status_code for _ in range(3)] == [200, 200, 429]


class TestSameSiteGuard:
    @staticmethod
    def _request(headers: dict[str, str], *, scheme: str = "http", root_path: str = "") -> Request:
        return Request(
            {
                "type": "http",
                "method": "POST",
                "scheme": scheme,
                "server": ("localhost", 8000),
                "path": f"{root_path}/api/query",
                "root_path": root_path,
                "query_string": b"",
                "headers": [
                    (key.lower().encode("ascii"), value.encode("ascii"))
                    for key, value in headers.items()
                ],
            }
        )

    @pytest.mark.parametrize("fetch_site", ["cross-site", "unknown", ""])
    async def test_unsafe_fetch_metadata_rejected(self, fetch_site) -> None:
        with pytest.raises(HTTPException) as exc:
            await security.require_same_site(self._request({"sec-fetch-site": fetch_site}))
        assert exc.value.status_code == 403

    @pytest.mark.parametrize("fetch_site", ["same-origin", "same-site", "none", "Same-Origin"])
    async def test_safe_fetch_metadata_without_origin_allowed(self, fetch_site) -> None:
        assert (
            await security.require_same_site(self._request({"sec-fetch-site": fetch_site})) is None
        )

    async def test_non_browser_client_allowed(self) -> None:
        assert await security.require_same_site(self._request({})) is None

    async def test_foreign_origin_fallback_rejected(self, monkeypatch) -> None:
        monkeypatch.setattr(security, "get_settings", lambda: _security_settings())
        with pytest.raises(HTTPException) as exc:
            await security.require_same_site(self._request({"origin": "http://evil.example"}))
        assert exc.value.status_code == 403

    async def test_allowlisted_origin_fallback_allowed(self, monkeypatch) -> None:
        monkeypatch.setattr(security, "get_settings", lambda: _security_settings())
        allowed = _security_settings().api.cors_origins[0]
        assert await security.require_same_site(self._request({"origin": allowed})) is None

    @pytest.mark.parametrize("fetch_site", ["same-origin", "same-site", "none"])
    async def test_safe_fetch_metadata_does_not_allow_foreign_origin(
        self, monkeypatch, fetch_site
    ) -> None:
        monkeypatch.setattr(security, "get_settings", lambda: _security_settings())
        with pytest.raises(HTTPException) as exc:
            await security.require_same_site(
                self._request({"sec-fetch-site": fetch_site, "origin": "http://evil.example"})
            )
        assert exc.value.status_code == 403

    async def test_unsafe_fetch_metadata_rejects_allowlisted_origin(self, monkeypatch) -> None:
        monkeypatch.setattr(security, "get_settings", lambda: _security_settings())
        with pytest.raises(HTTPException) as exc:
            await security.require_same_site(
                self._request(
                    {
                        "sec-fetch-site": "cross-site",
                        "origin": _security_settings().api.cors_origins[0],
                    }
                )
            )
        assert exc.value.status_code == 403

    @pytest.mark.parametrize(
        "scheme,host,root_path",
        [
            ("http", "localhost:9123", ""),
            ("https", "research.example", ""),
            ("https", "research.example:9443", "/research"),
            ("http", "[::1]:9123", "/research"),
        ],
    )
    async def test_request_origin_allowed_with_custom_host_and_mount(
        self, monkeypatch, scheme, host, root_path
    ) -> None:
        monkeypatch.setattr(security, "get_settings", lambda: _security_settings(cors_origins=[]))
        request = self._request(
            {"host": host, "origin": f"{scheme}://{host}"}, scheme=scheme, root_path=root_path
        )
        assert await security.require_same_site(request) is None

    @pytest.mark.parametrize(
        "origin", ["", "http://localhost:9123", "https://localhost:8000", "null"]
    )
    async def test_different_port_scheme_or_opaque_origin_rejected(self, monkeypatch, origin):
        monkeypatch.setattr(security, "get_settings", lambda: _security_settings(cors_origins=[]))
        with pytest.raises(HTTPException) as exc:
            await security.require_same_site(self._request({"origin": origin}))
        assert exc.value.status_code == 403


class TestValidateUpload:
    @pytest.fixture(autouse=True)
    def upload_settings(self, monkeypatch):
        monkeypatch.setattr(security, "get_settings", lambda: _security_settings())
        monkeypatch.setattr(security, "_MAGIC", None)

    @pytest.fixture
    def contents(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("a.xml", "<x/>")
        return {
            ".pdf": b"%PDF-1.4\nhi",
            ".docx": buffer.getvalue(),
            ".txt": b"Document text.",
            ".md": b"# Document\n\nDocument text.",
        }

    def test_pdf_magic_required_without_mime_detector(self) -> None:
        with pytest.raises(HTTPException) as exc:
            security.validate_upload(".pdf", b"not a pdf", "application/pdf")
        assert exc.value.status_code == 415

    @pytest.mark.parametrize(
        "extension,content_type",
        [
            (".pdf", "application/pdf"),
            (".docx", "application/zip"),
            (".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            (".txt", "text/plain"),
            (".txt", "text/markdown"),
            (".md", "text/plain"),
            (".md", "text/markdown"),
        ],
    )
    def test_matching_content_types_accepted(self, contents, extension, content_type):
        security.validate_upload(extension, contents[extension], content_type)

    @pytest.mark.parametrize("extension", [".pdf", ".docx", ".txt", ".md"])
    @pytest.mark.parametrize("content_type", [None, "", "application/octet-stream"])
    def test_generic_content_types_accepted(self, contents, extension, content_type):
        security.validate_upload(extension, contents[extension], content_type)

    @pytest.mark.parametrize(
        "extension,content_type",
        [
            (".pdf", "text/plain"),
            (".docx", "application/pdf"),
            (".txt", "application/zip"),
            (".md", "application/pdf"),
        ],
    )
    def test_content_type_for_another_extension_rejected(self, contents, extension, content_type):
        with pytest.raises(HTTPException) as exc:
            security.validate_upload(extension, contents[extension], content_type)
        assert exc.value.status_code == 415

    def test_content_type_parameters_and_case_accepted(self):
        security.validate_upload(".md", b"Document text.", " Text/Markdown ; charset=UTF-8 ")

    def test_extension_case_is_normalized(self, contents):
        security.validate_upload(".PDF", contents[".pdf"], "application/pdf")

    def test_unsupported_extension_rejected(self):
        with pytest.raises(HTTPException) as exc:
            security.validate_upload(".zip", b"Document text.", "application/octet-stream")
        assert exc.value.status_code == 415

    def test_disallowed_content_type(self) -> None:
        with pytest.raises(HTTPException) as exc:
            security.validate_upload(".txt", b"hello", "application/x-sh")
        assert exc.value.status_code == 415

    @pytest.mark.parametrize("extension", [".txt", ".md"])
    def test_text_must_be_utf8_without_mime_detector(self, extension) -> None:
        with pytest.raises(HTTPException) as exc:
            security.validate_upload(extension, b"\xff\xfe\x00binary", "text/plain")
        assert exc.value.status_code == 415

    @pytest.mark.parametrize("extra_byte", [False, True])
    def test_file_size_boundary_counts_utf8_bytes(self, monkeypatch, extra_byte) -> None:
        monkeypatch.setattr(security, "get_settings", lambda: _security_settings(max_upload_mb=1))
        text = "é" * (1024 * 1024 // 2) + ("x" if extra_byte else "")
        if extra_byte:
            with pytest.raises(HTTPException) as exc:
                security.validate_upload(".txt", text, "text/plain")
            assert exc.value.status_code == 413
        else:
            security.validate_upload(".txt", text, "text/plain")

    def test_invalid_docx_rejected_without_mime_detector(self) -> None:
        with pytest.raises(HTTPException) as exc:
            security.validate_upload(".docx", b"PK\x03\x04garbage", "application/zip")
        assert exc.value.status_code == 415

    @pytest.mark.parametrize(
        "extension,detected",
        [
            (".pdf", "text/plain"),
            (".docx", "application/pdf"),
            (".txt", "application/zip"),
            (".md", "application/pdf"),
        ],
    )
    def test_sniffed_content_mismatch_rejected(self, monkeypatch, contents, extension, detected):
        monkeypatch.setattr(security, "_MAGIC", SimpleNamespace(from_buffer=lambda data: detected))
        with pytest.raises(HTTPException) as exc:
            security.validate_upload(extension, contents[extension], "application/octet-stream")
        assert exc.value.status_code == 415

    @pytest.mark.parametrize("detected", ["text/x-rst", "TEXT/PLAIN", "inode/x-empty"])
    def test_sniffed_text_variants_accepted(self, monkeypatch, detected):
        monkeypatch.setattr(security, "_MAGIC", SimpleNamespace(from_buffer=lambda data: detected))
        data = b"" if detected == "inode/x-empty" else b"Document text."
        security.validate_upload(".md", data, "text/markdown")

    def test_mime_detector_failure_is_not_suppressed(self, monkeypatch):
        def fail(_data):
            raise OSError("MIME detector failed")

        monkeypatch.setattr(security, "_MAGIC", SimpleNamespace(from_buffer=fail))
        with pytest.raises(OSError, match="MIME detector failed"):
            security.validate_upload(".txt", b"Document text.", "text/plain")

    @pytest.mark.parametrize("extra_byte", [False, True])
    def test_docx_size_boundary_sums_archive_members(self, monkeypatch, extra_byte):
        monkeypatch.setattr(
            security, "get_settings", lambda: _security_settings(max_uncompressed_mb=1)
        )
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("a.xml", b"A" * (1024 * 1024 // 2))
            archive.writestr("b.xml", b"B" * (1024 * 1024 // 2 + int(extra_byte)))
        if extra_byte:
            with pytest.raises(HTTPException) as exc:
                security.validate_upload(".docx", buffer.getvalue(), "application/zip")
            assert exc.value.status_code == 413
        else:
            security.validate_upload(".docx", buffer.getvalue(), "application/zip")


class TestAppConfig:
    def test_cors_default_is_not_wildcard(self) -> None:

        origins = APISettings().cors_origins
        assert "*" not in origins
        assert origins and all(o.startswith("http://") for o in origins)

    def test_docs_disabled_when_api_key_set(self, monkeypatch) -> None:

        monkeypatch.setattr(api, "get_settings", lambda: _security_settings(api_key="k"))
        client = TestClient(api.create_app(), base_url="http://localhost")
        assert client.get("/docs").status_code == 404
        assert client.get("/openapi.json").status_code == 404

    def test_docs_enabled_in_local_mode(self, monkeypatch) -> None:

        monkeypatch.setattr(
            api, "get_settings", lambda: _security_settings(api_key="", docs_enabled=True)
        )
        client = TestClient(api.create_app(), base_url="http://localhost")
        assert client.get("/docs").status_code == 200


class TestIdSanitizer:
    def test_path_traversal_and_injection_rejected(self) -> None:
        from scinexusrag.storage.vector_store import _sanitize_id

        for bad in ["../../etc/passwd", "a/b", "a\\b", "id; DROP TABLE x", "a' OR '1'='1", "a b"]:
            with pytest.raises(ValueError):
                _sanitize_id(bad)

    def test_safe_ids_accepted(self) -> None:
        from scinexusrag.storage.vector_store import _sanitize_id

        for ok in ["doc_123", "chunk-abc", "AaZz09_-"]:
            assert _sanitize_id(ok) == ok

    @pytest.mark.parametrize("length", [129, 200_001])
    def test_oversized_ids_are_rejected_before_pattern_matching(self, monkeypatch, length):
        from scinexusrag.storage import vector_store

        def unexpected_match(_value):
            pytest.fail("Oversized IDs must be rejected before pattern matching.")

        monkeypatch.setattr(
            vector_store, "SAFE_ID_PATTERN", SimpleNamespace(fullmatch=unexpected_match)
        )

        with pytest.raises(ValueError, match="ID too long"):
            vector_store._sanitize_id("a" * (length - 1) + "!")


def _document(content, sectioned):
    document = ParsedDocument(id="limits", content=content, metadata={"filename": "limits.txt"})
    if sectioned:
        document.sections = [Section(title="Methods", content=content, level=1, page_number=4)]
    return document


def _assert_source(chunks, content, sectioned):
    prefix = "[Methods]\n\n" if sectioned else ""
    words = " ".join(chunk.content.removeprefix(prefix) for chunk in chunks).split()
    assert words == content.split()
    if sectioned:
        assert all(chunk.section_title == "Methods" and chunk.page_number == 4 for chunk in chunks)
        assert all(chunk.content.startswith(prefix) for chunk in chunks)
    assert all(chunk.document_name == "limits.txt" for chunk in chunks)


@pytest.mark.parametrize("sectioned", [False, True])
@pytest.mark.parametrize("target", [1200, 2000])
def test_small_leading_paragraph_cannot_merge_past_maximum(sectioned, target):
    content = "word " * 70 + "\n\n" + "long " * 295
    chunker = SemanticChunker(target_chunk_size=target, overlap_size=0, include_context=False)
    chunks = chunker.chunk(_document(content, sectioned))
    assert len(chunks) >= 2
    assert all(len(chunk.content) <= chunker.max_chunk_size for chunk in chunks)
    _assert_source(chunks, content, sectioned)


@pytest.mark.parametrize("sectioned", [False, True])
@pytest.mark.parametrize("leading_words", [400, 500])
def test_default_overlap_is_dropped_when_next_paragraph_needs_the_budget(sectioned, leading_words):
    content = "a " * leading_words + "\n\n" + "b " * 100 + "\n\n" + "c " * 700
    chunker = SemanticChunker(include_context=False)
    chunks = chunker.chunk(_document(content, sectioned))
    assert all(len(chunk.content) <= chunker.max_chunk_size for chunk in chunks)
    _assert_source(chunks, content, sectioned)


def test_overlap_is_retained_when_header_and_new_paragraph_fit():
    content = "alpha " * 10 + "\n\n" + "beta " * 10
    document = _document(content, True)
    chunker = SemanticChunker(
        min_chunk_size=20,
        target_chunk_size=80,
        max_chunk_size=150,
        overlap_size=60,
        include_context=False,
    )
    chunks = chunker.chunk(document)
    assert len(chunks) == 2
    assert all(len(chunk.content) <= chunker.max_chunk_size for chunk in chunks)
    assert (
        chunks[1].content
        == "[Methods]\n\n" + ("alpha " * 10).strip() + "\n\n" + ("beta " * 10).strip()
    )


@pytest.mark.parametrize("extra", [0, 1])
def test_exact_joined_paragraph_length_includes_header_and_separator(extra):
    content = "a" * 60 + "\n\n" + "b" * (77 + extra)
    chunker = SemanticChunker(
        min_chunk_size=20,
        target_chunk_size=200,
        max_chunk_size=150,
        overlap_size=0,
        include_context=False,
    )
    chunks = chunker.chunk(_document(content, True))
    assert len(chunks) == 1 + extra
    assert all(len(chunk.content) <= chunker.max_chunk_size for chunk in chunks)
    _assert_source(chunks, content, True)


@pytest.mark.parametrize("title", ["H" * 146, "H" * 147])
def test_section_header_without_any_content_budget_is_rejected(title):
    document = ParsedDocument(
        id="header", content="word", sections=[Section(title=title, content="word", level=1)]
    )
    chunker = SemanticChunker(
        min_chunk_size=20, target_chunk_size=100, max_chunk_size=150, overlap_size=0
    )
    with pytest.raises(ValueError, match="section header.*max_chunk_size"):
        chunker.chunk(document)


def test_header_and_word_can_use_exactly_the_maximum_budget():
    title = "H" * 145
    document = ParsedDocument(
        id="exact", content="x", sections=[Section(title=title, content="x", level=1)]
    )
    chunker = SemanticChunker(
        min_chunk_size=20, target_chunk_size=100, max_chunk_size=150, overlap_size=0
    )
    chunks = chunker.chunk(document)
    assert len(chunks) == 1
    assert len(chunks[0].content) == 150
    assert chunks[0].content == "[" + title + "]\n\nx"


@pytest.mark.parametrize("sectioned", [False, True])
def test_indivisible_oversized_word_is_rejected_without_returning_partial_chunks(sectioned):
    content = "Valid first paragraph.\n\n" + "x" * 1501
    chunker = SemanticChunker(include_context=False)
    with pytest.raises(ValueError, match="word length exceeds available chunk size"):
        chunker.chunk(_document(content, sectioned))


def test_header_reduces_the_available_word_budget():
    document = _document("x" * 141, True)
    chunker = SemanticChunker(
        min_chunk_size=20, target_chunk_size=100, max_chunk_size=150, overlap_size=0
    )
    with pytest.raises(ValueError, match="word length exceeds available chunk size"):
        chunker.chunk(document)


@pytest.fixture
def source_environment():
    source_dir = Path(__file__).resolve().parents[1] / "src"
    python_path = os.pathsep.join([str(source_dir), os.environ.get("PYTHONPATH", "")])
    return os.environ | {"PYTHONPATH": python_path}


class TestSettings:
    def test_api_keys_are_excluded_from_settings_representations(self, monkeypatch):

        key = "test-only-private-key"
        monkeypatch.setenv("NEXUSRAG_API_KEY", key)
        settings = Settings()

        assert settings.api.api_key == key
        assert key not in repr(settings)
        assert key not in repr(settings.api)
        assert key not in settings.model_dump_json()
        assert "api_key" not in settings.model_dump()["api"]

    def test_default_settings(self):

        settings = Settings()

        assert settings.llm.model == "llama3.2:3b"
        assert settings.llm.base_url == "http://localhost:11434"
        assert settings.embedding.model == "BAAI/bge-small-en-v1.5"
        assert settings.ingestion.chunk_size == 1200
        assert settings.retrieval.top_k == 8
        assert settings.self_correction.enabled is True
        assert settings.log_level == "INFO"

    def test_nested_settings_types(self):

        settings = Settings()

        assert isinstance(settings.llm, LLMSettings)
        assert isinstance(settings.embedding, EmbeddingSettings)
        assert isinstance(settings.ingestion, IngestionSettings)
        assert isinstance(settings.retrieval, RetrievalSettings)
        assert isinstance(settings.self_correction, SelfCorrectionSettings)
        assert isinstance(settings.storage, StorageSettings)
        assert isinstance(settings.api, APISettings)

    def test_env_override_llm(self, monkeypatch):
        monkeypatch.setenv("OLLAMA_BASE_URL", "http://custom:8080")
        monkeypatch.setenv("LLM_MODEL", "mistral:7b")

        llm = LLMSettings()

        assert llm.base_url == "http://custom:8080"
        assert llm.model == "mistral:7b"

    def test_env_override_embedding(self, monkeypatch):
        monkeypatch.setenv("EMBEDDING_MODEL", "custom-embedding-model")
        monkeypatch.setenv("EMBEDDING_DEVICE", "cuda")

        embedding = EmbeddingSettings()

        assert embedding.model == "custom-embedding-model"
        assert embedding.device == "cuda"

    def test_env_override_log_level(self, monkeypatch):
        monkeypatch.setenv("LOG_LEVEL", "DEBUG")

        settings = Settings()

        assert settings.log_level == "DEBUG"

    def test_storage_path_types(self):

        settings = Settings()

        assert isinstance(settings.storage.lancedb_path, Path)
        assert isinstance(settings.data_dir, Path)

    def test_api_settings(self):

        api = APISettings()

        assert api.host == "127.0.0.1"
        assert api.port == 8000

    def test_ingestion_settings_values(self):

        ingestion = IngestionSettings()

        assert ingestion.chunk_size > 0
        assert ingestion.chunk_overlap >= 0
        assert ingestion.chunk_overlap < ingestion.chunk_size
        assert ingestion.min_chunk_size > 0

    def test_retrieval_settings_values(self):

        retrieval = RetrievalSettings()

        assert retrieval.top_k > 0
        assert retrieval.max_query_length > 0

    def test_self_correction_settings(self):

        correction = SelfCorrectionSettings()

        assert correction.enabled is True
        assert 0 <= correction.confidence_tau <= 1
        assert correction.feedback_docs > 0 and correction.feedback_terms > 0

    def test_get_settings_caching(self):

        get_settings.cache_clear()

        settings1 = get_settings()
        settings2 = get_settings()

        assert settings1 is settings2

    def test_nested_llm_settings_read_environment(self, monkeypatch):

        monkeypatch.setenv("OLLAMA_BASE_URL", "http://env-file:1234")
        monkeypatch.setenv("LLM_MODEL", "from-env-file")

        settings = Settings()

        assert settings.llm.base_url == "http://env-file:1234"
        assert settings.llm.model == "from-env-file"


class TestSettingsPrecedence:
    def test_env_var_overrides_default(self, monkeypatch):

        monkeypatch.setenv("LLM_MODEL", "env-wins")
        assert Settings().llm.model == "env-wins"

    def test_llm_settings_ignore_bare_env_names(self, monkeypatch):

        monkeypatch.setenv("TEMPERATURE", "0.99")
        monkeypatch.setenv("TIMEOUT", "999")
        s = LLMSettings()
        assert s.temperature == 0.1
        assert s.timeout == 60

        monkeypatch.setenv("LLM_TEMPERATURE", "0.7")
        monkeypatch.setenv("LLM_TIMEOUT", "33")
        s = LLMSettings()
        assert s.temperature == 0.7
        assert s.timeout == 33


class TestSettingsValidation:
    @pytest.mark.parametrize(
        "env_name",
        [
            "LLM_MAX_TOKENS",
            "LLM_TIMEOUT",
            "EMBEDDING_BATCH_SIZE",
            "INGESTION_CHUNK_SIZE",
            "INGESTION_MIN_CHUNK_SIZE",
            "INGESTION_MAX_CHUNKS",
            "RETRIEVAL_TOP_K",
            "API_QUERY_RATE_PER_MINUTE",
            "API_UPLOAD_RATE_PER_MINUTE",
        ],
    )
    @pytest.mark.parametrize("value", ["0", "-1"])
    def test_invalid_runtime_limits_rejected_from_environment(self, monkeypatch, env_name, value):

        monkeypatch.setenv(env_name, value)
        with pytest.raises(ValidationError):
            Settings()

    @pytest.mark.parametrize(
        "overrides",
        [{"chunk_overlap": -1}, {"chunk_overlap": 100}, {"min_chunk_size": 101}],
    )
    def test_invalid_chunk_relationships_rejected(self, overrides):

        kwargs = {"chunk_size": 100, "chunk_overlap": 0, "min_chunk_size": 20, **overrides}
        with pytest.raises(ValidationError):
            IngestionSettings(**kwargs)

    def test_zero_overlap_and_equal_minimum_are_valid(self):

        assert IngestionSettings(chunk_size=100, chunk_overlap=0, min_chunk_size=100)
        assert LLMSettings(temperature=0).temperature == 0

    @pytest.mark.parametrize("field", ["confidence_tau", "grounding_threshold"])
    @pytest.mark.parametrize("value", [-0.1, 1.1])
    def test_invalid_probability_thresholds_rejected(self, field, value):

        with pytest.raises(ValidationError):
            SelfCorrectionSettings(**{field: value})

    def test_embedding_revision_matches_selected_model(self, monkeypatch):

        model = "sentence-transformers/all-MiniLM-L6-v2"
        monkeypatch.setenv("EMBEDDING_MODEL", model)
        monkeypatch.delenv("EMBEDDING_REVISION", raising=False)

        assert EmbeddingSettings().revision == HF_REVISIONS[model]
        assert EmbeddingSettings(model="custom-model").revision is None

    def test_explicit_embedding_revision_is_preserved(self, monkeypatch):

        monkeypatch.setenv("EMBEDDING_REVISION", "custom-sha")

        assert EmbeddingSettings(model="custom-model").revision == "custom-sha"

    def test_invalid_log_level(self, monkeypatch):
        monkeypatch.setenv("LOG_LEVEL", "INVALID")

        with pytest.raises(ValidationError):
            Settings()

    def test_invalid_device(self, monkeypatch):
        monkeypatch.setenv("EMBEDDING_DEVICE", "tpu")

        with pytest.raises(ValidationError):
            EmbeddingSettings()


class TestSettingsLoading:
    def test_env_file_in_cwd_is_loaded(self, tmp_path, source_environment):
        (tmp_path / ".env").write_text("LLM_MODEL=dotenv-model\n")
        env = {k: v for k, v in source_environment.items() if k != "LLM_MODEL"}
        out = subprocess.run(
            [
                sys.executable,
                "-c",
                "from scinexusrag.settings import Settings; print(Settings().llm.model)",
            ],
            cwd=tmp_path,
            env=env,
            capture_output=True,
            text=True,
            check=True,
        )
        assert out.stdout.strip() == "dotenv-model"

    def test_model_revisions_do_not_validate_unrelated_settings(self, tmp_path, source_environment):

        result = subprocess.run(
            [
                sys.executable,
                "-c",
                "from scinexusrag.settings import HF_REVISIONS; "
                "print(HF_REVISIONS['BAAI/bge-small-en-v1.5'])",
            ],
            cwd=tmp_path,
            env=source_environment | {"LLM_MAX_TOKENS": "0"},
            capture_output=True,
            text=True,
            check=True,
        )
        assert result.stdout.strip() == HF_REVISIONS["BAAI/bge-small-en-v1.5"]


def test_index_retains_its_own_corpus_when_the_input_list_changes():
    chunks = [
        Chunk(id="first", document_id="paper", content="kinase inhibition"),
        Chunk(id="second", document_id="paper", content="cardiac disease"),
    ]
    retriever = BM25Retriever()
    retriever.add(chunks)
    chunks.clear()

    assert retriever.count() == 2
    assert [result.chunk.id for result in retriever.retrieve("cardiac")] == ["second"]


def test_index_copies_passage_text_identity_and_nested_metadata():
    chunk = Chunk(
        id="original",
        document_id="paper",
        content="kinase inhibition",
        metadata={"units": ["mA"]},
    )
    retriever = BM25Retriever()
    retriever.add([chunk])
    chunk.id = "changed"
    chunk.content = "cardiac disease"
    chunk.metadata["units"].append("mV")

    result = retriever.retrieve("kinase")[0]

    assert result.chunk.id == "original"
    assert result.chunk.content == "kinase inhibition"
    assert result.chunk.metadata == {"units": ["mA"]}
    assert retriever.retrieve("cardiac") == []


def test_mutating_a_result_does_not_change_future_results_or_removal():
    retriever = BM25Retriever()
    retriever.add(
        [
            Chunk(
                id="original",
                document_id="paper",
                content="kinase inhibition",
                metadata={"units": ["mA"]},
            )
        ]
    )
    result = retriever.retrieve("kinase")[0]
    result.chunk.id = "changed"
    result.chunk.content = "cardiac disease"
    result.chunk.metadata["units"].append("mV")

    next_result = retriever.retrieve("kinase")[0]

    assert next_result.chunk.id == "original"
    assert next_result.chunk.content == "kinase inhibition"
    assert next_result.chunk.metadata == {"units": ["mA"]}
    assert retriever.remove({"original"}) == 1
    assert retriever.count() == 0


def test_incremental_rebuild_copies_new_chunks():
    retriever = BM25Retriever()
    retriever.add([Chunk(id="first", document_id="paper", content="kinase inhibition")])
    fresh = Chunk(id="second", document_id="paper", content="cardiac disease")
    retriever.add_incremental([fresh])
    fresh.id = "changed"
    fresh.content = "unrelated text"

    assert [result.chunk.id for result in retriever.retrieve("cardiac")] == ["second"]


def test_empty_corpus_replacement_clears_old_passages():
    retriever = BM25Retriever()
    retriever.add([Chunk(id="old", document_id="paper", content="kinase inhibition")])

    assert retriever.add([]) == 0
    assert retriever.count() == 0
    assert retriever.retrieve("kinase") == []


def test_empty_incremental_batch_keeps_existing_passages():
    retriever = BM25Retriever()
    retriever.add([Chunk(id="old", document_id="paper", content="kinase inhibition")])

    assert retriever.add_incremental([]) == 1
    assert retriever.count() == 1
    assert [result.chunk.id for result in retriever.retrieve("kinase")] == ["old"]


def test_duplicate_passage_ids_are_rejected_without_changing_the_old_index():
    retriever = BM25Retriever()
    retriever.add([Chunk(id="old", document_id="paper", content="kinase inhibition")])
    duplicate_ids = [
        Chunk(id="duplicate", document_id="paper", content="cardiac disease"),
        Chunk(id="duplicate", document_id="paper", content="neural response"),
    ]

    with pytest.raises(ValueError, match="IDs must be unique"):
        retriever.add(duplicate_ids)

    assert retriever.count() == 1
    assert [result.chunk.id for result in retriever.retrieve("kinase")] == ["old"]
    assert retriever.retrieve("cardiac") == []


@pytest.mark.parametrize("operation", ["replace", "incremental", "remove"])
def test_failed_rebuild_keeps_the_previous_index(monkeypatch, operation):
    retriever = BM25Retriever()
    retriever.add(
        [
            Chunk(id="first", document_id="paper", content="kinase inhibition"),
            Chunk(id="second", document_id="paper", content="cardiac disease"),
        ]
    )
    before = retriever.retrieve("kinase cardiac")
    new_chunks = [Chunk(id="new", document_id="paper", content="neural response")]

    def fail(_corpus):
        raise RuntimeError("index construction failed")

    monkeypatch.setattr("scinexusrag.retrieval.sparse.BM25Okapi", fail)

    with pytest.raises(RuntimeError, match="index construction failed"):
        if operation == "replace":
            retriever.add(new_chunks)
        elif operation == "incremental":
            retriever.add_incremental(new_chunks)
        else:
            retriever.remove({"second"})

    assert retriever.count() == 2
    assert retriever.retrieve("kinase cardiac") == before


@pytest.mark.parametrize("chunk_ids", [set(), {"absent"}])
def test_removing_no_passages_does_not_rebuild(monkeypatch, chunk_ids):
    retriever = BM25Retriever()
    retriever.add([Chunk(id="old", document_id="paper", content="kinase inhibition")])

    def fail(_corpus):
        raise AssertionError("unchanged index must not be rebuilt")

    monkeypatch.setattr("scinexusrag.retrieval.sparse.BM25Okapi", fail)

    assert retriever.remove(chunk_ids) == 0
    assert retriever.count() == 1
    assert [result.chunk.id for result in retriever.retrieve("kinase")] == ["old"]


def test_custom_stopword_mutation_does_not_change_an_existing_index():
    stopwords = set()
    retriever = BM25Retriever(stopwords=stopwords)
    retriever.add([Chunk(id="original", document_id="paper", content="kinase inhibition")])
    stopwords.add("kinase")

    assert [result.chunk.id for result in retriever.retrieve("kinase")] == ["original"]
    assert retriever.tokenize("kinase") == ["kinase"]


def test_custom_stopwords_are_copied_when_the_retriever_is_created():
    stopwords = {"kinase"}
    retriever = BM25Retriever(stopwords=stopwords)
    stopwords.clear()
    retriever.add([Chunk(id="original", document_id="paper", content="kinase inhibition")])

    assert retriever.tokenize("kinase inhibition") == ["inhibition"]
    assert retriever.retrieve("kinase") == []


def test_stopword_configuration_cannot_be_replaced_after_indexing():
    retriever = BM25Retriever(stopwords=frozenset())
    retriever.add([Chunk(id="original", document_id="paper", content="kinase inhibition")])

    with pytest.raises(AttributeError):
        retriever.stopwords = frozenset({"kinase"})

    assert [result.chunk.id for result in retriever.retrieve("kinase")] == ["original"]


@pytest.mark.parametrize("indexed", [False, True])
@pytest.mark.parametrize("top_k", [True, False, 1.5, "2", None])
def test_invalid_depth_is_rejected_even_with_an_empty_index(indexed, top_k):
    retriever = BM25Retriever()
    if indexed:
        retriever.add([Chunk(id="original", document_id="paper", content="kinase inhibition")])

    with pytest.raises(ValueError, match="top_k must be an integer"):
        retriever.retrieve("kinase", top_k=top_k)


@pytest.mark.parametrize("indexed", [False, True])
@pytest.mark.parametrize("query", [None, 42, b"kinase"])
def test_non_string_query_is_rejected_even_with_an_empty_index(indexed, query):
    retriever = BM25Retriever()
    if indexed:
        retriever.add([Chunk(id="original", document_id="paper", content="kinase inhibition")])

    with pytest.raises(ValueError, match="query must be a string"):
        retriever.retrieve(query)


@pytest.mark.parametrize("query,top_k", [("kinase", 0), ("kinase", -2), ("", 5), (" \n\t ", 5)])
def test_empty_requests_do_not_tokenize_or_score_the_index(monkeypatch, query, top_k):
    retriever = BM25Retriever()
    retriever.add([Chunk(id="original", document_id="paper", content="kinase inhibition")])

    def fail(_query):
        raise AssertionError("empty request must not be tokenized")

    monkeypatch.setattr(retriever, "tokenize", fail)

    assert retriever.retrieve(query, top_k=top_k) == []


@pytest.mark.parametrize("method", ["add", "add_incremental"])
@pytest.mark.parametrize(
    "chunks",
    [
        None,
        False,
        "",
        (),
        {},
        [None],
        ["passage"],
        [Chunk("first", None, "paper")],
        [Chunk(None, "passage", "paper")],
        [Chunk("", "passage", "paper")],
        [Chunk(" \t ", "passage", "paper")],
        [Chunk([], "passage", "paper")],
    ],
)
def test_malformed_corpus_inputs_never_change_the_existing_index(method, chunks):
    retriever = BM25Retriever()
    retriever.add(
        [
            Chunk("first", "kinase inhibition", "paper"),
            Chunk("second", "cardiac disease", "paper"),
        ]
    )
    before = retriever.retrieve("kinase cardiac")

    with pytest.raises(ValueError, match="chunks must be a list"):
        getattr(retriever, method)(chunks)

    assert retriever.count() == 2
    assert retriever.retrieve("kinase cardiac") == before


PIN = "49cf4c7b0db5b870a401ddf5e2669993ef3699c7"


@pytest.mark.parametrize("field", ["batch_size", "max_length"])
@pytest.mark.parametrize("value", [0, -1])
def test_invalid_encoding_parameters_fail_before_loading(monkeypatch, field, value):
    encode = Mock()
    monkeypatch.setattr(SpladeRetriever, "_encode", encode)

    with pytest.raises(ValueError, match="positive"):
        SpladeRetriever([], **{field: value})
    encode.assert_not_called()


@pytest.mark.parametrize("top_k", [0, -1, -99])
def test_nonpositive_depth_skips_query_encoding(monkeypatch, top_k):
    encode = Mock(return_value=sparse.csr_matrix([[1.0]]))
    monkeypatch.setattr(SpladeRetriever, "_encode", encode)
    retriever = SpladeRetriever([Chunk(id="doc", content="Evidence", document_id="doc")])

    assert retriever.retrieve("query", top_k=top_k) == []
    assert encode.call_count == 1


@pytest.mark.parametrize("top_k,expected", [(1, [1]), (2, [1, 2]), (99, [1, 2, 3, 0, 4])])
def test_tied_scores_keep_corpus_order_at_depth_boundary(monkeypatch, top_k, expected):
    encode = Mock(
        side_effect=[
            sparse.csr_matrix([[0.3], [0.7], [0.7], [0.7], [0.1]]),
            sparse.csr_matrix([[2.0]]),
        ]
    )
    monkeypatch.setattr(SpladeRetriever, "_encode", encode)
    chunks = [Chunk(id=str(i), content=str(i), document_id=str(i)) for i in range(5)]
    retriever = SpladeRetriever(chunks)
    results = retriever.retrieve("query", top_k=top_k)

    assert [result.chunk.id for result in results] == [str(i) for i in expected]
    assert results[0].score == pytest.approx(1.4)


@pytest.mark.parametrize(
    "model_name,revision,expected",
    [
        (DEFAULT_MODEL, None, PIN),
        (DEFAULT_MODEL, "explicit-pin", "explicit-pin"),
        ("custom/sparse-model", "custom-pin", "custom-pin"),
        ("custom/sparse-model", None, None),
    ],
)
def test_tokenizer_and_model_share_selected_revision(monkeypatch, model_name, revision, expected):
    tokenizer = Mock()
    model = Mock()
    model.to.return_value = model
    tokenizer_loader = Mock(return_value=tokenizer)
    model_loader = Mock(return_value=model)
    monkeypatch.setitem(
        sys.modules,
        "transformers",
        SimpleNamespace(
            AutoTokenizer=SimpleNamespace(from_pretrained=tokenizer_loader),
            AutoModelForMaskedLM=SimpleNamespace(from_pretrained=model_loader),
        ),
    )
    monkeypatch.setattr(SpladeRetriever, "_encode", lambda *_: sparse.csr_matrix((0, 0)))
    retriever = SpladeRetriever([], model_name=model_name, revision=revision)
    retriever._load()
    retriever._load()

    assert retriever.revision == expected
    tokenizer_loader.assert_called_once_with(model_name, revision=expected)
    model_loader.assert_called_once_with(model_name, revision=expected)
    model.to.assert_called_once_with("cpu")
    model.eval.assert_called_once()


@pytest.mark.parametrize(
    "model_name,revision,expected_model,expected_revision",
    [
        (None, None, DEFAULT_MODEL, PIN),
        ("custom/sparse-model", "custom-pin", "custom/sparse-model", "custom-pin"),
        ("custom/sparse-model", None, "custom/sparse-model", None),
    ],
)
def test_enabled_eval_records_revision_passed_to_retriever(
    monkeypatch, model_name, revision, expected_model, expected_revision
):
    dataset = datasets.IRDataset(
        name="fixture",
        corpus={"doc": {"title": "", "text": "Scientific evidence."}},
        queries={"query": "Which evidence?"},
        qrels={"query": {"doc": 1}},
    )
    monkeypatch.setattr(datasets, "load", lambda *a, **k: dataset)
    monkeypatch.setattr(systems, "ExactDenseRetriever", lambda *a: object())
    constructed = []

    def fake_splade(chunks: list[Chunk], **kwargs):
        constructed.append(kwargs)
        return SimpleNamespace(retrieve=lambda *a, **k: [RetrievalResult(chunks[0], 1.0, "splade")])

    monkeypatch.setattr("scinexusrag.retrieval.splade.SpladeRetriever", fake_splade)

    def splade_only(*args, **kwargs):
        built = systems.build_systems(*args, **kwargs)
        return {"SPLADE": built["SPLADE"]}

    monkeypatch.setattr(run, "build_systems", splade_only)
    output = run.evaluate(
        use_sample=True,
        include_splade=True,
        splade_model=model_name,
        splade_revision=revision,
    )

    assert constructed == [
        {"model_name": expected_model, "revision": expected_revision, "device": "cpu"}
    ]
    assert output["provenance"]["splade_model"] == expected_model
    assert output["provenance"]["splade_revision"] == expected_revision
    assert output["systems"]["SPLADE"]["means"]["nDCG@10"] == 1.0


def test_cli_forwards_splade_model_and_revision(monkeypatch, tmp_path):
    evaluate = Mock(return_value={"dataset": "fixture"})
    monkeypatch.setattr(run, "evaluate", evaluate)
    monkeypatch.setattr(run, "RESULTS_DIR", tmp_path)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "scinexusrag-eval",
            "--splade",
            "--splade-model",
            "custom/sparse-model",
            "--splade-revision",
            "custom-pin",
            "--out",
            str(tmp_path / "result.json"),
        ],
    )
    run.main()

    assert evaluate.call_args.kwargs["include_splade"] is True
    assert evaluate.call_args.kwargs["splade_model"] == "custom/sparse-model"
    assert evaluate.call_args.kwargs["splade_revision"] == "custom-pin"


@pytest.mark.parametrize(
    "text,expected",
    [
        (
            "First claim. Second one! Third? Yes.",
            ["First claim.", "Second one!", "Third?", "Yes."],
        ),
        (
            "We ran the test. 3 trials passed.",
            ["We ran the test.", "3 trials passed."],
        ),
    ],
)
def test_sentence_boundaries_preserve_the_complete_text(text, expected):
    assert split_sentences(text) == expected


def test_grounding_and_chunker_share_sentence_boundaries():
    from scinexusrag.generation import grounding
    from scinexusrag.ingestion import chunker

    assert grounding.split_sentences is split_sentences
    assert chunker.SENTENCE_BOUNDARY is SENTENCE_BOUNDARY


@pytest.mark.parametrize("text", ["42", "5.", "-5", "no", "是", "A.", "نعم"])
def test_short_claims_are_preserved(text):
    assert split_sentences(f"  {text}\n") == [text]


@pytest.mark.parametrize("text", ["", " \n\t", "...", "?!", "—", "()"])
def test_empty_and_punctuation_only_text_has_no_sentences(text):
    assert split_sentences(text) == []


@pytest.mark.parametrize("text", [None, 42, ["claim"]])
def test_nontext_input_has_clear_error(text):
    with pytest.raises(ValueError, match="text must be a string"):
        split_sentences(text)


def test_citations_remain_with_their_claims():
    assert split_sentences("First claim [1].\n\nNext claim [2].") == [
        "First claim [1].",
        "Next claim [2].",
    ]
    assert split_sentences("First claim. [1] Next claim. [2]") == [
        "First claim. [1] Next claim. [2]"
    ]


def test_short_numeric_claims_are_each_checked_for_grounding(monkeypatch):
    verifier = GroundingVerifier()
    scores = Mock(side_effect=[[1.0], [0.0]])
    monkeypatch.setattr(verifier, "entailment_scores", scores)

    report = verifier.verify("42. 5.", ["The measured value is 42."])

    assert [item.sentence for item in report.sentences] == ["42.", "5."]
    assert report.faithfulness == 0.5
    assert report.unsupported == ["5."]
    assert scores.call_args_list == [
        call([("The measured value is 42.", "42.")]),
        call([("The measured value is 42.", "5.")]),
    ]
    assert verifier._model is None


@pytest.mark.parametrize(
    "first,second",
    [
        ('"First claim."', "Next claim."),
        ("“First claim.”", "Next claim."),
        ("‘First claim!’", "Next claim?"),
        ("'First claim?'", "3 trials passed."),
        ("“First claim.”", "“Second claim.”"),
        ('"First claim."', '"Second claim."'),
    ],
)
def test_quoted_sentences_keep_their_quotes(first, second):
    text = first + " \n\t " + second

    assert split_sentences(text) == [first, second]
    assert all(match.group().isspace() for match in SENTENCE_BOUNDARY.finditer(text))


@pytest.mark.parametrize(
    "text",
    [
        "Smith et al. found an effect.",
        "Several methods, e.g. a baseline, were tested.",
        'Several methods, e.g. "a baseline", were tested.',
        "The measured value was 3.14 mg.",
        '"First claim." [1] Next claim. [2]',
    ],
)
def test_existing_abbreviation_decimal_and_citation_grouping_is_preserved(text):
    assert split_sentences(text) == [text]


def test_quoted_sentence_chunking_keeps_exact_source_offsets():
    first, second = '"First complete claim."', '"Second complete claim."'
    content = first + "  \n\t" + second
    chunker = SemanticChunker(
        target_chunk_size=32,
        max_chunk_size=32,
        min_chunk_size=1,
        overlap_size=0,
        context_chars=100,
    )

    chunks = chunker.chunk(ParsedDocument(id="quoted-claims", content=content))

    assert [chunk.content for chunk in chunks] == [first, second]
    assert [(chunk.start_offset, chunk.end_offset) for chunk in chunks] == [
        (0, len(first)),
        (content.index(second), len(content)),
    ]
    assert all(content[chunk.start_offset : chunk.end_offset] == chunk.content for chunk in chunks)
    assert chunks[0].context_before == ""
    assert chunks[0].context_after == second
    assert chunks[1].context_before == first
    assert chunks[1].context_after == ""


def _config(dim: int = 4, **changes) -> EmbeddingConfig:
    return replace(EmbeddingConfig("test-model", "a" * 40, "", "", True, dim), **changes)


def _store(tmp_path, dim: int = 4, **changes) -> VectorStore:
    return VectorStore(path=tmp_path / "lancedb", embedding_config=_config(dim, **changes))


@pytest.fixture
def vector_store(tmp_path) -> VectorStore:
    return _store(tmp_path)


class _ObservedLock:
    def __init__(self, lock):
        self.lock = lock
        self.contender = None
        self.attempted = Event()

    def __enter__(self):
        if get_ident() == self.contender:
            self.attempted.set()
        if not self.lock.acquire(timeout=5):
            raise TimeoutError("Vector store lock was not released")
        return self

    def __exit__(self, *_args):
        self.lock.release()

    def run_contender(self, operation):
        self.contender = get_ident()
        return operation()


def test_search_scores_are_cosine_similarity(vector_store) -> None:

    chunks = [
        Chunk(id="a", content="aligned unit", document_id="a"),
        Chunk(id="b", content="aligned scaled", document_id="b"),
        Chunk(id="c", content="orthogonal", document_id="c"),
    ]
    vectors = np.array(
        [[1, 0, 0, 0], [3, 0, 0, 0], [0, 1, 0, 0]],
        dtype=np.float32,
    )
    vector_store.add(chunks, vectors)

    results = {
        r.chunk.id: r.score
        for r in vector_store.search(np.array([1, 0, 0, 0], np.float32), top_k=3)
    }

    assert results["a"] > 0.99
    assert results["b"] > 0.99
    assert abs(results["c"]) < 0.01


def test_search_ranks_by_similarity(vector_store) -> None:
    chunks = [
        Chunk(id="near", content="near", document_id="near"),
        Chunk(id="far", content="far", document_id="far"),
    ]
    vectors = np.array([[1, 0.1, 0, 0], [0, 0, 1, 0]], dtype=np.float32)
    vector_store.add(chunks, vectors)

    ranked = vector_store.search(np.array([1, 0, 0, 0], np.float32), top_k=2)

    assert ranked[0].chunk.id == "near"
    assert ranked[0].score > ranked[1].score


def test_search_score_clamped_to_unit_interval(vector_store) -> None:

    chunks = [Chunk(id="opp", content="opposite", document_id="opp")]
    vector_store.add(chunks, np.array([[-1, 0, 0, 0]], dtype=np.float32))

    results = vector_store.search(np.array([1, 0, 0, 0], np.float32), top_k=1)

    assert results[0].score == 0.0


@pytest.mark.parametrize(
    "operation",
    [
        "delete",
        "delete_by_document",
        "search_by_document",
        "get_chunks_by_document",
        "get_chunk_ids_by_document",
    ],
)
def test_trailing_newline_id_is_rejected_before_database_access(vector_store, operation) -> None:

    with pytest.raises(ValueError, match="Invalid ID format"):
        if operation == "delete":
            vector_store.delete(["paper\n"])
        elif operation == "delete_by_document":
            vector_store.delete_by_document("paper\n")
        elif operation == "search_by_document":
            vector_store.search_by_document(np.ones(4, dtype=np.float32), "paper\n")
        else:
            getattr(vector_store, operation)("paper\n")

    assert vector_store._db is None


@pytest.mark.parametrize("access", ["search", "get_all_chunks", "get_chunks_by_document"])
def test_chunk_context_and_citation_metadata_survive_reload(tmp_path, access, vector_store) -> None:
    chunk = Chunk(
        id="context",
        content="The sample was measured.",
        document_id="paper",
        metadata={"file_type": "pdf"},
        document_name="Paper.pdf",
        section_title="Methods",
        page_number=3,
        chunk_index=7,
        context_before="The sample contained twelve participants.",
        context_after="All measurements were blinded.",
        start_offset=120,
        end_offset=144,
    )
    vector_store.add([chunk], np.array([[1, 0, 0, 0]], dtype=np.float32))

    reloaded = _store(tmp_path)
    if access == "search":
        restored = reloaded.search(np.array([1, 0, 0, 0], np.float32))[0].chunk
    elif access == "get_all_chunks":
        restored = reloaded.get_all_chunks()[0]
    else:
        restored = reloaded.get_chunks_by_document("paper")[0]

    assert restored.full_context == chunk.full_context
    assert restored.document_name == "Paper.pdf"
    assert restored.section_title == "Methods"
    assert restored.page_number == 3
    assert restored.chunk_index == 7
    assert (restored.start_offset, restored.end_offset) == (chunk.start_offset, chunk.end_offset)
    assert restored.metadata["start_offset"] == chunk.start_offset
    assert restored.metadata["end_offset"] == chunk.end_offset
    assert restored.metadata["file_type"] == "pdf"
    assert reloaded.list_documents() == ["paper"]


@pytest.mark.parametrize("existing", [False, True])
def test_legacy_connection_can_create_reopen_and_clear_tables(
    tmp_path, existing, vector_store
) -> None:
    vectors = np.array([[1, 0, 0, 0]], dtype=np.float32)
    if existing:
        vector_store.add(
            [Chunk(id="old", content="Original passage", document_id="paper")], vectors
        )

    class LegacyConnection:
        """Legacy interface."""

        def __init__(self, connection):
            self.connection = connection

        def table_names(self):
            return self.connection.list_tables().tables

        def __getattr__(self, name):
            if name == "list_tables":
                raise AttributeError(name)
            return getattr(self.connection, name)

    reloaded = _store(tmp_path)
    reloaded._db = LegacyConnection(vector_store.db)
    reloaded.add([Chunk(id="new", content="New passage", document_id="paper")], vectors)

    assert reloaded.count() == 1 + existing
    assert {result.chunk.id for result in reloaded.search(vectors[0])} == (
        {"old", "new"} if existing else {"new"}
    )
    reloaded.clear()
    assert "chunks" not in vector_store.db.list_tables().tables
    assert reloaded.count() == 0


@pytest.mark.parametrize("operation", ["open", "clear"])
def test_current_connection_finds_tables_on_later_pages(vector_store, operation) -> None:
    schema = vector_store.table.schema
    vector_store._table = None
    database = MagicMock(spec=["list_tables", "open_table", "create_table", "drop_table"])
    database.open_table.return_value.schema = schema
    database.open_table.return_value.count_rows.return_value = 0
    pages = {
        None: SimpleNamespace(tables=["another_table"], page_token="next_page"),
        "next_page": SimpleNamespace(tables=["chunks"], page_token=None),
    }
    database.list_tables.side_effect = lambda *, page_token: pages[page_token]
    vector_store._db = database

    if operation == "open":
        assert vector_store.table is database.open_table.return_value
        database.open_table.assert_called_once_with("chunks")
        database.create_table.assert_not_called()
    else:
        vector_store.clear()
        database.drop_table.assert_called_once_with("chunks")


def test_embedding_configuration_is_persisted_and_reopened(tmp_path) -> None:
    config = _config(query_prefix="query: ", doc_prefix="passage: ", normalize=False)
    vector_store = VectorStore(path=tmp_path / "lancedb", embedding_config=config)
    chunk = Chunk(id="stored", content="Stored passage", document_id="paper")
    vector_store.add([chunk], np.array([[1, 0, 0, 0]], dtype=np.float32))

    reloaded = VectorStore(path=vector_store.path, embedding_config=config)
    metadata = reloaded.table.schema.metadata

    assert json.loads(metadata[b"nexusrag.embedding_config"]) == {"version": 1, **asdict(config)}
    assert reloaded.embedding_dim == config.dimension
    assert reloaded.count() == 1
    assert reloaded.search(np.array([1, 0, 0, 0], dtype=np.float32))[0].chunk.id == "stored"


@pytest.mark.parametrize(
    "change",
    [
        {"model_name": "another-model"},
        {"revision": "b" * 40},
        {"query_prefix": "query: "},
        {"doc_prefix": "passage: "},
        {"normalize": False},
        {"dimension": 3},
    ],
    ids=["model", "revision", "query-prefix", "document-prefix", "normalization", "dimension"],
)
@pytest.mark.parametrize("operation", ["search", "add"])
def test_changed_configuration_rejects_before_vector_query_or_write(
    tmp_path, monkeypatch, change, operation
) -> None:
    original = _store(tmp_path)
    original.add(
        [Chunk(id="original", content="Original passage", document_id="paper")],
        np.ones((1, 4), dtype=np.float32),
    )
    table = original.table
    original_metadata = table.schema.metadata
    search = MagicMock(side_effect=AssertionError("queried incompatible vectors"))
    add = MagicMock(side_effect=AssertionError("wrote incompatible vectors"))
    monkeypatch.setattr(table, "search", search)
    monkeypatch.setattr(table, "add", add)
    reloaded = VectorStore(path=original.path, embedding_config=replace(_config(), **change))
    reloaded._db = MagicMock(wraps=original.db)
    reloaded._db.open_table.return_value = table

    for _ in range(2):
        with pytest.raises(EmbeddingConfigError):
            if operation == "search":
                reloaded.search(np.ones(reloaded.embedding_dim, dtype=np.float32))
            else:
                reloaded.add(
                    [Chunk(id="new", content="New passage", document_id="another")],
                    np.ones((1, reloaded.embedding_dim), dtype=np.float32),
                )
        assert reloaded._table is None

    search.assert_not_called()
    add.assert_not_called()
    assert table.count_rows() == 1
    assert table.schema.metadata == original_metadata


@pytest.mark.parametrize(
    "operation",
    [
        "count",
        "list_documents",
        "get_all_chunks",
        "get_chunks_by_document",
        "get_chunk_ids_by_document",
    ],
)
def test_read_helpers_propagate_configuration_errors(tmp_path, operation) -> None:
    original = _store(tmp_path)
    original.add(
        [Chunk(id="original", content="Original passage", document_id="paper")],
        np.ones((1, 4), dtype=np.float32),
    )
    changed = _store(tmp_path, revision="b" * 40)

    with pytest.raises(EmbeddingConfigError):
        if operation in {"get_chunks_by_document", "get_chunk_ids_by_document"}:
            getattr(changed, operation)("paper")
        else:
            getattr(changed, operation)()

    assert original.count() == 1


@pytest.mark.parametrize("error_type", [OSError, ValueError])
def test_count_propagates_storage_failure_instead_of_reporting_empty(
    vector_store, monkeypatch, error_type
) -> None:
    table = vector_store.table
    failure = error_type("Cannot read row count")
    monkeypatch.setattr(table, "count_rows", MagicMock(side_effect=failure))

    with pytest.raises(VectorStoreError, match="row count") as raised:
        vector_store.count()

    assert raised.value.__cause__ is failure


@pytest.mark.parametrize("chunk_count", [0, 3, 4, 5, 8])
def test_document_scans_cross_page_boundaries_without_losing_rows(
    vector_store, monkeypatch, chunk_count
) -> None:
    monkeypatch.setattr("scinexusrag.storage.vector_store.SCAN_PAGE_SIZE", 4)
    chunks = [
        Chunk(id=f"paper-{i}", content=f"Measurement {i}", document_id="paper")
        for i in range(chunk_count)
    ]
    unrelated = Chunk(id="other", content="Different document", document_id="other")
    vector_store.add([unrelated, *chunks], np.ones((chunk_count + 1, 4), dtype=np.float32))

    assert {chunk.id for chunk in vector_store.get_chunks_by_document("paper")} == {
        chunk.id for chunk in chunks
    }
    assert vector_store.get_chunk_ids_by_document("paper") == {chunk.id for chunk in chunks}
    assert {chunk.id for chunk in vector_store.get_all_chunks()} == {
        "other",
        *(c.id for c in chunks),
    }
    assert set(vector_store.list_documents()) == ({"other", "paper"} if chunks else {"other"})
    assert vector_store.get_chunks_by_document("missing") == []
    assert vector_store.get_chunk_ids_by_document("missing") == set()


def test_document_scans_include_chunks_beyond_the_previous_100000_limit(vector_store) -> None:
    count = 100_001
    ids = [f"chunk-{i}" for i in range(count)]
    vector_store.table.add(
        pa.table(
            {
                "id": ids,
                "content": ["Measured current is 42 mA."] * count,
                "document_id": ["large-paper"] * count,
                "metadata": ["{}"] * count,
                "vector": [[1.0, 0.0, 0.0, 0.0]] * count,
            },
            schema=vector_store.table.schema,
        )
    )

    assert vector_store.get_chunk_ids_by_document("large-paper") == set(ids)
    assert {chunk.id for chunk in vector_store.get_chunks_by_document("large-paper")} == set(ids)


@pytest.mark.parametrize(
    "operation",
    ["list_documents", "get_all_chunks", "get_chunks_by_document", "get_chunk_ids_by_document"],
)
def test_failed_scan_does_not_report_an_empty_or_partial_corpus(
    vector_store, monkeypatch, operation
) -> None:
    table = vector_store.table
    failure = OSError("Storage is unavailable")
    monkeypatch.setattr(table, "search", MagicMock(side_effect=failure))

    with pytest.raises(VectorStoreError, match="read passages") as raised:
        if operation in {"get_chunks_by_document", "get_chunk_ids_by_document"}:
            getattr(vector_store, operation)("paper")
        else:
            getattr(vector_store, operation)()

    assert raised.value.__cause__ is failure


def test_failure_on_a_later_scan_page_does_not_return_a_partial_document(
    vector_store, monkeypatch
) -> None:
    monkeypatch.setattr("scinexusrag.storage.vector_store.SCAN_PAGE_SIZE", 1)
    vector_store.add(
        [Chunk(id="a", content="First measurement", document_id="paper")],
        np.ones((1, 4), dtype=np.float32),
    )
    table = vector_store.table
    search = table.search
    calls = 0

    def fail_on_second_page(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise OSError("Cannot read the next page")
        return search(*args, **kwargs)

    monkeypatch.setattr(table, "search", fail_on_second_page)
    with pytest.raises(VectorStoreError, match="read passages"):
        vector_store.get_chunks_by_document("paper")

    assert calls == 2


def test_vector_search_failure_is_distinct_from_empty_results(vector_store, monkeypatch) -> None:
    failure = OSError("Storage is unavailable")
    monkeypatch.setattr(vector_store.table, "search", MagicMock(side_effect=failure))

    with pytest.raises(VectorStoreError, match="search the vector index") as raised:
        vector_store.search(np.ones(4, dtype=np.float32))

    assert raised.value.__cause__ is failure


@pytest.mark.parametrize("metadata", ["not JSON", "[]", "null"])
@pytest.mark.parametrize("operation", ["get_all_chunks", "get_chunks_by_document", "search"])
def test_corrupt_passage_metadata_is_reported_as_a_storage_error(
    vector_store, metadata, operation
) -> None:
    vector_store.table.add(
        [
            {
                "id": "corrupt",
                "content": "Stored measurement",
                "document_id": "paper",
                "metadata": metadata,
                "vector": [1, 0, 0, 0],
            }
        ]
    )

    with pytest.raises(VectorStoreError, match="decode passage metadata"):
        if operation == "get_chunks_by_document":
            vector_store.get_chunks_by_document("paper")
        elif operation == "search":
            vector_store.search(np.ones(4, dtype=np.float32))
        else:
            vector_store.get_all_chunks()


@pytest.mark.parametrize("populated", [False, True])
@pytest.mark.parametrize(
    "metadata_kind", ["missing", "incomplete", "corrupt", "unsupported-version"]
)
def test_unbound_or_invalid_tables_are_never_adopted(
    tmp_path, populated, metadata_kind, vector_store
) -> None:
    schema = vector_store.table.schema.remove_metadata()
    vector_store.clear()
    config_data = {"version": 1, **asdict(_config())}
    if metadata_kind == "incomplete":
        del config_data["doc_prefix"]
    elif metadata_kind == "unsupported-version":
        config_data["version"] = 2
    if metadata_kind != "missing":
        encoded = b"not JSON" if metadata_kind == "corrupt" else json.dumps(config_data).encode()
        schema = schema.with_metadata({b"nexusrag.embedding_config": encoded})
    raw_table = vector_store.db.create_table("chunks", schema=schema)
    if populated:
        raw_table.add(
            [
                {
                    "id": "old",
                    "content": "Old passage",
                    "document_id": "paper",
                    "metadata": "{}",
                    "vector": [1, 0, 0, 0],
                }
            ]
        )
    reloaded = _store(tmp_path)

    with pytest.raises(EmbeddingConfigError):
        reloaded.count()

    assert raw_table.count_rows() == int(populated)
    assert raw_table.schema.metadata == schema.metadata
    assert reloaded._table is None


def test_cached_table_configuration_is_checked_on_every_access(vector_store) -> None:
    table = vector_store.table
    changed_schema = table.schema.with_metadata(
        {
            b"nexusrag.embedding_config": json.dumps(
                {"version": 1, **asdict(_config(revision="b" * 40))}
            ).encode()
        }
    )
    cached = MagicMock(wraps=table)
    cached.schema = changed_schema
    vector_store._table = cached

    with pytest.raises(EmbeddingConfigError):
        vector_store.search(np.ones(4, dtype=np.float32))

    cached.search.assert_not_called()


def test_vector_schema_dimension_must_match_verified_metadata(vector_store) -> None:
    schema = vector_store.table.schema
    vector_store.clear()
    incompatible = schema.set(
        schema.get_field_index("vector"), pa.field("vector", pa.list_(pa.float32(), 2))
    )
    vector_store.db.create_table("chunks", schema=incompatible)

    with pytest.raises(EmbeddingConfigError):
        vector_store.count()

    assert vector_store._table is None


def test_active_embedder_configuration_must_match_store(vector_store) -> None:
    vector_store.validate_embedding_config(_config())

    with pytest.raises(EmbeddingConfigError):
        vector_store.validate_embedding_config(_config(query_prefix="different: "))

    assert vector_store.count() == 0


@pytest.mark.parametrize("document_id", [None, "paper"])
def test_dense_retrieval_checks_active_configuration_before_embedding(
    vector_store, document_id
) -> None:
    embedder = MagicMock()
    embedder.embedding_config = _config(query_prefix="another query prefix: ")
    retriever = DenseRetriever(embedder, vector_store)

    with pytest.raises(EmbeddingConfigError):
        retriever.retrieve("Which measurement?", document_id=document_id)

    embedder.embed_query.assert_not_called()
    assert vector_store._db is None


@pytest.mark.parametrize("legacy", [False, True])
@pytest.mark.parametrize("populated", [False, True])
def test_explicit_clear_recovers_rejected_table_with_a_new_dimension(
    tmp_path, legacy, populated
) -> None:
    original = _store(tmp_path)
    schema = original.table.schema
    if legacy:
        original.clear()
        raw_table = original.db.create_table("chunks", schema=schema.remove_metadata())
    else:
        raw_table = original.table
    if populated:
        raw_table.add(
            [
                {
                    "id": "old",
                    "content": "Old passage",
                    "document_id": "paper",
                    "metadata": "{}",
                    "vector": [1, 0, 0, 0],
                }
            ]
        )
    changed = _store(tmp_path, dim=2, model_name="new-model")

    with pytest.raises(EmbeddingConfigError):
        changed.count()

    changed.clear()
    changed.add(
        [Chunk(id="new", content="New passage", document_id="new-paper")],
        np.array([[1, 0]], dtype=np.float32),
    )
    reopened = _store(tmp_path, dim=2, model_name="new-model")
    assert reopened.count() == 1
    assert reopened.table.schema.field("vector").type.list_size == 2
    assert reopened.list_documents() == ["new-paper"]
    assert reopened.search(np.array([1, 0], dtype=np.float32))[0].chunk.id == "new"


@pytest.mark.parametrize("operation", ["count", "search", "add"])
@pytest.mark.parametrize("extra_old_write", [False, True], ids=["equal-version", "lower-version"])
def test_cached_handle_rejects_a_recreated_collection(tmp_path, operation, extra_old_write) -> None:
    original = _store(tmp_path)
    vector = np.array([[1, 0, 0, 0]], dtype=np.float32)
    original.add([Chunk(id="old", content="Old passage", document_id="old-paper")], vector)
    if extra_old_write:
        original.add(
            [Chunk(id="old-more", content="Another old passage", document_id="old-paper")], vector
        )
    original_version = original.table.version
    replacement = _store(tmp_path, model_name="another-model")
    replacement.clear()
    replacement.add([Chunk(id="new", content="New passage", document_id="new-paper")], vector)
    assert replacement.table.version <= original_version

    with pytest.raises(EmbeddingConfigError):
        if operation == "count":
            original.count()
        elif operation == "search":
            original.search(vector[0])
        else:
            original.add(
                [Chunk(id="stale", content="Stale writer", document_id="stale-paper")], vector
            )

    assert replacement.count() == 1
    assert replacement.list_documents() == ["new-paper"]


@pytest.mark.parametrize("bootstrap_stage", ["connect", "create_table"])
def test_concurrent_first_access_connects_and_creates_once(
    vector_store, monkeypatch, bootstrap_stage
):
    observed = _ObservedLock(vector_store._init_lock)
    vector_store._init_lock = observed
    entered = Event()
    release = Event()
    original = lancedb.connect if bootstrap_stage == "connect" else vector_store.db.create_table

    def held_bootstrap(*args, **kwargs):
        entered.set()
        if not release.wait(timeout=5):
            raise TimeoutError("Bootstrap was not released")
        return original(*args, **kwargs)

    bootstrap = MagicMock(side_effect=held_bootstrap)
    if bootstrap_stage == "connect":
        monkeypatch.setattr(lancedb, "connect", bootstrap)
    else:
        monkeypatch.setattr(vector_store.db, "create_table", bootstrap)

    with ThreadPoolExecutor(max_workers=2) as workers:
        first = workers.submit(vector_store.count)
        try:
            assert entered.wait(timeout=5)
            second = workers.submit(observed.run_contender, vector_store.count)
            assert observed.attempted.wait(timeout=5)
            assert bootstrap.call_count == 1
        finally:
            release.set()
        assert first.result(timeout=5) == second.result(timeout=5) == 0

    assert bootstrap.call_count == 1
    assert vector_store.db.list_tables().tables == ["chunks"]


@pytest.mark.parametrize("bootstrap_stage", ["connect", "create_table"])
def test_failed_bootstrap_releases_lock_and_can_retry(vector_store, monkeypatch, bootstrap_stage):
    original = lancedb.connect if bootstrap_stage == "connect" else vector_store.db.create_table

    def fail_once(*args, **kwargs):
        if bootstrap.call_count == 1:
            raise OSError("Storage temporarily unavailable")
        return original(*args, **kwargs)

    bootstrap = MagicMock(side_effect=fail_once)
    if bootstrap_stage == "connect":
        monkeypatch.setattr(lancedb, "connect", bootstrap)
    else:
        monkeypatch.setattr(vector_store.db, "create_table", bootstrap)

    with pytest.raises(VectorStoreError) as raised:
        vector_store.count()

    assert isinstance(raised.value.__cause__, OSError)
    assert vector_store._table is None
    if bootstrap_stage == "connect":
        assert vector_store._db is None
    assert vector_store.count() == 0
    assert bootstrap.call_count == 2
    assert (
        json.loads(vector_store.table.schema.metadata[b"nexusrag.embedding_config"])["version"] == 1
    )


def test_clear_waits_for_table_bootstrap_then_removes_the_published_table(
    vector_store, monkeypatch
):
    observed = _ObservedLock(vector_store._init_lock)
    vector_store._init_lock = observed
    entered = Event()
    release = Event()
    original_create = vector_store.db.create_table

    def held_create(*args, **kwargs):
        entered.set()
        if not release.wait(timeout=5):
            raise TimeoutError("Table creation was not released")
        return original_create(*args, **kwargs)

    monkeypatch.setattr(vector_store.db, "create_table", held_create)
    with ThreadPoolExecutor(max_workers=2) as workers:
        first = workers.submit(lambda: vector_store.table)
        try:
            assert entered.wait(timeout=5)
            cleared = workers.submit(observed.run_contender, vector_store.clear)
            assert observed.attempted.wait(timeout=5)
        finally:
            release.set()
        assert first.result(timeout=5) is not None
        assert cleared.result(timeout=5) is None

    assert vector_store._table is None
    assert "chunks" not in vector_store.db.list_tables().tables
    assert vector_store.count() == 0


def test_failed_clear_keeps_cached_data_and_allows_retry(vector_store, monkeypatch):
    vector_store.add(
        [Chunk(id="kept", content="Stored passage", document_id="paper")],
        np.ones((1, 4), dtype=np.float32),
    )
    original_drop = vector_store.db.drop_table
    attempts = 0

    def fail_once(*args, **kwargs):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise OSError("Cannot remove table")
        return original_drop(*args, **kwargs)

    monkeypatch.setattr(vector_store.db, "drop_table", fail_once)
    with pytest.raises(OSError):
        vector_store.clear()

    assert vector_store.count() == 1
    vector_store.clear()
    assert vector_store._table is None
    assert "chunks" not in vector_store.db.list_tables().tables


@pytest.fixture
def contract_vectors(tmp_path):
    return VectorStore(
        tmp_path / "vectors", embedding_config=EmbeddingConfig("model", "a" * 40, "", "", True, 2)
    )


@pytest.mark.parametrize(
    "vectors",
    [
        np.ones(2),
        np.ones((1, 3)),
        np.ones((2, 2)),
        np.array([[1.0, np.nan]]),
        np.array([[1.0, np.inf]]),
        np.array([[1.0, 1e100]]),
        np.zeros((1, 2)),
        np.array([[True, False]]),
        np.array([["1", "2"]]),
        np.ones((1, 2), dtype=complex),
    ],
)
def test_invalid_vectors_fail_before_opening_database(contract_vectors, vectors):
    with pytest.raises(ValueError):
        contract_vectors.add([Chunk("chunk", "passage", "paper")], vectors)

    assert contract_vectors._db is None


@pytest.mark.parametrize("field", ["id", "document_id"])
@pytest.mark.parametrize("value", [None, 1, "", "paper\n", "paper'", "a" * 129])
def test_invalid_chunk_ids_fail_before_opening_database(contract_vectors, field, value):
    chunk = Chunk("chunk", "passage", "paper")
    setattr(chunk, field, value)

    with pytest.raises(ValueError):
        contract_vectors.add([chunk], np.ones((1, 2), dtype=np.float32))

    assert contract_vectors._db is None


def test_duplicate_batch_ids_reject_entire_write(contract_vectors):
    chunks = [Chunk("same", "first", "paper"), Chunk("same", "second", "paper")]

    with pytest.raises(ValueError, match="Duplicate"):
        contract_vectors.add(chunks, np.ones((2, 2), dtype=np.float32))

    assert contract_vectors._db is None


@pytest.mark.parametrize("value", [np.nan, np.inf, -np.inf])
def test_nonfinite_metadata_is_not_persisted(contract_vectors, value):
    chunk = Chunk("chunk", "passage", "paper", metadata={"score": value})

    with pytest.raises(ValueError):
        contract_vectors.add([chunk], np.ones((1, 2), dtype=np.float32))

    assert contract_vectors._db is None


def test_empty_batch_does_not_open_database(contract_vectors):
    assert contract_vectors.add([], np.empty((0, 2), dtype=np.float32)) == 0
    assert contract_vectors._db is None


def test_vectors_without_chunks_are_rejected(contract_vectors):
    with pytest.raises(ValueError, match="shape"):
        contract_vectors.add([], np.ones((1, 2), dtype=np.float32))

    assert contract_vectors._db is None


def test_failed_vector_write_exposes_storage_error_with_original_cause(
    contract_vectors, monkeypatch
):
    failure = OSError("disk write failed")
    monkeypatch.setattr(contract_vectors.table, "add", Mock(side_effect=failure))

    with pytest.raises(VectorStoreError, match="write passages") as raised:
        contract_vectors.add(
            [Chunk("chunk", "passage", "paper")], np.ones((1, 2), dtype=np.float32)
        )

    assert raised.value.__cause__ is failure
    assert contract_vectors.count() == 0


@pytest.mark.parametrize(
    "vector",
    [
        np.ones((1, 2)),
        np.ones(3),
        np.array([np.nan, 1]),
        np.array([np.inf, 1]),
        np.array([1e100, 1]),
        np.zeros(2),
        np.array(["1", "2"]),
        np.array([True, False]),
    ],
)
def test_invalid_query_vectors_fail_before_opening_database(contract_vectors, vector):
    with pytest.raises(ValueError):
        contract_vectors.search(vector)

    assert contract_vectors._db is None


@pytest.mark.parametrize("top_k", [True, False, 1.5, "1", None])
def test_invalid_search_depth_fails_before_opening_database(contract_vectors, top_k):
    with pytest.raises(ValueError, match="top_k"):
        contract_vectors.search(np.ones(2, dtype=np.float32), top_k=top_k)

    assert contract_vectors._db is None


@pytest.mark.parametrize("top_k", [0, -1, -100])
def test_nonpositive_search_depth_does_not_open_database(contract_vectors, top_k):
    assert contract_vectors.search(np.ones(2, dtype=np.float32), top_k=top_k) == []
    assert contract_vectors._db is None


def search_with_row(contract_vectors, monkeypatch, row):
    query = Mock()
    query.metric.return_value = query
    query.limit.return_value = query
    query.to_list.return_value = [row]
    monkeypatch.setattr(contract_vectors.table, "search", Mock(return_value=query))
    return contract_vectors.search(np.ones(2, dtype=np.float32))


@pytest.mark.parametrize("distance", [None, True, "0.1", np.nan, np.inf, -np.inf])
def test_invalid_cosine_distance_is_not_a_confident_result(contract_vectors, monkeypatch, distance):
    row = {
        "id": "chunk",
        "content": "passage",
        "document_id": "paper",
        "metadata": "{}",
        "_distance": distance,
    }

    with pytest.raises(VectorStoreError, match="cosine distance"):
        search_with_row(contract_vectors, monkeypatch, row)


def test_missing_cosine_distance_is_a_storage_error(contract_vectors, monkeypatch):
    row = {"id": "chunk", "content": "passage", "document_id": "paper", "metadata": "{}"}

    with pytest.raises(VectorStoreError, match="cosine distance"):
        search_with_row(contract_vectors, monkeypatch, row)


@pytest.mark.parametrize("distance,expected", [(-1e-6, 1.0), (2.0, 0.0), (0.25, 0.75)])
def test_search_scores_stay_in_unit_interval(contract_vectors, monkeypatch, distance, expected):
    row = {
        "id": "chunk",
        "content": "passage",
        "document_id": "paper",
        "metadata": "{}",
        "_distance": distance,
    }

    results = search_with_row(contract_vectors, monkeypatch, row)

    assert results[0].score == expected
    assert type(results[0].score) is float


@pytest.mark.parametrize(
    "field,value", [("id", None), ("document_id", "invalid'"), ("content", None)]
)
def test_corrupt_stored_passage_fields_are_storage_errors(
    contract_vectors, monkeypatch, field, value
):
    row = {
        "id": "chunk",
        "content": "passage",
        "document_id": "paper",
        "metadata": "{}",
        "_distance": 0.0,
    }
    row[field] = value

    with pytest.raises(VectorStoreError, match="passage fields"):
        search_with_row(contract_vectors, monkeypatch, row)


def test_existing_ids_reject_entire_batch_without_overwriting(contract_vectors):
    contract_vectors.add(
        [Chunk("existing", "original passage", "original")], np.ones((1, 2), np.float32)
    )
    incoming = [Chunk("new", "new passage", "new"), Chunk("existing", "replacement", "new")]

    with pytest.raises(ValueError, match="already exists"):
        contract_vectors.add(incoming, np.ones((2, 2), np.float32))

    assert contract_vectors.count() == 1
    assert contract_vectors.get_all_chunks()[0].content == "original passage"


def test_existing_ids_are_checked_across_bounded_batches(contract_vectors, monkeypatch):
    monkeypatch.setattr("scinexusrag.storage.vector_store.SCAN_PAGE_SIZE", 2)
    contract_vectors.add(
        [Chunk("existing", "original passage", "original")], np.ones((1, 2), np.float32)
    )
    incoming = [Chunk(str(index), "new passage", "new") for index in range(4)]
    incoming.append(Chunk("existing", "replacement", "new"))

    with pytest.raises(ValueError, match="already exists"):
        contract_vectors.add(incoming, np.ones((5, 2), np.float32))

    assert contract_vectors.count() == 1
    assert contract_vectors.list_documents() == ["original"]


def test_failed_duplicate_check_never_writes_vectors(contract_vectors, monkeypatch):
    table = contract_vectors.table
    failure = OSError("cannot inspect existing ids")
    writer = Mock()
    monkeypatch.setattr(table, "search", Mock(side_effect=failure))
    monkeypatch.setattr(table, "add", writer)

    with pytest.raises(VectorStoreError, match="check existing IDs") as raised:
        contract_vectors.add([Chunk("new", "passage", "paper")], np.ones((1, 2), np.float32))

    assert raised.value.__cause__ is failure
    writer.assert_not_called()


@pytest.mark.parametrize("operation", ["delete", "delete_by_document"])
def test_deletion_count_ignores_unrelated_concurrent_writes(
    contract_vectors, monkeypatch, operation
):
    contract_vectors.add(
        [Chunk("target", "target passage", "paper"), Chunk("other", "other passage", "other")],
        np.ones((2, 2), dtype=np.float32),
    )
    writer = VectorStore(contract_vectors.path, embedding_config=contract_vectors.embedding_config)
    original_delete = contract_vectors.table.delete

    def delete_after_other_write(predicate):
        writer.add(
            [Chunk(f"new-{index}", "new passage", "other") for index in range(3)],
            np.ones((3, 2), dtype=np.float32),
        )
        return original_delete(predicate)

    monkeypatch.setattr(contract_vectors.table, "delete", delete_after_other_write)

    deleted = (
        contract_vectors.delete(["target"])
        if operation == "delete"
        else contract_vectors.delete_by_document("paper")
    )

    assert deleted == 1
    assert contract_vectors.count() == 4
    assert contract_vectors.get_chunk_ids_by_document("paper") == set()


def test_repeated_delete_ids_count_each_stored_row_once(contract_vectors):
    contract_vectors.add([Chunk("chunk", "passage", "paper")], np.ones((1, 2), dtype=np.float32))

    assert contract_vectors.delete(["chunk", "chunk", "missing"]) == 1
    assert contract_vectors.delete(["chunk"]) == 0
    assert contract_vectors.count() == 0


def test_partial_failed_deletion_reports_error_and_can_retry(contract_vectors, monkeypatch):
    contract_vectors.add(
        [Chunk("first", "first passage", "paper"), Chunk("second", "second passage", "paper")],
        np.ones((2, 2), dtype=np.float32),
    )
    table = contract_vectors.table
    original_delete = table.delete
    failure = OSError("delete interrupted")

    def partial_delete(_predicate):
        original_delete("id = 'first'")
        raise failure

    monkeypatch.setattr(table, "delete", partial_delete)
    with pytest.raises(VectorStoreError, match="delete passages") as raised:
        contract_vectors.delete_by_document("paper")

    assert raised.value.__cause__ is failure
    assert contract_vectors.get_chunk_ids_by_document("paper") == {"second"}
    monkeypatch.setattr(table, "delete", original_delete)
    assert contract_vectors.delete_by_document("paper") == 1
    assert contract_vectors.count() == 0


@pytest.mark.parametrize("mutation", ["delete", "add"])
def test_paginated_scans_keep_one_version_across_other_handles(
    contract_vectors, monkeypatch, mutation
):
    monkeypatch.setattr("scinexusrag.storage.vector_store.SCAN_PAGE_SIZE", 2)
    chunks = [Chunk(letter, f"passage {letter}", "paper") for letter in "abcd"]
    contract_vectors.add(chunks, np.ones((4, 2), dtype=np.float32))
    writer = VectorStore(contract_vectors.path, embedding_config=contract_vectors.embedding_config)
    scan = contract_vectors._read_rows(["id"])

    first = next(scan)["id"]
    if mutation == "delete":
        writer.delete([first])
    else:
        writer.add([Chunk("new", "new passage", "paper")], np.ones((1, 2), np.float32))
    remaining = [row["id"] for row in scan]

    assert [first, *remaining] == list("abcd")
    assert contract_vectors.count() == (3 if mutation == "delete" else 5)
    assert contract_vectors.get_chunk_ids_by_document("paper") == (
        set("bcd") if mutation == "delete" else {*"abcd", "new"}
    )


def test_closing_a_partial_scan_restores_writes(contract_vectors, monkeypatch):
    monkeypatch.setattr("scinexusrag.storage.vector_store.SCAN_PAGE_SIZE", 1)
    contract_vectors.add([Chunk("old", "old passage", "paper")], np.ones((1, 2), np.float32))
    scan = contract_vectors._read_rows(["id"])
    assert next(scan)["id"] == "old"

    scan.close()
    contract_vectors.add([Chunk("new", "new passage", "paper")], np.ones((1, 2), np.float32))

    assert contract_vectors.count() == 2


def test_snapshot_cleanup_failure_preserves_original_scan_error(contract_vectors, monkeypatch):
    table = contract_vectors.table
    original = OSError("scan failed")
    cleanup = OSError("restore failed")
    monkeypatch.setattr(table, "search", Mock(side_effect=original))
    monkeypatch.setattr(table, "checkout_latest", Mock(side_effect=cleanup))

    with pytest.raises(VectorStoreError, match="read passages") as raised:
        contract_vectors.get_all_chunks()

    assert raised.value.__cause__ is original
    assert contract_vectors._table is None
    assert contract_vectors.count() == 0


def test_snapshot_cleanup_failure_does_not_return_partial_success(contract_vectors, monkeypatch):
    table = contract_vectors.table
    cleanup = OSError("restore failed")
    monkeypatch.setattr(table, "checkout_latest", Mock(side_effect=cleanup))

    with pytest.raises(VectorStoreError, match="read passages") as raised:
        contract_vectors.get_all_chunks()

    assert raised.value.__cause__ is cleanup
    assert contract_vectors._table is None
    contract_vectors.add([Chunk("new", "new passage", "paper")], np.ones((1, 2), np.float32))
    assert contract_vectors.count() == 1


@pytest.mark.parametrize(
    "metadata",
    [
        {"context_before": None},
        {"context_after": ["passage"]},
        {"document_name": 7},
        {"section_title": False},
        {"page_number": "3"},
        {"page_number": -1},
        {"chunk_index": True},
        {"chunk_index": None},
        {"start_offset": -1},
        {"end_offset": 1.5},
        {"start_offset": 20, "end_offset": 10},
        {"nested": {"score": np.nan}},
        {"nested": {"score": np.inf}},
    ],
)
def test_invalid_persisted_chunk_metadata_is_a_storage_error(contract_vectors, metadata):
    contract_vectors.table.add(
        [
            {
                "id": "corrupt",
                "content": "passage",
                "document_id": "paper",
                "metadata": json.dumps(metadata),
                "vector": [1.0, 0.0],
            }
        ]
    )

    with pytest.raises(VectorStoreError, match="decode passage metadata"):
        contract_vectors.get_all_chunks()


@pytest.mark.parametrize(
    "field,value",
    [
        ("context_before", None),
        ("document_name", 3),
        ("page_number", "3"),
        ("chunk_index", True),
        ("start_offset", -1),
    ],
)
def test_invalid_chunk_metadata_never_opens_database(contract_vectors, field, value):
    chunk = Chunk("chunk", "passage", "paper")
    setattr(chunk, field, value)

    with pytest.raises(ValueError):
        contract_vectors.add([chunk], np.ones((1, 2), np.float32))

    assert contract_vectors._db is None


def test_zero_page_and_offsets_survive_round_trip(contract_vectors):
    chunk = Chunk("chunk", "passage", "paper", page_number=0, start_offset=0, end_offset=7)
    contract_vectors.add([chunk], np.ones((1, 2), np.float32))

    restored = contract_vectors.get_all_chunks()[0]

    assert restored.page_number == 0
    assert (restored.start_offset, restored.end_offset) == (0, 7)


@pytest.mark.parametrize("normalize", [False, True])
@pytest.mark.parametrize("stored_scale", [1e-30, 1.0, 1e20, 1e38])
@pytest.mark.parametrize("query_scale", [1e-30, 1.0, 1e20, 1e38])
def test_cosine_scores_do_not_depend_on_finite_vector_scale(
    contract_vectors, normalize, stored_scale, query_scale
):
    config = replace(contract_vectors.embedding_config, normalize=normalize)
    contract_vectors = VectorStore(contract_vectors.path, embedding_config=config)
    vectors = np.array([[stored_scale, 0], [0, stored_scale]], dtype=np.float32)
    query = np.array([query_scale, 0], dtype=np.float32)
    original_vectors, original_query = vectors.copy(), query.copy()
    chunks = [
        Chunk("aligned", "aligned passage", "paper"),
        Chunk("orthogonal", "other passage", "paper"),
    ]

    contract_vectors.add(chunks, vectors)
    results = {result.chunk.id: result.score for result in contract_vectors.search(query, top_k=2)}

    assert results["aligned"] == pytest.approx(1.0, abs=1e-6)
    assert results["orthogonal"] == pytest.approx(0.0, abs=1e-6)
    assert contract_vectors.embedding_config == config
    np.testing.assert_array_equal(vectors, original_vectors)
    np.testing.assert_array_equal(query, original_query)


@pytest.mark.parametrize("guard,status", [("auth", 401), ("origin", 403), ("rate", 429)])
async def test_upload_guards_run_before_multipart_parsing(
    local_settings, route_rag, monkeypatch, asgi_request, guard, status
):
    import starlette.formparsers

    local_settings.api.api_key = "test-only-key"
    local_settings.api.upload_rate_per_minute = 1
    request = httpx.Request(
        "POST", "http://localhost/api/ingest", files={"file": ("paper.txt", b"Evidence.")}
    )
    body = request.read()
    headers = [
        (b"host", b"localhost"),
        (b"content-type", request.headers["content-type"].encode()),
    ]
    if guard != "auth":
        headers.append((b"x-api-key", b"test-only-key"))
    if guard == "origin":
        headers.append((b"origin", b"https://foreign.example"))
    application = api.create_app()
    monkeypatch.setattr(security.limiter, "enabled", True)
    security.limiter.reset()
    try:
        if guard == "rate":
            route_rag.ingest_bytes.return_value = IngestResult("paper", "paper.txt", 1, 1, True)
            messages, consumed = await asgi_request(
                application, headers, [body], path="/api/ingest"
            )
            assert messages[0]["status"] == 200
            assert consumed == 1
            route_rag.ingest_bytes.reset_mock()
        spool = Mock(side_effect=AssertionError("Rejected upload created a temporary file"))
        monkeypatch.setattr(starlette.formparsers, "SpooledTemporaryFile", spool)
        messages, consumed = await asgi_request(application, headers, [body], path="/api/ingest")
        assert messages[0]["status"] == status
        assert consumed == 0
        spool.assert_not_called()
        route_rag.ingest_bytes.assert_not_called()
    finally:
        security.limiter.reset()


@pytest.mark.parametrize(
    "outcome,status",
    [
        ("success", 200),
        ("filename", 400),
        ("size", 413),
        ("failure", 500),
        ("field", 422),
        ("extra", 400),
    ],
)
def test_multipart_files_close_on_every_exit(
    application_client, route_rag, upload_limit_bytes, monkeypatch, outcome, status
):
    import starlette.formparsers

    temporary_files = []
    create_file = starlette.formparsers.SpooledTemporaryFile

    def track_file(*args, **kwargs):
        file = create_file(*args, **kwargs)
        temporary_files.append(file)
        return file

    monkeypatch.setattr(starlette.formparsers, "SpooledTemporaryFile", track_file)
    route_rag.ingest_bytes.return_value = IngestResult("paper", "paper.txt", 1, 1, True)
    if outcome == "failure":
        route_rag.ingest_bytes.side_effect = RuntimeError("Ingestion failed")
    content = b"x" * (upload_limit_bytes + 1) if outcome == "size" else b"Evidence."
    filename = "paper.xyz" if outcome == "filename" else "paper.txt"
    field = "other" if outcome == "field" else "file"
    files = [(field, (filename, content, "text/plain"))]
    if outcome == "extra":
        files.append(("file", ("extra.txt", b"More evidence.", "text/plain")))

    response = application_client.post("/api/ingest", files=files)

    assert response.status_code == status
    assert temporary_files
    assert all(file.closed for file in temporary_files)


def test_ingest_openapi_retains_multipart_upload(application_client):
    operation = application_client.get("/openapi.json").json()["paths"]["/api/ingest"]["post"]
    body = operation["requestBody"]
    schema = body["content"]["multipart/form-data"]["schema"]
    assert body["required"] is True
    assert schema["required"] == ["file"]
    assert schema["properties"]["file"]["type"] == "string"
    assert schema["properties"]["file"]["format"] == "binary"


@pytest.mark.parametrize("count,accepted", [(2, True), (3, False)])
def test_pdf_page_limit_precedes_extraction(tmp_path, monkeypatch, count, accepted):
    parser = DocumentParser()
    monkeypatch.setattr(parser, "MAX_PDF_PAGES", 2)
    extract = Mock(return_value="Evidence.")
    reader = SimpleNamespace(
        is_encrypted=False, pages=[SimpleNamespace(extract_text=extract)] * count
    )
    monkeypatch.setattr("scinexusrag.ingestion.parser.PdfReader", lambda path: reader)
    path = tmp_path / "pages.pdf"
    path.write_bytes(b"%PDF-1.4 fixture")
    if accepted:
        assert parser.parse(path).content == "Evidence.\n\nEvidence."
        assert extract.call_count == count
    else:
        with pytest.raises(DocumentParseError, match="page limit"):
            parser.parse(path)
        extract.assert_not_called()


@pytest.mark.parametrize(
    "texts,accepted,calls",
    [
        (["abc", "def", "g"], False, 2),
        (["      "], False, 1),
        (["a", "bc"], True, 2),
        (["ab", "cd"], False, 2),
    ],
)
def test_pdf_text_budget_counts_raw_text_and_join_separators(
    tmp_path, monkeypatch, texts, accepted, calls
):
    parser = DocumentParser()
    monkeypatch.setattr(parser, "MAX_EXTRACTED_CHARS", 5)
    extracts = [Mock(return_value=text) for text in texts]
    reader = SimpleNamespace(
        is_encrypted=False, pages=[SimpleNamespace(extract_text=extract) for extract in extracts]
    )
    monkeypatch.setattr("scinexusrag.ingestion.parser.PdfReader", lambda path: reader)
    path = tmp_path / "text.pdf"
    path.write_bytes(b"%PDF-1.4 fixture")
    if accepted:
        document = parser.parse(path)
        assert document.content == "a\n\nbc"
        _assert_section_spans(document)
    else:
        with pytest.raises(DocumentParseError, match="extracted text limit"):
            parser.parse(path)
    assert sum(extract.call_count for extract in extracts) == calls


@pytest.mark.parametrize("extension", [".txt", ".md", ".docx"])
@pytest.mark.parametrize("length", [5, 6])
def test_document_text_limit_applies_to_each_format(tmp_path, monkeypatch, extension, length):
    parser = DocumentParser()
    monkeypatch.setattr(parser, "MAX_EXTRACTED_CHARS", 5)
    path = tmp_path / ("text" + extension)
    if extension == ".docx":
        document = DocxDocument()
        document.add_paragraph("a" * length)
        document.save(path)
    else:
        path.write_text("a" * length, encoding="utf-8")
    if length == 5:
        assert parser.parse(path).content == "a" * length
    else:
        with pytest.raises(DocumentParseError, match="extracted text limit"):
            parser.parse(path)


def test_passage_limit_rejects_before_model_and_storage_initialization(test_settings):
    test_settings.ingestion.max_chunks = 1
    rag = NexusRAG(settings=test_settings)
    rag._chunker = Mock()
    rag._chunker.chunk.return_value = [Mock(), Mock()]

    result = rag.ingest_bytes(b"Evidence.", "paper.txt", ".txt")

    assert result.success is False
    assert "limit of 1 passages" in result.error
    assert rag._embedder is None
    assert rag._vector_store is None
    assert rag._document_store is None


def test_passage_limit_accepts_exact_boundary(pipeline_rag):
    pipeline_rag.settings.ingestion.max_chunks = 1
    result = pipeline_rag.ingest_bytes(b"Evidence.", "paper.txt", ".txt")
    assert result.success is True
    assert result.chunk_count == 1


@pytest.mark.parametrize("binary", [False, True])
def test_missing_file_error_excludes_unrelated_form_content(application_client, route_rag, binary):
    marker = "private-form-marker"
    response = application_client.post(
        "/api/ingest",
        files={"other": ("payload.bin", b"\xff" * (1024 * 1024 + 1))} if binary else None,
        data=None if binary else {"other": marker},
    )
    assert response.status_code == 422
    detail = response.json()["detail"]
    assert detail[0]["loc"] == ["body", "file"]
    assert all("input" not in error and "ctx" not in error for error in detail)
    assert marker not in response.text
    assert "SpooledTemporaryFile" not in response.text
    route_rag.ingest_bytes.assert_not_called()
